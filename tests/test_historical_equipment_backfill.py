from __future__ import annotations

import importlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from automacao_gd.infrastructure.persistence.atomic import atomic_copy_file


def _api():
    models = importlib.import_module("automacao_gd.domain.backfill_models")
    audit = importlib.import_module(
        "automacao_gd.application.historical_backfill.audit_service"
    )
    plan = importlib.import_module(
        "automacao_gd.application.historical_backfill.plan"
    )
    apply = importlib.import_module(
        "automacao_gd.application.historical_backfill.apply_service"
    )
    report = importlib.import_module(
        "automacao_gd.application.historical_backfill.report"
    )
    return models, audit, plan, apply, report


def _workbook(path: Path, rows_by_sheet: dict[str, list[tuple]], *, header_row: int = 1) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in rows_by_sheet.items():
        ws = wb.create_sheet(sheet_name)
        if header_row > 1:
            ws.cell(1, 1, "RELATÓRIO AUXILIAR")
        for column, value in enumerate(("Protocolo", "Placa", "Inversor", "Observação"), 1):
            ws.cell(header_row, column, value)
        for index, values in enumerate(rows, header_row + 1):
            for column, value in enumerate((*values, "PRESERVAR"), 1):
                ws.cell(index, column, value)
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:D{max(header_row + 1, ws.max_row)}"
        ws.column_dimensions["A"].width = 18
        ws.row_dimensions[header_row].height = 24
        ws["A1"].font = Font(bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        if rows:
            table = Table(
                displayName=f"Tabela{len(wb.sheetnames)}",
                ref=f"A{header_row}:D{header_row + len(rows)}",
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True
            )
            ws.add_table(table)
    wb.properties.title = "Fixture sintética"
    wb.save(path)


def _approved(module: str = "5x LEAPTON LP182", inverter: str = "1x HUAWEI SUN2000"):
    _, audit, _, _, _ = _api()
    return audit.TechnicalProposal(
        status="approved",
        module_text=module,
        inverter_text=inverter,
        source_hash="a" * 64,
        pdf_status="valid",
    )


def _resolver(mapping):
    return lambda protocol: mapping.get(
        protocol,
        _api()[1].TechnicalProposal(
            status="pdf_not_found", pdf_status="not_found"
        ),
    )


def test_audit_is_read_only_and_correct_rows_are_no_change(tmp_path: Path) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001012", "5x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    before = workbook.read_bytes()

    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001012": _approved()}),
    )

    assert workbook.read_bytes() == before
    assert result.items[0].action is models.BackfillAction.NO_CHANGE
    assert result.summary.total_updates == 0


def test_audit_proposes_both_equipment_cells_as_one_update(tmp_path: Path) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(workbook, {"2025": [("2600001013", "", "GOKIN AMBIGUO")]})

    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001013": _approved()}),
    )
    item = result.items[0]

    assert item.action is models.BackfillAction.PENDING_TECHNICAL_REVIEW
    assert item.proposed_module_text == "5x LEAPTON LP182"
    assert item.proposed_inverter_text == "1x HUAWEI SUN2000"
    assert {"MODULE_EMPTY", "GOKIN_IN_INVERTER", "EQUIPMENT_ITEM_LOST"}.issubset(
        item.reasons
    )
    assert item.expected_row_fingerprint


def test_audit_handles_multiple_sheets_header_offset_and_ignores_auxiliary(
    tmp_path: Path,
) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(
        workbook,
        {
            "2024": [("2600001014", "", "1x HUAWEI SUN2000")],
            "2025-2026": [("2600001015", "", "1x HUAWEI SUN2000")],
            "Dashboard": [("2600001016", "ANTIGA", "ANTIGO")],
        },
        header_row=3,
    )

    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver(
            {"2600001014": _approved(), "2600001015": _approved()}
        ),
    )

    assert result.summary.sheets_analyzed == 2
    assert {item.protocol for item in result.items} == {"2600001014", "2600001015"}
    assert all(
        item.action is models.BackfillAction.UPDATE_EQUIPMENT
        for item in result.items
    )


def test_missing_protocol_duplicate_and_missing_pdf_are_blocked(tmp_path: Path) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(
        workbook,
        {
            "2024": [("2600001017", "A", "B"), (None, "A", "B")],
            "2025": [("2600001017", "A", "B"), ("2600001018", "A", "B")],
        },
    )

    result = audit.audit_historical_workbook(workbook, technical_resolver=_resolver({}))
    actions = [item.action for item in result.items]

    assert actions.count(models.BackfillAction.DUPLICATE_PROTOCOL) == 2
    assert models.BackfillAction.PROTOCOL_MISSING in actions
    assert models.BackfillAction.PDF_NOT_FOUND in actions
    assert result.summary.total_updates == 0


@pytest.mark.parametrize("status", ["pending_review", "invalid_pdf", "ambiguous_pdf"])
def test_unapproved_pdf_never_proposes_update(tmp_path: Path, status: str) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / f"{status}.xlsx"
    _workbook(workbook, {"2025": [("2600001019", "A", "B")]})
    proposal = audit.TechnicalProposal(
        status=status,
        pdf_status=status,
        errors=("TECHNICAL_PENDING",),
    )

    result = audit.audit_historical_workbook(
        workbook, technical_resolver=_resolver({"2600001019": proposal})
    )

    assert result.items[0].action is models.BackfillAction.PENDING_TECHNICAL_REVIEW
    assert result.summary.total_updates == 0


def test_duplicate_solplanet_and_single_field_difference_are_updates(tmp_path: Path) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(
        workbook,
        {
            "2025": [
                (
                    "2600001020",
                    "5x LEAPTON LP182",
                    "SOLPLANET | ASW6000\nAISWEI | ASW6000",
                ),
                ("2600001021", "5x LEAPTON LP182", "INVERSOR ANTIGO"),
            ]
        },
    )
    resolver = _resolver(
        {
            "2600001020": _approved(inverter="1x SOLPLANET ASW6000"),
            "2600001021": _approved(),
        }
    )

    result = audit.audit_historical_workbook(workbook, technical_resolver=resolver)

    assert result.items[0].action is models.BackfillAction.UPDATE_EQUIPMENT
    assert result.items[1].action is models.BackfillAction.PENDING_TECHNICAL_REVIEW
    assert "SOLPLANET_ALIAS_NORMALIZED" in result.items[0].reasons
    assert (
        result.items[1].proposed_module_text
        == "LEAPTON | LP182\nQtd. total: 5 módulos"
    )


def test_audit_classifies_semantic_differences_without_using_them_as_source(
    tmp_path: Path,
) -> None:
    _, audit, _, _, _ = _api()
    workbook = tmp_path / "classificacoes.xlsx"
    _workbook(
        workbook,
        {
            "2025": [
                (
                    "2600001039",
                    "MICROINVERSOR | MODELO-ANTIGO",
                    "AISWEI | ASW5000\nAISWEI | ASW5000\nQtd. total: 2 inversores",
                )
            ]
        },
    )
    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver(
            {
                "2600001039": _approved(
                    module="5x LEAPTON LP182",
                    inverter="1x SOLPLANET ASW6000",
                )
            }
        ),
    )

    reasons = set(result.items[0].reasons)
    assert {
        "INVERTER_IN_MODULE",
        "SOLPLANET_ALIAS_NORMALIZED",
        "DUPLICATE_EQUIPMENT",
        "QUANTITY_DIVERGENT",
        "MODEL_DIVERGENT",
        "LEGACY_FORMAT",
    }.issubset(reasons)


def test_plan_is_deterministic_and_tampering_is_detected(tmp_path: Path) -> None:
    _, audit, plan, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(workbook, {"2025": [("2600001022", "A", "B")]})
    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001022": _approved()}),
    )

    first = plan.build_backfill_plan(result, created_at="2026-07-21T10:00:00Z")
    second = plan.build_backfill_plan(result, created_at="2026-07-21T10:00:00Z")
    assert first == second
    assert first["plan_hash"] == second["plan_hash"]
    assert plan.validate_backfill_plan(first) is None

    first["items"][0]["proposed_module_text"] = "ALTERADO"
    with pytest.raises(plan.BackfillPlanError, match="hash"):
        plan.validate_backfill_plan(first)


def test_plan_blocks_incompatible_versions_and_contains_no_paths(tmp_path: Path) -> None:
    _, audit, plan, _, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(workbook, {"2025": [("2600001023", "A", "B")]})
    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001023": _approved()}),
    )
    payload = plan.build_backfill_plan(result, created_at="2026-07-21T10:00:00Z")
    serialized = repr(payload)

    assert str(tmp_path) not in serialized
    assert ".xlsx" not in serialized
    payload["equipment_rules_version"] = "old"
    payload["plan_hash"] = plan.calculate_plan_hash(payload)
    with pytest.raises(plan.BackfillPlanError, match="regras"):
        plan.validate_backfill_plan(payload)


def _production_settings() -> SimpleNamespace:
    return SimpleNamespace(
        APP_ENV="production",
        DRY_RUN=False,
        APPLY_EXCEL=True,
        APPLY_ARCHIVE=False,
        BACKUP_EXCEL=True,
    )


class _ReadyPreflight:
    def raise_if_blocked(self) -> None:
        return None


def test_apply_requires_production_confirmation_and_valid_plan(tmp_path: Path) -> None:
    _, audit, plan, apply, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(workbook, {"2025": [("2600001024", "A", "B")]})
    audit_result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001024": _approved()}),
    )
    payload = plan.build_backfill_plan(audit_result, created_at="2026-07-21T10:00:00Z")

    bad_settings = _production_settings()
    bad_settings.APP_ENV = "development"
    with pytest.raises(apply.BackfillApplyError, match="production"):
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=bad_settings,
            confirmation=apply.required_apply_confirmation(payload),
        )
    with pytest.raises(apply.BackfillApplyError, match="(?i)confirmação"):
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation="SIM",
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )


def test_apply_creates_integral_backup_preserves_workbook_and_is_idempotent(
    tmp_path: Path,
) -> None:
    models, audit, plan, apply, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(
        workbook,
        {
            "2025": [
                ("2600001025", "1x LEAPTON LP182", "1x HUAWEI SUN2000"),
                ("2600001026", "OK", "OK"),
            ]
        },
    )
    wb = load_workbook(workbook)
    ws = wb["2025"]
    ws["D3"] = "=1+1"
    ws.merge_cells("E1:F1")
    ws["E1"] = "ESTRUTURA PRESERVADA"
    validation = DataValidation(type="list", formula1='"A,B"')
    ws.add_data_validation(validation)
    validation.add("D2:D3")
    ws.conditional_formatting.add(
        "D2:D3",
        CellIsRule(operator="equal", formula=['"PRESERVAR"']),
    )
    ws.print_area = "A1:F3"
    ws.page_setup.orientation = "landscape"
    ws["G1"], ws["G2"], ws["G3"] = "Série", 1, 2
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=7, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, "H2")
    wb.properties.subject = "Preservar propriedades"
    auxiliary = wb.create_sheet("Auxiliar")
    auxiliary.sheet_state = "hidden"
    auxiliary["A1"] = "NÃO ALTERAR"
    wb.save(workbook)
    wb.close()
    metadata_check = load_workbook(workbook, read_only=True)
    original_modified = metadata_check.properties.modified
    metadata_check.close()
    resolver = _resolver(
        {
            "2600001025": _approved(),
            "2600001026": _approved(module="OK", inverter="OK"),
        }
    )
    audit_result = audit.audit_historical_workbook(workbook, technical_resolver=resolver)
    payload = plan.build_backfill_plan(audit_result, created_at="2026-07-21T10:00:00Z")

    result = apply.apply_backfill_plan(
        workbook,
        payload,
        settings=_production_settings(),
        confirmation=apply.required_apply_confirmation(payload),
        preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
    )

    assert result.status is models.BackfillApplyStatus.APPLIED_SUCCESSFULLY
    assert result.backup_path is not None
    assert result.backup_path.read_bytes() != b""
    assert result.original_hash == result.backup_hash
    saved = load_workbook(workbook)
    assert saved["2025"]["B2"].value == "LEAPTON | LP182\nQtd. total: 5 módulos"
    assert saved["2025"]["C2"].value == "HUAWEI | SUN2000\nQtd. total: 1 inversor"
    assert saved["2025"]["B3"].value == "OK"
    assert saved["2025"]["D3"].value == "=1+1"
    assert saved["2025"].auto_filter.ref == "A1:D3"
    assert saved["2025"].column_dimensions["A"].width == 18
    assert saved.sheetnames == ["2025", "Auxiliar"]
    assert saved["Auxiliar"].sheet_state == "hidden"
    assert saved["Auxiliar"]["A1"].value == "NÃO ALTERAR"
    assert "E1:F1" in saved["2025"].merged_cells
    assert len(saved["2025"].data_validations.dataValidation) == 1
    assert len(saved["2025"].conditional_formatting) == 1
    assert saved["2025"].page_setup.orientation == "landscape"
    assert len(saved["2025"]._charts) == 1
    assert saved.properties.subject == "Preservar propriedades"
    assert saved.properties.modified == original_modified
    saved.close()

    second_audit = audit.audit_historical_workbook(workbook, technical_resolver=resolver)
    assert second_audit.items[0].action is models.BackfillAction.NO_CHANGE
    assert (
        second_audit.items[1].action
        is models.BackfillAction.PENDING_TECHNICAL_REVIEW
    )
    assert second_audit.summary.total_updates == 0


def test_row_conflict_is_not_modified(tmp_path: Path) -> None:
    models, audit, plan, apply, _ = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001027", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    audit_result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001027": _approved()}),
    )
    payload = plan.build_backfill_plan(audit_result, created_at="2026-07-21T10:00:00Z")
    wb = load_workbook(workbook)
    wb["2025"]["B2"] = "ALTERADA APÓS AUDITORIA"
    wb.save(workbook)

    with pytest.raises(apply.BackfillApplyError) as caught:
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )

    assert caught.value.result is not None
    assert caught.value.result.status is models.BackfillApplyStatus.ABORTED_PRE_FLIGHT
    assert caught.value.result.error_code == "WORKBOOK_CHANGED_AFTER_AUDIT"
    assert load_workbook(workbook)["2025"]["B2"].value == "ALTERADA APÓS AUDITORIA"


def test_public_report_is_allowlisted_and_redacts_paths(tmp_path: Path) -> None:
    _, audit, _, _, report = _api()
    workbook = tmp_path / "historico.xlsx"
    _workbook(workbook, {"2025": [("2600001028", r"C:\Pessoa\arquivo.pdf", "B")]})
    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001028": _approved()}),
    )

    public = report.build_audit_public_payload(result)
    serialized = repr(public)

    assert public["items"][0]["protocol"] == "2600001028"
    for forbidden in ("C:\\", "Z:\\", "\\\\servidor\\", "/home/", ".pdf", ".xlsx"):
        assert forbidden not in serialized


@pytest.mark.parametrize("protocol", ["2600001064", "2600001070", "2600001073"])
def test_known_anonymized_protocols_remain_pending_without_hardcode(
    tmp_path: Path, protocol: str
) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / f"{protocol}.xlsx"
    _workbook(workbook, {"2025": [(protocol, "A", "B")]})
    pending = audit.TechnicalProposal(
        status="pending_review",
        pdf_status="valid",
        errors=("TECHNICAL_VALIDATION_PENDING",),
    )

    result = audit.audit_historical_workbook(
        workbook, technical_resolver=_resolver({protocol: pending})
    )

    assert result.items[0].action is models.BackfillAction.PENDING_TECHNICAL_REVIEW
    assert result.summary.total_updates == 0


def test_cli_accepts_explicit_backfill_commands() -> None:
    from automacao_gd.presentation import cli

    audit_args = cli._parse_args(["backfill-audit"])
    apply_args = cli._parse_args(["backfill-apply", "--plan", "plano.json"])

    assert audit_args.command == "backfill-audit"
    assert apply_args.command == "backfill-apply"
    assert apply_args.plan == Path("plano.json")


def test_cli_audit_does_not_instantiate_portal_controller(monkeypatch) -> None:
    from automacao_gd.presentation import cli

    monkeypatch.setattr(cli, "ensure_directories", lambda: None)
    monkeypatch.setattr(cli, "setup_logger", lambda **kwargs: None)
    monkeypatch.setattr(
        cli,
        "ApplicationController",
        lambda: pytest.fail("a auditoria não pode criar o controlador do portal"),
    )
    monkeypatch.setattr(cli, "_run_backfill_audit", lambda: 0)

    assert cli.main(["backfill-audit"]) == 0


def test_apply_public_report_omits_backup_path(tmp_path: Path) -> None:
    models, _, _, _, report = _api()
    backup = tmp_path / "cliente" / "historico_backup.xlsx"
    result = models.BackfillApplyResult(
        status=models.BackfillApplyStatus.APPLIED_SUCCESSFULLY,
        items=(
            models.BackfillApplyItemResult(
                workbook_sheet="2025",
                workbook_row=2,
                protocol="2600001029",
                action=models.BackfillAction.UPDATE_EQUIPMENT,
                previous_module_text="ANTIGA",
                final_module_text="LEAPTON | LP182\nQtd. total: 5 módulos",
                previous_inverter_text="ANTIGO",
                final_inverter_text="HUAWEI | SUN2000\nQtd. total: 1 inversor",
            ),
        ),
        backup_path=backup,
        original_hash="a" * 64,
        backup_hash="a" * 64,
        final_hash="b" * 64,
        planned_updates=1,
        applied_updates=1,
        verification_passed=True,
        rollback_available=True,
        plan_hash="c" * 64,
    )

    public = report.build_apply_public_payload(result)

    assert public["backup_created"] is True
    assert str(tmp_path) not in repr(public)
    assert ".xlsx" not in repr(public)


def test_missing_or_ambiguous_header_is_reported_without_guessing(tmp_path: Path) -> None:
    _, audit, _, _, _ = _api()
    workbook = tmp_path / "cabecalhos.xlsx"
    wb = Workbook()
    ws_missing = wb.active
    ws_missing.title = "2024"
    ws_missing.append(("Protocolo", "Placa", "Observação"))
    ws_ambiguous = wb.create_sheet("2025")
    ws_ambiguous.append(("Protocolo", "Placa", "Placa", "Inversor"))
    wb.save(workbook)

    result = audit.audit_historical_workbook(workbook, technical_resolver=_resolver({}))

    assert result.summary.sheets_analyzed == 0
    assert result.summary.rows_analyzed == 0
    assert result.summary.unexpected_errors == 2


def test_only_module_difference_still_proposes_the_complete_pair(tmp_path: Path) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "somente_placa.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001030", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001030": _approved()}),
    )

    assert result.items[0].action is models.BackfillAction.UPDATE_EQUIPMENT
    assert result.items[0].proposed_module_text == "LEAPTON | LP182\nQtd. total: 5 módulos"
    assert (
        result.items[0].proposed_inverter_text
        == "HUAWEI | SUN2000\nQtd. total: 1 inversor"
    )


def test_preflight_block_prevents_backup_and_write(tmp_path: Path, monkeypatch) -> None:
    from automacao_gd.domain.errors import PreflightBlockedError

    _, audit, plan, apply, _ = _api()
    workbook = tmp_path / "bloqueada.xlsx"
    _workbook(workbook, {"2025": [("2600001031", "A", "B")]})
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001031": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    before = workbook.read_bytes()

    class _Blocked:
        def raise_if_blocked(self):
            raise PreflightBlockedError(
                code="WORKBOOK_LOCKED",
                user_message="Planilha bloqueada.",
                stage="pré-voo",
            )

    monkeypatch.setattr(
        apply,
        "_create_validated_backup",
        lambda *_: pytest.fail("backup não pode ocorrer com pré-voo bloqueado"),
    )
    with pytest.raises(apply.BackfillApplyError, match="Pré-voo"):
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _Blocked(),
        )
    assert workbook.read_bytes() == before


def test_save_failure_preserves_original_and_integral_backup(
    tmp_path: Path, monkeypatch
) -> None:
    models, audit, plan, apply, _ = _api()
    workbook = tmp_path / "falha_save.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001032", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001032": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    original = workbook.read_bytes()
    captured: dict[str, Path] = {}
    real_backup = apply._create_validated_backup

    def _backup(path, original_hash, equipment_rules_version=None):
        captured["path"] = real_backup(path, original_hash, equipment_rules_version)
        return captured["path"]

    monkeypatch.setattr(apply, "_create_validated_backup", _backup)
    monkeypatch.setattr(
        apply,
        "_save_temporary_workbook",
        lambda *args: (_ for _ in ()).throw(OSError("falha sintética")),
    )

    with pytest.raises(apply.BackfillApplyError) as caught:
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )
    assert workbook.read_bytes() == original
    assert captured["path"].read_bytes() == original
    assert caught.value.result is not None
    assert (
        caught.value.result.status
        is models.BackfillApplyStatus.ABORTED_UNEXPECTED_CHANGE
    )


def test_post_save_verification_failure_keeps_backup_for_manual_rollback(
    tmp_path: Path, monkeypatch
) -> None:
    _, audit, plan, apply, _ = _api()
    workbook = tmp_path / "falha_verificacao.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001033", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001033": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    original = workbook.read_bytes()
    captured: dict[str, Path] = {}
    real_backup = apply._create_validated_backup

    def _backup(path, original_hash, equipment_rules_version=None):
        captured["path"] = real_backup(path, original_hash, equipment_rules_version)
        return captured["path"]

    monkeypatch.setattr(apply, "_create_validated_backup", _backup)
    real_verify = apply._verify_saved_workbook
    calls = 0

    def _fail_final(path, prepared, preservation):
        nonlocal calls
        calls += 1
        return real_verify(path, prepared, preservation) if calls == 1 else False

    monkeypatch.setattr(apply, "_verify_saved_workbook", _fail_final)

    with pytest.raises(apply.BackfillApplyError, match="pós-gravação") as caught:
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )
    assert captured["path"].read_bytes() == original
    assert workbook.read_bytes() == original
    assert caught.value.result is not None
    assert caught.value.result.rollback_status == "CONFIRMED"


def test_post_save_verification_detects_unplanned_cell_change(
    tmp_path: Path, monkeypatch
) -> None:
    models, audit, plan, apply, _ = _api()
    workbook = tmp_path / "alteracao_extra.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001038", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001038": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    real_write = apply._write_prepared_updates

    def _write_with_contamination(wb, prepared):
        result = real_write(wb, prepared)
        wb["2025"]["D2"] = "ALTERAÇÃO NÃO PLANEJADA"
        return result

    monkeypatch.setattr(apply, "_write_prepared_updates", _write_with_contamination)

    original = workbook.read_bytes()
    with pytest.raises(apply.BackfillApplyError) as caught:
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )
    assert caught.value.result is not None
    assert (
        caught.value.result.status
        is models.BackfillApplyStatus.ABORTED_UNEXPECTED_CHANGE
    )
    assert workbook.read_bytes() == original


def test_documented_manual_rollback_restores_the_integral_backup(tmp_path: Path) -> None:
    _, audit, plan, apply, _ = _api()
    workbook = tmp_path / "rollback.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001034", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    original = workbook.read_bytes()
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001034": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    result = apply.apply_backfill_plan(
        workbook,
        payload,
        settings=_production_settings(),
        confirmation=apply.required_apply_confirmation(payload),
        preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
    )
    assert result.backup_path is not None

    atomic_copy_file(result.backup_path, workbook, private=True)

    assert workbook.read_bytes() == original
    restored = load_workbook(workbook)
    assert restored["2025"]["B2"].value == "1x LEAPTON LP182"
    assert restored["2025"]["C2"].value == "1x HUAWEI SUN2000"
    restored.close()


def test_audit_blocks_if_workbook_changes_during_read(tmp_path: Path, monkeypatch) -> None:
    _, audit, _, _, _ = _api()
    workbook = tmp_path / "concorrente.xlsx"
    _workbook(workbook, {"2025": [("2600001035", "A", "B")]})
    hashes = iter(("a" * 64, "b" * 64))
    monkeypatch.setattr(audit, "file_sha256", lambda path: next(hashes))

    with pytest.raises(audit.HistoricalAuditConsistencyError, match="alterada"):
        audit.audit_historical_workbook(
            workbook,
            technical_resolver=_resolver({"2600001035": _approved()}),
        )


def test_parser_failure_becomes_pending_and_does_not_abort_batch(
    tmp_path: Path, monkeypatch
) -> None:
    models, audit, _, _, _ = _api()
    workbook = tmp_path / "parser.xlsx"
    _workbook(
        workbook,
        {
            "2025": [
                ("2600001036", "A", "B"),
                ("2600001037", "A", "B"),
            ]
        },
    )
    pdf = tmp_path / "sintetico.pdf"
    pdf.write_bytes(b"%PDF-1.4 synthetic")
    monkeypatch.setattr(audit, "_pdf_candidates", lambda *args: (pdf,))
    monkeypatch.setattr(audit, "validate_pdf_input", lambda path: None)
    monkeypatch.setattr(audit, "extract_pdf_text", lambda path: "texto")
    monkeypatch.setattr(
        audit,
        "extract_protocol_from_pdf_text",
        lambda text: audit._CURRENT_TEST_PROTOCOL,
        raising=False,
    )
    monkeypatch.setattr(
        audit,
        "extract_generation_data",
        lambda *args, **kwargs: (_ for _ in ()).throw(ValueError("PDF sintético inválido")),
    )

    resolver = audit._default_resolver(tmp_path, None)
    items = []
    for protocol in ("2600001036", "2600001037"):
        monkeypatch.setattr(audit, "_CURRENT_TEST_PROTOCOL", protocol, raising=False)
        items.append(resolver(protocol))

    assert all(item.status == "pending_review" for item in items)
    assert all("TECHNICAL_EXTRACTION_FAILED" in item.errors for item in items)
    assert all(item.pdf_status == "valid" for item in items)
    assert models.BackfillAction.PENDING_TECHNICAL_REVIEW.value == "PENDING_TECHNICAL_REVIEW"


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        (lambda payload: payload.update({"unexpected": "value"}), "campos"),
        (
            lambda payload: payload["items"][0].update({"workbook_row": "inválida"}),
            "linha",
        ),
        (
            lambda payload: payload["items"][0].update(
                {"workbook_sheet": "C:\\dados\\2025"}
            ),
            "sensível",
        ),
    ],
)
def test_plan_rejects_unknown_fields_and_invalid_item_types(
    tmp_path: Path, mutation, message: str
) -> None:
    _, audit, plan, _, _ = _api()
    workbook = tmp_path / "plano_malformado.xlsx"
    _workbook(workbook, {"2025": [("2600001040", "A", "B")]})
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001040": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    mutation(payload)
    payload["plan_hash"] = plan.calculate_plan_hash(payload)

    with pytest.raises(plan.BackfillPlanError, match=message):
        plan.validate_backfill_plan(payload)


def test_backup_race_failure_is_controlled_and_does_not_modify_workbook(
    tmp_path: Path, monkeypatch
) -> None:
    _, audit, plan, apply, _ = _api()
    workbook = tmp_path / "falha_backup.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001041", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    before = workbook.read_bytes()
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001041": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    monkeypatch.setattr(
        apply,
        "_create_validated_backup",
        lambda *args: (_ for _ in ()).throw(PermissionError("bloqueio sintético")),
    )

    with pytest.raises(apply.BackfillApplyError, match="(?i)backup"):
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )
    assert workbook.read_bytes() == before


def test_workbook_open_race_failure_is_controlled(tmp_path: Path, monkeypatch) -> None:
    _, audit, plan, apply, _ = _api()
    workbook = tmp_path / "falha_abertura.xlsx"
    _workbook(workbook, {"2025": [("2600001042", "A", "B")]})
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001042": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    monkeypatch.setattr(
        apply,
        "load_workbook",
        lambda *args, **kwargs: (_ for _ in ()).throw(PermissionError("corrida sintética")),
    )

    with pytest.raises(apply.BackfillApplyError, match="abrir"):
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )


def test_post_save_reopen_error_has_failed_result_for_report(
    tmp_path: Path, monkeypatch
) -> None:
    models, audit, plan, apply, _ = _api()
    workbook = tmp_path / "falha_reabertura.xlsx"
    _workbook(
        workbook,
        {"2025": [("2600001043", "1x LEAPTON LP182", "1x HUAWEI SUN2000")]},
    )
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001043": _approved()}),
    )
    payload = plan.build_backfill_plan(audited, created_at="2026-07-21T10:00:00Z")
    real_verify = apply._verify_saved_workbook
    calls = 0

    def _fail_final(path, prepared, preservation):
        nonlocal calls
        calls += 1
        if calls == 1:
            return real_verify(path, prepared, preservation)
        raise OSError("reabertura sintética")

    monkeypatch.setattr(apply, "_verify_saved_workbook", _fail_final)

    with pytest.raises(apply.BackfillApplyError, match="pós-gravação") as caught:
        apply.apply_backfill_plan(
            workbook,
            payload,
            settings=_production_settings(),
            confirmation=apply.required_apply_confirmation(payload),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )
    assert caught.value.result is not None
    assert caught.value.result.status is models.BackfillApplyStatus.FAILED_ROLLED_BACK
    assert caught.value.result.rollback_available is True


def test_unsafe_technical_proposal_is_pending_and_never_reaches_public_report(
    tmp_path: Path,
) -> None:
    models, audit, _, _, report = _api()
    workbook = tmp_path / "proposta_insegura.xlsx"
    _workbook(workbook, {"2025": [("2600001044", "A", "B")]})
    unsafe = audit.TechnicalProposal(
        status="approved",
        module_text=r"C:\dados\cliente.pdf",
        inverter_text="1x HUAWEI SUN2000",
        source_hash="a" * 64,
        pdf_status="valid",
        warnings=("contato pessoa@example.com",),
    )

    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001044": unsafe}),
    )
    public = report.build_audit_public_payload(result)
    serialized = repr(public)

    assert result.items[0].action is models.BackfillAction.PENDING_TECHNICAL_REVIEW
    assert result.items[0].reasons == ("UNSAFE_TECHNICAL_PROPOSAL",)
    assert "C:\\" not in serialized
    assert ".pdf" not in serialized
    assert "@" not in serialized


def test_audit_markdown_includes_reasons_warnings_and_technical_status(
    tmp_path: Path,
) -> None:
    _, audit, _, _, report = _api()
    workbook = tmp_path / "markdown.xlsx"
    _workbook(workbook, {"2025": [("2600001045", "A", "B")]})
    proposal = audit.TechnicalProposal(
        status="approved",
        module_text="5x LEAPTON LP182",
        inverter_text="1x HUAWEI SUN2000",
        source_hash="a" * 64,
        pdf_status="valid",
        warnings=("AVISO_TECNICO",),
    )
    result = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver({"2600001045": proposal}),
    )

    markdown = report.build_audit_markdown(report.build_audit_public_payload(result))

    assert "Motivos" in markdown
    assert "Avisos" in markdown
    assert "Status técnico" in markdown
    assert "AVISO_TECNICO" in markdown


def test_audit_reports_share_plan_identity_and_versions(tmp_path: Path) -> None:
    _, audit, plan, _, report = _api()
    workbook = tmp_path / "linked_artifacts.xlsx"
    _workbook(
        workbook,
        {"2026": [("2600001085", "25x TSUN 615W N-TYPE", "1x SAJ 10K-R6")]},
    )
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver(
            {
                "2600001085": _approved(
                    module="25x TSUN 615W N-TYPE",
                    inverter="1x SAJ 10K-R6",
                )
            }
        ),
        created_at="2026-07-22T10:00:00+00:00",
    )
    planned = plan.build_backfill_plan(audited)

    public = report.build_audit_public_payload(audited, plan=planned)
    markdown = report.build_audit_markdown(public)
    shared = (
        "execution_id",
        "created_at",
        "plan_hash",
        "workbook_fingerprint",
        "plan_version",
        "equipment_format_version",
        "technical_processing_format_version",
        "equipment_rules_version",
    )

    assert all(public[field] == planned[field] for field in shared)
    assert all(str(public[field]) in markdown for field in shared)


def test_rules5_artifacts_use_one_timestamp_and_preserve_rules4(
    tmp_path: Path,
) -> None:
    _, audit, plan, _, report = _api()
    workbook = tmp_path / "artifact_names.xlsx"
    _workbook(
        workbook,
        {"2026": [("2600001084", "38x TSUN 610W N-TYPE", "1x SAJ 25K-R6")]},
    )
    audited = audit.audit_historical_workbook(
        workbook,
        technical_resolver=_resolver(
            {
                "2600001084": _approved(
                    module="38x TSUN 610W N-TYPE",
                    inverter="1x SAJ 25K-R6",
                )
            }
        ),
        created_at="2026-07-22T10:00:00+00:00",
    )
    planned = plan.build_backfill_plan(audited)
    legacy = tmp_path / "historical_equipment_backfill_plan_rules4.json"
    legacy.write_text("preservar", encoding="utf-8")
    json_path, markdown_path, plan_path = report.backfill_artifact_paths(
        tmp_path, planned["created_at"]
    )

    report.write_audit_reports(
        tmp_path,
        audited,
        plan=planned,
        json_path=json_path,
        markdown_path=markdown_path,
    )
    plan.write_backfill_plan(plan_path, planned)

    expected_suffixes = {
        "historical_equipment_backfill_audit_rules5_20260722T100000Z.json",
        "historical_equipment_backfill_audit_rules5_20260722T100000Z.md",
        "historical_equipment_backfill_plan_rules5_20260722T100000Z.json",
    }
    assert {json_path.name, markdown_path.name, plan_path.name} == expected_suffixes
    public = json.loads(json_path.read_text(encoding="utf-8"))
    saved_plan = json.loads(plan_path.read_text(encoding="utf-8"))
    assert public["execution_id"] == saved_plan["execution_id"]
    assert public["plan_hash"] == saved_plan["plan_hash"]
    assert legacy.read_text(encoding="utf-8") == "preservar"
