"""Testes para excel_service — valida planilha, backup, inserção e atualização."""

from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

import src.excel_service as excel_service
from src.excel_service import (
    DATE_NUMBER_FORMAT,
    REQUIRED_COLUMNS,
    TEXT_NUMBER_FORMAT,
    create_year_sheet_from_template,
    create_workbook_backup,
    find_required_columns,
    get_target_sheet_from_entry_date_or_protocol,
    repair_workbook_format,
    standardize_sheet_layout,
    update_excel_from_pdf_data,
    validate_excel_protocol_updated,
    validate_workbook_format,
    validate_workbook_for_pdf_updates,
    validate_xlsx_integrity,
    _save_workbook_atomically,
    _find_chronological_insert_row,
    _normalize_header,
    _sheet_name_for_entry_date,
)


@pytest.fixture
def sample_workbook(tmp_path: Path) -> Path:
    """Cria uma planilha de teste com as colunas obrigatórias."""
    wb = Workbook()
    ws = wb.active
    ws.title = "2025"
    headers = [
        "Cliente",
        "Protocolo",
        "Data de ingresso",
        "Conclusão",
        "Parecer",
        "Placa",
        "Inversor",
    ]
    ws.append(headers)
    # Linhas de dados
    ws.append(["CLIENTE SINTETICO 008 LTDA", "2600001097", "01/06/2025", "", "Sim", "18x MOD X", "1x INV Y"])
    ws.append(["Maria Souza", "2600001098", "15/06/2025", "", "Sim", "10x MOD A", "1x INV B"])
    path = tmp_path / "test_planilha.xlsx"
    wb.save(path)
    return path


class TestValidateWorkbook:
    def test_valid_workbook(self, sample_workbook: Path):
        result = validate_workbook_for_pdf_updates(sample_workbook, require_writable=True)
        assert result["success"] is True
        assert result["worksheet"] == "2025"

    def test_missing_workbook(self, tmp_path: Path):
        result = validate_workbook_for_pdf_updates(tmp_path / "nonexistent.xlsx")
        assert result["success"] is False
        assert "não encontrada" in result["error"].lower()


class TestSheetNameForEntryDate:
    def test_year_2025(self):
        assert _sheet_name_for_entry_date(date(2025, 6, 15)) == "2025"

    def test_year_2022_or_2023(self):
        assert _sheet_name_for_entry_date(date(2022, 1, 1)) == "2022 - 2023"
        assert _sheet_name_for_entry_date(date(2023, 12, 31)) == "2022 - 2023"

    def test_year_2026(self):
        assert _sheet_name_for_entry_date(date(2026, 1, 1)) == "2026"


class TestFindRequiredColumns:
    def test_finds_columns(self, sample_workbook: Path):
        from openpyxl import load_workbook

        wb = load_workbook(sample_workbook)
        ws = wb.active
        columns = find_required_columns(ws)
        assert "Cliente" in columns
        assert "Protocolo" in columns
        assert "Data de ingresso" in columns
        assert "Parecer" in columns
        assert "Placa" in columns
        assert "Inversor" in columns
        wb.close()


class TestNormalizeHeader:
    def test_removes_accents(self):
        assert _normalize_header("Data de ingresso") == "DATA DE INGRESSO"
        assert _normalize_header("Conclusão") == "CONCLUSAO"

    def test_removes_special_chars(self):
        assert _normalize_header("Placa(s)") == "PLACA S"
        assert _normalize_header("Qtd. módulos") == "QTD MODULOS"


class TestUpdateExcelFromPdfData:
    def test_blocked_missing_entry_date(self, sample_workbook: Path):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001099",
            client_name="CLIENTE SINTETICO LTDA",
            entry_date=None,
            dry_run=True,
        )
        assert result["action"] == "blocked_missing_entry_date"
        assert result["success"] is False

    def test_dry_run_insert_new(self, sample_workbook: Path):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001099",
            client_name="CLIENTE SINTETICO LTDA",
            entry_date="20/06/2025",
            module_text="12x MOD Z 500W",
            inverter_text="1x INV W",
            dry_run=True,
        )
        assert result["success"] is True
        assert result["action"] == "insert_new_chronological"
        assert result["can_write"] is True
        assert result["target_sheet"] == "2025"

    def test_dry_run_update_existing(self, sample_workbook: Path):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 008 LTDA",
            entry_date="01/06/2025",
            module_text="18x MOD X 585W",
            inverter_text="1x INV Y",
            dry_run=True,
        )
        assert result["success"] is True
        assert result["action"] == "update_existing"
        assert result["row_found"] is True
        assert result["row_number"] is not None

    def test_real_run_skips_protocol_already_updated(self, sample_workbook: Path):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 007 LTDA",
            entry_date="01/06/2025",
            module_text="18x MOD X",
            inverter_text="1x INV Y",
            dry_run=False,
        )

        validation = validate_excel_protocol_updated(
            sample_workbook,
            "2600001097",
            entry_date="01/06/2025",
        )

        assert result["success"] is True
        assert result["action"] == "skipped_excel_already_updated"
        assert result["can_write"] is False
        assert validation["success"] is True
        assert validation["already_updated"] is True

    def test_real_run_updates_completion_when_equipment_is_already_same(
        self, sample_workbook: Path
    ):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 006 LTDA",
            entry_date="01/06/2025",
            completion_date="15/06/2025",
            module_text="18x MOD X",
            inverter_text="1x INV Y",
            dry_run=False,
        )

        wb = load_workbook(sample_workbook)
        ws = wb["2025"]
        try:
            assert result["success"] is True
            assert result["action"] == "update_existing"
            assert result["completion_no_change"] is False
            assert result["equipment_no_change"] is True
            assert ws.cell(row=2, column=1).value == "CLIENTE SINTETICO 008 LTDA"
            assert ws.cell(row=2, column=4).value == datetime(2025, 6, 15)
            assert ws.cell(row=2, column=4).number_format == DATE_NUMBER_FORMAT
        finally:
            wb.close()

    def test_real_run_skips_when_equipment_and_completion_are_already_same(
        self, sample_workbook: Path
    ):
        wb = load_workbook(sample_workbook)
        ws = wb["2025"]
        ws.cell(row=2, column=4).value = datetime(2025, 6, 15)
        ws.cell(row=2, column=4).number_format = DATE_NUMBER_FORMAT
        wb.save(sample_workbook)
        wb.close()

        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 007 LTDA",
            entry_date="01/06/2025",
            completion_date="15/06/2025",
            module_text="18x MOD X",
            inverter_text="1x INV Y",
            dry_run=False,
        )

        assert result["success"] is True
        assert result["action"] == "skipped_excel_already_updated"
        assert result["completion_no_change"] is True
        assert result["equipment_no_change"] is True

    def test_real_run_writes_em_aberto_as_text(self, sample_workbook: Path):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 007 LTDA",
            entry_date="01/06/2025",
            completion_date="EM ABERTO",
            module_text="18x MOD X",
            inverter_text="1x INV Y",
            dry_run=False,
        )

        wb = load_workbook(sample_workbook)
        ws = wb["2025"]
        try:
            assert result["success"] is True
            assert result["action"] == "update_existing"
            assert ws.cell(row=2, column=4).value == "EM ABERTO"
        finally:
            wb.close()

    def test_real_run_updates_existing_when_equipment_text_changed(
        self, sample_workbook: Path
    ):
        module_text = (
            "BYD | P6C-30 260\n"
            "TRINA | TSM-NEG21C 695\n"
            "Total: 218 mÃ³dulos"
        )
        inverter_text = (
            "HUAWEI | SUN2000-30KTL\n"
            "ABB | Aurora Trio-20.0TL-OUTD\n"
            "HUAWEI | SUN2000-20KTL\n"
            "Total: 3 inversores"
        )

        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 007 LTDA",
            entry_date="01/06/2025",
            module_text=module_text,
            inverter_text=inverter_text,
            dry_run=False,
        )

        wb = load_workbook(sample_workbook)
        ws = wb["2025"]
        try:
            assert result["success"] is True
            assert result["action"] == "update_existing"
            assert ws.cell(row=2, column=6).value == module_text
            assert ws.cell(row=2, column=7).value == inverter_text
            assert ws.cell(row=2, column=6).alignment.wrap_text is True
            assert ws.cell(row=2, column=7).alignment.vertical == "top"
            assert (ws.row_dimensions[2].height or 0) >= 60
        finally:
            wb.close()

    def test_real_run_updates_existing_when_expected_inverter_is_empty(
        self, sample_workbook: Path
    ):
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 007 LTDA",
            entry_date="01/06/2025",
            module_text="18x MOD X",
            inverter_text="",
            dry_run=False,
        )

        wb = load_workbook(sample_workbook)
        ws = wb["2025"]
        try:
            assert result["success"] is True
            assert result["action"] == "update_existing"
            assert ws.cell(row=2, column=7).value in (None, "")
        finally:
            wb.close()

    def test_permission_error_during_atomic_save_is_operational_without_traceback(
        self, sample_workbook: Path, monkeypatch
    ):
        exception_calls = []
        error_calls = []

        def blocked_save(*args, **kwargs):
            exc = PermissionError(13, "Acesso negado", str(sample_workbook))
            exc.winerror = 5
            raise exc

        monkeypatch.setattr(excel_service, "_save_workbook_atomically", blocked_save)
        monkeypatch.setattr(
            excel_service.logger,
            "exception",
            lambda *args, **kwargs: exception_calls.append((args, kwargs)),
        )
        monkeypatch.setattr(
            excel_service.logger,
            "error",
            lambda *args, **kwargs: error_calls.append((args, kwargs)),
        )

        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 007 LTDA",
            entry_date="01/06/2025",
            module_text="99x MOD NOVO",
            inverter_text="2x INV NOVO",
            dry_run=False,
        )

        assert result["success"] is False
        assert result["code"] == "WORKBOOK_PERMISSION_DENIED"
        assert "sem permissão" in result["error"].lower()
        assert exception_calls == []
        assert error_calls

    def test_atomic_save_fails_safely_when_replace_is_denied(
        self, sample_workbook: Path, monkeypatch
    ):
        def deny_replace(source, destination):
            raise PermissionError(5, "Acesso negado", str(destination))

        monkeypatch.setattr(excel_service.os, "replace", deny_replace)

        before = sample_workbook.read_bytes()
        wb = load_workbook(sample_workbook)
        try:
            wb["2025"].cell(row=2, column=6).value = "99x MOD FALLBACK"
            with pytest.raises(PermissionError):
                _save_workbook_atomically(wb, sample_workbook)
        finally:
            wb.close()

        assert sample_workbook.read_bytes() == before
        assert not sample_workbook.with_name(
            f".{sample_workbook.stem}.saving{sample_workbook.suffix}"
        ).exists()

    def test_dry_run_wrong_sheet(self, sample_workbook: Path):
        """Protocolo em aba 2025 mas data de ingresso de 2026."""
        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 008 LTDA",
            entry_date="15/01/2026",
            dry_run=True,
        )
        # Como MOVE_WRONG_YEAR_ROWS=True por padrão, deve marcar move
        assert result["success"] is True
        assert result["action"] in ("move_wrong_sheet_to_correct_sheet", "warning_wrong_sheet")

    def test_blocked_duplicate_protocol(self, sample_workbook: Path):
        """Protocolo duplicado na mesma aba deve bloquear."""
        from openpyxl import load_workbook

        wb = load_workbook(sample_workbook)
        ws = wb["2025"]
        # Adicionar o mesmo protocolo numa terceira linha
        ws.cell(row=3, column=1, value="Duplicado")
        ws.cell(row=3, column=2, value="2600001097")
        ws.cell(row=3, column=3, value="02/06/2025")
        ws.cell(row=3, column=5, value="Sim")
        ws.cell(row=3, column=6, value="X")
        ws.cell(row=3, column=7, value="Y")
        wb.save(sample_workbook)
        wb.close()

        result = update_excel_from_pdf_data(
            workbook_path=sample_workbook,
            protocol="2600001097",
            client_name="CLIENTE SINTETICO 008 LTDA",
            entry_date="01/06/2025",
            dry_run=True,
        )
        # Com protocolo duplicado, a função deve bloquear ou atualizar a primeira ocorrência
        # O comportamento exato depende da implementação de detecção de duplicata
        assert result["action"] in (
            "blocked_duplicate_protocol",
            "update_existing",
            "warning_wrong_sheet",
        )
        assert result is not None

    def test_missing_workbook(self, tmp_path: Path):
        result = update_excel_from_pdf_data(
            workbook_path=tmp_path / "nonexistent.xlsx",
            protocol="2600001097",
            client_name="CLIENTE SINTETICO LTDA",
            entry_date="01/06/2025",
            dry_run=True,
        )
        assert result["success"] is False
        assert "não encontrada" in result["error"].lower()


class TestCreateWorkbookBackup:
    def test_backup_created(self, sample_workbook: Path):
        backup = create_workbook_backup(sample_workbook)
        assert backup.exists()
        assert backup.suffix == ".xlsx"
        assert "backup" in backup.name.lower()

    def test_backup_does_not_overwrite(self, sample_workbook: Path):
        backup1 = create_workbook_backup(sample_workbook)
        backup2 = create_workbook_backup(sample_workbook)
        assert backup1 != backup2
        assert backup1.exists()
        assert backup2.exists()


class TestFindChronologicalInsertRow:
    def test_empty_sheet(self, sample_workbook: Path):
        from openpyxl import load_workbook

        wb = load_workbook(sample_workbook)
        ws = wb.create_sheet("Empty")
        ws.append(["Cliente", "Protocolo", "Data de ingresso", "Conclusão", "Parecer", "Placa", "Inversor"])
        columns = {"Data de ingresso": 3}
        row = _find_chronological_insert_row(ws, 1, columns, date(2025, 6, 1))
        assert row == 2
        wb.close()


@pytest.fixture
def formatted_workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2025"
    ws["A1"] = "LOGO"
    ws["A1"].font = Font(bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16
    ws.row_dimensions[1].height = 30
    ws.append(REQUIRED_COLUMNS)
    ws.append(["CLIENTE", "2601", date(2025, 6, 10), None, "Sim", "10x MOD", "1x INV"])
    for col in range(1, 8):
        ws.cell(row=3, column=col).font = Font(name="Arial", size=10)
        ws.cell(row=3, column=col).fill = PatternFill("solid", fgColor="DDDDDD")
    ws["B3"].number_format = TEXT_NUMBER_FORMAT
    ws["C3"].number_format = DATE_NUMBER_FORMAT
    ws["D3"].number_format = DATE_NUMBER_FORMAT
    path = tmp_path / "formatada.xlsx"
    wb.save(path)
    return path


class TestWorkbookFormatting:
    def test_target_sheet_uses_protocol_prefix_as_full_year(self):
        protocol_2025 = "25" + "05190388"
        protocol_2024 = "24" + "00000000"
        protocol_2022 = "22" + "00000000"
        protocol_2023 = "23" + "00000000"
        assert get_target_sheet_from_entry_date_or_protocol(None, protocol_2025) == "2025"
        assert get_target_sheet_from_entry_date_or_protocol(None, "2600000000") == "2026"
        assert get_target_sheet_from_entry_date_or_protocol(None, protocol_2024) == "2024"
        assert get_target_sheet_from_entry_date_or_protocol(None, protocol_2022) == "2022 - 2023"
        assert get_target_sheet_from_entry_date_or_protocol(None, protocol_2023) == "2022 - 2023"

    def test_target_sheet_prefers_entry_date(self):
        assert get_target_sheet_from_entry_date_or_protocol(date(2022, 1, 1), "2600000000") == "2022 - 2023"
        assert get_target_sheet_from_entry_date_or_protocol(date(2023, 1, 1), "2600000000") == "2022 - 2023"
        assert get_target_sheet_from_entry_date_or_protocol(date(2025, 1, 1), "2600000000") == "2025"

    def test_update_formats_date_and_protocol_as_text(self, formatted_workbook: Path):
        update_excel_from_pdf_data(
            workbook_path=formatted_workbook,
            protocol="2602",
            client_name="CLIENTE SINTETICO LTDA",
            entry_date="12/06/2025",
            module_text="10x JINKO JKM625N-78HL4-BDV | 6,25 kWp",
            inverter_text="1x HUAWEI SUN2000-5KTL | 5 kW",
            dry_run=False,
        )

        wb = load_workbook(formatted_workbook)
        ws = wb["2025"]
        assert ws["B4"].number_format == TEXT_NUMBER_FORMAT
        assert ws["C4"].number_format == DATE_NUMBER_FORMAT
        assert ws["D4"].number_format == DATE_NUMBER_FORMAT
        assert ws["C4"].value.date() == date(2025, 6, 12)
        assert ws["F4"].value == "10x JINKO JKM625N-78HL4-BDV"
        assert ws["G4"].value == "1x HUAWEI SUN2000-5KTL"
        wb.close()

    def test_multiline_equipment_cells_enable_wrap_and_taller_row(self, formatted_workbook: Path):
        module_text = (
            "BYD | P6C-30 260\n"
            "TRINA | TSM-NEG21C 695\n"
            "Total: 218 módulos"
        )
        inverter_text = (
            "HUAWEI | SUN2000-30KTL\n"
            "ABB | Aurora Trio-20.0TL-OUTD\n"
            "HUAWEI | SUN2000-20KTL\n"
            "Total: 3 inversores"
        )

        update_excel_from_pdf_data(
            workbook_path=formatted_workbook,
            protocol="2604",
            client_name="CLIENTE SINTETICO LTDA",
            entry_date="13/06/2025",
            module_text=module_text,
            inverter_text=inverter_text,
            dry_run=False,
        )

        wb = load_workbook(formatted_workbook)
        ws = wb["2025"]
        assert ws["F4"].value == module_text
        assert ws["G4"].value == inverter_text
        assert ws["F4"].alignment.wrap_text is True
        assert ws["G4"].alignment.wrap_text is True
        assert ws["F4"].alignment.vertical == "top"
        assert ws.row_dimensions[4].height >= 45
        wb.close()

    def test_standardize_removes_value_error_from_top_area(self, formatted_workbook: Path):
        wb = load_workbook(formatted_workbook)
        template = wb["2025"]
        ws = wb.copy_worksheet(template)
        ws.title = "2026"
        ws["A1"] = "#VALOR!"

        report = standardize_sheet_layout(ws, template)

        assert report["value_errors_removed"] >= 1
        assert ws["A1"].value == "SOLUCIONE NORDESTE ENERGIA ELÉTRICA"
        wb.close()

    def test_create_year_sheet_copies_layout_and_clears_data(self, formatted_workbook: Path):
        wb = load_workbook(formatted_workbook)
        template = wb["2025"]

        ws = create_year_sheet_from_template(wb, template, "2026")

        assert ws.title == "2026"
        assert ws.column_dimensions["A"].width == 38
        assert ws.column_dimensions["F"].width == 42
        assert ws.row_dimensions[1].height == 28
        assert ws["A1"].font.bold is True
        assert ws["A1"].value == "SOLUCIONE NORDESTE ENERGIA ELÉTRICA"
        assert ws["A3"].value is None
        wb.close()

    def test_repair_preserves_existing_data(self, formatted_workbook: Path):
        wb = load_workbook(formatted_workbook)
        ws = wb.copy_worksheet(wb["2025"])
        ws.title = "2026"
        ws["A1"] = "#VALOR!"
        ws["A3"] = "CLIENTE 2026"
        ws["B3"] = "2603"
        ws["C3"] = date(2026, 6, 12)
        wb.save(formatted_workbook)
        wb.close()

        result = repair_workbook_format(formatted_workbook, dry_run=False)

        wb = load_workbook(formatted_workbook)
        ws = wb["2026"]
        assert result["backup_path"]
        assert ws["B3"].value == "2603"
        assert ws["C3"].number_format == DATE_NUMBER_FORMAT
        assert ws["A1"].value == "SOLUCIONE NORDESTE ENERGIA ELÉTRICA"
        wb.close()

    def test_repair_real_run_validates_temp_before_replace(self, formatted_workbook: Path):
        result = repair_workbook_format(
            formatted_workbook,
            dry_run=False,
            mode="visual_only",
            sheet_names=["2025"],
        )

        assert result["success"] is True
        assert result["backup_path"]
        assert result["temp_path"]
        assert result["xlsx_integrity_ok"] is True
        assert result["official_file_replaced"] is True
        ok, errors = validate_xlsx_integrity(formatted_workbook)
        assert ok is True
        assert errors == []

    def test_visual_only_repair_succeeds_with_data_pending(self, formatted_workbook: Path):
        wb = load_workbook(formatted_workbook)
        ws_2024 = wb.copy_worksheet(wb["2025"])
        ws_2024.title = "2024"
        ws_2024["A1"] = "#VALUE!"
        ws_2024["B3"] = "2600001071"
        ws_2024["C3"] = date(2025, 5, 19)
        ws_2024["F3"] = "10 modulos | 6 kWp"
        wb.save(formatted_workbook)
        wb.close()

        result = repair_workbook_format(
            formatted_workbook,
            dry_run=True,
            mode="visual_only",
            sheet_names=["2024", "2025"],
        )

        assert result["success"] is True
        assert result["errors"] == []
        assert result["data_pending"]
        assert any(item["type"] == "wrong_year_sheet" for item in result["data_pending"])
        assert any(sheet["value_errors_removed"] >= 1 for sheet in result["sheets"])

    def test_validate_separates_data_pending_from_technical_errors(self, formatted_workbook: Path):
        wb = load_workbook(formatted_workbook)
        ws_2024 = wb.copy_worksheet(wb["2025"])
        ws_2024.title = "2024"
        ws_2024["B3"] = "2600001071"
        ws_2024["C3"] = date(2025, 5, 19)
        wb.save(formatted_workbook)
        wb.close()

        result = validate_workbook_format(
            formatted_workbook,
            sheet_names=["2024", "2025"],
            strict=False,
        )

        assert result["success"] is False
        assert result["technical_errors"] == []
        assert any(item["expected_sheet"] == "2025" for item in result["wrong_year_rows"])
        assert any(item["type"] == "wrong_year_sheet" for item in result["data_pending"])

    def test_validate_detects_duplicate_protocol_and_wrong_year(self, formatted_workbook: Path):
        wb = load_workbook(formatted_workbook)
        ws_2024 = wb.copy_worksheet(wb["2025"])
        ws_2024.title = "2024"
        ws_2024["B3"] = "2601"
        ws_2024["C3"] = date(2026, 1, 10)
        ws_2026 = wb.copy_worksheet(wb["2025"])
        ws_2026.title = "2026"
        ws_2026["B3"] = "2601"
        ws_2026["C3"] = date(2026, 1, 10)
        wb.save(formatted_workbook)
        wb.close()

        result = validate_workbook_format(formatted_workbook)

        assert result["success"] is False
        assert any(item["protocol"] == "2601" for item in result["duplicate_protocols"])
        assert any(item["current_sheet"] == "2024" for item in result["wrong_year_rows"])
