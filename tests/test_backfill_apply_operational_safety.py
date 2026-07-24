from __future__ import annotations

import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from automacao_gd.application.historical_backfill import apply_service, audit_service
from automacao_gd.application.historical_backfill.plan import build_backfill_plan
from automacao_gd.application.historical_backfill.report import write_apply_reports
from automacao_gd.domain.backfill_models import BackfillApplyStatus
from automacao_gd.domain.equipment_semantics import (
    CanonicalEquipmentCollection,
    EquipmentSource,
    EquipmentSourceType,
    canonical_collection_to_dict,
    canonicalize_equipment,
)


RULES7_PLAN_PATH = (
    Path("data")
    / "logs"
    / "historical_equipment_backfill_plan_rules7_20260724T141444Z.json"
)


class _ReadyPreflight:
    def raise_if_blocked(self) -> None:
        return None


def _settings(environment: str | None = "production") -> SimpleNamespace:
    return SimpleNamespace(
        APP_ENV=environment,
        DRY_RUN=False,
        APPLY_EXCEL=True,
        APPLY_ARCHIVE=False,
        BACKUP_EXCEL=True,
    )


def _synthetic_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2025"
    sheet.append(("Protocolo", "Placa", "Inversor", "Observação"))
    sheet.append(
        (
            "2500001001",
            "1x LEAPTON LP182",
            "1x HUAWEI SUN2000",
            "PRESERVAR",
        )
    )
    sheet.auto_filter.ref = "A1:D2"
    sheet.column_dimensions["A"].width = 18
    workbook.save(path)
    workbook.close()


def _feature_rich_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2025"
    sheet.append(("Protocolo", "Placa", "Inversor", "Observação", "Status"))
    sheet.append(("2500001001", 0, 0, "=LEN(A2)", "OK"))
    sheet["A1"].font = Font(bold=True)
    sheet["B2"].fill = PatternFill("solid", fgColor="FFFF00")
    warning_fill = PatternFill("solid", fgColor="FFC7CE")
    sheet.conditional_formatting.add(
        "C2:C2",
        CellIsRule(operator="greaterThan", formula=["0"], fill=warning_fill),
    )
    sheet.merge_cells("D4:E4")
    sheet["D4"] = "MESCLADO"
    sheet.auto_filter.ref = "A1:E2"
    sheet.freeze_panes = "A2"
    sheet.page_margins.left = 0.5
    sheet.page_setup.orientation = "landscape"
    validation = DataValidation(type="list", formula1='"OK,PENDENTE"')
    validation.add(sheet["E2"])
    sheet.add_data_validation(validation)
    table = Table(displayName="TabelaBackfillTeste", ref="A1:E2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(path)
    workbook.close()


def _synthetic_plan(path: Path) -> dict:
    _synthetic_workbook(path)
    proposal = audit_service.TechnicalProposal(
        status="approved",
        module_text="5x LEAPTON LP182",
        inverter_text="1x HUAWEI SUN2000",
        source_hash="a" * 64,
        pdf_status="valid",
    )
    audit = audit_service.audit_historical_workbook(
        path,
        technical_resolver=lambda _protocol: proposal,
        created_at="2026-07-22T15:00:00+00:00",
    )
    return build_backfill_plan(audit)


def _prepare(path: Path, payload: dict):
    return apply_service.prepare_backfill_application(
        path,
        payload,
        settings=_settings(),
        preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
    )


def _canonical_collection(
    *,
    module: tuple[str, str, int, EquipmentSourceType],
    inverter: tuple[str, str, int, EquipmentSourceType],
) -> CanonicalEquipmentCollection:
    module_manufacturer, module_model, module_quantity, module_source = module
    inverter_manufacturer, inverter_model, inverter_quantity, inverter_source = inverter
    return CanonicalEquipmentCollection(
        modules=(
            canonicalize_equipment(
                equipment_type="module",
                manufacturer=module_manufacturer,
                model=module_model,
                quantity=module_quantity,
                source=EquipmentSource(module_source),
            ),
        ),
        conventional_inverters=(
            canonicalize_equipment(
                equipment_type="inverter",
                manufacturer=inverter_manufacturer,
                model=inverter_model,
                quantity=inverter_quantity,
                source=EquipmentSource(inverter_source),
            ),
        ),
    )


def _quality_gate_item(
    *,
    protocol: str = "2500000000",
    proposed_module: str,
    proposed_inverter: str,
    collection: CanonicalEquipmentCollection,
    warnings: tuple[str, ...] = (),
) -> dict:
    return {
        "protocol": protocol,
        "workbook_sheet": "2025",
        "workbook_row": 2,
        "proposed_module_text": proposed_module,
        "proposed_inverter_text": proposed_inverter,
        "proposed_canonical_collection": canonical_collection_to_dict(collection),
        "warnings": list(warnings),
        "semantic_validation_warnings": list(warnings),
        "technical_validation_status": "approved",
        "semantic_validation_status": "approved",
        "blocking_violations": [],
    }


def _load_rules7_plan() -> dict:
    return json.loads(RULES7_PLAN_PATH.read_text(encoding="utf-8"))


def test_row_quality_gate_accepts_valid_cross_field_transition() -> None:
    source = _canonical_collection(
        module=("LEAPTON", "LP182-M-72-MH 585W", 12, EquipmentSourceType.TABLE_CELL),
        inverter=("SOLPLANET", "ASW5000-S", 1, EquipmentSourceType.TABLE_CELL),
    )
    item = _quality_gate_item(
        proposed_module="12x LEAPTON LP182-M-72-MH 585W",
        proposed_inverter="1x SOLPLANET ASW5000-S",
        collection=source,
        warnings=("CROSS_FIELD_CONTAMINATION_RESOLVED",),
    )

    assert apply_service._row_quality_gate(
        item,
        "",
        (
            "Solplanet | LEAPTON ASW5000-S LP182-M-72-MH 585W\n"
            "Aiswei | LEAPTON ASW5000-S LP182-M-72-MH 585W\n"
            "Qtd. total: 1 inversor"
        ),
    )


def test_row_quality_gate_blocks_real_equipment_loss() -> None:
    source = _canonical_collection(
        module=("LEAPTON", "LP182-M-72-MH 585W", 12, EquipmentSourceType.TABLE_CELL),
        inverter=("SOLPLANET", "ASW5000-S", 1, EquipmentSourceType.TABLE_CELL),
    )
    item = _quality_gate_item(
        proposed_module="12x LEAPTON LP182-M-72-MH 585W",
        proposed_inverter="1x SOLPLANET ASW5000-S",
        collection=source,
    )

    assert not apply_service._row_quality_gate(
        item,
        "",
        "1x SOLPLANET ASW5000-S\n1x GROWATT MIC 3000TL-X\nLEAPTON LP182-M-72-MH 585W",
    )


def test_row_quality_gate_does_not_accept_warning_as_bypass() -> None:
    source = _canonical_collection(
        module=("LEAPTON", "LP182-M-72-MH 585W", 12, EquipmentSourceType.TABLE_CELL),
        inverter=("SOLPLANET", "ASW5000-S", 1, EquipmentSourceType.TABLE_CELL),
    )
    item = _quality_gate_item(
        proposed_module="12x LEAPTON LP182-M-72-MH 585W",
        proposed_inverter="1x SOLPLANET ASW5000-S",
        collection=source,
        warnings=("CROSS_FIELD_CONTAMINATION_RESOLVED",),
    )

    assert not apply_service._row_quality_gate(
        item,
        "",
        "1x SOLPLANET ASW5000-S\n1x GROWATT MIC 3000TL-X\nLEAPTON LP182-M-72-MH 585W",
    )


@pytest.mark.parametrize(
    ("current_module", "proposed_module", "collection"),
    [
        (
            "6x TSUN TSUN 605W BIFACIAL",
            "6x TSUN 605W BIFACIAL",
            _canonical_collection(
                module=("TSUN", "605W BIFACIAL", 6, EquipmentSourceType.TABLE_CELL),
                inverter=("GROWATT", "MIC 3000TL-X", 1, EquipmentSourceType.TABLE_CELL),
            ),
        ),
        (
            "7x DMEGC Dmegc 605W (Monocristalino/N- Type) Bifacial",
            "7x DMEGC 605W (Monocristalino/N- Type) Bifacial",
            _canonical_collection(
                module=(
                    "DMEGC",
                    "605W (Monocristalino/N- Type) Bifacial",
                    7,
                    EquipmentSourceType.TABLE_CELL,
                ),
                inverter=("GROWATT", "MIC 3000TL-X", 1, EquipmentSourceType.TABLE_CELL),
            ),
        ),
    ],
)
def test_row_quality_gate_accepts_duplicate_manufacturer_cleanup(
    current_module: str,
    proposed_module: str,
    collection: CanonicalEquipmentCollection,
) -> None:
    item = _quality_gate_item(
        proposed_module=proposed_module,
        proposed_inverter="1x GROWATT MIC 3000TL-X",
        collection=collection,
        warnings=("DUPLICATED_MANUFACTURER_IN_MODEL_REMOVED",),
    )

    assert apply_service._row_quality_gate(
        item,
        current_module,
        "1x GROWATT MIC 3000TL-X",
    )


def test_row_quality_gate_accepts_solplanet_alias_cleanup() -> None:
    source = _canonical_collection(
        module=("JINKO", "JKM625N-78HL4-BDV 625W", 8, EquipmentSourceType.TABLE_CELL),
        inverter=("SOLPLANET", "ASW4000-S", 1, EquipmentSourceType.TABLE_CELL),
    )
    item = _quality_gate_item(
        proposed_module="8x JINKO JKM625N-78HL4-BDV 625W",
        proposed_inverter="1x SOLPLANET ASW4000-S",
        collection=source,
        warnings=("SOLPLANET_ALIAS_NORMALIZED",),
    )

    assert apply_service._row_quality_gate(
        item,
        "8x JINKO JKM625N-78HL4-BDV 625W",
        "1x SOLPLANET/AISWEI ASW4000-S",
    )


def test_rules7_official_preflight_prepares_all_updates_without_conflicts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.setenv("DRY_RUN", "false")
    monkeypatch.setenv("APPLY_EXCEL", "true")
    monkeypatch.setenv("BACKUP_EXCEL", "true")
    from automacao_gd.infrastructure.config import get_settings

    settings = get_settings()
    plan = _load_rules7_plan()
    current_hash = audit_service.file_sha256(settings.planilha_path)

    if current_hash != plan["workbook_fingerprint"]:
        prepared, conflicts = apply_service._load_prepared_updates(
            settings.planilha_path, plan
        )
        assert len(prepared) == 0
        assert len(conflicts) == 30
        assert {reason for conflict in conflicts for reason in conflict.reasons} == {
            "ROW_FINGERPRINT_MISMATCH"
        }
        return

    prepared, conflicts = apply_service._load_prepared_updates(
        settings.planilha_path, plan
    )

    assert len(prepared) == 30
    assert conflicts == []
    assert {item["protocol"] for item in prepared}.issuperset(
        {
            "2502104207",
            "2503118917",
            "2504165265",
            "2506043351",
            "2507119114",
            "2508155042",
            "2510074020",
            "2510236705",
            "2601204137",
            "2602027219",
        }
    )


@pytest.mark.parametrize("environment", [None, "", "development", "test", "local", "other"])
def test_non_production_environment_aborts_before_confirmation(
    tmp_path: Path, environment: str | None
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)

    with pytest.raises(apply_service.BackfillApplyError) as caught:
        apply_service.prepare_backfill_application(
            workbook,
            payload,
            settings=_settings(environment),
            preflight_runner=lambda *args, **kwargs: _ReadyPreflight(),
        )

    assert caught.value.result is not None
    assert caught.value.result.status is BackfillApplyStatus.ABORTED_PRE_FLIGHT
    assert caught.value.result.error_code == "PRODUCTION_ENVIRONMENT_REQUIRED"


def test_production_preflight_builds_plan_specific_confirmation(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)

    preparation = _prepare(workbook, payload)

    assert preparation.required_confirmation == "APLICAR RULES-7 1 UPDATES"
    assert preparation.workbook_hash == payload["workbook_fingerprint"]
    assert preparation.updates_by_sheet == {"2025": 1}


def test_authorized_plan_confirmation_text_is_exact() -> None:
    assert apply_service.required_apply_confirmation(
        {
            "equipment_rules_version": "equipment-v2-technical-v6-rules-5",
            "total_updates": 120,
        }
    ) == "APLICAR RULES-5 120 UPDATES"


def test_old_or_partial_confirmation_is_rejected_exactly(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)

    for confirmation in ("APLICAR BACKFILL", "APLICAR RULES-6", " APLICAR RULES-7 1 UPDATES "):
        with pytest.raises(apply_service.BackfillApplyError) as caught:
            apply_service.execute_backfill_application(preparation, confirmation)
        assert caught.value.result is not None
        assert caught.value.result.status is BackfillApplyStatus.ABORTED_CONFIRMATION


def test_synthetic_application_validates_backup_temp_atomicity_and_idempotency(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)

    result = apply_service.execute_backfill_application(
        preparation, preparation.required_confirmation
    )

    assert result.status is BackfillApplyStatus.APPLIED_SUCCESSFULLY
    assert result.backup_path is not None
    assert result.backup_path.name.startswith("Planilha_pre_backfill_rules7_")
    assert "rules5" not in result.backup_path.name
    assert result.original_hash == result.backup_hash
    assert result.applied_updates == 1
    assert result.updates_by_sheet == {"2025": 1}
    assert result.idempotency_result == "NO_ADDITIONAL_CHANGES"
    saved = load_workbook(workbook, read_only=True)
    assert saved["2025"]["B2"].value == "5x LEAPTON LP182"
    assert saved["2025"]["C2"].value == "1x HUAWEI SUN2000"
    assert saved["2025"]["D2"].value == "PRESERVAR"
    saved.close()


def test_backup_hash_mismatch_aborts_before_original_write(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)
    original = workbook.read_bytes()
    monkeypatch.setattr(
        apply_service,
        "_validate_backup",
        lambda *_args, **_kwargs: False,
    )

    with pytest.raises(apply_service.BackfillApplyError) as caught:
        apply_service.execute_backfill_application(
            preparation, preparation.required_confirmation
        )

    assert caught.value.result is not None
    assert caught.value.result.status is BackfillApplyStatus.ABORTED_PRE_FLIGHT
    assert caught.value.result.error_code == "BACKUP_VALIDATION_FAILED"
    assert workbook.read_bytes() == original


def test_fingerprint_conflict_aborts_without_backup(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    changed = load_workbook(workbook)
    changed["2025"]["B2"] = "ALTERADO DEPOIS DA AUDITORIA"
    changed.save(workbook)
    changed.close()
    after_external_change = workbook.read_bytes()
    monkeypatch.setattr(
        apply_service,
        "file_sha256",
        lambda _path: payload["workbook_fingerprint"],
    )
    monkeypatch.setattr(
        apply_service,
        "_create_validated_backup",
        lambda *_args: pytest.fail("backup não pode anteceder fingerprints"),
    )

    with pytest.raises(apply_service.BackfillApplyError) as caught:
        _prepare(workbook, payload)

    assert caught.value.result is not None
    assert (
        caught.value.result.status
        is BackfillApplyStatus.ABORTED_FINGERPRINT_CONFLICT
    )
    assert caught.value.result.fingerprint_conflicts == 1
    assert workbook.read_bytes() == after_external_change


def test_invalid_temporary_never_replaces_original(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)
    original = workbook.read_bytes()
    monkeypatch.setattr(
        apply_service,
        "_validate_temporary_workbook",
        lambda *_args, **_kwargs: False,
    )

    with pytest.raises(apply_service.BackfillApplyError) as caught:
        apply_service.execute_backfill_application(
            preparation, preparation.required_confirmation
        )

    assert caught.value.result is not None
    assert caught.value.result.status is BackfillApplyStatus.ABORTED_UNEXPECTED_CHANGE
    assert workbook.read_bytes() == original


def test_authorized_cells_mask_value_and_type_changes(tmp_path: Path) -> None:
    path = tmp_path / "types.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2025"
    for column in range(1, 7):
        sheet.cell(1, column).value = "HEADER"
    for column in range(1, 6):
        sheet.cell(2, column).value = "PRESERVAR"
    sheet["F2"] = 0
    sheet["F2"].number_format = "General"
    workbook.save(path)
    workbook.close()

    workbook = load_workbook(path)
    sheet = workbook["2025"]
    target_values = {("2025", 2, 6)}
    before = apply_service._preservation_fingerprint(workbook, target_values)

    sheet["F2"] = "5x LEAPTON LP182"

    assert apply_service._preservation_fingerprint(workbook, target_values) == before
    workbook.close()


def test_non_authorized_cell_type_change_is_still_blocked(tmp_path: Path) -> None:
    path = tmp_path / "types.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2025"
    for column in range(1, 7):
        sheet.cell(1, column).value = "HEADER"
    for column in range(1, 6):
        sheet.cell(2, column).value = "PRESERVAR"
    sheet["D2"] = 0
    sheet["D2"].number_format = "General"
    workbook.save(path)
    workbook.close()

    workbook = load_workbook(path)
    sheet = workbook["2025"]
    before = apply_service._preservation_fingerprint(workbook, set())

    sheet["D2"] = "ALTERACAO INDEVIDA"

    assert apply_service._preservation_fingerprint(workbook, set()) != before
    workbook.close()


def test_save_without_updates_preserves_functional_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "features.xlsx"
    _feature_rich_workbook(workbook_path)
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        before = apply_service._preservation_fingerprint(workbook, set())
        temporary = apply_service._save_temporary_workbook(workbook, workbook_path)
    finally:
        workbook.close()

    assert apply_service._validate_temporary_workbook(temporary, [], before)


def test_authorized_cells_are_the_only_functional_changes(tmp_path: Path) -> None:
    workbook_path = tmp_path / "features.xlsx"
    _feature_rich_workbook(workbook_path)
    prepared = [
        {
            "workbook_sheet": "2025",
            "workbook_row": 2,
            "module_column": 2,
            "inverter_column": 3,
            "proposed_module_text": "5x LEAPTON LP182",
            "proposed_inverter_text": "1x HUAWEI SUN2000",
        }
    ]
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        target_values = {("2025", 2, 2), ("2025", 2, 3)}
        before = apply_service._preservation_fingerprint(workbook, target_values)
        sheet = workbook["2025"]
        sheet["B2"] = "5x LEAPTON LP182"
        sheet["C2"] = "1x HUAWEI SUN2000"
        temporary = apply_service._save_temporary_workbook(workbook, workbook_path)
    finally:
        workbook.close()

    assert apply_service._verify_saved_workbook(temporary, prepared, before)
    saved = load_workbook(temporary, data_only=False)
    try:
        sheet = saved["2025"]
        assert sheet["D2"].value == "=LEN(A2)"
        assert sheet.auto_filter.ref == "A1:E2"
        assert "D4:E4" in {str(item) for item in sheet.merged_cells.ranges}
        assert len(sheet.data_validations.dataValidation) == 1
        assert "TabelaBackfillTeste" in sheet.tables
    finally:
        saved.close()


def test_unexpected_functional_change_is_still_blocked(tmp_path: Path) -> None:
    workbook_path = tmp_path / "features.xlsx"
    _feature_rich_workbook(workbook_path)
    prepared = [
        {
            "workbook_sheet": "2025",
            "workbook_row": 2,
            "module_column": 2,
            "inverter_column": 3,
            "proposed_module_text": "5x LEAPTON LP182",
            "proposed_inverter_text": "1x HUAWEI SUN2000",
        }
    ]
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        target_values = {("2025", 2, 2), ("2025", 2, 3)}
        before = apply_service._preservation_fingerprint(workbook, target_values)
        sheet = workbook["2025"]
        sheet["B2"] = "5x LEAPTON LP182"
        sheet["C2"] = "1x HUAWEI SUN2000"
        sheet["D2"] = "=LEN(E2)"
        temporary = apply_service._save_temporary_workbook(workbook, workbook_path)
    finally:
        workbook.close()

    assert not apply_service._verify_saved_workbook(temporary, prepared, before)


def test_post_replace_failure_rolls_back_and_restores_original_hash(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)
    original_hash = preparation.workbook_hash
    real_verify = apply_service._verify_saved_workbook
    calls = 0

    def fail_only_final(path, prepared, preservation):
        nonlocal calls
        calls += 1
        return real_verify(path, prepared, preservation) if calls == 1 else False

    monkeypatch.setattr(apply_service, "_verify_saved_workbook", fail_only_final)

    with pytest.raises(apply_service.BackfillApplyError) as caught:
        apply_service.execute_backfill_application(
            preparation, preparation.required_confirmation
        )

    result = caught.value.result
    assert result is not None
    assert result.status is BackfillApplyStatus.FAILED_ROLLED_BACK
    assert result.rollback_executed is True
    assert result.rollback_status == "CONFIRMED"
    assert audit_service.file_sha256(workbook) == original_hash


def test_unconfirmed_rollback_has_critical_typed_status(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)
    real_verify = apply_service._verify_saved_workbook
    calls = 0

    def fail_only_final(path, prepared, preservation):
        nonlocal calls
        calls += 1
        return real_verify(path, prepared, preservation) if calls == 1 else False

    monkeypatch.setattr(apply_service, "_verify_saved_workbook", fail_only_final)
    monkeypatch.setattr(apply_service, "_restore_validated_backup", lambda *_: False)

    with pytest.raises(apply_service.BackfillApplyError) as caught:
        apply_service.execute_backfill_application(
            preparation, preparation.required_confirmation
        )

    assert caught.value.result is not None
    assert (
        caught.value.result.status
        is BackfillApplyStatus.FAILED_ROLLBACK_UNCONFIRMED
    )


def test_rules7_report_is_timestamped_allowlisted_and_path_free(tmp_path: Path) -> None:
    workbook = tmp_path / "synthetic.xlsx"
    payload = _synthetic_plan(workbook)
    preparation = _prepare(workbook, payload)
    result = apply_service.execute_backfill_application(
        preparation, preparation.required_confirmation
    )

    json_path, markdown_path = write_apply_reports(tmp_path, result)
    public = json.loads(json_path.read_text(encoding="utf-8"))
    serialized = json_path.read_text(encoding="utf-8") + markdown_path.read_text(
        encoding="utf-8"
    )

    assert json_path.name.startswith("historical_equipment_backfill_apply_rules7_")
    assert markdown_path.name.startswith("historical_equipment_backfill_apply_rules7_")
    assert "rules5" not in json_path.name
    assert "rules5" not in markdown_path.name
    assert public["status"] == "APPLIED_SUCCESSFULLY"
    assert public["execution_id"] == payload["execution_id"]
    assert {
        "execution_id",
        "plan_hash",
        "plan_version",
        "technical_processing_format_version",
        "equipment_rules_version",
        "workbook_hash_before",
        "backup_hash",
        "workbook_hash_after",
        "started_at",
        "finished_at",
        "status",
        "updates_planned",
        "updates_applied",
        "updates_by_sheet",
        "fingerprint_conflicts",
        "unexpected_changes",
        "confirmation_received",
        "rollback_executed",
        "rollback_status",
        "idempotency_result",
    }.issubset(public)
    assert str(tmp_path) not in serialized
    assert ".xlsx" not in serialized
    assert "1x LEAPTON" not in serialized


def test_cli_never_prompts_when_preflight_aborts(monkeypatch: pytest.MonkeyPatch) -> None:
    from automacao_gd.presentation import cli

    result = apply_service.BackfillApplyResult.preflight_abort(
        error_code="PRODUCTION_ENVIRONMENT_REQUIRED",
        message="Produção obrigatória.",
    )
    monkeypatch.setattr(cli, "load_backfill_plan", lambda _path: {})
    settings = _settings("development")
    settings.planilha_path = Path("synthetic.xlsx")
    settings.logs_dir_path = Path("synthetic-logs")
    monkeypatch.setattr(cli, "get_settings", lambda: settings)
    monkeypatch.setattr(cli, "write_apply_reports", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        cli,
        "prepare_backfill_application",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(
            apply_service.BackfillApplyError("bloqueado", result=result)
        ),
    )
    monkeypatch.setattr(
        "builtins.input",
        lambda *_args: pytest.fail("confirmação não pode ocorrer antes do pré-voo"),
    )

    assert cli._run_backfill_apply(Path("synthetic-plan.json")) == 2
