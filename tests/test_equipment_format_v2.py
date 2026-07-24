from __future__ import annotations

import ast
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from automacao_gd.domain.equipment import (
    EQUIPMENT_FORMAT_VERSION,
    format_inverters_for_excel_v2,
    format_modules_for_excel_v2,
    split_equipment_values,
)
from automacao_gd.domain.models import GenerationData, InverterEquipment
from automacao_gd.infrastructure.excel.service import update_excel_from_pdf_data
from automacao_gd.infrastructure.pdf.service import parse_generation_data_from_text


FUTURE_REASON = "Etapa 3/4 implementará equipment_format_v2"
CHECKPOINT_REASON = "Etapa 6 implementará checkpoint granular"
PROGRESS_REASON = "Etapa 6 implementará ProgressEvent"
ERROR_REASON = "Etapa 6 implementará erros estruturados"
CLEANUP_REASON = "Etapa 7 implementará limpeza segura de temporários"


def _budget_text(
    *,
    module_manufacturers: str = "RONMA",
    module_models: str = "RM182/144TB 585W",
    module_quantity: str = "14",
    module_total_kwp: str = "8,19",
    inverter_manufacturers: str | None = "HUAWEI",
    inverter_models: str | None = "SUN2000-6KTL",
    inverter_quantity: str | None = "1",
    inverter_total_kw: str | None = "6",
    micro_manufacturers: str | None = None,
    micro_models: str | None = None,
    micro_quantity: str | None = None,
    micro_total_kw: str | None = None,
) -> str:
    parts = [
        f"Fabricante(s) do(s) módulos(s): {module_manufacturers}",
        f"Modelo(s) do(s) módulos(s): {module_models}",
        f"Qtd módulos: {module_quantity}",
        f"Pot. total da(s) placa(s): {module_total_kwp} kWp",
    ]
    if inverter_manufacturers is not None or inverter_models is not None:
        parts.extend(
            [
                f"Fabricante(s) do(s) inversor(es): {inverter_manufacturers or ''}",
                f"Modelo(s) do(s) inversor(es): {inverter_models or ''}",
            ]
        )
        if inverter_quantity is not None:
            parts.append(f"Qtd inversores: {inverter_quantity}")
        if inverter_total_kw is not None:
            parts.append(f"Pot. total do(s) inversor(es): {inverter_total_kw} kW")
    if micro_manufacturers is not None or micro_models is not None:
        parts.extend(
            [
                f"Fabricante(s) do(s) micro-inversor(es): {micro_manufacturers or ''}",
                f"Modelo(s) do(s) micro-inversor(es): {micro_models or ''}",
            ]
        )
        if micro_quantity is not None:
            parts.append(f"Qtd micro-inversores: {micro_quantity}")
        if micro_total_kw is not None:
            parts.append(
                f"Pot. total do(s) micro-inversor(es): {micro_total_kw} kW"
            )
    return "\n".join(parts)


def _create_equipment_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    headers = [
        "Cliente",
        "Protocolo",
        "Data de ingresso",
        "Conclusão",
        "Parecer",
        "Placa",
        "Inversor",
    ]
    ws.append(headers)
    ws.append(
        [
            "Cliente Sintético",
            "2600000001",
            datetime(2026, 1, 10),
            "",
            "Sim",
            "14x OLD MODULE",
            "1x OLD INVERTER",
        ]
    )
    ws.auto_filter.ref = "A1:G2"
    ws.column_dimensions["F"].width = 42
    ws.column_dimensions["G"].width = 48
    fill = PatternFill(fill_type="solid", fgColor="FF7F00")
    for cell in ws[1]:
        cell.fill = fill
    wb.save(path)
    wb.close()


def _read_workbook_bytes(path: Path) -> bytes:
    return path.read_bytes()


def _formatted_equipment_from_text(text: str) -> tuple[str, str]:
    data = parse_generation_data_from_text(text)
    return (
        format_modules_for_excel_v2(data).text,
        format_inverters_for_excel_v2(data).text,
    )


def _write_equipment_to_synthetic_workbook(
    workbook_path: Path,
    *,
    module_text: str,
    inverter_text: str,
    dry_run: bool = False,
) -> dict:
    return update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol="2600000001",
        client_name="Cliente Sintético",
        entry_date="10/01/2026",
        module_text=module_text,
        inverter_text=inverter_text,
        dry_run=dry_run,
    )


def test_module_single_keeps_compact_format_without_total_power() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers="RONMA",
            module_models="RM182/144TB 585W",
            module_quantity="14",
            module_total_kwp="8,19",
        )
    )

    formatted = format_modules_for_excel_v2(data)
    result = formatted.text

    assert result == "14x RONMA RM182/144TB 585W"
    assert formatted.format_version == EQUIPMENT_FORMAT_VERSION
    assert "8,19" not in result
    assert "kWp" not in result
    assert "585W" in result


def test_inverter_single_keeps_compact_format_without_total_power() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="HUAWEI",
            inverter_models="SUN2000-6KTL",
            inverter_quantity="1",
            inverter_total_kw="6",
        )
    )

    formatted = format_inverters_for_excel_v2(data)
    result = formatted.text

    assert result == "1x HUAWEI SUN2000-6KTL"
    assert formatted.format_version == EQUIPMENT_FORMAT_VERSION
    assert "6 kW" not in result
    assert "kW" not in result


def test_solplanet_aiswei_is_canonical_inverter_manufacturer() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="SOLPLANET/AISWEI",
            inverter_models="ASW6000-S-G2",
            inverter_quantity="1",
            inverter_total_kw="6",
        )
    )

    formatted = format_inverters_for_excel_v2(data)

    assert formatted.text == "1x SOLPLANET ASW6000-S-G2"
    assert "SOLPLANET | ASW6000-S-G2" not in formatted.text
    assert "AISWEI | ASW6000-S-G2" not in formatted.text


def test_composite_manufacturer_split_normalizes_solplanet_aiswei() -> None:
    assert split_equipment_values(
        "SOLPLANET/AISWEI",
        context="manufacturer",
    ) == ["SOLPLANET"]


@pytest.mark.parametrize(
    "alias",
    [
        "Solplanet",
        "AISWEI",
        "Solplanet/AISWEI",
        "AISWEI/Solplanet",
        "Solplanet - AISWEI",
        "AISWEI - Solplanet",
        "SOLPLANET AISWEI",
        "AISWEI SOLPLANET",
    ],
)
def test_solplanet_aliases_use_one_canonical_name(alias: str) -> None:
    data = GenerationData(
        inverter_manufacturer=alias,
        inverter_model="ASW6000-S-G2",
        inverter_quantity=1,
    )

    assert format_inverters_for_excel_v2(data).text == (
        "1x SOLPLANET ASW6000-S-G2"
    )


def test_duplicate_solplanet_aliases_for_same_model_are_deduplicated() -> None:
    data = GenerationData(
        inverters=[
            InverterEquipment(manufacturer="SOLPLANET", model="ASW6000-S-G2"),
            InverterEquipment(manufacturer="AISWEI", model="ASW6000-S-G2"),
        ],
        inverter_total_quantity=1,
    )

    assert format_inverters_for_excel_v2(data).text == (
        "1x SOLPLANET ASW6000-S-G2"
    )


def test_solplanet_different_models_are_not_deduplicated() -> None:
    data = GenerationData(
        inverters=[
            InverterEquipment(manufacturer="SOLPLANET", model="ASW5000-S"),
            InverterEquipment(manufacturer="AISWEI", model="ASW6000-S"),
        ],
        inverter_total_quantity=2,
    )

    result = format_inverters_for_excel_v2(data)

    assert result.text == (
        "SOLPLANET | ASW5000-S\n"
        "SOLPLANET | ASW6000-S\n"
        "Qtd. total: 2 inversores"
    )
    assert any("não pôde ser distribuída" in warning for warning in result.warnings)


def test_gokin_is_module_manufacturer_and_not_inverter_manufacturer() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers="GOKIN",
            module_models="GKM-550M",
            module_quantity="20",
            inverter_manufacturers="GOKIN",
            inverter_models="",
            inverter_quantity="1",
            inverter_total_kw="5",
        )
    )

    module_text = format_modules_for_excel_v2(data).text
    inverter = format_inverters_for_excel_v2(data)

    assert module_text == "20x GOKIN GKM-550M"
    assert inverter.text == ""
    assert "GOKIN" not in inverter.text
    assert any("GOKIN" in warning for warning in inverter.warnings)


def test_gokin_is_rejected_as_microinverter_manufacturer() -> None:
    data = GenerationData(
        microinverters=[
            InverterEquipment(
                manufacturer="GOKIN",
                model="MICRO-1000",
                equipment_type="microinverter",
            )
        ],
        microinverter_total_quantity=1,
    )

    formatted = format_inverters_for_excel_v2(data)

    assert formatted.text == ""
    assert any("GOKIN" in warning for warning in formatted.warnings)


def test_modules_multiple_pipe_use_equipment_format_v2_total_label() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers="BYD | TRINA",
            module_models="P6C-30 260 | TSM-NEG21C 695",
            module_quantity="218",
            module_total_kwp="99,31",
        )
    )

    result = format_modules_for_excel_v2(data).text

    assert result == (
        "BYD | P6C-30 260\n"
        "TRINA | TSM-NEG21C 695\n"
        "Qtd. total: 218 módulos"
    )
    assert "99,31" not in result
    assert "kWp" not in result
    assert "109x" not in result


def test_modules_multiple_models_same_manufacturer_include_distributed_quantities() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers="GOKIN",
            module_models="GKM-550M | GKM-555M",
            module_quantity="10 | 12",
            module_total_kwp="12,16",
        )
    )

    result = format_modules_for_excel_v2(data).text

    assert result == (
        "10x GOKIN | GKM-550M\n"
        "12x GOKIN | GKM-555M\n"
        "Qtd. total: 22 m\u00f3dulos"
    )
    assert "12,16" not in result
    assert "kWp" not in result


def test_modules_multiple_slash_separator_is_supported_in_list_context() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers="DMEGC/INTELBRAS",
            module_models="EMSB-535M HC/ DMEGC BIFACE",
            module_quantity="305",
            module_total_kwp="175,43",
        )
    )

    result = format_modules_for_excel_v2(data).text

    assert result == (
        "DMEGC | EMSB-535M HC\n"
        "INTELBRAS | DMEGC BIFACE\n"
        "Qtd. total: 305 módulos"
    )
    assert "175,43" not in result
    assert "kWp" not in result


def test_modules_slash_inside_single_model_is_preserved() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers="RONMA",
            module_models="RM182/144TB 585W",
            module_quantity="14",
            module_total_kwp="8,19",
        )
    )

    result = format_modules_for_excel_v2(data).text

    assert result == "14x RONMA RM182/144TB 585W"


@pytest.mark.parametrize(
    ("manufacturers", "models"),
    [
        ("BYD; TRINA", "P6C-30 260; TSM-NEG21C 695"),
        ("BYD\nTRINA", "P6C-30 260\nTSM-NEG21C 695"),
    ],
)
def test_modules_multiple_semicolon_and_newline_separators(
    manufacturers: str, models: str
) -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            module_manufacturers=manufacturers,
            module_models=models,
            module_quantity="218",
        )
    )

    assert format_modules_for_excel_v2(data).text == (
        "BYD | P6C-30 260\n"
        "TRINA | TSM-NEG21C 695\n"
        "Qtd. total: 218 módulos"
    )


def test_inverters_duplicate_identity_is_consolidated_without_power() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="SOLIS | INTELBRAS | SOLIS",
            inverter_models="S5-GC25K | EGT 25000 MAX | S5-GC25K",
            inverter_quantity="3",
            inverter_total_kw="75",
        )
    )

    result = format_inverters_for_excel_v2(data).text

    assert result == (
        "SOLIS | S5-GC25K\n"
        "INTELBRAS | EGT 25000 MAX\n"
        "Qtd. total: 3 inversores"
    )
    assert "75" not in result
    assert "kW" not in result
    assert "1x" not in result


def test_solplanet_aiswei_from_pdf_path_does_not_create_two_inverters() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="Solplanet/Aiswei",
            inverter_models="ASW6000-S-G2",
            inverter_quantity="1",
            inverter_total_kw="6",
        )
    )

    assert len(data.inverters) == 1
    assert data.inverters[0].manufacturer == "SOLPLANET"
    assert format_inverters_for_excel_v2(data).text == (
        "1x SOLPLANET ASW6000-S-G2"
    )


def test_separate_solplanet_alias_rows_are_deduplicated_during_extraction() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="SOLPLANET | AISWEI",
            inverter_models="ASW6000-S-G2 | ASW6000-S-G2",
            inverter_quantity="1",
            inverter_total_kw="6",
        )
    )

    assert len(data.inverters) == 1
    assert data.inverters[0].manufacturer == "SOLPLANET"
    assert data.inverters[0].model == "ASW6000-S-G2"
    assert format_inverters_for_excel_v2(data).text == "1x SOLPLANET ASW6000-S-G2"


def test_inverters_mismatched_manufacturer_model_counts_warns() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="SOLIS | INTELBRAS | SOLIS",
            inverter_models="S5-GC25K | EGT 25000 MAX",
            inverter_quantity="3",
        )
    )

    formatted = format_inverters_for_excel_v2(data)

    assert formatted.text == (
        "SOLIS | S5-GC25K\n"
        "INTELBRAS | EGT 25000 MAX\n"
        "SOLIS\n"
        "Qtd. total: 3 inversores"
    )
    assert formatted.warnings
    assert (
        "Quantidade de fabricantes e modelos de inversores divergente: "
        "3 fabricantes, 2 modelos."
    ) in formatted.warnings


def test_microinverter_single_keeps_compact_format() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers=None,
            inverter_models=None,
            inverter_quantity=None,
            inverter_total_kw=None,
            micro_manufacturers="APSYSTEMS",
            micro_models="DS3D",
            micro_quantity="1",
        )
    )

    assert format_inverters_for_excel_v2(data).text == "1x MICROINVERSOR APSYSTEMS DS3D"


def test_microinverters_multiple_use_type_prefix_and_v2_total_label() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers=None,
            inverter_models=None,
            inverter_quantity=None,
            inverter_total_kw=None,
            micro_manufacturers="APSYSTEMS | HOYMYLES",
            micro_models="DS3D | HMS-2000DW-4T",
            micro_quantity="11",
        )
    )

    assert format_inverters_for_excel_v2(data).text == (
        "MICROINVERSOR — APSYSTEMS | DS3D\n"
        "MICROINVERSOR — HOYMYLES | HMS-2000DW-4T\n"
        "Qtd. total: 11 microinversores"
    )


def test_conventional_inverter_and_microinverter_are_both_identified_v2() -> None:
    data = parse_generation_data_from_text(
        _budget_text(
            inverter_manufacturers="SOLIS",
            inverter_models="S5-GC25K",
            inverter_quantity="1",
            inverter_total_kw="25",
            micro_manufacturers="APSYSTEMS",
            micro_models="DS3D",
            micro_quantity="10",
            micro_total_kw="20",
        )
    )

    result = format_inverters_for_excel_v2(data).text

    assert result == (
        "INVERSOR — SOLIS | S5-GC25K\n"
        "MICROINVERSOR — APSYSTEMS | DS3D\n"
        "Qtd. total: 1 inversor + 10 microinversores"
    )
    assert "25 kW" not in result
    assert "20" not in result
    assert "kW" not in result


def test_excel_multiline_equipment_applies_wrap_top_alignment_and_row_height(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)
    module_text = "BYD | P6C-30 260\nTRINA | TSM-NEG21C 695\nQtd. total: 218 módulos"
    inverter_text = (
        "SOLIS | S5-GC25K\n"
        "INTELBRAS | EGT 25000 MAX\n"
        "SOLIS | S5-GC25K\n"
        "Qtd. total: 3 inversores"
    )

    result = update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol="2600000001",
        client_name="Cliente Sintético",
        entry_date="10/01/2026",
        module_text=module_text,
        inverter_text=inverter_text,
        dry_run=False,
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        assert result["action"] == "update_existing"
        assert ws.cell(row=2, column=6).alignment.wrap_text is True
        assert ws.cell(row=2, column=7).alignment.wrap_text is True
        assert ws.cell(row=2, column=6).alignment.vertical == "top"
        assert ws.cell(row=2, column=7).alignment.vertical == "top"
        assert (ws.row_dimensions[2].height or 0) >= 45
    finally:
        wb.close()


def test_excel_update_preserves_header_filter_color_and_width(tmp_path: Path) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)

    update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol="2600000001",
        client_name="Cliente Sintético",
        entry_date="10/01/2026",
        module_text="BYD | P6C-30 260\nTRINA | TSM-NEG21C 695\nQtd. total: 218 módulos",
        inverter_text="1x HUAWEI SUN2000-6KTL",
        dry_run=False,
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        assert [cell.value for cell in ws[1]][:7] == [
            "Cliente",
            "Protocolo",
            "Data de ingresso",
            "Conclusão",
            "Parecer",
            "Placa",
            "Inversor",
        ]
        assert ws.auto_filter.ref == "A1:G2"
        assert ws.column_dimensions["F"].width == 42
        assert ws.column_dimensions["G"].width == 48
        assert ws.cell(row=1, column=1).fill.fgColor.rgb == "00FF7F00"
    finally:
        wb.close()


def test_excel_dry_run_does_not_save_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)
    before = _read_workbook_bytes(workbook_path)

    result = update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol="2600000001",
        client_name="Cliente Sintético",
        entry_date="10/01/2026",
        module_text="BYD | P6C-30 260\nTRINA | TSM-NEG21C 695\nQtd. total: 218 módulos",
        inverter_text="1x HUAWEI SUN2000-6KTL",
        dry_run=True,
    )

    assert result["action"] == "update_existing"
    assert _read_workbook_bytes(workbook_path) == before


def test_excel_writes_single_module_and_single_inverter_compact_values(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)
    module_text, inverter_text = _formatted_equipment_from_text(
        _budget_text(
            module_manufacturers="RONMA",
            module_models="RM182/144TB 585W",
            module_quantity="14",
            module_total_kwp="8,19",
            inverter_manufacturers="HUAWEI",
            inverter_models="SUN2000-6KTL",
            inverter_quantity="1",
            inverter_total_kw="6",
        )
    )

    result = _write_equipment_to_synthetic_workbook(
        workbook_path,
        module_text=module_text,
        inverter_text=inverter_text,
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        assert result["action"] == "update_existing"
        assert ws.cell(row=2, column=6).value == "14x RONMA RM182/144TB 585W"
        assert ws.cell(row=2, column=7).value == "1x HUAWEI SUN2000-6KTL"
        assert "8,19" not in ws.cell(row=2, column=6).value
        assert "kWp" not in ws.cell(row=2, column=6).value
        assert "6 kW" not in ws.cell(row=2, column=7).value
    finally:
        wb.close()


def test_excel_writes_multiple_modules_and_inverters_v2_multiline(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)
    module_text, inverter_text = _formatted_equipment_from_text(
        _budget_text(
            module_manufacturers="DMEGC/INTELBRAS",
            module_models="EMSB-535M HC/ DMEGC BIFACE",
            module_quantity="305",
            module_total_kwp="175,43",
            inverter_manufacturers="SOLIS | INTELBRAS | SOLIS",
            inverter_models="S5-GC25K | EGT 25000 MAX | S5-GC25K",
            inverter_quantity="3",
            inverter_total_kw="75",
        )
    )

    _write_equipment_to_synthetic_workbook(
        workbook_path,
        module_text=module_text,
        inverter_text=inverter_text,
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        assert ws.cell(row=2, column=6).value == (
            "DMEGC | EMSB-535M HC\n"
            "INTELBRAS | DMEGC BIFACE\n"
            "Qtd. total: 305 módulos"
        )
        assert ws.cell(row=2, column=7).value == (
            "SOLIS | S5-GC25K\n"
            "INTELBRAS | EGT 25000 MAX\n"
            "Qtd. total: 3 inversores"
        )
        assert "175,43" not in ws.cell(row=2, column=6).value
        assert "kWp" not in ws.cell(row=2, column=6).value
        assert "75 kW" not in ws.cell(row=2, column=7).value
        assert ws.cell(row=2, column=6).alignment.wrap_text is True
        assert ws.cell(row=2, column=7).alignment.vertical == "top"
        assert (ws.row_dimensions[2].height or 0) >= 45
    finally:
        wb.close()


def test_excel_writes_isolated_microinverter_in_inverter_column(tmp_path: Path) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)
    module_text, inverter_text = _formatted_equipment_from_text(
        _budget_text(
            inverter_manufacturers=None,
            inverter_models=None,
            inverter_quantity=None,
            inverter_total_kw=None,
            micro_manufacturers="APSYSTEMS",
            micro_models="DS3D",
            micro_quantity="1",
            micro_total_kw="2",
        )
    )

    _write_equipment_to_synthetic_workbook(
        workbook_path,
        module_text=module_text,
        inverter_text=inverter_text,
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        assert ws.cell(row=2, column=7).value == "1x MICROINVERSOR APSYSTEMS DS3D"
        assert "2 kW" not in ws.cell(row=2, column=7).value
    finally:
        wb.close()


def test_excel_writes_conventional_inverter_and_microinverter_together(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)
    module_text, inverter_text = _formatted_equipment_from_text(
        _budget_text(
            inverter_manufacturers="SOLIS",
            inverter_models="S5-GC25K",
            inverter_quantity="1",
            inverter_total_kw="25",
            micro_manufacturers="APSYSTEMS",
            micro_models="DS3D",
            micro_quantity="10",
            micro_total_kw="20",
        )
    )

    _write_equipment_to_synthetic_workbook(
        workbook_path,
        module_text=module_text,
        inverter_text=inverter_text,
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        assert ws.cell(row=2, column=7).value == (
            "INVERSOR — SOLIS | S5-GC25K\n"
            "MICROINVERSOR — APSYSTEMS | DS3D\n"
            "Qtd. total: 1 inversor + 10 microinversores"
        )
        assert "25 kW" not in ws.cell(row=2, column=7).value
        assert "20 kW" not in ws.cell(row=2, column=7).value
    finally:
        wb.close()


def test_excel_existing_protocol_update_preserves_non_equipment_cell_values(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)

    before_wb = load_workbook(workbook_path)
    before_ws = before_wb["2026"]
    before_values = [before_ws.cell(row=2, column=column).value for column in range(1, 8)]
    before_wb.close()

    _write_equipment_to_synthetic_workbook(
        workbook_path,
        module_text="BYD | P6C-30 260\nTRINA | TSM-NEG21C 695\nQtd. total: 218 módulos",
        inverter_text="1x HUAWEI SUN2000-6KTL",
    )

    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    try:
        after_values = [ws.cell(row=2, column=column).value for column in range(1, 8)]
        assert after_values[:5] == before_values[:5]
        assert after_values[5] != before_values[5]
        assert after_values[6] != before_values[6]
    finally:
        wb.close()


def test_excel_workbook_remains_valid_after_equipment_update(tmp_path: Path) -> None:
    workbook_path = tmp_path / "planilha_sintetica.xlsx"
    _create_equipment_workbook(workbook_path)

    _write_equipment_to_synthetic_workbook(
        workbook_path,
        module_text="BYD | P6C-30 260\nTRINA | TSM-NEG21C 695\nQtd. total: 218 módulos",
        inverter_text="1x HUAWEI SUN2000-6KTL",
    )

    wb = load_workbook(workbook_path)
    try:
        assert "2026" in wb.sheetnames
        assert wb["2026"].cell(row=2, column=2).value == "2600000001"
    finally:
        wb.close()


def test_checkpoint_granular_state_contract_contains_expected_steps() -> None:
    from automacao_gd.infrastructure.state.pipeline_state import PipelineStateStore

    store = PipelineStateStore(path=Path("synthetic_state.json"), resume=False)
    entry = store.protocol_entry("2600000001")
    expected_steps = [
        "discovered",
        "selected",
        "detail_opened",
        "pdf_downloaded",
        "pdf_validated",
        "pdf_parsed",
        "client_folder_resolved",
        "excel_updated",
        "pdf_archived",
        "completed",
    ]

    assert [step["name"] for step in entry["steps"]] == expected_steps
    for step in entry["steps"]:
        assert {
            "status",
            "started_at",
            "finished_at",
            "input_fingerprint",
            "output_fingerprint",
            "error",
            "equipment_format_version",
        } <= set(step)


def test_progress_event_contract_and_monotonic_sequence() -> None:
    from automacao_gd.application.contracts import ProgressEvent

    events = [
        ProgressEvent(overall_percent=0, stage_percent=0, protocol_percent=0, stage="preflight"),
        ProgressEvent(overall_percent=35, stage_percent=100, protocol_percent=20, stage="download"),
        ProgressEvent(overall_percent=100, stage_percent=100, protocol_percent=100, stage="cleanup"),
    ]

    assert [event.overall_percent for event in events] == sorted(
        event.overall_percent for event in events
    )
    assert events[0].overall_percent == 0
    assert events[-1].overall_percent == 100
    for event in events:
        assert hasattr(event, "stage")
        assert hasattr(event, "protocol")
        assert hasattr(event, "message")


def test_protocol_progress_stage_order_contract() -> None:
    from automacao_gd.application.contracts import protocol_progress_stages

    assert protocol_progress_stages() == [
        "localizar_pagina_protocolo",
        "abrir_detalhe",
        "baixar_ou_reutilizar_pdf",
        "extrair_pdf",
        "localizar_pasta",
        "atualizar_planilha",
        "arquivar_documento",
        "salvar_checkpoint",
    ]


@pytest.mark.parametrize(
    ("code", "stage", "suggested_action"),
    [
        ("EXCEL_LOCKED", "excel_update", "Fechar a planilha e executar novamente."),
        ("CLIENT_FOLDER_AMBIGUOUS", "archive", "pendência de conferência"),
        ("PDF_INVALID_SIGNATURE", "pdf_validation", "conferência"),
    ],
)
def test_structured_operation_error_contract(
    code: str, stage: str, suggested_action: str
) -> None:
    from automacao_gd.application.contracts import OperationError

    error = OperationError(
        run_id="run-sintetico",
        protocol="2600000001",
        stage=stage,
        code=code,
        user_message="Mensagem operacional sintética.",
        technical_cause="Causa técnica sintética.",
        recoverable=True,
        action_taken="Operação interrompida com segurança.",
        suggested_action=suggested_action,
        traceback_ref="logs/sintetico.json#1",
    )

    assert error.code == code
    assert error.stage == stage
    assert suggested_action in error.suggested_action


def test_cleanup_dry_run_lists_temp_files_without_deleting(tmp_path: Path) -> None:
    from automacao_gd.infrastructure.files.cleanup import cleanup_temp_files

    temp_file = tmp_path / ".pipeline.tmp"
    temp_file.write_text("temp", encoding="utf-8")

    report = cleanup_temp_files(tmp_path, dry_run=True)

    assert temp_file in report["would_remove"]
    assert temp_file.exists()


def test_cleanup_apply_removes_only_allowed_temp_files(tmp_path: Path) -> None:
    from automacao_gd.infrastructure.files.cleanup import cleanup_temp_files

    removable = [
        tmp_path / ".a.tmp",
        tmp_path / ".b.temp",
        tmp_path / "download.crdownload",
        tmp_path / "download.part",
        tmp_path / "~$planilha.xlsx",
    ]
    protected = [
        tmp_path / "orcamento.pdf",
        tmp_path / "planilha.xlsx",
        tmp_path / "pipeline_cdp_state.json",
        tmp_path / "client_folder_cache.json",
        tmp_path / "relatorio.md",
        tmp_path / "backup.bak",
    ]
    for path in [*removable, *protected]:
        path.write_text("synthetic", encoding="utf-8")

    cleanup_temp_files(tmp_path, dry_run=False)

    assert all(not path.exists() for path in removable)
    assert all(path.exists() for path in protected)


def test_cleanup_rejects_path_traversal_and_respects_allowed_root(tmp_path: Path) -> None:
    from automacao_gd.infrastructure.files.cleanup import cleanup_temp_files

    outside = tmp_path.parent / "outside.tmp"
    outside.write_text("synthetic", encoding="utf-8")

    with pytest.raises(ValueError):
        cleanup_temp_files(tmp_path, candidates=[outside], dry_run=False)

    assert outside.exists()


def test_domain_layer_does_not_import_infrastructure_or_ui_frameworks() -> None:
    forbidden_roots = {
        "automacao_gd.infrastructure",
        "automacao_gd.presentation",
        "openpyxl",
        "playwright",
        "PySide6",
        "tkinter",
    }
    domain_root = Path("automacao_gd/domain")

    for path in domain_root.glob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            imported = None
            if isinstance(node, ast.ImportFrom):
                imported = node.module
            elif isinstance(node, ast.Import):
                imported = node.names[0].name
            if imported:
                assert not any(
                    imported == forbidden or imported.startswith(f"{forbidden}.")
                    for forbidden in forbidden_roots
                ), f"{path} importa dependência proibida na camada domain: {imported}"


def test_future_apps_desktop_does_not_import_infrastructure_directly_if_present() -> None:
    apps_desktop = Path("apps/desktop")
    if not apps_desktop.exists():
        return

    forbidden_imports = {
        "automacao_gd.infrastructure.portal.cdp_service",
        "automacao_gd.infrastructure.excel.service",
        "playwright",
        "openpyxl",
    }
    for path in apps_desktop.rglob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            imported = None
            if isinstance(node, ast.ImportFrom):
                imported = node.module
            elif isinstance(node, ast.Import):
                imported = node.names[0].name
            if imported:
                assert not any(
                    imported == forbidden or imported.startswith(f"{forbidden}.")
                    for forbidden in forbidden_imports
                ), f"{path} importa infraestrutura proibida para apps/desktop: {imported}"
