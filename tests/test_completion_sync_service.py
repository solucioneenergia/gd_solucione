from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from automacao_gd.application.completion_sync_service import (
    CompletionPortalRecord,
    CompletionSyncAction,
    CompletionSyncStatus,
    COMPLETION_APPLY_REPORT_PREFIX,
    OPTION5_COMPLETED_PIPELINE_SOURCE,
    run_completion_sync,
    run_completion_sync_from_records,
)
from automacao_gd.application.contracts import OperationResult, OperationStatus
from automacao_gd.infrastructure.config import Settings
from automacao_gd.presentation.cli import (
    COMPLETION_SYNC_STRONG_CONFIRMATION,
    OPTION5_COMPLETION_STRONG_CONFIRMATION,
    _confirmed_completion_sync,
    _confirmed_pipeline,
    main as cli_main,
)


HEADERS = [
    "Cliente",
    "Protocolo",
    "Data de ingresso",
    "Conclusão",
    "Parecer",
    "Placa",
    "Inversor",
]


def _workbook(path: Path, rows: list[tuple[str, object, object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    ws.append(HEADERS)
    for protocol, entry_date, completion in rows:
        ws.append(
            [
                "Cliente teste",
                protocol,
                entry_date,
                completion,
                True,
                "1x MODULO TESTE",
                "1x INVERSOR TESTE",
            ]
        )
    wb.save(path)
    wb.close()
    return path


def _cell(path: Path, coord: str):
    wb = load_workbook(path, data_only=False)
    try:
        return wb["2026"][coord].value, wb["2026"][coord].number_format
    finally:
        wb.close()


def test_completed_project_writes_excel_date_with_format(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000001", date(2026, 7, 1), None)])

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000001", status="Concluído", completion_date="15/07/2026")],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["total_updated_dates"] == 1
    assert result["protocol_results"][0]["action"] == CompletionSyncAction.COMPLETION_DATE_UPDATED.value
    value, number_format = _cell(workbook, "D2")
    assert value == datetime(2026, 7, 15)
    assert number_format == "dd/mm/yyyy"


def test_open_project_marks_em_aberto(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000002", date(2026, 7, 1), None)])

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000002", status="Em andamento", completion_date=None)],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["protocol_results"][0]["action"] == CompletionSyncAction.MARKED_AS_OPEN.value
    assert _cell(workbook, "D2")[0] == "EM ABERTO"


def test_open_project_that_becomes_completed_replaces_em_aberto(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000003", date(2026, 7, 1), "EM ABERTO")])

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000003", status="Concluído", completion_date="20/07/2026")],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["protocol_results"][0]["action"] == CompletionSyncAction.COMPLETION_DATE_UPDATED.value
    assert _cell(workbook, "D2")[0] == datetime(2026, 7, 20)


def test_same_date_is_no_change(tmp_path: Path) -> None:
    workbook = _workbook(
        tmp_path / "planilha.xlsx",
        [("2600000004", date(2026, 7, 1), datetime(2026, 7, 15))],
    )

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000004", status="Concluído", completion_date="15/07/2026")],
        workbook_path=workbook,
        dry_run=True,
        apply_changes=False,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["protocol_results"][0]["action"] == CompletionSyncAction.NO_CHANGE.value


def test_divergent_date_is_pending_and_does_not_write(tmp_path: Path) -> None:
    workbook = _workbook(
        tmp_path / "planilha.xlsx",
        [("2600000005", date(2026, 7, 1), datetime(2026, 7, 14))],
    )

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000005", status="Concluído", completion_date="15/07/2026")],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["protocol_results"][0]["status"] == CompletionSyncStatus.PENDING_REVIEW.value
    assert result["protocol_results"][0]["reason"] == "COMPLETION_DATE_CONFLICT"
    assert _cell(workbook, "D2")[0] == datetime(2026, 7, 14)


def test_status_regression_preserves_existing_date(tmp_path: Path) -> None:
    workbook = _workbook(
        tmp_path / "planilha.xlsx",
        [("2600000006", date(2026, 7, 1), datetime(2026, 7, 15))],
    )

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000006", status="Em andamento", completion_date=None)],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["protocol_results"][0]["reason"] == "PORTAL_COMPLETION_STATUS_REGRESSION"
    assert _cell(workbook, "D2")[0] == datetime(2026, 7, 15)


def test_completed_without_date_is_pending(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000007", date(2026, 7, 1), None)])

    result = run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000007", status="Concluído", completion_date=None)],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["protocol_results"][0]["reason"] == "PORTAL_COMPLETION_DATE_MISSING"
    assert _cell(workbook, "D2")[0] is None


def test_option5_completed_without_date_is_marked_open_proposal(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000017", date(2026, 7, 1), None)])

    result = run_completion_sync_from_records(
        [
            CompletionPortalRecord(
                protocol="2600000017",
                status="Solicitação Concluída",
                completion_date=None,
                source=f"{OPTION5_COMPLETED_PIPELINE_SOURCE}:portal_detail_page_1",
            )
        ],
        workbook_path=workbook,
        dry_run=True,
        apply_changes=False,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    row = result["protocol_results"][0]
    assert row["action"] == CompletionSyncAction.MARKED_AS_OPEN.value
    assert row["reason"] == "POINT_OF_CONNECTION_COMPLETION_DATE_NOT_AVAILABLE"
    assert row["proposed_completion"] == "EM ABERTO"
    assert _cell(workbook, "D2")[0] is None


def test_allowlist_preserves_other_columns(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000008", date(2026, 7, 1), None)])

    run_completion_sync_from_records(
        [CompletionPortalRecord(protocol="2600000008", status="Concluído", completion_date="15/07/2026")],
        workbook_path=workbook,
        dry_run=False,
        apply_changes=True,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    wb = load_workbook(workbook, data_only=False)
    try:
        ws = wb["2026"]
        assert ws["A2"].value == "Cliente teste"
        assert ws["B2"].value == "2600000008"
        assert ws["C2"].value == datetime(2026, 7, 1)
        assert ws["E2"].value is True
        assert ws["F2"].value == "1x MODULO TESTE"
        assert ws["G2"].value == "1x INVERSOR TESTE"
    finally:
        wb.close()


def test_rollback_preserves_original_on_replace_failure(tmp_path: Path) -> None:
    workbook = _workbook(tmp_path / "planilha.xlsx", [("2600000009", date(2026, 7, 1), None)])

    def fail_replace(src: Path, dst: Path) -> None:
        raise PermissionError("simulated replace failure")

    with pytest.raises(PermissionError):
        run_completion_sync_from_records(
            [CompletionPortalRecord(protocol="2600000009", status="Concluído", completion_date="15/07/2026")],
            workbook_path=workbook,
            dry_run=False,
            apply_changes=True,
            max_protocols=5,
            today=date(2026, 7, 27),
            replace_workbook=fail_replace,
        )

    assert _cell(workbook, "D2")[0] is None


def test_limit_caps_protocols_analyzed(tmp_path: Path) -> None:
    workbook = _workbook(
        tmp_path / "planilha.xlsx",
        [(f"26000000{i:02d}", date(2026, 7, 1), None) for i in range(10)],
    )
    records = [
        CompletionPortalRecord(protocol=f"26000000{i:02d}", status="Em andamento", completion_date=None)
        for i in range(10)
    ]

    result = run_completion_sync_from_records(
        records,
        workbook_path=workbook,
        dry_run=True,
        apply_changes=False,
        max_protocols=5,
        today=date(2026, 7, 27),
    )

    assert result["total_protocols_available"] == 10
    assert result["total_protocols_analyzed"] == 5
    assert len(result["protocol_results"]) == 5


def test_run_sync_uses_only_option5_protocols_and_dry_run_prevents_write(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = _workbook(
        tmp_path / "planilha.xlsx",
        [(f"26000010{i}", date(2026, 7, 1), None) for i in range(7)],
    )
    logs_dir = tmp_path / "logs"
    logs_dir.mkdir()
    (logs_dir / "downloads_orcamentos_concluidos_cdp.json").write_text(
        """
        {
          "finished_at": "2026-07-27T10:05:22",
          "results": [
            {"protocol": "260000100", "status": "Solicitação Concluída", "download_status": "downloaded"},
            {"protocol": "260000101", "status": "Solicitação Concluída", "download_status": "existing_pdf_after_skip"},
            {"protocol": "260000102", "status": "Solicitação Concluída", "download_status": "existing_pdf_after_skip"},
            {"protocol": "260000103", "status": "Solicitação Concluída", "download_status": "existing_pdf_after_skip"},
            {"protocol": "260000104", "status": "Solicitação Concluída", "download_status": "existing_pdf_after_skip"},
            {"protocol": "260000105", "status": "Solicitação Concluída", "download_status": "existing_pdf_after_skip"},
            {"protocol": "260000999", "status": "Em andamento", "download_status": "ignored"}
          ]
        }
        """,
        encoding="utf-8",
    )
    (logs_dir / "pipeline_cdp_completo.json").write_text(
        '{"finished_at": "2026-07-27T10:05:22"}',
        encoding="utf-8",
    )
    (logs_dir / "processamento_pdfs_planilha_clientes.json").write_text(
        '{"results": []}',
        encoding="utf-8",
    )

    class ReadyPreflight:
        ready = True

        def raise_if_blocked(self) -> None:
            raise AssertionError("preflight should be ready")

    captured_targets: list[str] = []

    def fake_preflight(*args, **kwargs):
        return ReadyPreflight()

    def fake_collect(settings: Settings, *, target_records: list[CompletionPortalRecord] | None = None):
        assert target_records is not None
        captured_targets.extend(record.protocol for record in target_records)
        return [
            CompletionPortalRecord(
                protocol=record.protocol,
                status="Solicitação Concluída",
                completion_date=None,
                source=f"{OPTION5_COMPLETED_PIPELINE_SOURCE}:portal_detail_page_1",
                option5_run_at=record.option5_run_at,
                option5_action=record.option5_action,
                pdf_source=record.pdf_source,
            )
            for record in target_records
        ]

    monkeypatch.setattr(
        "automacao_gd.application.completion_sync_service.run_preflight",
        fake_preflight,
    )
    monkeypatch.setattr(
        "automacao_gd.application.completion_sync_service.collect_completion_records_from_portal",
        fake_collect,
    )

    settings = Settings(
        APP_ENV="production",
        PLANILHA_PATH=workbook,
        LOGS_DIR=logs_dir,
        SYNC_COMPLETION_STATUS=True,
        DRY_RUN=True,
        APPLY_COMPLETION_STATUS=False,
        APPLY_EXCEL=True,
        MAX_COMPLETION_PROTOCOLS_PER_RUN=5,
    )

    result = run_completion_sync(settings)

    assert captured_targets == [
        "260000100",
        "260000101",
        "260000102",
        "260000103",
        "260000104",
    ]
    assert result["origin"] == OPTION5_COMPLETED_PIPELINE_SOURCE
    assert result["total_option5_protocols_available"] == 6
    assert result["total_protocols_analyzed"] == 5
    assert result["protocols_outside_option5"] == 0
    assert result["protocols_added_after_limit"] == 0
    assert result["protocols_duplicated"] == 0
    assert result["total_marked_open"] == 5
    assert _cell(workbook, "D2")[0] is None


def test_real_sync_uses_approved_dry_run_report_without_portal(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = _workbook(
        tmp_path / "planilha.xlsx",
        [(f"26000020{i}", date(2026, 7, 1), None) for i in range(5)],
    )
    logs_dir = tmp_path / "logs"
    logs_dir.mkdir()
    dry_run_rows = [
        {
            "protocol": f"26000020{i}",
            "portal_status": "Concluído",
            "portal_completion_date": None,
            "source": f"{OPTION5_COMPLETED_PIPELINE_SOURCE}:portal_detail_page_1",
            "origin": OPTION5_COMPLETED_PIPELINE_SOURCE,
            "option5_run_at": "2026-07-27T10:05:22",
            "pdf_source": "reutilizado",
            "worksheet": "2026",
            "row": i + 2,
            "current_completion": None,
            "proposed_completion": "EM ABERTO",
            "action": CompletionSyncAction.MARKED_AS_OPEN.value,
            "status": CompletionSyncStatus.MARKED_AS_OPEN.value,
            "reason": "POINT_OF_CONNECTION_COMPLETION_DATE_NOT_AVAILABLE",
            "write_permission": True,
        }
        for i in range(5)
    ]
    (logs_dir / "completion_status_sync_20260727T122433.json").write_text(
        __import__("json").dumps(
            {
                "status": "SUCESSO",
                "dry_run": True,
                "origin": OPTION5_COMPLETED_PIPELINE_SOURCE,
                "protocols_outside_option5": 0,
                "protocols_duplicated": 0,
                "protocol_results": dry_run_rows,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    class ReadyPreflight:
        ready = True

        def raise_if_blocked(self) -> None:
            raise AssertionError("preflight should be ready")

    def fake_preflight(*args, **kwargs):
        assert kwargs["real_run"] is True
        assert kwargs["require_cdp"] is False
        return ReadyPreflight()

    def fail_portal_collect(*args, **kwargs):
        raise AssertionError("real apply must not collect portal records")

    monkeypatch.setattr(
        "automacao_gd.application.completion_sync_service.run_preflight",
        fake_preflight,
    )
    monkeypatch.setattr(
        "automacao_gd.application.completion_sync_service.collect_completion_records_from_portal",
        fail_portal_collect,
    )
    settings = Settings(
        APP_ENV="production",
        PLANILHA_PATH=workbook,
        LOGS_DIR=logs_dir,
        SYNC_COMPLETION_STATUS=True,
        DRY_RUN=False,
        APPLY_COMPLETION_STATUS=True,
        MAX_COMPLETION_PROTOCOLS_PER_RUN=5,
    )

    result = run_completion_sync(settings)

    assert result["mode"] == "APPLY"
    assert result["source_dry_run_report"] == "completion_status_sync_20260727T122433.json"
    assert result["total_protocols_analyzed"] == 5
    assert result["total_marked_open"] == 5
    assert result["total_updates_applied"] == 5
    assert result["backup_created"] is True
    assert result["backup_validated"] is True
    assert result["atomic_replace"] == "SUCCESS"
    assert result["rollback"] == "NOT_REQUIRED"
    assert Path(result["json_report_path"]).name.startswith(
        f"{COMPLETION_APPLY_REPORT_PREFIX}_"
    )
    for row in range(2, 7):
        assert _cell(workbook, f"D{row}")[0] == "EM ABERTO"


def test_completion_sync_real_write_rejects_plain_sim(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class FakeSettings:
        DRY_RUN = False
        APPLY_COMPLETION_STATUS = True

    class FakeController:
        settings = FakeSettings()

        def sync_completion_status(self):
            raise AssertionError("sync must not run without strong confirmation")

    monkeypatch.setattr("builtins.input", lambda _: "SIM")

    result = _confirmed_completion_sync(FakeController())

    assert result.status == "BLOQUEADO"
    assert result.payload["code"] == "STRONG_CONFIRMATION_REQUIRED"
    assert result.payload["confirmation_required"] == COMPLETION_SYNC_STRONG_CONFIRMATION


def test_option5_real_pipeline_rejects_plain_sim(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    calls: list[str] = []

    class FakeSettings:
        DRY_RUN = False
        MAX_COMPLETED_TO_PROCESS = 5

    class FakeController:
        settings = FakeSettings()

        def preflight(self, require_cdp: bool = False):
            calls.append(f"preflight:{require_cdp}")
            raise AssertionError("preflight must not run without strong confirmation")

        def run_pipeline(self):
            calls.append("run_pipeline")
            raise AssertionError("pipeline must not run without strong confirmation")

    monkeypatch.setattr("builtins.input", lambda _: "SIM")

    result = _confirmed_pipeline(FakeController())

    assert result.status == OperationStatus.BLOQUEADO
    assert result.payload["cancelled_by_user"] is True
    assert result.payload["confirmation_required"] == OPTION5_COMPLETION_STRONG_CONFIRMATION
    assert "confirmação forte inválida ou ausente" in result.message
    assert calls == []


def test_cli_menu_no_longer_exposes_completion_option_7(
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr("builtins.input", lambda _: "0")

    exit_code = cli_main([])

    output = capsys.readouterr().out
    assert exit_code == 0
    assert "5 - Executar pipeline CDP completo" in output
    assert "7 - Sincronizar datas de conclusão" not in output
