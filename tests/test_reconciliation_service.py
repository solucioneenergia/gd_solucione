from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from automacao_gd.application.reconciliation_service import (
    ReconciliationIssueType,
    WorkbookProtocolRecord,
    build_workbook_protocol_index,
    normalize_protocol,
    reconcile_portal_workbook,
    save_reconciliation_reports,
)
from automacao_gd.domain.models import PortalSolicitation


def _portal_records(*protocols: str) -> list[PortalSolicitation]:
    return [
        PortalSolicitation(
            protocol=protocol,
            status="Solicitação Concluída",
            entry_date="01/01/2026",
            page_number=(index // 50) + 1,
            row_index=index + 1,
        )
        for index, protocol in enumerate(protocols)
    ]


def _workbook(path: Path, rows: list[tuple[str, object, object, object, object, object]]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    ws.append(["Cliente", "Protocolo", "Data de ingresso", "Conclusão", "Parecer", "Placa", "Inversor"])
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def test_reconciliation_audits_all_portal_protocols_before_operational_limit(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [(f"Cliente {index}", f"2600{index:04d}", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV") for index in range(100)],
    )
    portal_records = _portal_records(*(f"2600{index:04d}" for index in range(100)))

    result = reconcile_portal_workbook(
        portal_records,
        workbook_path,
        operational_selected_protocols=[f"2600{index:04d}" for index in range(5)],
    )

    assert result.portal_summary["portal_concluded_unique"] == 100
    assert result.operational_batch["operational_protocols_selected"] == 5
    assert result.operational_batch["reconciliation_protocols_audited"] == 100


def test_missing_protocol_is_reported_without_creating_workbook_row(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV")],
    )

    result = reconcile_portal_workbook(_portal_records("2600000001", "2600000002"), workbook_path)

    assert [item["protocol"] for item in result.missing_in_workbook] == ["2600000002"]
    assert result.missing_in_workbook[0]["issue_type"] == ReconciliationIssueType.PORTAL_COMPLETED_MISSING_IN_WORKBOOK
    wb = load_workbook(workbook_path, read_only=True)
    assert wb["2026"].max_row == 2
    wb.close()


def test_duplicate_protocols_are_reported_in_same_sheet_and_across_sheets(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [
            ("Cliente", "2600000001", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV"),
            ("Cliente", "2600000001", datetime(2026, 1, 2), None, True, "1x MOD", "1x INV"),
        ],
    )
    wb = load_workbook(workbook_path)
    ws = wb.create_sheet("2025")
    ws.append(["Cliente", "Protocolo", "Data de ingresso", "Conclusão", "Parecer", "Placa", "Inversor"])
    ws.append(["Cliente", "2600000001", datetime(2025, 1, 1), None, True, "1x MOD", "1x INV"])
    wb.save(workbook_path)

    result = reconcile_portal_workbook(_portal_records("2600000001"), workbook_path)

    duplicates = result.duplicate_workbook_protocols
    assert len(duplicates) == 1
    assert duplicates[0]["issue_type"] == ReconciliationIssueType.DUPLICATE_PROTOCOL_IN_WORKBOOK
    assert {(item["sheet_name"], item["row_number"]) for item in duplicates[0]["occurrences"]} == {
        ("2026", 2),
        ("2026", 3),
        ("2025", 2),
    }


def test_numeric_protocol_is_compared_canonically_without_modifying_cell(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", 2600000001, datetime(2026, 1, 1), None, True, "1x MOD", "1x INV")],
    )

    result = reconcile_portal_workbook(_portal_records("2600000001"), workbook_path)

    record = result.workbook_index.records_by_protocol["2600000001"][0]
    assert record.protocol_cell_type == "numeric"
    assert result.set_reconciliation["matched_unique"] == 1
    wb = load_workbook(workbook_path, read_only=True)
    assert isinstance(wb["2026"]["B2"].value, int)
    wb.close()


def test_wrong_year_sheet_and_grouped_year_sheet_are_classified(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", datetime(2025, 1, 1), None, True, "1x MOD", "1x INV")],
    )
    wb = load_workbook(workbook_path)
    ws = wb.create_sheet("2022 - 2023")
    ws.append(["Cliente", "Protocolo", "Data de ingresso", "Conclusão", "Parecer", "Placa", "Inversor"])
    ws.append(["Cliente", "2300000001", datetime(2023, 1, 1), None, True, "1x MOD", "1x INV"])
    wb.save(workbook_path)

    result = reconcile_portal_workbook(_portal_records("2600000001", "2300000001"), workbook_path)

    assert [item["protocol"] for item in result.wrong_year_sheet] == ["2600000001"]
    assert result.wrong_year_sheet[0]["issue_type"] == ReconciliationIssueType.PROTOCOL_IN_WRONG_YEAR_SHEET


def test_invalid_ingress_date_disables_year_inference(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", "sem data", None, True, "1x MOD", "1x INV")],
    )

    result = reconcile_portal_workbook(_portal_records("2600000001"), workbook_path)

    assert result.structural_findings[0]["issue_type"] == ReconciliationIssueType.YEAR_SHEET_VALIDATION_UNAVAILABLE


def test_completion_and_equipment_empty_are_classified_by_status(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [
            ("Cliente", "2600000001", datetime(2026, 1, 1), None, True, "", ""),
            ("Cliente", "2600000002", datetime(2026, 1, 1), None, "Cancelado", "", ""),
            ("Cliente", "2600000003", datetime(2026, 1, 1), "texto invalido", None, "", ""),
        ],
    )

    result = reconcile_portal_workbook(_portal_records("2600000001", "2600000002", "2600000003"), workbook_path)

    completion_types = {item["protocol"]: item["issue_type"] for item in result.completion_audit}
    assert completion_types["2600000001"] == ReconciliationIssueType.COMPLETION_EMPTY
    equipment_types = {item["protocol"]: item["issue_type"] for item in result.equipment_field_audit}
    assert equipment_types["2600000001"] == ReconciliationIssueType.EQUIPMENT_EMPTY_REQUIRES_REVIEW
    assert equipment_types["2600000002"] == ReconciliationIssueType.EQUIPMENT_EMPTY_EXPECTED_BY_STATUS
    incomplete = {item["protocol"]: item["missing_fields"] for item in result.incomplete_records}
    assert "parecer" in incomplete["2600000003"]
    assert "completion" in incomplete["2600000003"]


def test_trailing_empty_rows_are_structural_not_project_rows(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV")],
    )
    wb = load_workbook(workbook_path)
    ws = wb["2026"]
    ws.cell(row=5, column=1, value="")
    wb.save(workbook_path)

    result = reconcile_portal_workbook(_portal_records("2600000001"), workbook_path)

    assert result.workbook_summary["workbook_project_rows"] == 1
    assert result.structural_findings[-1]["issue_type"] == ReconciliationIssueType.TRAILING_MATERIALIZED_EMPTY_ROW


def test_workbook_only_protocol_is_informative(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV")],
    )

    result = reconcile_portal_workbook(_portal_records("2600000002"), workbook_path)

    assert result.workbook_only_protocols[0]["issue_type"] == ReconciliationIssueType.WORKBOOK_PROTOCOL_NOT_IN_CURRENT_PORTAL_COMPLETED_SET


def test_read_only_reconciliation_preserves_sha_and_generates_private_reports(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente Nome Sensivel", "2600000001", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV")],
    )
    before = workbook_path.read_bytes()

    result = reconcile_portal_workbook(_portal_records("2600000001"), workbook_path)
    json_path, md_path = save_reconciliation_reports(result, tmp_path)

    assert workbook_path.read_bytes() == before
    assert result.safety["workbook_sha_before_reconciliation"] == result.safety["workbook_sha_after_reconciliation"]
    assert "Cliente Nome Sensivel" not in json_path.read_text(encoding="utf-8")
    assert "Cliente Nome Sensivel" not in md_path.read_text(encoding="utf-8")


def test_partial_pagination_marks_reconciliation_not_authoritative(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", datetime(2026, 1, 1), None, True, "1x MOD", "1x INV")],
    )

    result = reconcile_portal_workbook(
        _portal_records("2600000001"),
        workbook_path,
        pagination={
            "pagination_complete": False,
            "last_page_confirmed": False,
            "last_page_number": None,
            "next_page_available_after_stop": True,
            "pagination_stop_reason": "safety_cap_reached_with_next_page",
            "pagination_safety_cap": 11,
            "pages_visited": list(range(1, 12)),
        },
    )

    assert result.metrics_scope == "partial"
    assert result.authoritative_status["set_reconciliation_authoritative"] is False
    assert result.decision == "STAGE2_BLOCKED - PORTAL_PAGINATION_INCOMPLETE"


def test_complete_pagination_marks_reconciliation_authoritative(tmp_path: Path) -> None:
    workbook_path = _workbook(
        tmp_path / "planilha.xlsx",
        [("Cliente", "2600000001", datetime(2026, 1, 1), datetime(2026, 2, 1), True, "1x MOD", "1x INV")],
    )

    result = reconcile_portal_workbook(
        _portal_records("2600000001"),
        workbook_path,
        pagination={
            "pagination_complete": True,
            "last_page_confirmed": True,
            "last_page_number": 14,
            "next_page_available_after_stop": False,
            "pagination_stop_reason": "last_page_reached",
            "pagination_safety_cap": 20,
            "pages_visited": list(range(1, 15)),
        },
    )

    assert result.metrics_scope == "global"
    assert result.authoritative_status["set_reconciliation_authoritative"] is True
    assert result.pagination["last_page_number"] == 14


@pytest.mark.parametrize(
    ("raw", "expected", "cell_type"),
    [
        ("260 000 0001", "2600000001", "text"),
        (2600000001, "2600000001", "numeric"),
        (2600000001.0, "2600000001", "numeric"),
    ],
)
def test_protocol_normalization(raw: object, expected: str, cell_type: str) -> None:
    assert normalize_protocol(raw).normalized == expected
    assert normalize_protocol(raw).cell_type == cell_type


def test_invalid_protocol_normalization() -> None:
    assert normalize_protocol("2.601E+09").normalized is None
    assert normalize_protocol(2600000001.5).normalized is None
