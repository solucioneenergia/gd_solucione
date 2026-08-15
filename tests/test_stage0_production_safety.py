from __future__ import annotations

import errno
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, mock_open
import hashlib
import json

import pytest
from openpyxl import Workbook

from automacao_gd.application import full_pipeline
from automacao_gd.application import processing_service
from automacao_gd.application.contracts import OperationResult, OperationStatus
from automacao_gd.application.preflight import run_preflight
from automacao_gd.domain.errors import PreflightBlockedError
from automacao_gd.infrastructure.config import Settings
from automacao_gd.infrastructure.excel import availability
from automacao_gd.infrastructure.excel.availability import (
    WorkbookAvailabilityCode,
    validate_workbook_availability,
)
from automacao_gd.presentation import cli
from automacao_gd.presentation.cli import exit_code_for_status
from automacao_gd.presentation.controller import ApplicationController
from automacao_gd.presentation.operational_output import format_operation_summary
from tests._operational_auth import authorize_synthetic_pdfs


def _workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    ws.append(
        [
            "Cliente",
            "Protocolo",
            "Data de ingresso",
            "Conclusão",
            "Parecer",
            "Placa",
            "Inversor",
        ]
    )
    wb.save(path)
    wb.close()
    return path


def _settings(tmp_path: Path, **overrides) -> Settings:
    workbook = _workbook(tmp_path / "planilha.xlsx")
    clients = tmp_path / "clientes"
    clients.mkdir(exist_ok=True)
    values = {
        "APP_ENV": "production",
        "PORTAL_GD_URL": "https://example.com/",
        "PLANILHA_PATH": workbook,
        "CLIENTES_ROOT": clients,
        "DOWNLOADS_DIR": tmp_path / "downloads",
        "LOGS_DIR": tmp_path / "logs",
        "AUTH_STATE_PATH": tmp_path / "auth" / "state.json",
        "BROWSER_PROFILE_DIR": tmp_path / "browser",
        "CDP_ENDPOINT": "http://127.0.0.1:9222",
        "DRY_RUN": False,
        "APPLY_EXCEL": True,
        "APPLY_ARCHIVE": True,
    }
    values.update(overrides)
    return Settings(_env_file=None, **values)


@pytest.mark.parametrize("app_env", ["development", "test"])
def test_real_run_requires_production(tmp_path: Path, app_env: str) -> None:
    report = run_preflight(_settings(tmp_path, APP_ENV=app_env), real_run=True)

    assert report.ready is False
    assert any(check.code == "ENVIRONMENT_NOT_PRODUCTION" for check in report.checks)
    assert "DRY_RUN=false exige APP_ENV=production" in " ".join(report.blocking_errors)


def test_production_real_run_and_test_simulation_are_allowed(tmp_path: Path) -> None:
    production = run_preflight(_settings(tmp_path), real_run=True)
    simulation = run_preflight(
        _settings(tmp_path, APP_ENV="test", DRY_RUN=True), real_run=False
    )

    assert production.ready is True
    assert simulation.ready is True


def test_real_excel_application_outside_production_is_blocked(tmp_path: Path) -> None:
    settings = _settings(
        tmp_path,
        APP_ENV="development",
        DRY_RUN=False,
        APPLY_EXCEL=True,
        APPLY_ARCHIVE=False,
    )

    assert run_preflight(settings, real_run=True).ready is False


def test_valid_workbook_preflight_is_byte_for_byte_non_destructive(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "available.xlsx")
    before = path.read_bytes()

    result = validate_workbook_availability(path, require_writable=True)

    assert result.ok is True
    assert path.read_bytes() == before
    assert list(tmp_path.glob(".*availability*")) == []


def test_workbook_missing_and_parent_missing_have_stable_codes(tmp_path: Path) -> None:
    missing = validate_workbook_availability(tmp_path / "missing.xlsx")
    missing_parent = validate_workbook_availability(
        tmp_path / "missing-parent" / "planilha.xlsx"
    )

    assert missing.code is WorkbookAvailabilityCode.WORKBOOK_NOT_FOUND
    assert missing_parent.code is WorkbookAvailabilityCode.WORKBOOK_TEMPORARILY_UNAVAILABLE


def test_disconnected_windows_drive_is_not_reported_as_missing_file(monkeypatch) -> None:
    monkeypatch.setattr(availability.os.path, "exists", lambda value: False)

    result = validate_workbook_availability(Path(r"Z:\Clientes\planilha.xlsx"))

    assert result.code is WorkbookAvailabilityCode.NETWORK_DRIVE_UNAVAILABLE
    assert "Z:\\" in result.user_message


def test_unc_and_local_paths_do_not_require_drive_letter_probe(
    tmp_path: Path, monkeypatch
) -> None:
    local = _workbook(tmp_path / "local.xlsx")
    probes: list[str] = []
    original_exists = availability.os.path.exists

    def record_exists(value) -> bool:
        probes.append(str(value))
        return original_exists(value)

    monkeypatch.setattr(availability.os.path, "exists", record_exists)
    assert validate_workbook_availability(local).ok is True
    assert "Z:\\" not in probes
    assert availability.windows_drive_root(r"\\SERVIDOR\SINTETICO\folder\book.xlsx") is None


def test_valid_unc_path_is_checked_without_mapped_drive(monkeypatch) -> None:
    path = Path(r"\\SERVIDOR\SINTETICO\folder\book.xlsx")
    workbook = SimpleNamespace(close=Mock())
    monkeypatch.setattr(Path, "exists", lambda _self: True)
    monkeypatch.setattr(Path, "is_dir", lambda self: self != path)
    monkeypatch.setattr(Path, "is_file", lambda self: self == path)
    monkeypatch.setattr(Path, "open", lambda *_args, **_kwargs: mock_open()())
    monkeypatch.setattr(availability.os, "access", lambda *_args: True)
    monkeypatch.setattr(availability, "load_workbook", lambda *_args, **_kwargs: workbook)

    result = validate_workbook_availability(path)

    assert result.ok is True
    workbook.close.assert_called_once()


def test_locked_workbook_is_distinct_from_permission_denied(
    tmp_path: Path, monkeypatch
) -> None:
    path = _workbook(tmp_path / "locked.xlsx")
    real_open = Path.open

    def locked_open(self: Path, mode="r", *args, **kwargs):
        if self == path and mode == "r+b":
            exc = PermissionError(errno.EACCES, "sharing violation", str(self))
            exc.winerror = 32
            raise exc
        return real_open(self, mode, *args, **kwargs)

    monkeypatch.setattr(Path, "open", locked_open)
    locked = validate_workbook_availability(path, require_writable=True)
    monkeypatch.setattr(
        availability.os,
        "access",
        lambda candidate, _mode: str(candidate) not in {str(path), str(path.parent)},
    )
    denied = validate_workbook_availability(path, require_writable=True)

    assert locked.code is WorkbookAvailabilityCode.WORKBOOK_LOCKED
    assert denied.code is WorkbookAvailabilityCode.WORKBOOK_PERMISSION_DENIED


def test_corrupt_workbook_and_open_failure_are_classified(
    tmp_path: Path, monkeypatch
) -> None:
    corrupt_path = tmp_path / "corrupt.xlsx"
    corrupt_path.write_bytes(b"not an xlsx")
    corrupt = validate_workbook_availability(corrupt_path)
    valid_path = _workbook(tmp_path / "valid.xlsx")
    monkeypatch.setattr(
        availability,
        "load_workbook",
        Mock(side_effect=RuntimeError("openpyxl failed")),
    )
    open_failed = validate_workbook_availability(valid_path)

    assert corrupt.code is WorkbookAvailabilityCode.WORKBOOK_INVALID
    assert open_failed.code is WorkbookAvailabilityCode.WORKBOOK_OPEN_FAILED


def test_temp_creation_and_atomic_replace_failures_are_classified(
    tmp_path: Path, monkeypatch
) -> None:
    path = _workbook(tmp_path / "available.xlsx")
    monkeypatch.setattr(
        availability.tempfile,
        "mkstemp",
        Mock(side_effect=PermissionError(errno.EACCES, "denied")),
    )
    create_failed = validate_workbook_availability(path, require_writable=True)
    monkeypatch.undo()
    monkeypatch.setattr(
        availability.os,
        "replace",
        Mock(side_effect=PermissionError(errno.EACCES, "denied")),
    )
    replace_failed = validate_workbook_availability(path, require_writable=True)

    assert create_failed.code is WorkbookAvailabilityCode.WORKBOOK_TEMP_CREATE_FAILED
    assert replace_failed.code is WorkbookAvailabilityCode.WORKBOOK_ATOMIC_REPLACE_FAILED
    assert path.exists()
    assert not list(tmp_path.glob(".*availability*"))


def test_temp_write_failure_is_classified_and_cleaned(
    tmp_path: Path, monkeypatch
) -> None:
    path = _workbook(tmp_path / "available.xlsx")
    monkeypatch.setattr(
        availability.os,
        "fsync",
        Mock(side_effect=OSError(errno.EIO, "sync failed")),
    )

    result = validate_workbook_availability(path, require_writable=True)

    assert result.code is WorkbookAvailabilityCode.WORKBOOK_TEMP_CREATE_FAILED
    assert not list(tmp_path.glob(".*availability*"))


def test_temporary_unavailability_has_specific_code(tmp_path: Path, monkeypatch) -> None:
    path = _workbook(tmp_path / "available.xlsx")
    real_open = Path.open

    def unavailable_open(self: Path, mode="r", *args, **kwargs):
        if self == path and mode == "rb":
            raise OSError(errno.EBUSY, "temporarily busy")
        return real_open(self, mode, *args, **kwargs)

    monkeypatch.setattr(Path, "open", unavailable_open)

    result = validate_workbook_availability(path)

    assert result.code is WorkbookAvailabilityCode.WORKBOOK_TEMPORARILY_UNAVAILABLE


def test_preflight_block_occurs_before_state_and_portal(monkeypatch, tmp_path: Path) -> None:
    settings = _settings(tmp_path, APP_ENV="development")
    download = Mock()
    state_store = Mock()
    playwright = Mock()
    processing = Mock()
    monkeypatch.setattr(full_pipeline, "_run_download_step", download)
    monkeypatch.setattr(full_pipeline, "PipelineStateStore", state_store)
    monkeypatch.setattr(full_pipeline, "sync_playwright", playwright)
    monkeypatch.setattr(full_pipeline, "process_downloaded_pdfs", processing)

    with pytest.raises(PreflightBlockedError):
        full_pipeline.run_full_cdp_pipeline(settings)

    download.assert_not_called()
    state_store.assert_not_called()
    playwright.assert_not_called()
    processing.assert_not_called()


def test_option5_real_run_reuses_frozen_dry_run_plan_without_cdp_selection(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=False,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=2,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    pdfs: list[Path] = []
    artifacts: list[dict[str, str]] = []
    results: list[dict[str, object]] = []
    selected_protocols: list[dict[str, str]] = []
    for offset, protocol in enumerate(("2600000000", "2600000001"), start=1):
        pdf = settings.downloads_dir_path / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
        pdf.parent.mkdir(parents=True, exist_ok=True)
        pdf.write_bytes(b"%PDF-1.4 synthetic " + protocol.encode("ascii"))
        sha256 = hashlib.sha256(pdf.read_bytes()).hexdigest()
        pdfs.append(pdf)
        artifacts.append(
            {
                "protocol": protocol,
                "path": str(pdf),
                "sha256": sha256,
            }
        )
        selected_protocols.append(
            {
                "protocol": protocol,
                "client_name": f"CLIENTE SINTETICO {offset}",
            }
        )
        results.append(
            {
                "protocol": protocol,
                "client_name": f"CLIENTE SINTETICO {offset}",
                "download_status": "existing_pdf_after_skip",
                "process_pdf_path": str(pdf),
                "selected_for_processing": True,
                "selected_by_global_limit": True,
                "global_limit_status": "selected",
            }
        )
    digest_source = "\n".join(f"{item['protocol']}:{item['sha256']}" for item in artifacts)
    dry_run_payload = {
        "dry_run": True,
        "status": "SUCESSO",
        "requested_batch_limit": 2,
        "authorized_batch_limit": 60,
        "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
        "total_errors": 0,
        "run_error": None,
        "download": {
            "run_error": None,
            "total_selected": 2,
            "total_for_processing": 2,
            "total_sent_to_processing": 2,
            "total_existing_reused": 2,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": selected_protocols,
            "results": results,
            "frozen_batch_created": True,
            "frozen_batch": {
                "requested_limit": 2,
                "authorized_limit": 60,
                "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
                "protocols": ["2600000000", "2600000001"],
                "unique_before_limit": 2,
                "dropped_by_limit": 0,
                "duplicate_protocols_in_frozen_batch": 0,
                "protocols_added_after_freeze": 0,
            },
            "frozen_pdf_scope": {
                "digest": hashlib.sha256(digest_source.encode("utf-8")).hexdigest(),
                "artifacts": artifacts,
            },
        },
    }
    settings.logs_dir_path.mkdir(parents=True, exist_ok=True)
    (settings.logs_dir_path / full_pipeline.PIPELINE_JSON_REPORT_NAME).write_text(
        json.dumps(dry_run_payload),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    monkeypatch.setattr(
        full_pipeline,
        "_run_download_step",
        Mock(side_effect=AssertionError("CDP selection must not run")),
    )
    received: dict[str, object] = {}

    def fake_process(**kwargs):
        received.update(kwargs)
        return {
            "total_pdfs": 2,
            "total_success": 2,
            "total_errors": 0,
            "total_excel_updated": 2,
            "results": [
                {
                    "protocol": "2600000000",
                    "success": True,
                    "action": "update_existing",
                    "excel_status": {"success": True, "action": "update_existing"},
                    "archive_status": {"success": True, "action": "simulation_only"},
                },
                {
                    "protocol": "2600000001",
                    "success": True,
                    "action": "insert_new_chronological",
                    "excel_status": {"success": True, "action": "insert_new_chronological"},
                    "archive_status": {"success": True, "action": "simulation_only"},
                },
            ],
        }

    monkeypatch.setattr(full_pipeline, "process_downloaded_pdfs", fake_process)

    result = full_pipeline.run_full_cdp_pipeline(
        settings,
        confirmation=full_pipeline.build_option5_strong_confirmation(2),
    )

    assert result["download"]["reused_from_dry_run_plan"] is True
    assert result["download"]["dry_run_plan_source"].endswith(
        full_pipeline.PIPELINE_JSON_REPORT_NAME
    )
    assert received["pdf_paths"] == pdfs
    assert received["allowed_protocols"] == {"2600000000", "2600000001"}


def test_option5_real_run_blocks_without_frozen_dry_run_plan_before_cdp(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=False,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=2,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    download = Mock(side_effect=AssertionError("CDP selection must not run"))
    processing = Mock(side_effect=AssertionError("real processing must not run"))
    monkeypatch.setattr(full_pipeline, "_run_download_step", download)
    monkeypatch.setattr(full_pipeline, "process_downloaded_pdfs", processing)

    with pytest.raises(PreflightBlockedError) as exc:
        full_pipeline.run_full_cdp_pipeline(
            settings,
            confirmation=full_pipeline.build_option5_strong_confirmation(2),
        )

    assert exc.value.code == "FROZEN_DRY_RUN_PLAN_REQUIRED"
    download.assert_not_called()
    processing.assert_not_called()


def test_option5_dry_run_records_frozen_pdf_scope_for_real_reuse(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=True,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=1,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        OP5_RECONCILIATION_MODE="batch_fast",
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    protocol = "2600000000"
    pdf = settings.downloads_dir_path / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.write_bytes(b"%PDF-1.4 synthetic frozen scope")
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    monkeypatch.setattr(
        full_pipeline,
        "_run_download_step",
        lambda *_args: {
            "run_error": None,
            "total_selected": 1,
            "total_for_processing": 1,
            "total_sent_to_processing": 1,
            "total_existing_reused": 1,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": [{"protocol": protocol, "client_name": "CLIENTE SINTETICO"}],
            "results": [
                {
                    "protocol": protocol,
                    "client_name": "CLIENTE SINTETICO",
                    "download_status": "existing_pdf_after_skip",
                    "process_pdf_path": str(pdf),
                    "selected_for_processing": True,
                }
            ],
        },
    )
    monkeypatch.setattr(
        full_pipeline,
        "process_downloaded_pdfs",
        lambda **_kwargs: {
            "total_pdfs": 1,
            "total_success": 1,
            "total_errors": 0,
            "total_excel_updated": 0,
            "results": [{"protocol": protocol, "success": True, "action": "simulation_only"}],
        },
    )

    result = full_pipeline.run_full_cdp_pipeline(
        settings,
        confirmation=full_pipeline.build_option5_strong_confirmation(1),
    )
    report = json.loads(
        (settings.logs_dir_path / full_pipeline.PIPELINE_JSON_REPORT_NAME).read_text(
            encoding="utf-8"
        )
    )

    assert result["download"]["frozen_pdf_scope"]["artifacts"][0]["protocol"] == protocol
    assert report["download"]["frozen_pdf_scope"]["artifacts"][0]["sha256"] == hashlib.sha256(
        pdf.read_bytes()
    ).hexdigest()


def test_option5_dry_run_persists_explicit_frozen_plan(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=True,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=1,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        OP5_RECONCILIATION_MODE="batch_fast",
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    protocol = "2600000000"
    pdf = settings.downloads_dir_path / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.write_bytes(b"%PDF-1.4 synthetic op5 plan")
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    monkeypatch.setattr(
        full_pipeline,
        "_run_download_step",
        lambda *_args: {
            "run_error": None,
            "total_selected": 1,
            "total_for_processing": 1,
            "total_sent_to_processing": 1,
            "total_existing_reused": 1,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": [{"protocol": protocol, "client_name": "CLIENTE SINTETICO"}],
            "results": [
                {
                    "protocol": protocol,
                    "client_name": "CLIENTE SINTETICO",
                    "download_status": "existing_pdf_after_skip",
                    "process_pdf_path": str(pdf),
                    "selected_for_processing": True,
                }
            ],
        },
    )
    monkeypatch.setattr(
        full_pipeline,
        "process_downloaded_pdfs",
        lambda **_kwargs: {
            "total_pdfs": 1,
            "total_success": 1,
            "total_errors": 0,
            "total_excel_updated": 0,
            "results": [
                {
                    "protocol": protocol,
                    "success": True,
                    "action": "insert_new_chronological",
                    "excel_status": {
                        "success": True,
                        "action": "insert_new_chronological",
                    },
                }
            ],
        },
    )

    result = full_pipeline.run_full_cdp_pipeline(
        settings,
        confirmation=full_pipeline.build_option5_strong_confirmation(1),
    )
    plan_path = settings.logs_dir_path / "op5_plan_latest.json"
    plan = json.loads(plan_path.read_text(encoding="utf-8"))

    assert result["op5_plan_path"].endswith("op5_plan_latest.json")
    assert plan["schema_version"] == 1
    assert plan["dry_run"] is True
    assert plan["requested_batch_limit"] == 1
    assert plan["workbook_sha256"] == hashlib.sha256(
        settings.planilha_path.read_bytes()
    ).hexdigest()
    assert plan["frozen_pdf_scope"]["artifacts"][0]["protocol"] == protocol
    assert plan["planned_excel_actions"] == [
        {"protocol": protocol, "action": "insert_new_chronological"}
    ]


def test_option5_dry_run_reports_accepted_plan_limited_to_requested_batch(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=True,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=1,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        OP5_RECONCILIATION_MODE="batch_fast",
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    protocols = ["2600000001", "2600000002"]
    pdfs: dict[str, Path] = {}
    download_call: dict[str, object] = {}
    processing_call: dict[str, object] = {}
    for protocol in protocols:
        pdf = settings.downloads_dir_path / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
        pdf.parent.mkdir(parents=True, exist_ok=True)
        pdf.write_bytes(f"%PDF-1.4 synthetic {protocol}".encode("utf-8"))
        pdfs[protocol] = pdf
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    def fake_run_download_step(effective_settings, *_args) -> dict:
        download_call["max_completed_to_process"] = (
            effective_settings.MAX_COMPLETED_TO_PROCESS
        )
        return {
            "run_error": None,
            "total_selected": 2,
            "total_for_processing": 2,
            "total_sent_to_processing": 2,
            "total_existing_reused": 2,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": [
                {"protocol": protocol, "client_name": f"CLIENTE {protocol}"}
                for protocol in protocols
            ],
            "results": [
                {
                    "protocol": protocol,
                    "client_name": f"CLIENTE {protocol}",
                    "download_status": "existing_pdf_after_skip",
                    "process_pdf_path": str(pdfs[protocol]),
                    "selected_for_processing": True,
                }
                for protocol in protocols
            ],
        }

    monkeypatch.setattr(
        full_pipeline,
        "_run_download_step",
        fake_run_download_step,
    )
    def fake_process_downloaded_pdfs(**kwargs) -> dict:
        processing_call.update(kwargs)
        return {
            "total_pdfs": 2,
            "total_success": 2,
            "total_errors": 0,
            "total_updates_planned": 2,
            "total_updates_applied": 0,
            "results": [
                {
                    "protocol": protocol,
                    "success": True,
                    "excel_status": {
                        "success": True,
                        "action": "update_existing",
                    },
                }
                for protocol in protocols
            ],
        }

    monkeypatch.setattr(
        full_pipeline,
        "process_downloaded_pdfs",
        fake_process_downloaded_pdfs,
    )

    result = full_pipeline.run_full_cdp_pipeline(
        settings,
        confirmation=full_pipeline.build_option5_strong_confirmation(1),
    )
    report = json.loads(
        (settings.logs_dir_path / full_pipeline.PIPELINE_JSON_REPORT_NAME).read_text(
            encoding="utf-8"
        )
    )

    assert result["total_updates_planned"] == 1
    assert result["total_selected"] == 1
    assert result["download"]["total_selected"] == 1
    assert download_call["max_completed_to_process"] == 1
    assert processing_call["allowed_protocols"] == {"2600000001"}
    assert processing_call["pdf_paths"] == [pdfs["2600000001"]]
    assert result["planned_excel_actions"] == [
        {"protocol": "2600000001", "action": "update_existing"}
    ]
    assert report["total_updates_planned"] == 1
    assert report["download"]["frozen_batch"]["protocols"] == ["2600000001"]


def test_failed_op5_plan_does_not_leave_previous_latest_plan_usable(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=True,
        APPLY_ARCHIVE=True,
        MAX_COMPLETED_TO_PROCESS=50,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    settings.logs_dir_path.mkdir(parents=True, exist_ok=True)
    plan_path = settings.logs_dir_path / "op5_plan_latest.json"
    plan_path.write_text(
        json.dumps(
            {
                "schema_version": 1,
                "dry_run": True,
                "status": "SUCESSO",
                "requested_batch_limit": 50,
                "authorized_batch_limit": 60,
                "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
                "total_errors": 0,
                "planned_excel_actions": [
                    {"protocol": "2600001048", "action": "update_existing"}
                ],
            }
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    monkeypatch.setattr(
        full_pipeline,
        "_run_download_step",
        lambda *_args: {
            "run_error": "failed_return_to_listing",
            "aborted": True,
            "abort_reason": "failed_return_to_listing",
            "total_selected": 0,
            "total_for_processing": 0,
            "total_sent_to_processing": 0,
            "total_existing_reused": 0,
            "total_downloaded": 0,
            "total_cdp_errors": 1,
            "total_download_errors": 0,
            "total_errors": 1,
            "selected_protocols": [],
            "results": [],
        },
    )
    processing = Mock(side_effect=AssertionError("processing must not run after CDP failure"))
    monkeypatch.setattr(full_pipeline, "process_downloaded_pdfs", processing)

    result = full_pipeline.run_full_cdp_pipeline(
        settings,
        confirmation=full_pipeline.build_option5_strong_confirmation(50),
    )
    latest_plan = json.loads(plan_path.read_text(encoding="utf-8"))

    assert result["status"] != "SUCESSO"
    assert latest_plan["stale_after_failed_plan"] is True
    assert latest_plan["status"] == result["status"]
    assert latest_plan["run_error"] == "failed_return_to_listing"
    assert latest_plan["planned_excel_actions"] == []
    assert latest_plan["op5_workbook_coverage"]["planned_excel_actions"] == 0


def test_option5_real_run_blocks_when_workbook_changed_after_explicit_plan(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=False,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=1,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
    )
    protocol = "2600000000"
    pdf = settings.downloads_dir_path / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.write_bytes(b"%PDF-1.4 synthetic op5 apply plan")
    sha256 = hashlib.sha256(pdf.read_bytes()).hexdigest()
    digest = hashlib.sha256(f"{protocol}:{sha256}".encode("utf-8")).hexdigest()
    plan = {
        "schema_version": 1,
        "created_at": "2026-08-10T00:00:00",
        "source_report_path": "pipeline_cdp_completo.json",
        "dry_run": True,
        "status": "SUCESSO",
        "requested_batch_limit": 1,
        "authorized_batch_limit": 60,
        "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
        "strong_confirmation_contract": full_pipeline.build_option5_strong_confirmation(1),
        "workbook_sha256": "0" * 64,
        "workbook_path": str(settings.planilha_path),
        "apply_excel": True,
        "apply_archive": False,
        "total_errors": 0,
        "frozen_batch": {
            "requested_limit": 1,
            "authorized_limit": 60,
            "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
            "protocols": [protocol],
            "unique_before_limit": 1,
            "dropped_by_limit": 0,
            "duplicate_protocols_in_frozen_batch": 0,
            "protocols_added_after_freeze": 0,
        },
        "frozen_pdf_scope": {
            "digest": digest,
            "artifacts": [{"protocol": protocol, "path": str(pdf), "sha256": sha256}],
        },
        "planned_excel_actions": [{"protocol": protocol, "action": "update_existing"}],
        "download": {
            "run_error": None,
            "total_selected": 1,
            "total_for_processing": 1,
            "total_sent_to_processing": 1,
            "total_existing_reused": 1,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": [{"protocol": protocol, "client_name": "CLIENTE SINTETICO"}],
            "results": [
                {
                    "protocol": protocol,
                    "client_name": "CLIENTE SINTETICO",
                    "download_status": "existing_pdf_after_skip",
                    "process_pdf_path": str(pdf),
                    "selected_for_processing": True,
                    "selected_by_global_limit": True,
                    "global_limit_status": "selected",
                }
            ],
            "frozen_batch_created": True,
            "frozen_batch": {
                "requested_limit": 1,
                "authorized_limit": 60,
                "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
                "protocols": [protocol],
                "unique_before_limit": 1,
                "dropped_by_limit": 0,
                "duplicate_protocols_in_frozen_batch": 0,
                "protocols_added_after_freeze": 0,
            },
            "frozen_pdf_scope": {
                "digest": digest,
                "artifacts": [{"protocol": protocol, "path": str(pdf), "sha256": sha256}],
            },
        },
    }
    settings.logs_dir_path.mkdir(parents=True, exist_ok=True)
    (settings.logs_dir_path / "op5_plan_latest.json").write_text(
        json.dumps(plan),
        encoding="utf-8",
    )
    monkeypatch.setattr(
        full_pipeline,
        "global_execution_lock_path",
        lambda _settings: tmp_path / "locks" / "real_run_execution.lock",
    )
    monkeypatch.setattr(
        full_pipeline,
        "_run_download_step",
        Mock(side_effect=AssertionError("CDP selection must not run")),
    )
    monkeypatch.setattr(
        full_pipeline,
        "process_downloaded_pdfs",
        Mock(side_effect=AssertionError("real processing must not run")),
    )

    with pytest.raises(PreflightBlockedError) as exc:
        full_pipeline.run_full_cdp_pipeline(
            settings,
            confirmation=full_pipeline.build_option5_strong_confirmation(1),
        )

    assert exc.value.code == "OP5_PLAN_WORKBOOK_CHANGED"


def test_option5_real_run_uses_configured_explicit_plan_path(
    tmp_path: Path,
) -> None:
    settings = _settings(
        tmp_path,
        DRY_RUN=False,
        APPLY_ARCHIVE=False,
        MAX_COMPLETED_TO_PROCESS=1,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        LOGS_DIR=tmp_path / "logs",
        DOWNLOADS_DIR=tmp_path / "downloads",
        OP5_PLAN_PATH=tmp_path / "logs" / "op5_plan_custom.json",
    )
    protocol = "2600000000"
    pdf = settings.downloads_dir_path / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.write_bytes(b"%PDF-1.4 synthetic explicit op5 plan")
    pdf_sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
    digest = hashlib.sha256(f"{protocol}:{pdf_sha}".encode("utf-8")).hexdigest()
    valid_plan = {
        "schema_version": 1,
        "created_at": "2026-08-10T00:00:00",
        "source_report_path": "pipeline_cdp_completo.json",
        "dry_run": True,
        "status": "SUCESSO",
        "requested_batch_limit": 1,
        "authorized_batch_limit": 60,
        "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
        "strong_confirmation_contract": full_pipeline.build_option5_strong_confirmation(1),
        "workbook_sha256": hashlib.sha256(settings.planilha_path.read_bytes()).hexdigest(),
        "workbook_path": str(settings.planilha_path),
        "apply_excel": True,
        "apply_archive": False,
        "total_errors": 0,
        "frozen_batch": {
            "requested_limit": 1,
            "authorized_limit": 60,
            "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
            "protocols": [protocol],
            "unique_before_limit": 1,
            "dropped_by_limit": 0,
            "duplicate_protocols_in_frozen_batch": 0,
            "protocols_added_after_freeze": 0,
        },
        "frozen_pdf_scope": {
            "digest": digest,
            "artifacts": [{"protocol": protocol, "path": str(pdf), "sha256": pdf_sha}],
        },
        "planned_excel_actions": [{"protocol": protocol, "action": "update_existing"}],
        "download": {
            "run_error": None,
            "total_selected": 1,
            "total_for_processing": 1,
            "total_sent_to_processing": 1,
            "total_existing_reused": 1,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": [{"protocol": protocol, "client_name": "CLIENTE SINTETICO"}],
            "results": [
                {
                    "protocol": protocol,
                    "client_name": "CLIENTE SINTETICO",
                    "download_status": "existing_pdf_after_skip",
                    "process_pdf_path": str(pdf),
                    "selected_for_processing": True,
                    "selected_by_global_limit": True,
                    "global_limit_status": "selected",
                }
            ],
            "frozen_batch_created": True,
            "frozen_batch": {
                "requested_limit": 1,
                "authorized_limit": 60,
                "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
                "protocols": [protocol],
                "unique_before_limit": 1,
                "dropped_by_limit": 0,
                "duplicate_protocols_in_frozen_batch": 0,
                "protocols_added_after_freeze": 0,
            },
            "frozen_pdf_scope": {
                "digest": digest,
                "artifacts": [{"protocol": protocol, "path": str(pdf), "sha256": pdf_sha}],
            },
        },
    }
    settings.logs_dir_path.mkdir(parents=True, exist_ok=True)
    (settings.logs_dir_path / "op5_plan_latest.json").write_text(
        json.dumps({**valid_plan, "workbook_sha256": "0" * 64}),
        encoding="utf-8",
    )
    Path(settings.OP5_PLAN_PATH).write_text(json.dumps(valid_plan), encoding="utf-8")
    authorization = full_pipeline.validate_requested_batch_limit(
        1,
        full_pipeline.batch_authorization_policy_from_settings(settings),
    )

    plan = full_pipeline.load_frozen_dry_run_plan(settings, authorization)

    assert plan.source_path == Path(settings.OP5_PLAN_PATH)
    assert plan.summary["dry_run_plan_source_kind"] == "explicit_op5_plan"
    assert plan.summary["cdp_selection_skipped"] is True


def test_controller_catches_operational_block_without_traceback(
    tmp_path: Path, monkeypatch
) -> None:
    controller = ApplicationController(_settings(tmp_path, APP_ENV="test"))
    exception_log = Mock()
    monkeypatch.setattr(
        "automacao_gd.application.use_cases.run_pipeline.run_full_cdp_pipeline",
        Mock(side_effect=PreflightBlockedError(
            code="ENVIRONMENT_NOT_PRODUCTION",
            user_message="Execução real bloqueada.",
            stage="pré-voo",
        )),
    )
    monkeypatch.setattr("automacao_gd.presentation.controller.logger.exception", exception_log)

    result = controller.run_pipeline()

    assert result.status is OperationStatus.BLOQUEADO
    assert result.payload["total_pages_read"] == 0
    assert result.payload["total_downloaded"] == 0
    exception_log.assert_not_called()


def test_second_validation_blocks_all_real_writes_and_keeps_pdfs_pending(
    tmp_path: Path, monkeypatch
) -> None:
    pdf = tmp_path / "Orcamento_de_Conexao_2600000000.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    simulated = processing_service._empty_result(pdf)
    simulated.update(
        {
            "success": True,
            "protocol": "2600000000",
            "client_name": "Cliente",
            "client_folder_match_type": "protocol",
        }
    )
    simulated["excel_status"].update({"success": True, "can_write": True})
    simulated["archive_status"].update({"success": True, "simulated": True})
    validation = Mock(
        side_effect=[
            {"success": True, "error": None, "code": "WORKBOOK_AVAILABLE"},
            {
                "success": False,
                "error": "A planilha está aberta ou bloqueada pelo Excel.",
                "code": "WORKBOOK_LOCKED",
                "technical_cause": "winerror=32",
            },
        ]
    )
    process_one = Mock(return_value=simulated)
    backup = Mock()
    exception_log = Mock()
    monkeypatch.setattr(processing_service, "ensure_directories", lambda: None)
    monkeypatch.setattr(processing_service, "clear_folder_cache", lambda: None)
    monkeypatch.setattr(
        processing_service,
        "get_settings",
        lambda: SimpleNamespace(logs_dir_path=tmp_path / "logs", BACKUP_EXCEL=True),
    )
    monkeypatch.setattr(processing_service, "_real_run_preflight_result", validation)
    monkeypatch.setattr(processing_service, "_process_single_pdf", process_one)
    monkeypatch.setattr(processing_service, "create_workbook_backup", backup)
    monkeypatch.setattr(processing_service, "atomic_write_json", lambda *_a, **_k: None)
    monkeypatch.setattr(processing_service, "atomic_write_text", lambda *_a, **_k: None)
    monkeypatch.setattr(processing_service.logger, "exception", exception_log)

    payload = processing_service.process_downloaded_pdfs(
        downloads_root=tmp_path,
        workbook_path=tmp_path / "planilha.xlsx",
        clientes_root=tmp_path / "clientes",
        dry_run=False,
        pdf_paths=[pdf],
        apply_excel=True,
        apply_archive=True,
        allowed_protocols={"2600000000"},
        authorization=authorize_synthetic_pdfs(tmp_path, [pdf]),
    )

    assert validation.call_count == 2
    assert process_one.call_count == 1
    assert process_one.call_args.args[3] is True
    backup.assert_not_called()
    exception_log.assert_not_called()
    assert payload["status"] == OperationStatus.BLOQUEADO.value
    assert payload["real_run_block_code"] == "WORKBOOK_LOCKED"
    assert payload["pending_pdf_paths"] == [str(pdf)]
    assert payload["total_excel_updated"] == 0
    assert payload["total_archived"] == 0


def test_unexpected_controller_exception_is_failed(tmp_path: Path, monkeypatch) -> None:
    controller = ApplicationController(_settings(tmp_path, DRY_RUN=True))
    monkeypatch.setattr(
        "automacao_gd.application.use_cases.process_downloads.ProcessDownloadedPdfsUseCase.execute",
        Mock(side_effect=RuntimeError("fatal")),
    )

    result = controller.process_downloads(dry_run=True)

    assert result.status is OperationStatus.FALHOU


@pytest.mark.parametrize(
    ("payload", "expected_status", "expected_message"),
    [
        (
            {
                "total_downloaded": 2,
                "total_existing_reused": 0,
                "total_processed_success": 1,
                "total_excel_updated": 1,
                "total_errors": 1,
                "total_selected": 2,
                "total_eligible_after_skip": 2,
                "processing": {},
                "dry_run": False,
            },
            OperationStatus.PARCIAL,
            "parciais",
        ),
        (
            {
                "total_downloaded": 1,
                "total_existing_reused": 0,
                "total_processed_success": 1,
                "total_excel_updated": 1,
                "total_errors": 0,
                "total_selected": 1,
                "total_eligible_after_skip": 1,
                "processing": {},
                "dry_run": False,
            },
            OperationStatus.SUCESSO,
            "sucesso",
        ),
        (
            {
                "total_downloaded": 0,
                "total_existing_reused": 0,
                "total_processed_success": 0,
                "total_excel_updated": 0,
                "total_errors": 0,
                "total_selected": 0,
                "total_eligible_after_skip": 0,
                "processing": {},
                "dry_run": False,
            },
            OperationStatus.SUCESSO,
            "Nenhuma atualização necessária.",
        ),
        (
            {
                "total_downloaded": 1,
                "total_existing_reused": 0,
                "total_processed_success": 0,
                "total_excel_updated": 0,
                "total_errors": 1,
                "total_selected": 1,
                "total_eligible_after_skip": 1,
                "processing": {},
                "dry_run": False,
            },
            OperationStatus.FALHOU,
            "Nenhum PDF",
        ),
    ],
)
def test_pipeline_status_semantics(payload, expected_status, expected_message) -> None:
    status, message = full_pipeline._classify_pipeline_result(payload)

    assert status is expected_status
    assert expected_message in message


@pytest.mark.parametrize(
    ("status", "exit_code"),
    [
        (OperationStatus.SUCESSO, 0),
        (OperationStatus.FALHOU, 1),
        (OperationStatus.BLOQUEADO, 2),
        (OperationStatus.PARCIAL, 3),
    ],
)
def test_status_exit_codes(status: OperationStatus, exit_code: int) -> None:
    assert exit_code_for_status(status) == exit_code


class _MenuOption5Settings(SimpleNamespace):
    def model_copy(self, *, update: dict):
        data = dict(self.__dict__)
        data.update(update)
        return _MenuOption5Settings(**data)


def _menu_option5_settings(tmp_path: Path) -> _MenuOption5Settings:
    workbook = tmp_path / "planilha.xlsx"
    workbook.write_bytes(b"workbook sintetico")
    return _MenuOption5Settings(
        DRY_RUN=False,
        APPLY_EXCEL=True,
        APPLY_ARCHIVE=True,
        MAX_COMPLETED_TO_PROCESS=5,
        OPTION5_AUTHORIZED_MAX_PROTOCOLS=60,
        OP5_RECONCILIATION_MODE="inline_global",
        planilha_path=workbook,
        logs_dir_path=tmp_path / "logs",
    )


def _write_valid_menu_op5_plan(
    settings: _MenuOption5Settings,
    *,
    protocol: str = "2600001048",
    action: str = "update_existing",
    requested_limit: int = 1,
) -> Path:
    pdf = settings.logs_dir_path.parent / "downloads" / protocol / f"Orcamento_de_Conexao_{protocol}.pdf"
    pdf.parent.mkdir(parents=True, exist_ok=True)
    pdf.write_bytes(b"%PDF-1.4 synthetic menu op5 plan")
    pdf_sha = hashlib.sha256(pdf.read_bytes()).hexdigest()
    digest = hashlib.sha256(f"{protocol}:{pdf_sha}".encode("utf-8")).hexdigest()
    settings.logs_dir_path.mkdir(parents=True, exist_ok=True)
    plan_path = settings.logs_dir_path / "op5_plan_latest.json"
    plan = {
        "schema_version": 1,
        "created_at": "2026-08-13T00:00:00",
        "source_report_path": "pipeline_cdp_completo.json",
        "dry_run": True,
        "status": OperationStatus.SUCESSO.value,
        "requested_batch_limit": requested_limit,
        "authorized_batch_limit": 60,
        "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
        "strong_confirmation_contract": full_pipeline.build_option5_strong_confirmation(requested_limit),
        "workbook_sha256": hashlib.sha256(Path(settings.planilha_path).read_bytes()).hexdigest(),
        "workbook_path": str(settings.planilha_path),
        "apply_excel": True,
        "apply_archive": True,
        "total_selected": requested_limit,
        "total_updates_planned": requested_limit,
        "total_updates_applied": 0,
        "total_errors": 0,
        "frozen_batch": {
            "requested_limit": requested_limit,
            "authorized_limit": 60,
            "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
            "protocols": [protocol],
            "unique_before_limit": requested_limit,
            "dropped_by_limit": 0,
            "duplicate_protocols_in_frozen_batch": 0,
            "protocols_added_after_freeze": 0,
        },
        "frozen_pdf_scope": {
            "digest": digest,
            "artifacts": [{"protocol": protocol, "path": str(pdf), "sha256": pdf_sha}],
        },
        "planned_excel_actions": [{"protocol": protocol, "action": action}],
        "download": {
            "run_error": None,
            "total_selected": requested_limit,
            "total_for_processing": requested_limit,
            "total_sent_to_processing": requested_limit,
            "total_existing_reused": requested_limit,
            "total_downloaded": 0,
            "total_cdp_errors": 0,
            "total_errors": 0,
            "selected_protocols": [{"protocol": protocol, "client_name": "CLIENTE SINTETICO"}],
            "results": [
                {
                    "protocol": protocol,
                    "client_name": "CLIENTE SINTETICO",
                    "download_status": "existing_pdf_after_skip",
                    "process_pdf_path": str(pdf),
                    "selected_for_processing": True,
                    "selected_by_global_limit": True,
                    "global_limit_status": "selected",
                }
            ],
            "frozen_batch_created": True,
            "frozen_batch": {
                "requested_limit": requested_limit,
                "authorized_limit": 60,
                "authorization_scope": full_pipeline.CONTROLLED_PRODUCTION_UP_TO_60_AUTHORIZATION_SCOPE,
                "protocols": [protocol],
                "unique_before_limit": requested_limit,
                "dropped_by_limit": 0,
                "duplicate_protocols_in_frozen_batch": 0,
                "protocols_added_after_freeze": 0,
            },
            "frozen_pdf_scope": {
                "digest": digest,
                "artifacts": [{"protocol": protocol, "path": str(pdf), "sha256": pdf_sha}],
            },
        },
    }
    plan_path.write_text(json.dumps(plan), encoding="utf-8")
    return plan_path


def test_interactive_option5_orchestrates_plan_backup_and_frozen_apply(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = _menu_option5_settings(tmp_path)
    plan_path = tmp_path / "logs" / "op5_plan_latest.json"
    backup_path = tmp_path / "planilha_backup.xlsx"
    calls: list[tuple[object, str | None]] = []
    summaries: list[str] = []

    class FakeController:
        def __init__(self, active_settings=None):
            self.settings = active_settings or settings

        def run_pipeline(self, *, confirmation: str | None = None):
            calls.append((self.settings, confirmation))
            if self.settings.DRY_RUN:
                return OperationResult(
                    True,
                    "plano gerado",
                    {
                        "status": OperationStatus.SUCESSO.value,
                        "dry_run": True,
                        "apply_archive": True,
                        "requested_batch_limit": 2,
                        "authorized_batch_limit": 60,
                        "total_updates_planned": 2,
                        "total_updates_applied": 0,
                        "total_errors": 0,
                        "op5_plan_path": str(plan_path),
                    },
                    status=OperationStatus.SUCESSO,
                )
            return OperationResult(
                True,
                "apply executado",
                {"status": OperationStatus.SUCESSO.value, "total_updates_applied": 2},
                status=OperationStatus.SUCESSO,
            )

    answers = iter(["2", full_pipeline.build_option5_strong_confirmation(2)])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    monkeypatch.setattr(cli, "ApplicationController", FakeController)
    monkeypatch.setattr(cli, "print_operation_summary", lambda label, _result: summaries.append(label))
    monkeypatch.setattr(
        cli,
        "validate_workbook_availability",
        lambda *_args, **_kwargs: SimpleNamespace(ok=True, user_message=""),
    )
    monkeypatch.setattr(cli, "create_workbook_backup", lambda _path: backup_path, raising=False)
    monkeypatch.setattr(
        cli,
        "_sha256_file",
        lambda _path: "a" * 64,
        raising=False,
    )

    result = cli._interactive_option5_plan_apply(FakeController(settings))

    assert result.status is OperationStatus.SUCESSO
    assert summaries == ["pipeline"]
    assert len(calls) == 2
    plan_settings, plan_confirmation = calls[0]
    apply_settings, apply_confirmation = calls[1]
    assert plan_settings.DRY_RUN is True
    assert plan_settings.MAX_COMPLETED_TO_PROCESS == 2
    assert plan_settings.OP5_RECONCILIATION_MODE == "batch_fast"
    assert plan_confirmation == full_pipeline.build_option5_strong_confirmation(2)
    assert apply_settings.DRY_RUN is False
    assert apply_settings.OP5_PLAN_PATH == plan_path
    assert apply_confirmation == full_pipeline.build_option5_strong_confirmation(2)


def test_interactive_option5_reuses_valid_pending_plan_without_new_cdp_dry_run(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = _menu_option5_settings(tmp_path)
    plan_path = _write_valid_menu_op5_plan(settings, requested_limit=1)
    backup_path = tmp_path / "planilha_backup.xlsx"
    calls: list[tuple[object, str | None]] = []

    class FakeController:
        def __init__(self, active_settings=None):
            self.settings = active_settings or settings

        def run_pipeline(self, *, confirmation: str | None = None):
            calls.append((self.settings, confirmation))
            if self.settings.DRY_RUN:
                raise AssertionError("dry-run/CDP nao deve ser reexecutado")
            return OperationResult(
                True,
                "apply executado",
                {"status": OperationStatus.SUCESSO.value, "total_updates_applied": 1},
                status=OperationStatus.SUCESSO,
            )

    answers = iter(["1", full_pipeline.build_option5_strong_confirmation(1)])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    monkeypatch.setattr(cli, "ApplicationController", FakeController)
    monkeypatch.setattr(cli, "print_operation_summary", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        cli,
        "validate_workbook_availability",
        lambda *_args, **_kwargs: SimpleNamespace(ok=True, user_message=""),
    )
    monkeypatch.setattr(cli, "create_workbook_backup", lambda _path: backup_path, raising=False)
    monkeypatch.setattr(cli, "_sha256_file", lambda _path: "a" * 64, raising=False)

    result = cli._interactive_option5_plan_apply(FakeController(settings))

    assert result.status is OperationStatus.SUCESSO
    assert len(calls) == 1
    apply_settings, apply_confirmation = calls[0]
    assert apply_settings.DRY_RUN is False
    assert apply_settings.OP5_PLAN_PATH == plan_path
    assert apply_confirmation == full_pipeline.build_option5_strong_confirmation(1)


def test_interactive_option5_rejects_generic_confirmation_before_backup_or_apply(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    settings = _menu_option5_settings(tmp_path)
    plan_path = tmp_path / "logs" / "op5_plan_latest.json"
    calls: list[tuple[object, str | None]] = []
    backup = Mock(side_effect=AssertionError("backup nao deve ser criado"))

    class FakeController:
        def __init__(self, active_settings=None):
            self.settings = active_settings or settings

        def run_pipeline(self, *, confirmation: str | None = None):
            calls.append((self.settings, confirmation))
            return OperationResult(
                True,
                "plano gerado",
                {
                    "status": OperationStatus.SUCESSO.value,
                    "dry_run": True,
                    "apply_archive": True,
                    "requested_batch_limit": 2,
                    "authorized_batch_limit": 60,
                    "total_updates_planned": 2,
                    "total_updates_applied": 0,
                    "total_errors": 0,
                    "op5_plan_path": str(plan_path),
                },
                status=OperationStatus.SUCESSO,
            )

    answers = iter(["2", "Confirmar"])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    monkeypatch.setattr(cli, "ApplicationController", FakeController)
    monkeypatch.setattr(cli, "print_operation_summary", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(cli, "create_workbook_backup", backup, raising=False)

    result = cli._interactive_option5_plan_apply(FakeController(settings))

    assert result.status is OperationStatus.BLOQUEADO
    assert result.payload["code"] == "STRONG_CONFIRMATION_MISMATCH"
    assert len(calls) == 1
    backup.assert_not_called()


def test_interactive_option5_prints_exact_apply_confirmation_on_dedicated_line(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    settings = _menu_option5_settings(tmp_path)
    plan_path = tmp_path / "logs" / "op5_plan_latest.json"
    confirmation = full_pipeline.build_option5_strong_confirmation(2)

    class FakeController:
        def __init__(self, active_settings=None):
            self.settings = active_settings or settings

        def run_pipeline(self, *, confirmation: str | None = None):
            return OperationResult(
                True,
                "plano gerado",
                {
                    "status": OperationStatus.SUCESSO.value,
                    "dry_run": True,
                    "apply_archive": True,
                    "requested_batch_limit": 2,
                    "authorized_batch_limit": 60,
                    "total_updates_planned": 2,
                    "total_updates_applied": 0,
                    "total_errors": 0,
                    "op5_plan_path": str(plan_path),
                },
                status=OperationStatus.SUCESSO,
            )

    answers = iter(["2", "Confirmar"])
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))
    monkeypatch.setattr(cli, "ApplicationController", FakeController)
    monkeypatch.setattr(cli, "print_operation_summary", lambda *_args, **_kwargs: None)

    result = cli._interactive_option5_plan_apply(FakeController(settings))

    output_lines = capsys.readouterr().out.splitlines()
    assert result.status is OperationStatus.BLOQUEADO
    assert "Para aplicar este plano, digite exatamente:" in output_lines
    assert confirmation in output_lines
    assert all(
        line == confirmation or confirmation not in line
        for line in output_lines
    )


def test_partial_listing_recovery_batch_is_not_classified_as_success() -> None:
    payload = {
        "dry_run": True,
        "requested_batch_limit": 55,
        "download": {
            "partial_batch_due_to_listing_recovery": True,
            "abort_reason": "stopped_after_failed_return_to_listing",
        },
        "total_downloaded": 0,
        "total_existing_reused": 14,
        "total_processed_success": 14,
        "total_excel_updated": 0,
        "total_errors": 0,
        "total_selected": 55,
        "total_eligible_after_skip": 97,
        "total_updates_planned": 14,
        "processing": {
            "total_updates_planned": 14,
            "blocked_real_run": False,
        },
    }

    status, message = full_pipeline._classify_pipeline_result(payload)

    assert status is OperationStatus.PARCIAL
    assert "listagem" in message.lower()


def test_interactive_option5_rejects_partial_listing_recovery_plan_even_if_counts_match() -> None:
    block = cli._validate_interactive_option5_plan(
        {
            "status": OperationStatus.SUCESSO.value,
            "dry_run": True,
            "apply_archive": True,
            "total_updates_planned": 10,
            "total_updates_applied": 0,
            "total_errors": 0,
            "download": {
                "partial_batch_due_to_listing_recovery": True,
                "abort_reason": "stopped_after_failed_return_to_listing",
            },
        },
        requested_limit=10,
    )

    assert block is not None
    assert block.status is OperationStatus.BLOQUEADO
    assert "partial_listing_recovery" in block.payload["blockers"]


def test_interactive_option5_accepts_safe_partial_plan_with_technical_pending_review() -> None:
    block = cli._validate_interactive_option5_plan(
        {
            "status": OperationStatus.SUCESSO.value,
            "dry_run": True,
            "apply_archive": True,
            "requested_batch_limit": 40,
            "total_updates_planned": 39,
            "total_updates_applied": 0,
            "total_errors": 0,
            "total_pending_review": 1,
            "planned_excel_actions": [
                {"protocol": f"250000{i:04d}", "action": "update_existing"}
                for i in range(39)
            ],
            "frozen_batch": {
                "requested_limit": 40,
                "protocols": [f"250000{i:04d}" for i in range(39)],
            },
            "processing": {
                "total_updates_planned": 39,
                "total_pending_review": 1,
                "results": [
                    {
                        "protocol": "2600001107",
                        "action": "pending_technical_review",
                        "technical_validation_status": "pending_review",
                    }
                ],
            },
            "download": {
                "partial_batch_due_to_listing_recovery": False,
                "abort_reason": None,
            },
        },
        requested_limit=40,
    )

    assert block is None


@pytest.mark.parametrize(
    ("status", "expected"),
    [
        (OperationStatus.SUCESSO, 0),
        (OperationStatus.FALHOU, 1),
        (OperationStatus.BLOQUEADO, 2),
        (OperationStatus.PARCIAL, 3),
    ],
)
def test_interactive_cli_returns_last_operation_exit_code(
    status: OperationStatus, expected: int, monkeypatch
) -> None:
    answers = iter(["5", "0"])

    class FakeController:
        settings = SimpleNamespace(DRY_RUN=True)

    monkeypatch.setattr(cli, "ensure_directories", lambda: None)
    monkeypatch.setattr(cli, "setup_logger", lambda **_kwargs: None)
    monkeypatch.setattr(cli, "ApplicationController", FakeController)
    monkeypatch.setattr(
        cli,
        "_interactive_option5_plan_apply",
        lambda _controller: OperationResult(
            status is OperationStatus.SUCESSO,
            "resultado",
            {},
            status=status,
        ),
    )
    monkeypatch.setattr("builtins.input", lambda _prompt="": next(answers))

    assert cli.main([]) == expected


def test_operational_summary_starts_with_explicit_status() -> None:
    result = OperationResult(
        False,
        "A planilha está aberta ou bloqueada pelo Excel.",
        {
            "stage": "pré-voo",
            "code": "WORKBOOK_LOCKED",
            "total_pages_read": 0,
            "total_downloaded": 0,
            "total_excel_updated": 0,
        },
        status=OperationStatus.BLOQUEADO,
    )

    output = format_operation_summary("pipeline", result)

    assert output.startswith("Status: BLOQUEADO")
    assert "Pipeline CDP concluído" not in output
    assert "Páginas lidas: 0" in output
