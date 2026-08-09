from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.worksheet.datavalidation import DataValidation

from automacao_gd.application.workbook_visual_audit_service import (
    audit_workbook_visual_structure,
    save_workbook_visual_audit_artifacts,
    serialize_excel_color,
)


HEADERS = [
    "Cliente",
    "Protocolo",
    "Data de ingresso",
    "Conclusão",
    "Parecer",
    "Placa",
    "Inversor",
]


def _base_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "2025"
    wb.create_sheet("2026")
    for sheet in wb.worksheets:
        sheet["A1"] = "SOLUCIONE NORDESTE ENERGIA ELÉTRICA"
        for index, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(2, index)
            cell.value = header
            cell.font = Font(bold=True, name="Arial", size=11)
            cell.fill = PatternFill("solid", fgColor="D9EAD3")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.column_dimensions[cell.column_letter].width = 18
        _row(sheet, 3, "Cliente 1", "2600001012", "2025-01-10", None, True, "1x MOD A", "1x INV A")
        sheet.auto_filter.ref = "A2:G3"
    wb.save(path)
    wb.close()
    return path


def _row(sheet, row: int, client, protocol, ingress, completion, parecer, placa, inversor) -> None:
    values = [client, protocol, ingress, completion, parecer, placa, inversor]
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column)
        cell.value = value
        cell.font = Font(name="Arial", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.cell(row, 2).number_format = "@"
    sheet.cell(row, 3).number_format = "dd/mm/yyyy"
    sheet.cell(row, 4).number_format = "dd/mm/yyyy"


def _issue_types(audit: dict) -> set[str]:
    return {
        finding["issue_type"]
        for finding in audit.get("findings", [])
        + audit.get("informational_findings", [])
    }


def test_style_id_different_but_same_appearance_has_no_visual_mismatch(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")

    audit = audit_workbook_visual_structure(path)

    assert "DATA_CELL_STYLE_MISMATCH" not in _issue_types(audit)


def test_header_visual_difference_is_reported(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"]["A2"].font = Font(bold=False, name="Arial", size=11)
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "HEADER_STYLE_MISMATCH" in _issue_types(audit)


def test_header_only_filter_is_reported(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"].auto_filter.ref = "A2:G2"
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "FILTER_HEADER_ONLY" in _issue_types(audit)


def test_correct_filter_is_reported(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")

    audit = audit_workbook_visual_structure(path)

    assert audit["sheet_comparisons"]["2026"]["filter"]["issue_type"] == "FILTER_RANGE_CORRECT"


def test_trailing_styled_empty_row_safe_to_remove(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"]["A4"].font = Font(name="Arial")
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)
    finding = next(f for f in audit["findings"] if f["issue_type"] == "TRAILING_MATERIALIZED_EMPTY_ROW")

    assert finding["safe_to_apply"] is True


def test_trailing_empty_row_with_validation_is_not_safe_to_remove(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2026"]
    ws["A4"].font = Font(name="Arial")
    dv = DataValidation(type="list", formula1='"A,B"')
    ws.add_data_validation(dv)
    dv.add(ws["A4"])
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)
    finding = next(f for f in audit["findings"] if f["issue_type"] == "TRAILING_MATERIALIZED_EMPTY_ROW")

    assert finding["safe_to_apply"] is False


def test_numeric_integer_protocol_safe_to_convert(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"]["B3"].value = 2600001083
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "PROTOCOL_NUMERIC_SAFE_TO_CONVERT" in _issue_types(audit)


def test_numeric_ambiguous_protocol_requires_review(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"]["B3"].value = 2600001083.5
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "PROTOCOL_NUMERIC_REQUIRES_REVIEW" in _issue_types(audit)


def test_multimanufacturer_with_wrap_and_height_is_canonical(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2026"]
    ws["F3"] = "TRINA | A\nOSDA | B\nQtd. total: 2 módulos"
    ws["G3"] = "HUAWEI | X\nOMNIK | Y\nQtd. total: 2 inversores"
    ws["F3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["G3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 60
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "MULTI_MANUFACTURER_STYLE_CANONICAL" in _issue_types(audit)
    assert all(
        action["action_type"] != "MULTI_MANUFACTURER_STYLE_CANONICAL"
        for action in audit["proposed_actions"]
    )


def test_multimanufacturer_without_wrap_is_reported(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2026"]
    ws["F3"] = "TRINA | A\nOSDA | B\nQtd. total: 2 módulos"
    ws["F3"].alignment = Alignment(wrap_text=False)
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "MULTI_MANUFACTURER_WRAP_MISSING" in _issue_types(audit)


def test_multimanufacturer_height_insufficient_is_reported(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2026"]
    ws["F3"] = "TRINA | A\nOSDA | B\nQtd. total: 2 módulos"
    ws["F3"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[3].height = 15
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "ROW_HEIGHT_ADJUSTMENT_SAFE" in _issue_types(audit)


def test_content_hash_is_preserved_in_read_only_audit(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")

    audit = audit_workbook_visual_structure(path)

    assert audit["safety"]["workbook_content_hash_before"] == audit["safety"]["workbook_content_hash_after"]
    assert audit["safety"]["workbook_sha_before"] == audit["safety"]["workbook_sha_after"]


def test_formulas_are_protected_by_content_hash(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"]["H3"] = "=1+1"
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert audit["content_invariants"]["protected_formula_cells_count"] == 1


def test_content_change_attempt_rejects_plan() -> None:
    audit = {
        "proposed_actions": [
            {"changes_content": True, "issue_type": "CONTENT_CHANGE_OUTSIDE_SCOPE"}
        ]
    }

    with pytest.raises(ValueError, match="CONTENT_CHANGE_OUTSIDE_SCOPE"):
        audit_workbook_visual_structure.validate_plan_safety(audit)  # type: ignore[attr-defined]


def test_last_real_row_differs_from_max_row(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"]["A7"].font = Font(name="Arial")
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    sheet = audit["sheet_comparisons"]["2026"]["structure"]
    assert sheet["last_real_data_row"] == 3
    assert sheet["physical_max_row"] >= 7


def test_adaptive_height_archetypes_are_distinct(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2025"]
    _row(ws, 4, "Cliente 2", "2600001013", "2025-01-11", None, True, "A\nB", "I")
    ws.row_dimensions[4].height = 36
    _row(ws, 5, "Cliente 3", "2600001014", "2025-01-12", None, True, "A\nB\nC", "I")
    ws.row_dimensions[5].height = 54
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    archetypes = audit["canonical_fingerprint"]["row_archetypes"]
    assert "SINGLE_LINE_STANDARD" in archetypes
    assert "TWO_LINE_EQUIPMENT" in archetypes
    assert "THREE_LINE_EQUIPMENT" in archetypes


def test_visual_difference_in_canonical_reference_becomes_safe_style_action(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2025"]
    _row(ws, 4, "Cliente 2", "2600001013", "2025-01-11", None, True, "1x MOD B", "1x INV B")
    ws["A4"].font = Font(name="Calibri", size=10)
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert audit["metrics"]["canonical_reference_inconsistencies"] == 0
    assert any(
        action["action_type"] == "STYLE_STANDARDIZATION"
        for action in audit["proposed_actions"]
    )


def test_read_only_does_not_save_backup_temp_or_replace(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")

    audit = audit_workbook_visual_structure(path)

    assert audit["safety"]["workbook_save_called"] == 0
    assert audit["safety"]["backup_created"] == 0
    assert audit["safety"]["temporary_file_created"] == 0
    assert audit["safety"]["os_replace_called"] == 0


def test_row_visual_hash_ignores_height_for_canonical_consistency(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2025"]
    _row(ws, 4, "Cliente 2", "2600001013", "2025-01-11", None, True, "1x MOD B", "1x INV B")
    ws.row_dimensions[3].height = 27.95
    ws.row_dimensions[4].height = 45
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    archetype = audit["canonical_fingerprint"]["row_archetypes"]["SINGLE_LINE_STANDARD"]
    assert len(archetype["row_visual_style_hashes"]) == 1
    assert audit["metrics"]["canonical_reference_inconsistencies"] == 0


def test_protected_multimanufacturer_content_is_not_action(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2026"]
    ws["F3"] = "FAB A | MODELO A\nFAB B | MODELO B\nQtd. total: 2 módulos"
    ws["G3"] = "INV A"
    ws["F3"].alignment = Alignment(wrap_text=True, vertical="center")
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert any(target["protection_reason"] == "MULTI_MANUFACTURER_CONTENT_PROTECTED" for target in audit["protected_targets"])
    assert all(action["action_type"] != "MULTI_MANUFACTURER_CONTENT_PROTECTED" for action in audit["proposed_actions"])
    assert audit["blocked_actions"] == []


def test_excel_color_serialization_supports_rgb_indexed_theme_and_none() -> None:
    rgb = serialize_excel_color(Color(rgb="FF595959"))
    indexed = serialize_excel_color(Color(indexed=64))
    theme = serialize_excel_color(Color(theme=1, tint=0.5))
    none = serialize_excel_color(None)

    assert rgb["type"] == "rgb"
    assert rgb["rgb"] == "FF595959"
    assert indexed["type"] == "indexed"
    assert indexed["indexed"] == 64
    assert indexed["rgb"] is None
    assert theme["type"] == "theme"
    assert theme["theme"] == 1
    assert theme["tint"] == 0.5
    assert none == {"type": None, "rgb": None, "indexed": None, "theme": None, "tint": None, "auto": None}


def test_single_visual_line_with_multiple_equipment_entries_has_adequate_height(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    ws = wb["2026"]
    ws["F3"] = "FAB A | MOD A / FAB B | MOD B"
    ws["F3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 27.95
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "ROW_HEIGHT_ALREADY_ADEQUATE" in _issue_types(audit)
    assert "ROW_HEIGHT_ADJUSTMENT_SAFE" not in _issue_types(audit)


def test_freeze_panes_a3_is_informational_and_mismatch_is_safe_action(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2025"].freeze_panes = "A3"
    wb["2026"].freeze_panes = "A78"
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert "FREEZE_PANES_CORRECT" in _issue_types(audit)
    assert "FREEZE_PANES_MISMATCH" in _issue_types(audit)
    assert any(action["action_type"] == "FREEZE_PANES_STANDARDIZATION" for action in audit["proposed_actions"])


def test_final_action_model_has_zero_blocked_actions_and_informational_outside_equation(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    wb = load_workbook(path)
    wb["2026"].auto_filter.ref = "A2:G2"
    wb["2026"].freeze_panes = "A9"
    wb.save(path)
    wb.close()

    audit = audit_workbook_visual_structure(path)

    assert audit["blocked_actions"] == []
    assert audit["metrics"]["blocked_proposed_actions"] == 0
    assert audit["metrics"]["safe_proposed_actions"] == audit["metrics"]["proposed_actions"]
    assert audit["metrics"]["informational_findings_total"] > 0
    assert audit["metrics"]["action_equation_valid"] is True


def test_artifacts_do_not_contain_openpyxl_error_strings(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")
    audit = audit_workbook_visual_structure(path)
    outputs = save_workbook_visual_audit_artifacts(audit, tmp_path / "logs", timestamp="20260101T000000Z")

    forbidden = ["Values must be of type", "<class", "TypeError", "ValueError", "Descriptor"]
    for output in outputs.values():
        text = output.read_text(encoding="utf-8")
        assert not any(item in text for item in forbidden)
    assert audit["metrics"]["artifact_error_strings"] == 0


def test_final_decision_requires_no_blockers(tmp_path: Path) -> None:
    path = _base_workbook(tmp_path / "book.xlsx")

    audit = audit_workbook_visual_structure(path)

    assert audit["metrics"]["canonical_reference_inconsistencies"] == 0
    assert audit["metrics"]["unresolved_canonical_decisions"] == 0
    assert audit["metrics"]["blocked_proposed_actions"] == 0
    assert audit["metrics"]["content_changes_proposed"] == 0
    assert audit["metrics"]["apply_now_true"] == 0
    assert audit["decision"] == "STAGE4_1_COMPLETE - CANONICAL_REFERENCE_STABILIZED_AND_READ_ONLY_PLAN_APPROVED"
