from __future__ import annotations

import importlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from openpyxl import load_workbook

from automacao_gd.infrastructure.excel import service as excel_service
from automacao_gd.infrastructure.state import pipeline_state
from automacao_gd.application import processing_service
from automacao_gd.application import full_pipeline
from automacao_gd.application.contracts import OperationStatus


def test_replace_denied_returns_typed_error_and_preserves_original(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = tmp_path / "synthetic_workbook.xlsx"
    original = Workbook()
    original.active["A1"] = "ORIGINAL"
    original.save(workbook_path)
    original_bytes = workbook_path.read_bytes()

    changed = Workbook()
    changed.active["A1"] = "CHANGED"

    def deny_replace(_source: Path, _destination: Path) -> None:
        raise PermissionError("synthetic replace denial")

    monkeypatch.setattr(excel_service.os, "replace", deny_replace)

    with pytest.raises(PermissionError) as exc_info:
        excel_service._save_workbook_atomically(changed, workbook_path)

    assert getattr(exc_info.value, "code", None) == "WORKBOOK_ATOMIC_REPLACE_DENIED"
    assert workbook_path.read_bytes() == original_bytes
    assert not workbook_path.with_name(".synthetic_workbook.saving.xlsx").exists()


def test_shareable_report_is_aggregate_only_and_recursively_safe() -> None:
    report_module = importlib.import_module(
        "automacao_gd.application.shareable_reports"
    )
    forbidden = {
        "windows_user_path": "C:" + "\\Users\\OPERADOR\\arquivo.xlsx",
        "network_path": "Z:" + "\\CLIENTES\\arquivo.pdf",
        "posix_path": "/CAMINHO/SINTETICO",
        "client_name": "CLIENTE " + "NAO COMPARTILHAVEL",
        "protocol": "2600001107",
        "uc": "2600001003",
        "address": "RUA " + "NAO COMPARTILHAVEL",
        "token": "tok" + "_nao_compartilhavel",
        "cookie": "cook" + "ie_nao_compartilhavel",
        "portal_raw": "PORTAL_" + "RAW_NAO_COMPARTILHAVEL",
    }
    payload = {
        "status": "PARCIAL",
        "dry_run": True,
        "total_pdfs": 2,
        "total_success": 1,
        "total_errors": 1,
        "blocked_real_run": False,
        "results": [forbidden],
        "nested": {"leak": forbidden},
    }

    report = report_module.build_shareable_report(
        payload,
        report_type="processing",
    )
    serialized = json.dumps(report, ensure_ascii=False)

    assert report == {
        "classification": "SHAREABLE",
        "schema_version": 1,
        "report_type": "processing",
        "status": "PARCIAL",
        "mode": "dry_run",
        "blocked": False,
        "totals": {
            "total_pdfs": 2,
            "total_success": 1,
            "total_errors": 1,
        },
    }
    for value in forbidden.values():
        assert value not in serialized


def test_rollback_trace_preserves_archive_effect_and_requires_manual_action() -> None:
    results = [
        {
            "success": True,
            "excel_status": {"can_write": True, "skipped": False, "success": True},
            "excel_effect": "applied",
            "archive_effect": "archived",
            "state_effect": "persisted",
            "report_effect": "not_applied",
            "manual_action_required": False,
            "rollback_possible": True,
        }
    ]

    marked = processing_service._mark_results_after_rollback(
        results,
        reason="synthetic systemic failure",
    )

    assert marked[0]["excel_effect"] == "rolled_back"
    assert marked[0]["archive_effect"] == "archived"
    assert marked[0]["state_effect"] == "persisted"
    assert marked[0]["manual_action_required"] is True
    assert marked[0]["rollback_possible"] is False


def test_restoration_is_confirmed_only_for_matching_valid_workbook(tmp_path: Path) -> None:
    backup = tmp_path / "synthetic_backup.xlsx"
    restored = tmp_path / "synthetic_restored.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "SYNTHETIC"
    workbook.save(backup)
    restored.write_bytes(backup.read_bytes())

    assert processing_service._verify_restored_workbook(backup, restored) is True

    different = Workbook()
    different.active["A1"] = "DIFFERENT"
    different.save(restored)
    assert processing_service._verify_restored_workbook(backup, restored) is False


def test_report_failure_is_returned_as_effect_instead_of_reusing_success(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = {
        "status": "SUCESSO",
        "operation_message": "synthetic success",
        "dry_run": False,
        "total_pdfs": 1,
        "total_success": 1,
        "total_errors": 0,
        "blocked_real_run": False,
        "results": [
            {
                "success": True,
                "report_effect": "not_applied",
                "manual_action_required": False,
            }
        ],
    }

    def fail_write(*_args, **_kwargs) -> None:
        raise PermissionError("synthetic report denial")

    monkeypatch.setattr(processing_service, "atomic_write_json", fail_write)

    persisted = processing_service._persist_processing_reports(tmp_path, payload)

    assert persisted["status"] == "FALHOU"
    assert persisted["report_effect"] == "failed"
    assert persisted["report_error_code"] == "REPORT_PERSISTENCE_FAILED"
    assert persisted["results"][0]["report_effect"] == "failed"
    assert persisted["results"][0]["manual_action_required"] is True


def test_pipeline_persistence_emits_aggregate_shareable_artifacts(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = {
        "status": "SUCESSO",
        "dry_run": True,
        "blocked_real_run": False,
        "total_processed": 2,
        "total_excel_updated": 0,
        "total_archived": 0,
        "protocols": [
            {
                "protocol": "2600001107",
                "client_name": "CLIENTE NAO COMPARTILHAVEL",
                "pdf_path": "C:" + "\\Users\\OPERADOR\\arquivo.pdf",
            }
        ],
    }
    monkeypatch.setattr(full_pipeline, "_build_markdown_report", lambda _payload: "private")

    full_pipeline._save_pipeline_reports(tmp_path, payload)

    shareable_json = tmp_path / "pipeline_cdp_shareable.json"
    shareable_markdown = tmp_path / "pipeline_cdp_shareable.md"
    assert shareable_json.is_file()
    assert shareable_markdown.is_file()
    report = json.loads(shareable_json.read_text(encoding="utf-8"))
    assert report["classification"] == "SHAREABLE"
    assert "protocols" not in report
    serialized = json.dumps(report, ensure_ascii=False)
    assert "2600001107" not in serialized
    assert "OPERADOR" not in serialized


def test_full_pipeline_report_failure_becomes_traceable_partial(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = {
        "status": OperationStatus.SUCESSO.value,
        "operation_message": "synthetic success",
        "total_downloaded": 1,
    }
    monkeypatch.setattr(
        full_pipeline,
        "_save_pipeline_reports",
        lambda *_: (_ for _ in ()).throw(PermissionError("synthetic denied")),
    )

    json_path, markdown_path = full_pipeline._persist_pipeline_reports(
        tmp_path,
        payload,
    )

    assert json_path is None
    assert markdown_path is None
    assert payload["status"] == OperationStatus.PARCIAL.value
    assert payload["report_effect"] == "failed"
    assert payload["manual_action_required"] is True
    assert "synthetic denied" not in payload["operation_message"]


def test_integrated_synthetic_dry_run_has_no_real_write(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    protocol = "2600000000"
    workbook_path = tmp_path / "synthetic_workbook.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2026"
    sheet.append(
        [
            "Cliente",
            "Protocolo",
            "Data de ingresso",
            "Conclusão",
            "Parecer",
            "Placa",
            "Inversor",
        ]
    )
    workbook.save(workbook_path)
    original_workbook = workbook_path.read_bytes()
    pdf_path = tmp_path / f"Orcamento_de_Conexao_{protocol}.pdf"
    pdf_path.write_bytes(b"SYNTHETIC PDF PLACEHOLDER")
    client_folder = tmp_path / "CLIENTE_SINTETICO"
    client_folder.mkdir()
    logs_dir = tmp_path / "logs"

    validation = SimpleNamespace(
        approved=True,
        status="approved",
        errors=[],
        warnings=[],
        module_source="synthetic",
        inverter_source="synthetic",
    )
    monkeypatch.setattr(
        processing_service,
        "_load_or_extract_technical_data",
        lambda *_args, **_kwargs: (
            protocol,
            "CLIENTE SINTETICO LTDA",
            None,
            None,
            "1x MODULO SINTETICO",
            "1x INVERSOR SINTETICO",
            validation,
        ),
    )
    monkeypatch.setattr(
        processing_service,
        "load_portal_metadata",
        lambda *_args, **_kwargs: ({"entry_date": "01/06/2026"}, "synthetic"),
    )
    monkeypatch.setattr(
        processing_service,
        "find_client_folder",
        lambda *_args, **_kwargs: SimpleNamespace(
            match_type="protocol",
            matched_path=str(client_folder),
            confidence=100,
            cache_hit=False,
            cache_key="synthetic",
            reason=None,
            found_by="protocol",
            protocol_search_hit=True,
            search_elapsed_seconds=0,
        ),
    )
    monkeypatch.setattr(
        processing_service,
        "resolve_archive_destination_folder",
        lambda **_kwargs: SimpleNamespace(
            destination_folder=client_folder,
            match_type="protocol",
            reason=None,
            should_create_folder=False,
            fallback_mode=None,
            legacy_gd_ignored=False,
        ),
    )
    monkeypatch.setattr(
        processing_service,
        "get_settings",
        lambda: SimpleNamespace(
            logs_dir_path=logs_dir,
            downloads_dir_path=tmp_path,
            BACKUP_EXCEL=True,
        ),
    )

    payload = processing_service.process_downloaded_pdfs(
        downloads_root=tmp_path,
        workbook_path=workbook_path,
        clientes_root=tmp_path,
        dry_run=True,
        pdf_paths=[pdf_path],
        apply_excel=True,
        apply_archive=True,
        allowed_protocols={protocol},
    )

    assert payload["status"] == "SUCESSO"
    assert payload["total_excel_updated"] == 0
    assert payload["total_archived"] == 0
    assert workbook_path.read_bytes() == original_workbook
    assert pdf_path.is_file()
    assert (logs_dir / "processamento_pdfs_shareable.json").is_file()


def test_synthetic_restore_and_resume_do_not_duplicate_effects(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    protocol = "2600000000"
    workbook_path = tmp_path / "synthetic_workbook.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2026"
    sheet.append(
        [
            "Cliente",
            "Protocolo",
            "Data de ingresso",
            "Conclusão",
            "Parecer",
            "Placa",
            "Inversor",
        ]
    )
    workbook.save(workbook_path)
    backup_path = excel_service.create_workbook_backup(workbook_path)

    first = excel_service.update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol=protocol,
        client_name="CLIENTE SINTETICO LTDA",
        entry_date="01/06/2026",
        module_text="1x MODULO SINTETICO",
        inverter_text="1x INVERSOR SINTETICO",
        parecer="Sim",
        dry_run=False,
    )
    assert first["success"] is True

    processing_service.atomic_copy_file(backup_path, workbook_path, private=True)
    assert processing_service._verify_restored_workbook(backup_path, workbook_path)

    resumed = excel_service.update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol=protocol,
        client_name="CLIENTE SINTETICO LTDA",
        entry_date="01/06/2026",
        module_text="1x MODULO SINTETICO",
        inverter_text="1x INVERSOR SINTETICO",
        parecer="Sim",
        dry_run=False,
    )
    repeated = excel_service.update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol=protocol,
        client_name="CLIENTE SINTETICO LTDA",
        entry_date="01/06/2026",
        module_text="1x MODULO SINTETICO",
        inverter_text="1x INVERSOR SINTETICO",
        parecer="Sim",
        dry_run=False,
    )
    assert resumed["success"] is True
    assert repeated["success"] is True

    reloaded = load_workbook(workbook_path, data_only=False)
    occurrences = sum(
        1
        for worksheet in reloaded.worksheets
        for row in worksheet.iter_rows(min_row=2, values_only=True)
        if str(row[1] or "") == protocol
    )
    reloaded.close()
    assert occurrences == 1

    source_pdf = tmp_path / "synthetic_source.pdf"
    destination_pdf = tmp_path / "synthetic_archive" / "Orcamento_de_Conexao.pdf"
    source_pdf.write_bytes(b"SYNTHETIC PDF PLACEHOLDER")
    processing_service.atomic_copy_file(source_pdf, destination_pdf, private=True)
    processing_service.atomic_copy_file(source_pdf, destination_pdf, private=True)
    assert list(destination_pdf.parent.glob("*.pdf")) == [destination_pdf]

    monkeypatch.setattr(
        pipeline_state,
        "get_settings",
        lambda: SimpleNamespace(resolve_path=lambda path: Path(path)),
    )
    state_path = tmp_path / "synthetic_state.json"
    state = pipeline_state.PipelineStateStore(path=state_path, resume=False)
    state.mark_completed(protocol)
    state.mark_completed(protocol)
    resumed_state = pipeline_state.PipelineStateStore(path=state_path, resume=True)
    assert list(resumed_state.state["protocols"]) == [protocol]


def test_excel_update_can_mark_boolean_parecer_false_for_unavailable_budget(
    tmp_path: Path,
) -> None:
    protocol = "2600000001"
    workbook_path = tmp_path / "synthetic_workbook.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "2026"
    sheet.append(
        [
            "Cliente",
            "Protocolo",
            "Data de ingresso",
            "ConclusÃ£o",
            "Parecer",
            "Placa",
            "Inversor",
        ]
    )
    sheet.append(
        [
            "CLIENTE SINTETICO LTDA",
            protocol,
            "01/06/2026",
            None,
            True,
            "1x MODULO EXISTENTE",
            "1x INVERSOR EXISTENTE",
        ]
    )
    workbook.save(workbook_path)

    result = excel_service.update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol=protocol,
        client_name="CLIENTE SINTETICO LTDA",
        entry_date="01/06/2026",
        completion_date="10/06/2026",
        module_text=None,
        inverter_text=None,
        parecer="False",
        dry_run=False,
    )

    assert result["success"] is True
    assert result["action"] == "update_existing"

    reloaded = load_workbook(workbook_path, data_only=False)
    try:
        row = next(reloaded["2026"].iter_rows(min_row=2, max_row=2, values_only=True))
    finally:
        reloaded.close()
    assert row[4] is False
    assert row[5] == "1x MODULO EXISTENTE"
    assert row[6] == "1x INVERSOR EXISTENTE"
