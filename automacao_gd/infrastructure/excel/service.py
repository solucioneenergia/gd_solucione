import copy
import os
import re
import shutil
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from automacao_gd.infrastructure.config import get_settings
from automacao_gd.infrastructure.dates import date_to_excel_datetime, parse_date
from automacao_gd.infrastructure.excel.availability import (
    classify_workbook_write_error,
    validate_workbook_availability,
)
from automacao_gd.infrastructure.excel.excel_update_helpers import (
    OPEN_COMPLETION_TEXT,
    base_excel_result as _base_excel_result,
    completion_value_matches as _completion_value_matches,
    equipment_text_matches as _equipment_text_matches,
    has_text as _has_text,
    is_open_completion_value as _is_open_completion_value,
    normalize_cell as _normalize_cell,
    normalize_equipment_text as _normalize_equipment_text,
    payload_value as _payload_value,
    set_payload_value as _set_payload_value,
    sheet_name_for_entry_date as _sheet_name_for_entry_date,
    sheet_name_for_year as _sheet_name_for_year,
    status_payload as _status_payload,
)
from automacao_gd.infrastructure.logging import logger
from automacao_gd.infrastructure.persistence.atomic import atomic_copy_file
from automacao_gd.domain.models import GenerationData, PortalSolicitation


REQUIRED_COLUMNS = [
    "Cliente",
    "Protocolo",
    "Data de ingresso",
    "Conclusão",
    "Parecer",
    "Placa",
    "Inversor",
]
MAIN_WORKSHEET_NAMES = ["2022 - 2023", "2024", "2025", "2026"]
DATE_NUMBER_FORMAT = "dd/mm/yyyy"
TEXT_NUMBER_FORMAT = "@"
TOP_LAYOUT_SCAN_ROWS = 5
SAFE_LOGO_TEXT = "SOLUCIONE NORDESTE ENERGIA ELÉTRICA"
SAFE_COLUMN_WIDTHS = {
    "A": 38,
    "B": 16,
    "C": 16,
    "D": 16,
    "E": 13,
    "F": 42,
    "G": 38,
}


class WorkbookAtomicReplaceError(PermissionError):
    """Falha fechada quando o arquivo oficial nao pode ser substituido atomicamente."""

    code = "WORKBOOK_ATOMIC_REPLACE_DENIED"


def load_workbook_safe(path: Path):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {path}")
    logger.info(f"Carregando planilha: {path}")
    return load_workbook(path)


def find_required_columns(ws: Worksheet) -> dict[str, int]:
    _, columns = _find_required_columns_with_header(ws)
    return columns


def validate_workbook_for_pdf_updates(
    workbook_path: Path, require_writable: bool = False
) -> dict:
    workbook_path = Path(workbook_path)
    result: dict[str, Any] = {
        "success": False,
        "worksheet": None,
        "header_row": None,
        "error": None,
        "code": None,
        "technical_cause": None,
    }

    availability = validate_workbook_availability(
        workbook_path,
        require_writable=require_writable,
        stage="validação da planilha",
    )
    if not availability.ok:
        result["error"] = availability.user_message
        result["code"] = availability.code.value
        result["technical_cause"] = availability.technical_cause
        return result

    wb: Workbook | None = None
    try:
        wb = load_workbook_safe(workbook_path)
        ws, header_row, _ = _find_workbook_sheet_with_required_columns(wb)
        result.update(
            {
                "success": True,
                "worksheet": ws.title,
                "header_row": header_row,
                "code": availability.code.value,
            }
        )
        return result
    except Exception as exc:
        result["error"] = str(exc)
        result["code"] = "WORKBOOK_INVALID"
        result["technical_cause"] = f"type={type(exc).__name__}"
        return result
    finally:
        if wb is not None:
            wb.close()


def plan_excel_update_from_pdf_data(
    workbook_path: Path,
    protocol: str,
    entry_date: Any,
    dry_run: bool = True,
) -> dict:
    return update_excel_from_pdf_data(
        workbook_path=workbook_path,
        protocol=protocol,
        client_name="",
        entry_date=entry_date,
        completion_date=None,
        module_text="",
        inverter_text="",
        dry_run=True,
        apply_changes=False,
    )


def validate_excel_protocol_updated(
    workbook_path: Path,
    protocol: str,
    entry_date: Any = None,
    completion_date: Any = None,
    module_text: str | None = None,
    inverter_text: str | None = None,
    require_completion_present: bool = False,
) -> dict:
    workbook_path = Path(workbook_path)
    result: dict[str, Any] = {
        "success": False,
        "already_updated": False,
        "protocol": protocol,
        "worksheet": None,
        "row": None,
        "target_sheet": None,
        "error": None,
    }

    entry_dt = parse_date(entry_date)
    if entry_dt is not None:
        result["target_sheet"] = _sheet_name_for_entry_date(entry_dt)

    if not workbook_path.exists():
        result["error"] = f"Planilha não encontrada: {workbook_path}"
        return result

    wb: Workbook | None = None
    try:
        wb = load_workbook_safe(workbook_path)
        sheet_map = _workbook_required_sheets(wb)
        occurrences = _find_protocol_occurrences(sheet_map, protocol)
        if len(occurrences) != 1:
            result["error"] = (
                "Protocolo não encontrado na planilha"
                if not occurrences
                else "Protocolo duplicado na planilha"
            )
            return result

        occurrence = occurrences[0]
        info = sheet_map[occurrence["sheet"]]
        row = occurrence["row"]
        result.update(
            {
                "success": True,
                "worksheet": occurrence["sheet"],
                "row": row,
            }
        )
        if result["target_sheet"] and occurrence["sheet"] != result["target_sheet"]:
            return result

        if require_completion_present and completion_date is None:
            completion_col = _column_by_normalized_name(info["columns"], "CONCLUSAO")
            if not _has_text(info["worksheet"].cell(row=row, column=completion_col).value):
                result["already_updated"] = False
                result["error"] = "Conclusão vazia na planilha."
                return result

        result["already_updated"] = _row_has_required_update_values(
            info["worksheet"],
            row,
            info["columns"],
            entry_dt,
            completion_date,
            module_text,
            inverter_text,
        )
        return result
    except Exception as exc:
        result["error"] = str(exc)
        return result
    finally:
        if wb is not None:
            wb.close()


def protocol_row_has_required_values(workbook_path: Path, protocol: str) -> dict:
    workbook_path = Path(workbook_path)
    result: dict[str, Any] = {
        "success": False,
        "complete": False,
        "protocol": protocol,
        "worksheet": None,
        "row": None,
        "missing_columns": [],
        "error": None,
    }
    if not workbook_path.exists():
        result["error"] = f"Planilha não encontrada: {workbook_path}"
        return result

    wb: Workbook | None = None
    try:
        wb = load_workbook_safe(workbook_path)
        sheet_map = _workbook_required_sheets(wb)
        occurrences = _find_protocol_occurrences(sheet_map, protocol)
        if len(occurrences) != 1:
            result["error"] = (
                "Protocolo não encontrado na planilha"
                if not occurrences
                else "Protocolo duplicado na planilha"
            )
            return result

        occurrence = occurrences[0]
        info = sheet_map[occurrence["sheet"]]
        row = occurrence["row"]
        required_columns = _required_columns_for_completed_protocol_row(
            info["worksheet"],
            row,
            info["columns"],
        )
        missing = [
            column_name
            for column_name in required_columns
            if not _has_text(
                info["worksheet"].cell(
                    row=row,
                    column=info["columns"][column_name],
                ).value
            )
        ]
        result.update(
            {
                "success": True,
                "complete": not missing,
                "worksheet": occurrence["sheet"],
                "row": row,
                "missing_columns": missing,
            }
        )
        return result
    except Exception as exc:
        result["error"] = str(exc)
        return result
    finally:
        if wb is not None:
            wb.close()


def _required_columns_for_completed_protocol_row(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
) -> list[str]:
    parecer = ws.cell(row=row, column=columns["Parecer"]).value
    if isinstance(parecer, bool) and parecer is False:
        return [
            column
            for column in REQUIRED_COLUMNS
            if column not in {"Placa", "Inversor"}
        ]
    if _normalize_cell(parecer).upper() in {
        "FALSE",
        "FALSO",
        "NAO",
        "NÃƒO",
        "NÃO",
        "NO",
        "0",
    }:
        return [
            column
            for column in REQUIRED_COLUMNS
            if column not in {"Placa", "Inversor"}
        ]
    return REQUIRED_COLUMNS


def update_excel_equipment_columns(
    workbook_path: Path,
    protocol: str,
    module_text: str | None,
    inverter_text: str | None,
    *,
    dry_run: bool = True,
    backup_path: Path | None = None,
    apply_changes: bool = True,
) -> dict:
    """Atualiza somente as colunas Placa e Inversor de um protocolo existente."""
    workbook_path = Path(workbook_path)
    result = _base_excel_result(protocol, dry_run)
    result["action"] = "reformat_existing_excel_row"

    if not workbook_path.exists():
        result["error"] = f"Planilha nÃ£o encontrada: {workbook_path}"
        return result

    if not dry_run:
        try:
            _ensure_workbook_writable(workbook_path)
        except PermissionError as exc:
            _set_workbook_operational_error(result, exc)
            return result

    wb: Workbook | None = None
    try:
        wb = load_workbook_safe(workbook_path)
        sheet_map = _workbook_required_sheets(wb)
        occurrences = _find_protocol_occurrences(sheet_map, protocol)

        if len(occurrences) != 1:
            result.update(
                {
                    "action": "blocked_missing_protocol"
                    if not occurrences
                    else "blocked_duplicate_protocol",
                    "error": (
                        "Protocolo nÃ£o encontrado na planilha"
                        if not occurrences
                        else "Protocolo duplicado na planilha"
                    ),
                }
            )
            return result

        occurrence = occurrences[0]
        ws_info = sheet_map[occurrence["sheet"]]
        ws = ws_info["worksheet"]
        row = occurrence["row"]
        columns = ws_info["columns"]
        placa_col = columns["Placa"]
        inversor_col = columns["Inversor"]
        old_module = ws.cell(row=row, column=placa_col).value
        old_inverter = ws.cell(row=row, column=inversor_col).value
        cleaned_module = _clean_equipment_text(module_text)
        cleaned_inverter = _clean_equipment_text(inverter_text)

        result.update(
            {
                "success": True,
                "can_write": True,
                "row_found": True,
                "row_number": row,
                "worksheet": occurrence["sheet"],
                "target_sheet": occurrence["sheet"],
                "existing_sheet": occurrence["sheet"],
                "source_sheet": occurrence["sheet"],
                "source_row": row,
                "target_row": row,
                "existing_row": row,
                "new_row": row,
                "created_new_row": False,
                "old_module_text": old_module,
                "old_inverter_text": old_inverter,
                "new_module_text": cleaned_module,
                "new_inverter_text": cleaned_inverter,
                "updated_columns": ["Placa", "Inversor"],
            }
        )

        if not dry_run and apply_changes:
            result["backup_path"] = str(backup_path) if backup_path else None
            ws.cell(row=row, column=placa_col).value = cleaned_module
            ws.cell(row=row, column=inversor_col).value = cleaned_inverter
            _apply_equipment_cell_layout(ws, row, columns, cleaned_module, cleaned_inverter)
            _save_workbook_atomically(wb, workbook_path)
        return result
    except PermissionError as exc:
        _set_workbook_operational_error(result, exc)
        result["success"] = False
        result["can_write"] = False
        logger.error(result["error"])
        return result
    except Exception as exc:
        result["error"] = str(exc)
        logger.exception(
            f"Falha ao reformatar equipamentos na planilha para {protocol}: {exc}"
        )
        return result
    finally:
        if wb is not None:
            wb.close()


def update_excel_from_pdf_data(
    workbook_path: Path,
    protocol: str,
    client_name: str,
    entry_date: Any = None,
    completion_date: Any = None,
    module_text: str | None = "",
    inverter_text: str | None = "",
    parecer: str = "Sim",
    dry_run: bool = True,
    backup_path: Path | None = None,
    apply_changes: bool = True,
    require_completion_date: bool = False,
) -> dict:
    workbook_path = Path(workbook_path)
    entry_dt = parse_date(entry_date)
    result = _base_excel_result(protocol, dry_run)

    if entry_dt is None:
        result.update(
            {
                "action": "blocked_missing_entry_date",
                "error": "Data de ingresso não encontrada para o protocolo",
            }
        )
        return result

    result["entry_date"] = entry_dt.isoformat()
    result["target_sheet"] = _sheet_name_for_entry_date(entry_dt)

    if not workbook_path.exists():
        result["error"] = f"Planilha não encontrada: {workbook_path}"
        return result

    if not dry_run:
        try:
            _ensure_workbook_writable(workbook_path)
        except PermissionError as exc:
            _set_workbook_operational_error(result, exc)
            return result

    wb: Workbook | None = None
    try:
        wb = load_workbook_safe(workbook_path)
        sheet_map = _workbook_required_sheets(wb)
        occurrences = _find_protocol_occurrences(sheet_map, protocol)

        if len(occurrences) > 1:
            result.update(
                {
                    "action": "blocked_duplicate_protocol",
                    "error": (
                        "Protocolo duplicado na planilha: "
                        + ", ".join(
                            f"{item['sheet']}!{item['row']}"
                            for item in occurrences
                        )
                    ),
                }
            )
            return result

        if occurrences:
            return _handle_existing_row(
                result=result,
                wb=wb,
                sheet_map=sheet_map,
                occurrence=occurrences[0],
                entry_dt=entry_dt,
                protocol=protocol,
                client_name=client_name,
                completion_date=completion_date,
                module_text=module_text,
                inverter_text=inverter_text,
                parecer=parecer,
                dry_run=dry_run,
                backup_path=backup_path,
                apply_changes=apply_changes,
                workbook_path=workbook_path,
                require_completion_date=require_completion_date,
            )

        return _handle_new_row(
            result=result,
            wb=wb,
            sheet_map=sheet_map,
            entry_dt=entry_dt,
            protocol=protocol,
            client_name=client_name,
            completion_date=completion_date,
            module_text=module_text,
            inverter_text=inverter_text,
            parecer=parecer,
            dry_run=dry_run,
            backup_path=backup_path,
            apply_changes=apply_changes,
            workbook_path=workbook_path,
            require_completion_date=require_completion_date,
        )
    except PermissionError as exc:
        _set_workbook_operational_error(result, exc)
        result["success"] = False
        result["can_write"] = False
        logger.error(result["error"])
        return result
    except Exception as exc:
        result["error"] = str(exc)
        logger.exception(f"Falha ao mapear/atualizar planilha para {protocol}: {exc}")
        return result
    finally:
        if wb is not None:
            wb.close()


def _handle_existing_row(
    result: dict,
    wb: Workbook,
    sheet_map: dict[str, dict],
    occurrence: dict,
    entry_dt: date,
    protocol: str,
    client_name: str,
    completion_date: Any,
    module_text: str | None,
    inverter_text: str | None,
    parecer: str,
    dry_run: bool,
    backup_path: Path | None,
    apply_changes: bool,
    workbook_path: Path,
    require_completion_date: bool,
) -> dict:
    existing_sheet = occurrence["sheet"]
    existing_row = occurrence["row"]
    result.update(
        {
            "row_found": True,
            "row_number": existing_row,
            "worksheet": existing_sheet,
            "existing_sheet": existing_sheet,
            "source_sheet": existing_sheet,
            "source_row": existing_row,
            "existing_row": existing_row,
            "created_new_row": False,
        }
    )
    ws_info = sheet_map[existing_sheet]
    if require_completion_date and completion_date is None:
        completion_col = _column_by_normalized_name(ws_info["columns"], "CONCLUSAO")
        current_completion = ws_info["worksheet"].cell(
            row=existing_row,
            column=completion_col,
        ).value
        if not _has_text(current_completion):
            result.update(
                {
                    "success": False,
                    "can_write": False,
                    "action": "blocked_missing_completion_date",
                    "target_row": existing_row,
                    "new_row": existing_row,
                    "moved_from": f"{existing_sheet}!{existing_row}",
                    "moved_to": f"{existing_sheet}!{existing_row}",
                    "completion_no_change": False,
                    "completion_action": "MISSING_COMPLETION_DATE",
                    "error": (
                        "Conclusão vazia na planilha e data de conclusão não "
                        "disponível no plano OP5."
                    ),
                }
            )
            return result

    if existing_sheet != result["target_sheet"]:
        return _handle_wrong_sheet(
            result=result,
            wb=wb,
            sheet_map=sheet_map,
            existing_sheet=existing_sheet,
            existing_row=existing_row,
            entry_dt=entry_dt,
            protocol=protocol,
            client_name=client_name,
            completion_date=completion_date,
            module_text=module_text,
            inverter_text=inverter_text,
            parecer=parecer,
            dry_run=dry_run,
            backup_path=backup_path,
            apply_changes=apply_changes,
            workbook_path=workbook_path,
        )

    update_state = _row_required_update_state(
        ws_info["worksheet"],
        existing_row,
        ws_info["columns"],
        entry_dt,
        completion_date,
        module_text,
        inverter_text,
        parecer,
    )
    result.update(update_state)
    if update_state["all_no_change"]:
        result.update(
            {
                "success": True,
                "can_write": False,
                "action": "skipped_excel_already_updated",
                "target_row": existing_row,
                "new_row": existing_row,
                "moved_from": f"{existing_sheet}!{existing_row}",
                "moved_to": f"{existing_sheet}!{existing_row}",
                "warning": "Protocolo ja estava atualizado na planilha.",
            }
        )
        return result

    result.update(
        {
            "success": True,
            "can_write": True,
            "action": "update_existing",
            "target_row": existing_row,
            "new_row": existing_row,
            "moved_from": f"{existing_sheet}!{existing_row}",
            "moved_to": f"{existing_sheet}!{existing_row}",
        }
    )
    if not dry_run and apply_changes:
        result["backup_path"] = str(backup_path) if backup_path else None
        _write_existing_excel_row_updates(
            ws_info["worksheet"],
            occurrence["row"],
            ws_info["columns"],
            completion_date,
            module_text,
            inverter_text,
            parecer,
            update_state,
        )
        _save_workbook_atomically(wb, workbook_path)
    return result


def _handle_wrong_sheet(
    result: dict,
    wb: Workbook,
    sheet_map: dict[str, dict],
    existing_sheet: str,
    existing_row: int,
    entry_dt: date,
    protocol: str,
    client_name: str,
    completion_date: Any,
    module_text: str | None,
    inverter_text: str | None,
    parecer: str,
    dry_run: bool,
    backup_path: Path | None,
    apply_changes: bool,
    workbook_path: Path,
) -> dict:
    settings = get_settings()
    if not settings.MOVE_WRONG_YEAR_ROWS:
        result.update(
            {
                "success": True,
                "can_write": False,
                "action": "warning_wrong_sheet",
                "warning": (
                    "Protocolo encontrado em aba diferente da data de ingresso"
                ),
            }
        )
        return result

    target_info = _get_or_prepare_target_sheet(
        wb=wb,
        sheet_map=sheet_map,
        target_sheet=str(result["target_sheet"]),
        dry_run=dry_run,
    )
    if not target_info["success"]:
        result.update(target_info)
        return result

    if target_info.get("sheet_would_be_created"):
        target_row = target_info["header_row"] + 1
    else:
        target_row = _find_chronological_insert_row(
            target_info["worksheet"],
            target_info["header_row"],
            target_info["columns"],
            entry_dt,
        )

    result.update(
        {
            "success": True,
            "can_write": True,
            "action": "move_wrong_sheet_to_correct_sheet",
            "worksheet": result["target_sheet"],
            "row_number": target_row,
            "target_row": target_row,
            "new_row": target_row,
            "moved_from": f"{existing_sheet}!{existing_row}",
            "moved_to": f"{result['target_sheet']}!{target_row}",
            "created_new_row": True,
            "warning": _movement_message(
                existing_sheet,
                existing_row,
                str(result["target_sheet"]),
                target_row,
                dry_run,
                target_info.get("warning"),
            ),
        }
    )

    if not dry_run and apply_changes:
        source_ws = sheet_map[existing_sheet]["worksheet"]
        target_ws = target_info["worksheet"]
        _insert_row_preserving_style(target_ws, target_row)
        _write_excel_row(
            target_ws,
            target_row,
            target_info["columns"],
            protocol,
            client_name,
            entry_dt,
            completion_date,
            module_text,
            inverter_text,
            parecer,
            target_info.get("parecer_boolean"),
        )
        _validate_protocol_at_row(
            target_ws,
            target_row,
            target_info["columns"]["Protocolo"],
            protocol,
        )
        source_ws.delete_rows(existing_row)
        result["backup_path"] = str(backup_path) if backup_path else None
        result["warning"] = _movement_message(
            existing_sheet,
            existing_row,
            str(result["target_sheet"]),
            target_row,
            False,
            None,
        )
        _save_workbook_atomically(wb, workbook_path)
    return result


def _handle_new_row(
    result: dict,
    wb: Workbook,
    sheet_map: dict[str, dict],
    entry_dt: date,
    protocol: str,
    client_name: str,
    completion_date: Any,
    module_text: str | None,
    inverter_text: str | None,
    parecer: str,
    dry_run: bool,
    backup_path: Path | None,
    apply_changes: bool,
    workbook_path: Path,
    require_completion_date: bool,
) -> dict:
    if require_completion_date and completion_date is None:
        result.update(
            {
                "success": False,
                "can_write": False,
                "action": "blocked_missing_completion_date",
                "completion_no_change": False,
                "completion_action": "MISSING_COMPLETION_DATE",
                "error": (
                    "Data de conclusão não disponível para criar linha OP5 "
                    "com Conclusão preenchida."
                ),
            }
        )
        return result

    target_info = _get_or_prepare_target_sheet(
        wb=wb,
        sheet_map=sheet_map,
        target_sheet=str(result["target_sheet"]),
        dry_run=dry_run,
    )
    if not target_info["success"]:
        result.update(target_info)
        return result

    result["worksheet"] = result["target_sheet"]
    result["warning"] = target_info.get("warning")
    if target_info.get("sheet_would_be_created"):
        target_row = target_info["header_row"] + 1
    else:
        target_row = _find_chronological_insert_row(
            target_info["worksheet"],
            target_info["header_row"],
            target_info["columns"],
            entry_dt,
        )
    result.update(
        {
            "success": True,
            "can_write": True,
            "action": "insert_new_chronological",
            "row_number": target_row,
            "target_row": target_row,
            "new_row": target_row,
            "created_new_row": True,
        }
    )

    if not dry_run and apply_changes:
        ws = target_info["worksheet"]
        _insert_row_preserving_style(ws, target_row)
        _write_excel_row(
            ws,
            target_row,
            target_info["columns"],
            protocol,
            client_name,
            entry_dt,
            completion_date,
            module_text,
            inverter_text,
            parecer,
            target_info.get("parecer_boolean"),
        )
        result["backup_path"] = str(backup_path) if backup_path else None
        _save_workbook_atomically(wb, workbook_path)
    return result


def _save_workbook_atomically(wb: Workbook, workbook_path: Path) -> None:
    """Salva no mesmo volume, valida o XLSX e substitui o arquivo original."""
    workbook_path = Path(workbook_path)
    temp_path = workbook_path.with_name(
        f".{workbook_path.stem}.saving{workbook_path.suffix}"
    )
    temp_path.unlink(missing_ok=True)
    try:
        wb.save(temp_path)
        ok, errors = _validate_xlsx_zip_xml_only(temp_path)
        if not ok:
            raise ValueError("XLSX temporário inválido: " + "; ".join(errors))
        try:
            os.replace(temp_path, workbook_path)
        except PermissionError as exc:
            raise WorkbookAtomicReplaceError(
                "A substituicao atomica do workbook foi negada. "
                "O arquivo original foi preservado; verifique bloqueio, permissao ou sincronizacao."
            ) from exc
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    finally:
        temp_path.unlink(missing_ok=True)


def _workbook_save_permission_message(workbook_path: Path) -> str:
    return (
        "A planilha não pôde ser salva ou substituída. "
        "Verifique se o arquivo está aberto, se há permissão de escrita "
        "na pasta ou se a unidade de rede/sincronização bloqueou a operação. "
        f"Arquivo: {workbook_path}"
    )


def create_workbook_backup(workbook_path: Path) -> Path:
    return _backup_workbook(Path(workbook_path))


def validate_xlsx_integrity(path: Path) -> tuple[bool, list[str]]:
    path = Path(path)
    errors: list[str] = []
    if not path.exists():
        return False, [f"Arquivo nao encontrado: {path}"]
    if not zipfile.is_zipfile(path):
        return False, [f"Arquivo nao e um ZIP/XLSX valido: {path}"]

    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            required = {"xl/workbook.xml"}
            for item in required:
                if item not in names:
                    errors.append(f"XML obrigatorio ausente: {item}")
            xml_names = [
                name
                for name in archive.namelist()
                if name == "xl/workbook.xml"
                or name == "xl/styles.xml"
                or (name.startswith("xl/worksheets/") and name.endswith(".xml"))
            ]
            for name in xml_names:
                try:
                    ElementTree.fromstring(archive.read(name))
                except Exception as exc:
                    errors.append(f"XML invalido em {name}: {exc}")
    except Exception as exc:
        errors.append(f"Falha ao abrir XLSX como ZIP: {exc}")

    if errors:
        return False, errors

    wb: Workbook | None = None
    try:
        wb = load_workbook(path)
        resave_path = path.with_name(f"{path.stem}_integrity_resave{path.suffix}")
        wb.save(resave_path)
        wb.close()
        wb = None
        ok, resave_errors = _validate_xlsx_zip_xml_only(resave_path)
        if not ok:
            errors.extend(f"Resave invalido: {error}" for error in resave_errors)
        try:
            resave_path.unlink(missing_ok=True)
        except OSError:
            pass
    except Exception as exc:
        errors.append(f"openpyxl nao conseguiu abrir/salvar o XLSX: {exc}")
    finally:
        if wb is not None:
            wb.close()

    return not errors, errors


def _validate_xlsx_zip_xml_only(path: Path) -> tuple[bool, list[str]]:
    errors: list[str] = []
    if not zipfile.is_zipfile(path):
        return False, [f"Arquivo nao e ZIP/XLSX valido: {path}"]
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name == "xl/workbook.xml" or name == "xl/styles.xml" or (
                    name.startswith("xl/worksheets/") and name.endswith(".xml")
                ):
                    try:
                        ElementTree.fromstring(archive.read(name))
                    except Exception as exc:
                        errors.append(f"XML invalido em {name}: {exc}")
    except Exception as exc:
        errors.append(str(exc))
    return not errors, errors


def _save_repaired_workbook_safely(wb: Workbook, workbook_path: Path) -> dict:
    temp_path = workbook_path.with_name(f"{workbook_path.stem}.tmp{workbook_path.suffix}")
    result: dict[str, Any] = {
        "temp_path": str(temp_path),
        "xlsx_integrity_ok": False,
        "official_file_replaced": False,
        "technical_errors": [],
    }
    try:
        if temp_path.exists():
            temp_path.unlink()
        wb.save(temp_path)
        ok, integrity_errors = validate_xlsx_integrity(temp_path)
        result["xlsx_integrity_ok"] = ok
        if not ok:
            result["technical_errors"].extend(integrity_errors)
            result["technical_errors"].append(
                f"Arquivo temporario preservado para diagnostico: {temp_path}"
            )
            return result
        shutil.move(str(temp_path), str(workbook_path))
        result["official_file_replaced"] = True
        return result
    except Exception as exc:
        result["technical_errors"].append(f"Falha ao salvar/substituir XLSX: {exc}")
        return result


def repair_workbook_format(
    workbook_path: Path,
    dry_run: bool = True,
    template_sheet_name: str | None = None,
    sheet_names: list[str] | None = None,
    mode: str | None = None,
) -> dict:
    workbook_path = Path(workbook_path)
    settings = get_settings()
    template_sheet_name = template_sheet_name or settings.EXCEL_DEFAULT_TEMPLATE_SHEET
    sheet_names = sheet_names or MAIN_WORKSHEET_NAMES
    mode = (mode or settings.WORKBOOK_REPAIR_MODE or "visual_only").strip().lower()
    result: dict[str, Any] = {
        "success": False,
        "dry_run": dry_run,
        "mode": mode,
        "workbook_path": str(workbook_path),
        "backup_path": None,
        "temp_path": None,
        "xlsx_integrity_ok": None,
        "official_file_replaced": False,
        "template_sheet": template_sheet_name,
        "sheets": [],
        "warnings": [],
        "errors": [],
        "visual_errors": [],
        "technical_errors": [],
        "data_pending": [],
    }

    if not workbook_path.exists():
        message = f"Planilha nao encontrada: {workbook_path}"
        result["errors"].append(message)
        result["technical_errors"].append(message)
        return result

    if not dry_run:
        try:
            _ensure_workbook_writable(workbook_path)
        except PermissionError:
            message = "A planilha parece estar aberta no Excel. Feche o arquivo e tente novamente."
            result["errors"].append(message)
            result["technical_errors"].append(message)
            return result

    wb: Workbook | None = None
    try:
        backup_path = None if dry_run else create_workbook_backup(workbook_path)
        result["backup_path"] = str(backup_path) if backup_path else None
        wb = load_workbook_safe(workbook_path)
        if template_sheet_name not in wb.sheetnames:
            message = f"Aba modelo nao encontrada: {template_sheet_name}"
            result["errors"].append(message)
            result["technical_errors"].append(message)
            return result
        template_ws = wb[template_sheet_name]

        if mode != "validate_only":
            for sheet_name in sheet_names:
                if sheet_name not in wb.sheetnames:
                    if dry_run:
                        result["sheets"].append(
                            {
                                "sheet": sheet_name,
                                "would_create": True,
                                "formatted_rows": 0,
                                "value_errors_removed": 0,
                                "logo_copied": False,
                                "warnings": [f"Aba {sheet_name} seria criada."],
                            }
                        )
                        continue
                    ws = create_year_sheet_from_template(wb, template_ws, sheet_name)
                else:
                    ws = wb[sheet_name]

                sheet_report = standardize_sheet_layout(ws, template_ws)
                result["sheets"].append(sheet_report)

        validation = validate_workbook_format(wb=wb, sheet_names=sheet_names)
        result["validation"] = validation
        result["warnings"].extend(validation.get("warnings", []))
        result["visual_errors"].extend(validation.get("visual_errors", []))
        result["technical_errors"].extend(validation.get("technical_errors", []))
        result["data_pending"].extend(validation.get("data_pending", []))
        if mode == "visual_only":
            result["errors"].extend(result["technical_errors"])
            result["errors"].extend(result["visual_errors"])
        else:
            result["errors"].extend(validation.get("errors", []))
        result["success"] = not result["errors"]
        if not dry_run and mode != "validate_only":
            save_result = _save_repaired_workbook_safely(wb, workbook_path)
            result.update(save_result)
            result["technical_errors"].extend(save_result.get("technical_errors", []))
            result["errors"].extend(save_result.get("technical_errors", []))
            result["success"] = result["success"] and bool(
                save_result.get("official_file_replaced")
            )
        return result
    except Exception as exc:
        logger.exception(f"Falha ao reparar formatacao da planilha: {exc}")
        result["errors"].append(str(exc))
        result["technical_errors"].append(str(exc))
        return result
    finally:
        if wb is not None:
            wb.close()


def validate_workbook_format(
    workbook_path: Path | None = None,
    wb: Workbook | None = None,
    sheet_names: list[str] | None = None,
    strict: bool | None = None,
) -> dict:
    sheet_names = sheet_names or MAIN_WORKSHEET_NAMES
    settings = get_settings()
    strict = settings.STRICT_WORKBOOK_VALIDATION if strict is None else strict
    result: dict[str, Any] = {
        "success": False,
        "strict": strict,
        "workbook_path": str(workbook_path) if workbook_path else None,
        "sheets": [],
        "errors": [],
        "visual_errors": [],
        "technical_errors": [],
        "data_pending": [],
        "warnings": [],
        "duplicate_protocols": [],
        "wrong_year_rows": [],
    }
    close_workbook = False

    try:
        if wb is None:
            if workbook_path is None or not Path(workbook_path).exists():
                _add_validation_issue(
                    result,
                    "technical_errors",
                    f"Planilha nao encontrada: {workbook_path}",
                )
                return result
            wb = load_workbook_safe(Path(workbook_path))
            close_workbook = True

        if "2026" in sheet_names and "2026" not in wb.sheetnames:
            _add_validation_issue(result, "visual_errors", "Aba 2026 nao encontrada.")

        sheet_map = _workbook_required_sheets(wb)
        protocol_locations: dict[str, list[str]] = {}
        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                result["warnings"].append(f"Aba ausente: {sheet_name}")
                continue
            ws = wb[sheet_name]
            sheet_result = _validate_sheet_format(ws, sheet_map.get(sheet_name))
            result["sheets"].append(sheet_result)
            result["visual_errors"].extend(sheet_result.get("visual_errors", []))
            result["technical_errors"].extend(sheet_result.get("technical_errors", []))
            result["data_pending"].extend(sheet_result.get("data_pending", []))
            result["errors"].extend(sheet_result["errors"])
            result["warnings"].extend(sheet_result["warnings"])

            info = sheet_map.get(sheet_name)
            if not info:
                continue
            for row in range(info["header_row"] + 1, ws.max_row + 1):
                protocol = _normalize_cell(ws.cell(row=row, column=info["columns"]["Protocolo"]).value)
                if not protocol:
                    continue
                protocol_locations.setdefault(protocol, []).append(f"{sheet_name}!{row}")
                entry_dt = parse_date(ws.cell(row=row, column=info["columns"]["Data de ingresso"]).value)
                expected_sheet = get_target_sheet_from_entry_date_or_protocol(
                    entry_dt,
                    protocol,
                )
                if expected_sheet and expected_sheet != sheet_name:
                    pending = {
                        "protocol": protocol,
                        "current_sheet": sheet_name,
                        "expected_sheet": expected_sheet,
                        "row": row,
                        "sheet": sheet_name,
                        "cell": f"{sheet_name}!B{row}",
                        "type": "wrong_year_sheet",
                        "description": (
                            f"Protocolo {protocol} esta na aba {sheet_name}, "
                            f"mas a aba esperada e {expected_sheet}."
                        ),
                        "suggested_action": "Revisar manualmente antes de mover registro historico.",
                    }
                    result["wrong_year_rows"].append(pending)

            if sheet_name == "2026":
                order_pending = _validate_chronological_order(ws, info)
                result["data_pending"].extend(order_pending)
                sheet_result["data_pending"].extend(order_pending)

        result["duplicate_protocols"] = [
            {"protocol": protocol, "locations": locations}
            for protocol, locations in protocol_locations.items()
            if len(locations) > 1
        ]
        for duplicate in result["duplicate_protocols"]:
            result["data_pending"].append(
                {
                    "type": "duplicate_protocol",
                    "protocol": duplicate["protocol"],
                    "cell": ", ".join(duplicate["locations"]),
                    "description": (
                        f"Protocolo duplicado {duplicate['protocol']}: "
                        + ", ".join(duplicate["locations"])
                    ),
                    "suggested_action": "Revisar duplicidade manualmente.",
                }
            )
        for wrong in result["wrong_year_rows"]:
            result["data_pending"].append(wrong)

        result["errors"] = (
            result["technical_errors"]
            + result["visual_errors"]
            + [item["description"] for item in result["data_pending"] if item.get("description")]
        )
        result["success"] = not (
            result["technical_errors"]
            or result["visual_errors"]
            or result["data_pending"]
        )
        if strict:
            result["success"] = not result["errors"]
        return result
    except Exception as exc:
        logger.exception(f"Falha ao validar formatacao da planilha: {exc}")
        _add_validation_issue(result, "technical_errors", str(exc))
        return result
    finally:
        if close_workbook and wb is not None:
            wb.close()


def create_year_sheet_from_template(
    wb: Workbook, template_ws: Worksheet, target_sheet: str
) -> Worksheet:
    new_ws = wb.create_sheet(target_sheet)
    _apply_safe_workbook_sheet_layout(new_ws)
    standardize_sheet_layout(new_ws, template_ws)
    return new_ws


def standardize_sheet_layout(ws: Worksheet, template_ws: Worksheet) -> dict:
    report = {
        "sheet": ws.title,
        "formatted_rows": 0,
        "value_errors_removed": 0,
        "logo_copied": False,
        "official_file_replaced": False,
        "warnings": [],
    }
    _remove_unsafe_sheet_objects(ws)
    report["value_errors_removed"] = _clear_value_errors(ws)
    _apply_safe_workbook_sheet_layout(ws)
    report["value_errors_removed"] += _clear_value_errors(ws)
    report["warnings"].append("Logomarca em imagem nao foi copiada por seguranca; cabecalho textual aplicado.")
    header_row, columns = _find_required_columns_with_header(ws)
    format_metrics = _format_data_rows(ws, header_row, columns)
    report["formatted_rows"] = _count_data_rows(ws, header_row, columns)
    report.update(format_metrics)
    return report


def project_dry_run_target_rows(results: list[dict], workbook_path: Path) -> list[dict]:
    workbook_path = Path(workbook_path)
    if not workbook_path.exists():
        return results

    wb: Workbook | None = None
    try:
        wb = load_workbook(workbook_path, read_only=True)
        existing_sheets = set(wb.sheetnames)
    except Exception as exc:
        logger.warning(
            f"Não foi possível projetar linhas de destino no dry-run: {exc}"
        )
        return results
    finally:
        if wb is not None:
            wb.close()

    grouped: dict[str, list[dict]] = {}
    for item in results:
        status = _status_payload(item)
        action = _payload_value(item, status, "action")
        target_sheet = _payload_value(item, status, "target_sheet")
        if (
            action in {"insert_new_chronological", "move_wrong_sheet_to_correct_sheet"}
            and target_sheet
            and target_sheet not in existing_sheets
        ):
            grouped.setdefault(str(target_sheet), []).append(item)

    for target_sheet, items in grouped.items():
        base_row = min(
            (
                int(row)
                for row in (
                    _payload_value(item, _status_payload(item), "target_row")
                    or _payload_value(item, _status_payload(item), "row_number")
                    for item in items
                )
                if row
            ),
            default=2,
        )
        ordered = sorted(
            enumerate(items),
            key=lambda pair: (
                parse_date(pair[1].get("entry_date")) or date.max,
                pair[0],
            ),
        )
        for offset, (_, item) in enumerate(ordered):
            target_row = base_row + offset
            status = _status_payload(item)
            _set_payload_value(item, status, "target_row", target_row)
            _set_payload_value(item, status, "row_number", target_row)
            _set_payload_value(item, status, "new_row", target_row)
            source_sheet = _payload_value(item, status, "source_sheet")
            source_row = _payload_value(item, status, "source_row")
            if source_sheet and source_row:
                _set_payload_value(
                    item, status, "moved_to", f"{target_sheet}!{target_row}"
                )
                _set_payload_value(
                    item,
                    status,
                    "warning",
                    _movement_message(
                        str(source_sheet),
                        int(source_row),
                        target_sheet,
                        target_row,
                        True,
                        f"A aba {target_sheet} seria criada a partir da aba "
                        f"{get_settings().EXCEL_DEFAULT_TEMPLATE_SHEET}",
                    ),
                )

    return results


def update_row_by_protocol(
    workbook_path: Path,
    protocol: str,
    solicitation: PortalSolicitation | None = None,
    generation_data: GenerationData | None = None,
    conclusao: str | None = None,
    parecer: str | None = None,
    save: bool = True,
) -> int:
    workbook_path = Path(workbook_path)
    wb = load_workbook_safe(workbook_path)
    ws = wb.active
    header_row, columns = _find_required_columns_with_header(ws)
    protocol_col = columns["Protocolo"]

    target_row = _find_row_by_protocol(ws, header_row, protocol_col, protocol)
    if target_row is None:
        raise ValueError(f"Protocolo não encontrado na planilha: {protocol}")

    if solicitation:
        _set_cell_if_value(ws, target_row, columns["Cliente"], solicitation.client_name)
        _set_cell_if_value(ws, target_row, columns["Data de ingresso"], solicitation.entry_date)

    _set_cell_if_value(
        ws,
        target_row,
        _column_by_normalized_name(columns, "CONCLUSAO"),
        conclusao,
    )
    _set_cell_if_value(ws, target_row, columns["Parecer"], parecer)

    if generation_data:
        module_text = generation_data.format_module_for_planilha()
        inverter_text = generation_data.format_inverter_for_planilha()
        _set_cell_if_value(ws, target_row, columns["Placa"], module_text)
        _set_cell_if_value(ws, target_row, columns["Inversor"], inverter_text)
        _adjust_equipment_row_layout(ws, target_row, columns, module_text, inverter_text)

    if save:
        _save_workbook_atomically(wb, workbook_path)
        logger.info(f"Planilha atualizada na linha {target_row}: {workbook_path}")

    return target_row


def _set_workbook_operational_error(result: dict, exc: PermissionError) -> None:
    classified = classify_workbook_write_error(exc)
    result["error"] = classified.user_message
    result["code"] = classified.code.value
    result["technical_cause"] = classified.technical_cause


def _apply_safe_workbook_sheet_layout(ws: Worksheet) -> None:
    _remove_conflicting_merges(ws, min_row=1, max_row=1, min_col=1, max_col=7)
    for row in range(1, min(ws.max_row, TOP_LAYOUT_SCAN_ROWS) + 1):
        for column in range(1, 8):
            cell = ws.cell(row=row, column=column)
            text = str(cell.value or "").strip().upper()
            if text in {"#VALOR!", "#VALUE!"} or str(cell.value or "").startswith("="):
                cell.value = None

    ws.merge_cells("A1:G1")
    logo_cell = ws["A1"]
    logo_cell.value = SAFE_LOGO_TEXT
    logo_cell.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    logo_cell.fill = PatternFill("solid", fgColor="F58220")
    logo_cell.alignment = Alignment(horizontal="center", vertical="center")
    logo_cell.border = _thin_border()
    ws.row_dimensions[1].height = 28

    for index, header in enumerate(REQUIRED_COLUMNS, start=1):
        cell = ws.cell(row=2, column=index)
        cell.value = header
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="595959")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()
    ws.row_dimensions[2].height = 24

    for column_letter, width in SAFE_COLUMN_WIDTHS.items():
        ws.column_dimensions[column_letter].width = width

    ws.auto_filter.ref = "A2:G2"
    ws.freeze_panes = "A3"


def _remove_conflicting_merges(
    ws: Worksheet,
    *,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        if (
            merged_range.min_row <= max_row
            and merged_range.max_row >= min_row
            and merged_range.min_col <= max_col
            and merged_range.max_col >= min_col
        ):
            ws.unmerge_cells(str(merged_range))


def _remove_unsafe_sheet_objects(ws: Worksheet) -> None:
    if hasattr(ws, "_images"):
        ws._images = []
    if hasattr(ws, "_charts"):
        ws._charts = []


def _thin_border() -> Border:
    side = Side(style="thin", color="D9D9D9")
    return Border(left=side, right=side, top=side, bottom=side)


def _apply_safe_data_row_style(ws: Worksheet, row: int, columns: dict[str, int]) -> None:
    ws.row_dimensions[row].height = 28
    border = _thin_border()
    for column in range(1, 8):
        cell = ws.cell(row=row, column=column)
        cell.font = Font(name="Arial", size=9, color="000000")
        cell.fill = PatternFill("solid", fgColor="FFFFFF")
        cell.border = border
        horizontal = "left" if column == columns["Cliente"] else "center"
        cell.alignment = Alignment(
            horizontal=horizontal,
            vertical="center",
            wrap_text=True,
        )


def _clear_value_errors(ws: Worksheet) -> int:
    removed = 0
    for row in range(1, min(ws.max_row, TOP_LAYOUT_SCAN_ROWS) + 1):
        for column in range(1, min(ws.max_column, 7) + 1):
            cell = ws.cell(row=row, column=column)
            value_text = str(cell.value or "").strip()
            if value_text.upper() in {"#VALOR!", "#VALUE!"} or value_text.startswith("="):
                cell.value = None
                removed += 1
    return removed


def _format_data_rows(ws: Worksheet, header_row: int, columns: dict[str, int]) -> dict:
    metrics = {"date_cells_formatted": 0, "protocol_cells_as_text": 0}
    for row in range(header_row + 1, ws.max_row + 1):
        if _is_empty_data_row(ws, row, columns):
            continue
        row_metrics = _format_data_row(ws, row, columns)
        metrics["date_cells_formatted"] += row_metrics["date_cells_formatted"]
        metrics["protocol_cells_as_text"] += row_metrics["protocol_cells_as_text"]
    return metrics


def _format_data_row(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    *,
    clean_equipment: bool = False,
) -> dict:
    metrics = {"date_cells_formatted": 0, "protocol_cells_as_text": 0}
    _apply_safe_data_row_style(ws, row, columns)

    protocol_cell = ws.cell(row=row, column=columns["Protocolo"])
    if protocol_cell.value not in (None, ""):
        protocol_cell.value = str(protocol_cell.value).strip()
    protocol_cell.number_format = TEXT_NUMBER_FORMAT
    metrics["protocol_cells_as_text"] += 1

    _normalize_date_cell(ws.cell(row=row, column=columns["Data de ingresso"]))
    _normalize_date_cell(
        ws.cell(row=row, column=_column_by_normalized_name(columns, "CONCLUSAO"))
    )
    metrics["date_cells_formatted"] += 2
    if clean_equipment:
        ws.cell(row=row, column=columns["Placa"]).value = _clean_equipment_text(
            ws.cell(row=row, column=columns["Placa"]).value
        )
        ws.cell(row=row, column=columns["Inversor"]).value = _clean_equipment_text(
            ws.cell(row=row, column=columns["Inversor"]).value
        )
    return metrics


def _normalize_date_cell(cell) -> None:
    if cell.value in (None, ""):
        cell.number_format = DATE_NUMBER_FORMAT
        return
    parsed = parse_date(cell.value)
    if parsed:
        cell.value = date_to_excel_datetime(parsed)
    cell.number_format = DATE_NUMBER_FORMAT


def _data_style_source_row(
    ws: Worksheet, target_row: int, columns: dict[str, int]
) -> int | None:
    for row in range(target_row + 1, ws.max_row + 1):
        if not _is_empty_data_row(ws, row, columns):
            return row
    for row in range(target_row - 1, 0, -1):
        if not _row_is_completely_empty(ws, row):
            return row
    return None


def _copy_row_style(ws: Worksheet, source_row: int, target_row: int) -> None:
    for column in range(1, ws.max_column + 1):
        _copy_cell_style(
            ws.cell(row=source_row, column=column),
            ws.cell(row=target_row, column=column),
        )
    if source_row in ws.row_dimensions:
        ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy.copy(source._style)
    target.number_format = source.number_format
    target.alignment = copy.copy(source.alignment)
    target.border = copy.copy(source.border)
    target.fill = copy.copy(source.fill)
    target.font = copy.copy(source.font)
    target.protection = copy.copy(source.protection)


def _count_data_rows(ws: Worksheet, header_row: int, columns: dict[str, int]) -> int:
    return sum(
        1
        for row in range(header_row + 1, ws.max_row + 1)
        if not _is_empty_data_row(ws, row, columns)
    )


def _clean_equipment_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if "\n" in text:
        return "\n".join(
            re.sub(r"\s+", " ", line).strip()
            for line in text.splitlines()
            if line.strip()
        )
    text = text.split("|", 1)[0]
    text = re.sub(r"\b(?:kWp|kW)\b", "", text, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", text).strip()


def _adjust_equipment_row_layout(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    module_text: str | None,
    inverter_text: str | None,
) -> None:
    line_count = max(
        len(str(module_text or "").splitlines()),
        len(str(inverter_text or "").splitlines()),
    )
    if line_count <= 1:
        return
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 28, 15 * line_count)
    for column_name in ["Placa", "Inversor"]:
        cell = ws.cell(row=row, column=columns[column_name])
        cell.alignment = Alignment(
            horizontal=cell.alignment.horizontal or "center",
            vertical="top",
            wrap_text=True,
        )


def _apply_equipment_cell_layout(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    module_text: str,
    inverter_text: str,
) -> None:
    line_count = max(
        1,
        len(str(module_text or "").splitlines()),
        len(str(inverter_text or "").splitlines()),
    )
    ws.row_dimensions[row].height = max(ws.row_dimensions[row].height or 18, 15 * line_count)
    for column_name in ["Placa", "Inversor"]:
        cell = ws.cell(row=row, column=columns[column_name])
        cell.alignment = Alignment(
            horizontal=cell.alignment.horizontal or "center",
            vertical="top",
            wrap_text=True,
        )


def _validate_sheet_format(ws: Worksheet, info: dict | None) -> dict:
    report = {
        "sheet": ws.title,
        "errors": [],
        "visual_errors": [],
        "technical_errors": [],
        "data_pending": [],
        "warnings": [],
    }
    for row in range(1, min(ws.max_row, TOP_LAYOUT_SCAN_ROWS) + 1):
        for column in range(1, min(ws.max_column, 7) + 1):
            value = str(ws.cell(row=row, column=column).value or "").strip().upper()
            if value in {"#VALOR!", "#VALUE!"}:
                _add_validation_issue(
                    report,
                    "visual_errors",
                    f"{ws.title}!{row}:{column} contem {value}.",
                )

    if not info:
        _add_validation_issue(
            report,
            "technical_errors",
            f"Aba {ws.title} nao contem cabecalhos obrigatorios.",
        )
        return report

    columns = info["columns"]
    header_row = info["header_row"]
    for column_name in REQUIRED_COLUMNS:
        value = ws.cell(
            row=header_row,
            column=_column_by_normalized_name(columns, _normalize_header(column_name)),
        ).value
        if _normalize_header(value) != _normalize_header(column_name):
            _add_validation_issue(
                report,
                "visual_errors",
                f"Cabecalho invalido em {ws.title}: {column_name}.",
            )

    for row in range(header_row + 1, ws.max_row + 1):
        if _is_empty_data_row(ws, row, columns):
            continue
        protocol_cell = ws.cell(row=row, column=columns["Protocolo"])
        if protocol_cell.number_format != TEXT_NUMBER_FORMAT:
            report["warnings"].append(f"{ws.title}!B{row} protocolo nao esta como texto.")

        for normalized_column in ["DATA DE INGRESSO", "CONCLUSAO"]:
            cell = ws.cell(
                row=row,
                column=_column_by_normalized_name(columns, normalized_column),
            )
            if cell.value in (None, ""):
                continue
            if normalized_column == "CONCLUSAO" and _is_open_completion_value(cell.value):
                continue
            if not isinstance(cell.value, (date, datetime)):
                _add_data_pending(
                    report,
                    ws.title,
                    f"{ws.title}!{cell.coordinate}",
                    "invalid_date",
                    f"{ws.title}!{cell.coordinate} nao contem data real.",
                    "Converter manualmente se for data historica valida.",
                )
                continue
            if parse_date(cell.value) is None:
                _add_data_pending(
                    report,
                    ws.title,
                    f"{ws.title}!{cell.coordinate}",
                    "invalid_date",
                    f"{ws.title}!{cell.coordinate} nao contem data real.",
                    "Converter manualmente se for data historica valida.",
                )
            if cell.number_format.lower() != DATE_NUMBER_FORMAT:
                report["warnings"].append(
                    f"{ws.title}!{cell.coordinate} formato de data nao e {DATE_NUMBER_FORMAT}."
                )

        for column_name in ["Placa", "Inversor"]:
            value = str(ws.cell(row=row, column=columns[column_name]).value or "")
            if "\n" not in value and ("|" in value or re.search(
                r"\bkWp\b|\bkW\b|m[oó]dulos?|inversor",
                value,
                flags=re.IGNORECASE,
            )):
                _add_data_pending(
                    report,
                    ws.title,
                    f"{ws.title}!{row}",
                    "legacy_equipment_text",
                    f"{ws.title}!{row} {column_name} contem unidade ou separador indevido.",
                    "Revisar texto legado; novos registros ja sao normalizados.",
                )

    return report


def _validate_chronological_order(ws: Worksheet, info: dict) -> list[dict[str, Any]]:
    pending: list[dict] = []
    previous_date: date | None = None
    date_col = info["columns"]["Data de ingresso"]
    for row in range(info["header_row"] + 1, ws.max_row + 1):
        if _is_empty_data_row(ws, row, info["columns"]):
            continue
        current_date = parse_date(ws.cell(row=row, column=date_col).value)
        if current_date is None:
            continue
        if previous_date and current_date < previous_date:
            pending.append(
                {
                    "sheet": ws.title,
                    "cell": f"{ws.title}!{row}",
                    "type": "chronological_order",
                    "description": f"Aba 2026 fora de ordem cronologica na linha {row}.",
                    "suggested_action": "Revisar ordem cronologica manualmente.",
                }
            )
        previous_date = current_date
    return pending


def _add_validation_issue(report: dict, bucket: str, message: str) -> None:
    report.setdefault(bucket, []).append(message)
    report.setdefault("errors", []).append(message)


def _add_data_pending(
    report: dict,
    sheet: str,
    cell: str,
    issue_type: str,
    description: str,
    suggested_action: str,
) -> None:
    report.setdefault("data_pending", []).append(
        {
            "sheet": sheet,
            "cell": cell,
            "type": issue_type,
            "description": description,
            "suggested_action": suggested_action,
        }
    )


def get_target_sheet_from_entry_date_or_protocol(
    entry_date: Any = None,
    protocol: str | None = None,
) -> str | None:
    parsed = parse_date(entry_date)
    if parsed:
        year = parsed.year
        if 0 <= year < 100:
            year += 2000
        return _sheet_name_for_year(year)

    protocol_text = re.sub(r"\D+", "", str(protocol or ""))
    if len(protocol_text) >= 2:
        prefix = protocol_text[:2]
        if prefix in {"22", "23"}:
            return "2022 - 2023"
        if prefix in {"24", "25", "26"}:
            return f"20{prefix}"
    return None


def _get_or_prepare_target_sheet(
    wb: Workbook, sheet_map: dict[str, dict], target_sheet: str, dry_run: bool
) -> dict:
    settings = get_settings()
    if target_sheet in sheet_map:
        return {"success": True, **sheet_map[target_sheet]}

    warning = (
        f"A aba {target_sheet} seria criada a partir da aba "
        f"{settings.EXCEL_DEFAULT_TEMPLATE_SHEET}"
    )
    if dry_run:
        template = sheet_map.get(settings.EXCEL_DEFAULT_TEMPLATE_SHEET)
        if not template:
            return {
                "success": False,
                "action": "blocked_missing_year_sheet",
                "error": (
                    f"Aba {target_sheet} não existe e a aba modelo "
                    f"{settings.EXCEL_DEFAULT_TEMPLATE_SHEET} não foi encontrada."
                ),
            }
        return {
            "success": True,
            "worksheet": template["worksheet"],
            "header_row": template["header_row"],
            "columns": template["columns"],
            "warning": warning,
            "sheet_would_be_created": True,
        }

    if not settings.CREATE_YEAR_SHEET_IF_MISSING:
        return {
            "success": False,
            "action": "blocked_missing_year_sheet",
            "error": f"Aba {target_sheet} não existe e criação automática está desativada.",
        }

    template_ws = wb[settings.EXCEL_DEFAULT_TEMPLATE_SHEET] if settings.EXCEL_DEFAULT_TEMPLATE_SHEET in wb.sheetnames else None
    if template_ws is None:
        return {
            "success": False,
            "action": "blocked_missing_year_sheet",
            "error": (
                f"Aba {target_sheet} não existe e a aba modelo "
                f"{settings.EXCEL_DEFAULT_TEMPLATE_SHEET} não foi encontrada."
            ),
        }

    new_ws = create_year_sheet_from_template(wb, template_ws, target_sheet)
    header_row, columns = _find_required_columns_with_header(new_ws)
    parecer_boolean = _column_uses_boolean(new_ws, columns["Parecer"])
    logger.info(f"Aba {target_sheet} criada a partir da aba {template_ws.title}.")
    sheet_map[target_sheet] = {
        "worksheet": new_ws,
        "header_row": header_row,
        "columns": columns,
        "parecer_boolean": parecer_boolean,
    }
    return {
        "success": True,
        "worksheet": new_ws,
        "header_row": header_row,
        "columns": columns,
        "warning": None,
        "parecer_boolean": parecer_boolean,
    }


def _find_required_columns_with_header(ws: Worksheet) -> tuple[int, dict[str, int]]:
    expected = _required_column_aliases()

    for row in range(1, min(ws.max_row, 10) + 1):
        found: dict[str, int] = {}
        for cell in ws[row]:
            normalized = _normalize_header(cell.value)
            if normalized in expected:
                found[expected[normalized]] = cell.column

        if all(column in found for column in REQUIRED_COLUMNS):
            logger.debug(f"Cabeçalho encontrado na linha {row}: {found}")
            return row, found

    missing = ", ".join(REQUIRED_COLUMNS)
    raise ValueError(f"Não foi possível localizar todas as colunas exigidas: {missing}")


def _required_column_aliases() -> dict[str, str]:
    aliases = {_normalize_header(name): name for name in REQUIRED_COLUMNS}
    aliases["CONCLUS O"] = "Conclusão"
    return aliases


def _find_workbook_sheet_with_required_columns(
    wb: Workbook,
) -> tuple[Worksheet, int, dict[str, int]]:
    sheet_map = _workbook_required_sheets(wb)
    if not sheet_map:
        raise ValueError(
            "Nenhuma aba da planilha contém todas as colunas obrigatórias: "
            + ", ".join(REQUIRED_COLUMNS)
        )
    first = next(iter(sheet_map.values()))
    logger.info(f"Aba da base localizada: '{first['worksheet'].title}'.")
    return first["worksheet"], first["header_row"], first["columns"]


def _workbook_required_sheets(wb: Workbook) -> dict[str, dict]:
    sheet_map: dict[str, dict] = {}
    for ws in wb.worksheets:
        try:
            header_row, columns = _find_required_columns_with_header(ws)
            sheet_map[ws.title] = {
                "worksheet": ws,
                "header_row": header_row,
                "columns": columns,
                "parecer_boolean": _column_uses_boolean(ws, columns["Parecer"]),
            }
        except ValueError:
            continue
    return sheet_map


def _find_protocol_occurrences(sheet_map: dict[str, dict], protocol: str) -> list[dict]:
    occurrences: list[dict] = []
    for sheet_name, info in sheet_map.items():
        row_number = _find_row_by_protocol(
            info["worksheet"],
            info["header_row"],
            info["columns"]["Protocolo"],
            protocol,
        )
        if row_number is not None:
            occurrences.append({"sheet": sheet_name, "row": row_number})
    return occurrences


def _find_row_by_protocol(
    ws: Worksheet, header_row: int, protocol_col: int, protocol: str
) -> int | None:
    wanted = _normalize_cell(protocol)
    for row in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=protocol_col).value
        if _normalize_cell(value) == wanted:
            return row
    return None


def _find_chronological_insert_row(
    ws: Worksheet, header_row: int, columns: dict[str, int], entry_dt: date
) -> int:
    date_col = columns["Data de ingresso"]
    last_data_row = header_row
    insert_row = ws.max_row + 1

    for row in range(header_row + 1, ws.max_row + 1):
        row_date = parse_date(ws.cell(row=row, column=date_col).value)
        if _is_empty_data_row(ws, row, columns):
            continue
        last_data_row = row
        if row_date and row_date > entry_dt:
            return row
        if row_date and row_date <= entry_dt:
            insert_row = row + 1

    if last_data_row == header_row:
        return header_row + 1
    return max(insert_row, last_data_row + 1)


def _is_empty_data_row(ws: Worksheet, row: int, columns: dict[str, int]) -> bool:
    return all(ws.cell(row=row, column=col).value in (None, "") for col in columns.values())


def _insert_row_preserving_style(ws: Worksheet, row: int) -> None:
    if row <= ws.max_row and not _row_is_completely_empty(ws, row):
        ws.insert_rows(row)
        source_row = row + 1 if row + 1 <= ws.max_row else max(row - 1, 1)
    elif row <= ws.max_row:
        source_row = row
    else:
        source_row = max(row - 1, 1)

    _copy_row_style(ws, source_row, row)


def _row_is_completely_empty(ws: Worksheet, row: int) -> bool:
    return all(ws.cell(row=row, column=col).value in (None, "") for col in range(1, ws.max_column + 1))


def _write_excel_row(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    protocol: str,
    client_name: str,
    entry_dt: date,
    completion_date: Any,
    module_text: str | None,
    inverter_text: str | None,
    parecer: str,
    parecer_boolean: bool | None = None,
) -> None:
    _format_data_row(ws, row, columns, clean_equipment=True)
    _set_cell_if_value(ws, row, columns["Cliente"], client_name)
    protocol_cell = ws.cell(row=row, column=columns["Protocolo"])
    protocol_cell.value = str(protocol)
    protocol_cell.number_format = TEXT_NUMBER_FORMAT
    entry_cell = ws.cell(row=row, column=columns["Data de ingresso"])
    entry_cell.value = date_to_excel_datetime(entry_dt)
    entry_cell.number_format = DATE_NUMBER_FORMAT
    completion_col = _column_by_normalized_name(columns, "CONCLUSAO")
    completion_cell = ws.cell(row=row, column=completion_col)
    if _is_open_completion_value(completion_date):
        completion_cell.value = OPEN_COMPLETION_TEXT
        completion_cell.number_format = TEXT_NUMBER_FORMAT
    else:
        completion_dt = date_to_excel_datetime(completion_date)
        if completion_dt:
            completion_cell.value = completion_dt
            completion_cell.number_format = DATE_NUMBER_FORMAT
    ws.cell(row=row, column=columns["Parecer"]).value = _parecer_value(
        ws, columns["Parecer"], parecer, parecer_boolean
    )
    cleaned_module = _clean_equipment_text(module_text)
    cleaned_inverter = _clean_equipment_text(inverter_text)
    _set_cell_if_value(ws, row, columns["Placa"], cleaned_module)
    _set_cell_if_value(ws, row, columns["Inversor"], cleaned_inverter)
    _adjust_equipment_row_layout(ws, row, columns, cleaned_module, cleaned_inverter)


def _write_existing_excel_row_updates(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    completion_date: Any,
    module_text: str | None,
    inverter_text: str | None,
    parecer: str,
    update_state: dict[str, Any],
) -> None:
    completion_col = _column_by_normalized_name(columns, "CONCLUSAO")
    completion_cell = ws.cell(row=row, column=completion_col)
    if completion_date is not None and not update_state.get("completion_no_change"):
        if _is_open_completion_value(completion_date):
            completion_cell.value = OPEN_COMPLETION_TEXT
            completion_cell.number_format = TEXT_NUMBER_FORMAT
        else:
            completion_dt = date_to_excel_datetime(completion_date)
            if completion_dt:
                completion_cell.value = completion_dt
                completion_cell.number_format = DATE_NUMBER_FORMAT

    module_changed = not update_state.get("module_no_change")
    inverter_changed = not update_state.get("inverter_no_change")
    cleaned_module = _clean_equipment_text(module_text)
    cleaned_inverter = _clean_equipment_text(inverter_text)
    if module_changed:
        _set_cell_if_value(ws, row, columns["Placa"], cleaned_module)
    if inverter_changed:
        _set_cell_if_value(ws, row, columns["Inversor"], cleaned_inverter)
    if module_changed or inverter_changed:
        _adjust_equipment_row_layout(ws, row, columns, cleaned_module, cleaned_inverter)
    if not update_state.get("parecer_no_change"):
        ws.cell(row=row, column=columns["Parecer"]).value = _parecer_value(
            ws,
            columns["Parecer"],
            parecer,
            _column_uses_boolean(ws, columns["Parecer"]),
        )


def _row_has_required_update_values(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    entry_dt: date | None,
    completion_date: Any = None,
    module_text: str | None = None,
    inverter_text: str | None = None,
) -> bool:
    return _row_required_update_state(
        ws,
        row,
        columns,
        entry_dt,
        completion_date,
        module_text,
        inverter_text,
    )["all_no_change"]


def _row_required_update_state(
    ws: Worksheet,
    row: int,
    columns: dict[str, int],
    entry_dt: date | None,
    completion_date: Any = None,
    module_text: str | None = None,
    inverter_text: str | None = None,
    parecer: str | None = None,
) -> dict[str, Any]:
    row_entry_dt = parse_date(ws.cell(row=row, column=columns["Data de ingresso"]).value)
    ingress_no_change = bool(
        row_entry_dt is not None and (entry_dt is None or row_entry_dt == entry_dt)
    )

    current_module = ws.cell(row=row, column=columns["Placa"]).value
    current_inverter = ws.cell(row=row, column=columns["Inversor"]).value
    module_no_change = _equipment_text_matches(current_module, module_text)
    inverter_no_change = _equipment_text_matches(current_inverter, inverter_text)
    equipment_no_change = module_no_change and inverter_no_change
    current_parecer = ws.cell(row=row, column=columns["Parecer"]).value
    expected_parecer = (
        None
        if parecer is None
        else _parecer_value(
            ws,
            columns["Parecer"],
            parecer,
            _column_uses_boolean(ws, columns["Parecer"]),
        )
    )
    parecer_no_change = (
        True if expected_parecer is None else current_parecer == expected_parecer
    )

    completion_col = _column_by_normalized_name(columns, "CONCLUSAO")
    current_completion = ws.cell(row=row, column=completion_col).value
    completion_no_change = _completion_value_matches(current_completion, completion_date)
    completion_action = None
    if completion_date is not None:
        if completion_no_change:
            completion_action = "NO_CHANGE"
        elif _is_open_completion_value(completion_date):
            completion_action = "MARKED_AS_OPEN"
        else:
            completion_action = "COMPLETION_DATE_UPDATED"

    return {
        "ingress_no_change": ingress_no_change,
        "module_no_change": module_no_change,
        "inverter_no_change": inverter_no_change,
        "equipment_no_change": equipment_no_change,
        "parecer_no_change": parecer_no_change,
        "completion_no_change": completion_no_change,
        "completion_action": completion_action,
        "all_no_change": (
            ingress_no_change
            and equipment_no_change
            and parecer_no_change
            and completion_no_change
        ),
    }


def _copy_row_values(
    source_ws: Worksheet, source_row: int, target_ws: Worksheet, target_row: int
) -> None:
    max_column = max(source_ws.max_column, target_ws.max_column)
    for column in range(1, max_column + 1):
        target_ws.cell(row=target_row, column=column).value = source_ws.cell(
            row=source_row, column=column
        ).value


def _validate_protocol_at_row(
    ws: Worksheet, row: int, protocol_col: int, protocol: str
) -> None:
    value = ws.cell(row=row, column=protocol_col).value
    if _normalize_cell(value) != _normalize_cell(protocol):
        raise RuntimeError(
            "Falha ao validar a nova linha antes de remover a antiga: "
            f"esperado protocolo {protocol}, encontrado {value!r}."
        )


def _movement_message(
    source_sheet: str,
    source_row: int,
    target_sheet: str,
    target_row: int,
    dry_run: bool,
    sheet_warning: str | None = None,
) -> str:
    action = "Moveria" if dry_run else "Linha movida"
    message = (
        f"{action} da aba {source_sheet}, linha {source_row}, "
        f"para a aba {target_sheet}, linha {target_row}"
    )
    if sheet_warning:
        message = f"{message}. {sheet_warning}"
    return message


def _parecer_value(
    ws: Worksheet, parecer_col: int, parecer: str, parecer_boolean: bool | None = None
) -> bool | str:
    if parecer_boolean is True:
        return _parecer_to_bool(parecer)
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=parecer_col).value
        if isinstance(value, bool):
            return _parecer_to_bool(parecer)
    return parecer


def _parecer_to_bool(parecer: str | bool | None) -> bool:
    if isinstance(parecer, bool):
        return parecer
    normalized = _normalize_cell(parecer).upper()
    return normalized not in {"FALSE", "FALSO", "NAO", "NÃO", "NO", "0"}


def _column_uses_boolean(ws: Worksheet, column: int) -> bool:
    for row in range(1, ws.max_row + 1):
        if isinstance(ws.cell(row=row, column=column).value, bool):
            return True
    return False


def _clear_sheet_data_rows(ws: Worksheet, header_row: int) -> None:
    if ws.max_row > header_row + 1:
        ws.delete_rows(header_row + 2, ws.max_row - header_row - 1)
    for row in range(header_row + 1, ws.max_row + 1):
        for column in range(1, ws.max_column + 1):
            ws.cell(row=row, column=column).value = None


def _normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    if "Ã" in text or "Â" in text:
        try:
            text = text.encode("latin1").decode("utf-8")
        except UnicodeError:
            pass
    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = text.upper()
    text = re.sub(r"[^A-Z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _column_by_normalized_name(columns: dict[str, int], normalized_name: str) -> int:
    wanted = re.sub(r"\s+", "", normalized_name)
    for name, column in columns.items():
        candidate = re.sub(r"\s+", "", _normalize_header(name))
        if candidate == wanted:
            return column
        if wanted == "CONCLUSAO" and candidate.startswith("CONCLUSA"):
            return column
    raise KeyError(normalized_name)


def _set_cell_if_value(ws: Worksheet, row: int, column: int, value: str | None) -> None:
    if value is not None:
        ws.cell(row=row, column=column).value = value


def _ensure_workbook_writable(workbook_path: Path) -> None:
    try:
        with workbook_path.open("r+b"):
            return
    except PermissionError:
        raise


def _backup_workbook(workbook_path: Path) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = workbook_path.with_name(
        f"{workbook_path.stem}_backup_{timestamp}{workbook_path.suffix}"
    )
    version = 2
    while backup_path.exists():
        backup_path = workbook_path.with_name(
            f"{workbook_path.stem}_backup_{timestamp}_v{version}{workbook_path.suffix}"
        )
        version += 1

    atomic_copy_file(workbook_path, backup_path, private=True)
    logger.info(f"Backup da planilha criado: {backup_path}")
    return backup_path
