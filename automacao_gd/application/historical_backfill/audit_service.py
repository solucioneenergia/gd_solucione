from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Callable
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from automacao_gd.application.historical_backfill.comparison import (
    comparison_reasons,
    contains_sensitive_or_path_text,
    equipment_texts_equal,
    normalize_protocol,
    row_fingerprint,
)
from automacao_gd.domain.backfill_models import (
    BackfillAction,
    HistoricalAuditResult,
    HistoricalAuditSummary,
    HistoricalEquipmentAuditItem,
)
from automacao_gd.domain.equipment_semantics import (
    CanonicalEquipmentCollection,
    EquipmentSourceType,
    SemanticComparison,
    canonical_collection_from_formatted_cells,
    canonical_collection_from_generation_data,
    canonical_collection_to_dict,
    compare_equipment_row_transition,
    compare_canonical_collections,
    format_canonical_collection,
    row_text_cleanup_warnings,
    semantic_compare_equipment_cells,
)
from automacao_gd.domain.equipment_validation import validate_canonical_equipment
from automacao_gd.infrastructure.pdf.service import (
    extract_generation_data,
    extract_pdf_text,
    extract_protocol_from_pdf_text,
    validate_pdf_input,
)


TechnicalResolver = Callable[[str], "TechnicalProposal"]

_MANDATORY_TEXTUAL_CLEANUP_REASONS = frozenset(
    {
        "DUPLICATED_MANUFACTURER_IN_MODEL_REMOVED",
        "SOLPLANET_ALIAS_NORMALIZED",
        "PROVEN_STRUCTURAL_CONTAMINATION_REMOVED",
        "CROSS_FIELD_CONTAMINATION_RESOLVED",
        "GENERIC_LABEL_REMOVED",
    }
)


class HistoricalAuditConsistencyError(RuntimeError):
    """A origem mudou enquanto a fotografia somente leitura era construída."""


@dataclass(frozen=True, slots=True)
class TechnicalProposal:
    status: str
    module_text: str | None = None
    inverter_text: str | None = None
    source_hash: str | None = None
    pdf_status: str = "not_found"
    errors: tuple[str, ...] = ()
    warnings: tuple[str, ...] = ()
    canonical_collection: CanonicalEquipmentCollection | None = None


@dataclass(frozen=True, slots=True)
class _WorkbookRow:
    sheet: str
    row: int
    protocol: str | None
    module_text: str
    inverter_text: str


def audit_historical_workbook(
    workbook_path: Path,
    *,
    technical_resolver: TechnicalResolver | None = None,
    downloads_root: Path | None = None,
    state_path: Path | None = None,
    created_at: str | None = None,
) -> HistoricalAuditResult:
    """Audit operational year sheets without saving or mutating the workbook."""
    path = Path(workbook_path)
    workbook_hash = file_sha256(path)
    resolver = technical_resolver or _default_resolver(downloads_root, state_path)
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        rows, sheets_analyzed, invalid_operational_sheets = _read_operational_rows(wb)
    finally:
        wb.close()

    protocol_locations: dict[str, list[_WorkbookRow]] = defaultdict(list)
    for row in rows:
        if row.protocol:
            protocol_locations[row.protocol].append(row)

    items: list[HistoricalEquipmentAuditItem] = []
    for row in rows:
        fingerprint = row_fingerprint(
            row.sheet, row.row, row.protocol, row.module_text, row.inverter_text
        )
        if row.protocol is None:
            items.append(
                _item(
                    row,
                    BackfillAction.PROTOCOL_MISSING,
                    "not_checked",
                    "invalid",
                    ("PROTOCOL_MISSING",),
                    fingerprint,
                )
            )
            continue
        if len(protocol_locations[row.protocol]) > 1:
            locations = tuple(
                f"{entry.sheet}:{entry.row}"
                for entry in protocol_locations[row.protocol]
            )
            items.append(
                _item(
                    row,
                    BackfillAction.DUPLICATE_PROTOCOL,
                    "not_checked",
                    "invalid",
                    ("DUPLICATE_PROTOCOL", *locations),
                    fingerprint,
                )
            )
            continue
        if contains_sensitive_or_path_text(row.module_text) or contains_sensitive_or_path_text(
            row.inverter_text
        ):
            items.append(
                _item(
                    row,
                    BackfillAction.INVALID_WORKBOOK_ROW,
                    "not_checked",
                    "invalid",
                    ("UNSAFE_TECHNICAL_CELL_CONTENT",),
                    fingerprint,
                )
            )
            continue

        proposal = resolver(row.protocol)
        if _technical_proposal_is_unsafe(proposal):
            items.append(
                _item(
                    row,
                    BackfillAction.PENDING_TECHNICAL_REVIEW,
                    "invalid",
                    "pending_review",
                    ("UNSAFE_TECHNICAL_PROPOSAL",),
                    fingerprint,
                )
            )
            continue
        if proposal.status == "pdf_not_found":
            items.append(
                _item(
                    row,
                    BackfillAction.PDF_NOT_FOUND,
                    proposal.pdf_status,
                    "not_checked",
                    ("PDF_NOT_FOUND",),
                    fingerprint,
                )
            )
            continue
        if (
            proposal.status != "approved"
            or not proposal.module_text
            or not proposal.inverter_text
        ):
            items.append(
                _item(
                    row,
                    BackfillAction.PENDING_TECHNICAL_REVIEW,
                    proposal.pdf_status,
                    "pending_review",
                    proposal.errors or ("TECHNICAL_VALIDATION_PENDING",),
                    fingerprint,
                    warnings=proposal.warnings,
                    source_hash=proposal.source_hash,
                )
            )
            continue

        current_collection = canonical_collection_from_formatted_cells(
            row.module_text, row.inverter_text
        )
        proposed_collection = proposal.canonical_collection or (
            canonical_collection_from_formatted_cells(
                proposal.module_text,
                proposal.inverter_text,
                source=EquipmentSourceType.STRUCTURED_CACHE,
            )
        )
        if proposal.canonical_collection is not None:
            semantic = compare_equipment_row_transition(
                row.module_text,
                row.inverter_text,
                proposal.module_text,
                proposal.inverter_text,
                proposal.canonical_collection,
            )
        else:
            semantic = compare_canonical_collections(
                current_collection, proposed_collection
            )
            text_semantic = semantic_compare_equipment_cells(
                row.module_text,
                row.inverter_text,
                proposal.module_text,
                proposal.inverter_text,
            )
            semantic = _merge_semantic_results(semantic, text_semantic)
        exact_text_no_change = equipment_texts_equal(
            row.module_text, proposal.module_text
        ) and equipment_texts_equal(row.inverter_text, proposal.inverter_text)
        if exact_text_no_change and set(semantic.errors).issubset(
            {"CANONICAL_STRUCTURE_INCOMPLETE", "SEMANTIC_REGRESSION"}
        ):
            semantic = SemanticComparison(
                status="approved",
                equivalent=True,
                errors=(),
                warnings=semantic.warnings,
                current_fingerprint=semantic.current_fingerprint,
                proposed_fingerprint=semantic.proposed_fingerprint,
            )
        if semantic.status != "approved":
            semantic_reasons = tuple(
                dict.fromkeys(
                    (
                        *semantic.errors,
                        *comparison_reasons(
                            row.module_text,
                            row.inverter_text,
                            proposal.module_text,
                            proposal.inverter_text,
                        ),
                    )
                )
            )
            items.append(
                _item(
                    row,
                    BackfillAction.PENDING_TECHNICAL_REVIEW,
                    proposal.pdf_status,
                    "pending_review",
                    semantic_reasons,
                    fingerprint,
                    proposed_module=proposal.module_text,
                    proposed_inverter=proposal.inverter_text,
                    warnings=proposal.warnings,
                    source_hash=proposal.source_hash,
                    semantic_status=semantic.status,
                    semantic_errors=semantic.errors,
                    semantic_warnings=semantic.warnings,
                    current_semantic_fingerprint=semantic.current_fingerprint,
                    proposed_semantic_fingerprint=semantic.proposed_fingerprint,
                    current_collection=current_collection,
                    proposed_collection=proposed_collection,
                )
            )
            continue
        cleanup_warnings = (
            row_text_cleanup_warnings(
                row.module_text,
                row.inverter_text,
                proposal.canonical_collection,
            )
            if proposal.canonical_collection is not None
            else ()
        )
        comparison_reasons_ = comparison_reasons(
            row.module_text,
            row.inverter_text,
            proposal.module_text,
            proposal.inverter_text,
        )
        mandatory_cleanup_reasons = _mandatory_textual_cleanup_reasons(
            exact_text_no_change=exact_text_no_change,
            comparison_reasons=comparison_reasons_,
            semantic_warnings=semantic.warnings,
            cleanup_warnings=cleanup_warnings,
            current_module=row.module_text,
            current_inverter=row.inverter_text,
            proposed_module=proposal.module_text,
            proposed_inverter=proposal.inverter_text,
            proposed_collection=proposed_collection,
        )
        no_change = exact_text_no_change or (
            semantic.equivalent and not mandatory_cleanup_reasons
        )
        reasons = () if no_change else comparison_reasons_
        warnings = tuple(
            dict.fromkeys(
                (*proposal.warnings, *cleanup_warnings, *mandatory_cleanup_reasons)
            )
        )
        items.append(
            _item(
                row,
                BackfillAction.NO_CHANGE if no_change else BackfillAction.UPDATE_EQUIPMENT,
                proposal.pdf_status,
                "approved",
                reasons,
                fingerprint,
                proposed_module=proposal.module_text,
                proposed_inverter=proposal.inverter_text,
                warnings=warnings,
                source_hash=proposal.source_hash,
                semantic_status=semantic.status,
                semantic_warnings=tuple(
                    dict.fromkeys((*semantic.warnings, *cleanup_warnings))
                ),
                current_semantic_fingerprint=semantic.current_fingerprint,
                proposed_semantic_fingerprint=semantic.proposed_fingerprint,
                current_collection=current_collection,
                proposed_collection=proposed_collection,
            )
        )

    summary = _build_summary(items, sheets_analyzed, invalid_operational_sheets)
    if file_sha256(path) != workbook_hash:
        raise HistoricalAuditConsistencyError(
            "A planilha foi alterada durante a auditoria; gere uma nova simulação."
        )
    return HistoricalAuditResult(
        workbook_fingerprint=workbook_hash,
        items=tuple(items),
        summary=summary,
        created_at=created_at or datetime.now(timezone.utc).isoformat(),
    )


def _merge_semantic_results(
    canonical: SemanticComparison, textual: SemanticComparison
) -> SemanticComparison:
    errors = tuple(dict.fromkeys((*canonical.errors, *textual.errors)))
    warnings = tuple(dict.fromkeys((*canonical.warnings, *textual.warnings)))
    return SemanticComparison(
        status="pending_review" if errors else "approved",
        equivalent=canonical.equivalent and textual.equivalent,
        errors=errors,
        warnings=warnings,
        current_fingerprint=canonical.current_fingerprint,
        proposed_fingerprint=canonical.proposed_fingerprint,
    )


def _mandatory_textual_cleanup_reasons(
    *,
    exact_text_no_change: bool,
    comparison_reasons: tuple[str, ...],
    semantic_warnings: tuple[str, ...],
    cleanup_warnings: tuple[str, ...],
    current_module: str,
    current_inverter: str,
    proposed_module: str,
    proposed_inverter: str,
    proposed_collection: CanonicalEquipmentCollection,
) -> tuple[str, ...]:
    if exact_text_no_change:
        return ()
    evidence = (*comparison_reasons, *semantic_warnings, *cleanup_warnings)
    if (
        "DUPLICATED_MANUFACTURER_IN_MODEL_REMOVED" in evidence
        and not _manufacturer_occurrence_removed(
            current_module,
            current_inverter,
            proposed_module,
            proposed_inverter,
            proposed_collection,
        )
    ):
        evidence = tuple(
            reason
            for reason in evidence
            if reason != "DUPLICATED_MANUFACTURER_IN_MODEL_REMOVED"
        )
    return tuple(
        dict.fromkeys(
            reason
            for reason in evidence
            if reason in _MANDATORY_TEXTUAL_CLEANUP_REASONS
        )
    )


def _manufacturer_occurrence_removed(
    current_module: str,
    current_inverter: str,
    proposed_module: str,
    proposed_inverter: str,
    proposed_collection: CanonicalEquipmentCollection,
) -> bool:
    current_text = f"{current_module}\n{current_inverter}"
    proposed_text = f"{proposed_module}\n{proposed_inverter}"
    for manufacturer in {
        item.canonical_manufacturer for item in proposed_collection.all_items()
    }:
        if (
            _word_occurrence_count(current_text, manufacturer)
            > _word_occurrence_count(proposed_text, manufacturer)
        ):
            return True
    return False


def _word_occurrence_count(text: str, token: str) -> int:
    words = re.findall(r"[A-Za-z0-9]+", _strip_accents(text).lower())
    token_words = re.findall(r"[A-Za-z0-9]+", _strip_accents(token).lower())
    if not words or not token_words:
        return 0
    size = len(token_words)
    return sum(
        tuple(words[index : index + size]) == tuple(token_words)
        for index in range(0, len(words) - size + 1)
    )


def _strip_accents(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    return "".join(char for char in normalized if not unicodedata.combining(char))


def _read_operational_rows(workbook) -> tuple[list[_WorkbookRow], int, int]:
    rows: list[_WorkbookRow] = []
    sheets_analyzed = 0
    invalid_operational_sheets = 0
    for ws in workbook.worksheets:
        if not _is_operational_sheet(ws.title):
            continue
        header = _find_header(ws)
        if header is None:
            invalid_operational_sheets += 1
            continue
        header_row, columns = header
        sheets_analyzed += 1
        for row_number in range(header_row + 1, ws.max_row + 1):
            protocol_value = ws.cell(row_number, columns["Protocolo"]).value
            module_value = ws.cell(row_number, columns["Placa"]).value
            inverter_value = ws.cell(row_number, columns["Inversor"]).value
            if all(value is None or str(value).strip() == "" for value in (
                protocol_value,
                module_value,
                inverter_value,
            )):
                continue
            rows.append(
                _WorkbookRow(
                    sheet=ws.title,
                    row=row_number,
                    protocol=normalize_protocol(protocol_value),
                    module_text=str(module_value or "").strip(),
                    inverter_text=str(inverter_value or "").strip(),
                )
            )
    return rows, sheets_analyzed, invalid_operational_sheets


def _is_operational_sheet(name: str) -> bool:
    return bool(re.fullmatch(r"\s*\d{4}(?:\s*-\s*\d{4})?\s*", name))


def _find_header(ws) -> tuple[int, dict[str, int]] | None:
    expected = {"protocolo": "Protocolo", "placa": "Placa", "inversor": "Inversor"}
    for row_number in range(1, min(ws.max_row, 50) + 1):
        found: dict[str, list[int]] = defaultdict(list)
        for cell in ws[row_number]:
            key = _header_key(cell.value)
            if key in expected:
                found[expected[key]].append(cell.column)
        if set(found) == set(expected.values()) and all(
            len(columns) == 1 for columns in found.values()
        ):
            return row_number, {name: columns[0] for name, columns in found.items()}
    return None


def _header_key(value: object) -> str:
    normalized = unicodedata.normalize("NFKD", str(value or ""))
    without_accents = "".join(
        char for char in normalized if not unicodedata.combining(char)
    )
    return re.sub(r"\s+", " ", without_accents).strip().casefold()


def _item(
    row: _WorkbookRow,
    action: BackfillAction,
    pdf_status: str,
    technical_status: str,
    reasons: tuple[str, ...],
    fingerprint: str,
    *,
    proposed_module: str | None = None,
    proposed_inverter: str | None = None,
    warnings: tuple[str, ...] = (),
    source_hash: str | None = None,
    semantic_status: str = "not_checked",
    semantic_errors: tuple[str, ...] = (),
    semantic_warnings: tuple[str, ...] = (),
    current_semantic_fingerprint: str | None = None,
    proposed_semantic_fingerprint: str | None = None,
    current_collection: CanonicalEquipmentCollection | None = None,
    proposed_collection: CanonicalEquipmentCollection | None = None,
) -> HistoricalEquipmentAuditItem:
    return HistoricalEquipmentAuditItem(
        workbook_sheet=row.sheet,
        workbook_row=row.row,
        protocol=row.protocol,
        current_module_text=row.module_text,
        current_inverter_text=row.inverter_text,
        proposed_module_text=proposed_module,
        proposed_inverter_text=proposed_inverter,
        technical_validation_status=technical_status,
        action=action,
        reasons=reasons,
        warnings=warnings,
        pdf_status=pdf_status,
        source_hash=source_hash,
        expected_row_fingerprint=fingerprint,
        semantic_validation_status=semantic_status,
        semantic_validation_errors=semantic_errors,
        semantic_validation_warnings=semantic_warnings,
        blocking_violations=semantic_errors,
        current_semantic_fingerprint=current_semantic_fingerprint,
        proposed_semantic_fingerprint=proposed_semantic_fingerprint,
        current_canonical_collection=(
            canonical_collection_to_dict(current_collection)
            if current_collection is not None
            else None
        ),
        proposed_canonical_collection=(
            canonical_collection_to_dict(proposed_collection)
            if proposed_collection is not None
            else None
        ),
    )


def _build_summary(
    items: list[HistoricalEquipmentAuditItem],
    sheets_analyzed: int,
    invalid_operational_sheets: int,
) -> HistoricalAuditSummary:
    actions = Counter(item.action for item in items)
    classifications = Counter(reason for item in items for reason in item.reasons)
    return HistoricalAuditSummary(
        sheets_analyzed=sheets_analyzed,
        rows_analyzed=len(items),
        valid_protocols=sum(item.protocol is not None for item in items),
        no_change=actions[BackfillAction.NO_CHANGE],
        total_updates=actions[BackfillAction.UPDATE_EQUIPMENT],
        pending_review=actions[BackfillAction.PENDING_TECHNICAL_REVIEW],
        pdf_not_found=actions[BackfillAction.PDF_NOT_FOUND],
        protocol_missing=actions[BackfillAction.PROTOCOL_MISSING],
        duplicate_protocols=actions[BackfillAction.DUPLICATE_PROTOCOL],
        conflicts=actions[BackfillAction.ROW_CONFLICT],
        unexpected_errors=invalid_operational_sheets,
        classifications=dict(classifications),
    )


def _technical_proposal_is_unsafe(proposal: TechnicalProposal) -> bool:
    values = (
        proposal.module_text,
        proposal.inverter_text,
        proposal.pdf_status,
        *proposal.errors,
        *proposal.warnings,
    )
    if any(contains_sensitive_or_path_text(value) for value in values):
        return True
    return proposal.status == "approved" and (
        not isinstance(proposal.source_hash, str)
        or len(proposal.source_hash) != 64
        or any(char not in "0123456789abcdef" for char in proposal.source_hash)
    )


def _default_resolver(
    downloads_root: Path | None, state_path: Path | None
) -> TechnicalResolver:
    root = Path(downloads_root) if downloads_root else Path("data/downloads")
    state_candidates = _state_pdf_index(state_path)

    def resolve(protocol: str) -> TechnicalProposal:
        candidates = _pdf_candidates(protocol, root, state_candidates.get(protocol, ()))
        valid: list[tuple[Path, str, str]] = []
        invalid_seen = False
        for candidate in candidates:
            try:
                validate_pdf_input(candidate)
                text = extract_pdf_text(candidate)
                if extract_protocol_from_pdf_text(text) != protocol:
                    invalid_seen = True
                    continue
                valid.append((candidate, text, file_sha256(candidate)))
            except (OSError, ValueError):
                invalid_seen = True
        unique_hashes = {entry[2] for entry in valid}
        if len(unique_hashes) > 1:
            return TechnicalProposal(
                status="ambiguous_pdf",
                pdf_status="ambiguous",
                errors=("MULTIPLE_PDF_CANDIDATES",),
            )
        if not valid:
            if invalid_seen:
                return TechnicalProposal(
                    status="invalid_pdf",
                    pdf_status="invalid",
                    errors=("PDF_INVALID_OR_PROTOCOL_MISMATCH",),
                )
            return TechnicalProposal(status="pdf_not_found", pdf_status="not_found")

        path, text, source_hash = valid[0]
        try:
            data = extract_generation_data(path, text=text)
            collection = canonical_collection_from_generation_data(data)
            module_text, inverter_text = format_canonical_collection(collection)
            validation = validate_canonical_equipment(
                collection,
                module_text,
                inverter_text,
                module_source=data.module_source,
                inverter_source=data.inverter_source,
            )
        except (OSError, ValueError, TypeError, KeyError, AttributeError):
            return TechnicalProposal(
                status="pending_review",
                pdf_status="valid",
                source_hash=source_hash,
                errors=("TECHNICAL_EXTRACTION_FAILED",),
            )
        if not validation.approved:
            return TechnicalProposal(
                status="pending_review",
                pdf_status="valid",
                source_hash=source_hash,
                errors=tuple(validation.errors),
                warnings=tuple(validation.warnings),
            )
        return TechnicalProposal(
            status="approved",
            module_text=module_text,
            inverter_text=inverter_text,
            source_hash=source_hash,
            pdf_status="valid",
            warnings=tuple(validation.warnings),
            canonical_collection=collection,
        )

    return resolve


def _state_pdf_index(state_path: Path | None) -> dict[str, tuple[Path, ...]]:
    if state_path is None or not Path(state_path).is_file():
        return {}
    try:
        payload = json.loads(Path(state_path).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    result: dict[str, tuple[Path, ...]] = {}
    for protocol, entry in (payload.get("protocols") or {}).items():
        paths = []
        for section, field in (("download", "pdf_path"), ("archive", "archived_pdf_path")):
            value = (entry.get(section) or {}).get(field)
            if value:
                paths.append(Path(value))
        result[str(protocol)] = tuple(paths)
    return result


def _pdf_candidates(
    protocol: str, downloads_root: Path, state_candidates: tuple[Path, ...]
) -> tuple[Path, ...]:
    candidates: list[Path] = list(state_candidates)
    protocol_dir = downloads_root / protocol
    if protocol_dir.is_dir():
        candidates.extend(sorted(protocol_dir.glob("*.pdf")))
    if downloads_root.is_dir():
        candidates.extend(sorted(downloads_root.rglob(f"*{protocol}*.pdf")))
    unique: dict[str, Path] = {}
    for path in candidates:
        unique[str(path.resolve(strict=False)).casefold()] = path
    return tuple(unique.values())


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
