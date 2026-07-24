from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
import zipfile
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, NoReturn

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from automacao_gd.application.historical_backfill.audit_service import (
    _find_header,
    file_sha256,
)
from automacao_gd.application.historical_backfill.comparison import row_fingerprint
from automacao_gd.application.historical_backfill.plan import (
    BackfillPlanError,
    validate_backfill_plan,
)
from automacao_gd.application.historical_backfill.report import rules_artifact_suffix
from automacao_gd.application.preflight import run_preflight
from automacao_gd.domain.backfill_models import (
    BackfillAction,
    BackfillApplyItemResult,
    BackfillApplyResult,
    BackfillApplyStatus,
)
from automacao_gd.domain.errors import OperationalBlockError
from automacao_gd.domain.equipment_semantics import (
    canonical_collection_from_dict,
    compare_equipment_row_transition,
    format_canonical_collection,
)
from automacao_gd.infrastructure.excel.availability import validate_workbook_availability
from automacao_gd.infrastructure.persistence.atomic import atomic_copy_file


_MINIMUM_FREE_SPACE_MULTIPLIER = 3


@dataclass(frozen=True, slots=True)
class BackfillApplicationPreparation:
    workbook_path: Path
    payload: dict[str, Any]
    workbook_hash: str
    required_confirmation: str
    updates_by_sheet: dict[str, int]
    started_at: str


class BackfillApplyError(RuntimeError):
    """Bloqueio operacional esperado antes ou durante o backfill."""

    def __init__(
        self, message: str, *, result: BackfillApplyResult | None = None
    ) -> None:
        super().__init__(message)
        self.result = result


def required_apply_confirmation(payload: dict[str, Any]) -> str:
    rules_label = str(payload["equipment_rules_version"]).rsplit("-", 2)[-2:]
    return f"APLICAR {'-'.join(rules_label).upper()} {payload['total_updates']} UPDATES"


def prepare_backfill_application(
    workbook_path: Path,
    payload: dict[str, Any],
    *,
    settings: Any,
    preflight_runner=run_preflight,
) -> BackfillApplicationPreparation:
    """Validate the entire operation without creating backups or changing the workbook."""
    started_at = _utc_now()
    try:
        validate_backfill_plan(payload)
    except BackfillPlanError as exc:
        _raise_abort(
            payload,
            started_at,
            "PLAN_VALIDATION_FAILED",
            "Plano bloqueado pela validação operacional.",
            cause=exc,
        )
    _validate_production_environment(settings, payload, started_at)
    try:
        preflight_runner(settings, real_run=True, require_cdp=False).raise_if_blocked()
    except OperationalBlockError as exc:
        _raise_abort(
            payload,
            started_at,
            str(exc.code),
            "Pré-voo operacional bloqueou a aplicação.",
            cause=exc,
        )

    path = Path(workbook_path)
    availability = validate_workbook_availability(
        path,
        require_writable=True,
        stage=f"pré-voo do backfill {payload['equipment_rules_version']}",
    )
    if not availability.ok:
        _raise_abort(
            payload,
            started_at,
            availability.code.value,
            availability.user_message,
        )
    try:
        free_space = shutil.disk_usage(path.parent).free
        required_space = path.stat().st_size * _MINIMUM_FREE_SPACE_MULTIPLIER
        workbook_hash = file_sha256(path)
    except OSError as exc:
        _raise_abort(
            payload,
            started_at,
            "WORKBOOK_TEMPORARILY_UNAVAILABLE",
            "A planilha ficou indisponível durante o pré-voo.",
            cause=exc,
        )
    if free_space <= required_space:
        _raise_abort(
            payload,
            started_at,
            "INSUFFICIENT_DISK_SPACE",
            "Espaço insuficiente para backup e arquivo temporário.",
        )
    if workbook_hash != payload["workbook_fingerprint"]:
        _raise_abort(
            payload,
            started_at,
            "WORKBOOK_CHANGED_AFTER_AUDIT",
            "A planilha mudou depois da auditoria.",
            original_hash=workbook_hash,
        )
    try:
        prepared, conflicts = _load_prepared_updates(path, payload)
    except BackfillApplyError as exc:
        _raise_abort(
            payload,
            started_at,
            "WORKBOOK_OPEN_FAILED",
            "Não foi possível abrir a planilha durante o pré-voo.",
            cause=exc,
            original_hash=workbook_hash,
        )
    if conflicts:
        _raise_conflicts(payload, started_at, workbook_hash, conflicts)
    _validate_prepared_updates(payload, prepared, started_at, workbook_hash)
    updates_by_sheet = dict(
        sorted(Counter(str(item["workbook_sheet"]) for item in prepared).items())
    )
    return BackfillApplicationPreparation(
        workbook_path=path,
        payload=payload,
        workbook_hash=workbook_hash,
        required_confirmation=required_apply_confirmation(payload),
        updates_by_sheet=updates_by_sheet,
        started_at=started_at,
    )


def execute_backfill_application(
    preparation: BackfillApplicationPreparation,
    confirmation: str,
) -> BackfillApplyResult:
    """Create validated artifacts and replace the workbook only after confirmation."""
    if confirmation != preparation.required_confirmation:
        result = _base_result(
            preparation.payload,
            preparation.started_at,
            status=BackfillApplyStatus.ABORTED_CONFIRMATION,
            original_hash=preparation.workbook_hash,
            confirmation_received=False,
            error_code="CONFIRMATION_MISMATCH",
            message="Confirmação forte divergente.",
        )
        raise BackfillApplyError("Confirmação forte inválida.", result=result)

    path = preparation.workbook_path
    payload = preparation.payload
    try:
        current_hash = file_sha256(path)
    except OSError as exc:
        _raise_abort(
            payload,
            preparation.started_at,
            "WORKBOOK_TEMPORARILY_UNAVAILABLE",
            "A planilha ficou indisponível depois da confirmação.",
            cause=exc,
            confirmation_received=True,
        )
    if current_hash != preparation.workbook_hash:
        _raise_abort(
            payload,
            preparation.started_at,
            "WORKBOOK_CHANGED_AFTER_PREFLIGHT",
            "A planilha mudou depois do pré-voo.",
            original_hash=current_hash,
            confirmation_received=True,
        )
    try:
        prepared, conflicts = _load_prepared_updates(path, payload)
    except BackfillApplyError as exc:
        _raise_abort(
            payload,
            preparation.started_at,
            "WORKBOOK_OPEN_FAILED",
            "Não foi possível reabrir a planilha antes do backup.",
            cause=exc,
            original_hash=preparation.workbook_hash,
            confirmation_received=True,
        )
    if conflicts:
        _raise_conflicts(
            payload,
            preparation.started_at,
            preparation.workbook_hash,
            conflicts,
            confirmation_received=True,
        )
    _validate_prepared_updates(
        payload, prepared, preparation.started_at, preparation.workbook_hash
    )

    backup_path: Path | None = None
    temporary_path: Path | None = None
    applied: list[BackfillApplyItemResult] = []
    replaced = False
    preservation_before = ""
    backup_hash: str | None = None
    try:
        backup_path = _create_validated_backup(
            path,
            preparation.workbook_hash,
            payload.get("equipment_rules_version"),
        )
        backup_hash = file_sha256(backup_path)
        workbook = load_workbook(path, data_only=False)
        try:
            target_values = {
                (item["workbook_sheet"], item["workbook_row"], column)
                for item in prepared
                for column in (item["module_column"], item["inverter_column"])
            }
            preservation_before = _preservation_fingerprint(workbook, target_values)
            applied = _write_prepared_updates(workbook, prepared)
            if len(applied) != len(prepared):
                raise BackfillApplyError(
                    "Quantidade aplicada divergiu da quantidade preparada."
                )
            temporary_path = _save_temporary_workbook(workbook, path)
        finally:
            workbook.close()

        if not _validate_temporary_workbook(
            temporary_path, prepared, preservation_before
        ):
            result = _base_result(
                payload,
                preparation.started_at,
                status=BackfillApplyStatus.ABORTED_UNEXPECTED_CHANGE,
                items=tuple(applied),
                backup_path=backup_path,
                original_hash=preparation.workbook_hash,
                backup_hash=backup_hash,
                unexpected_changes=1,
                confirmation_received=True,
                error_code="UNEXPECTED_WORKBOOK_CHANGE",
                message="O temporário alterou conteúdo fora da allowlist.",
            )
            raise BackfillApplyError(result.message or "Temporário inválido.", result=result)
        if not _validate_backup(backup_path, preparation.workbook_hash):
            _raise_backup_failure(preparation, backup_path, backup_hash)

        os.replace(temporary_path, path)
        temporary_path = None
        replaced = True
        _sync_file(path)
        if not _verify_saved_workbook(path, prepared, preservation_before):
            _raise_post_replace_failure(
                preparation, applied, backup_path, backup_hash
            )
        idempotency = _validate_idempotency(path, prepared)
        if idempotency != "NO_ADDITIONAL_CHANGES":
            _raise_post_replace_failure(
                preparation, applied, backup_path, backup_hash
            )
        return _base_result(
            payload,
            preparation.started_at,
            status=BackfillApplyStatus.APPLIED_SUCCESSFULLY,
            items=tuple(applied),
            backup_path=backup_path,
            original_hash=preparation.workbook_hash,
            backup_hash=backup_hash,
            final_hash=file_sha256(path),
            applied_updates=len(applied),
            updates_by_sheet=preparation.updates_by_sheet,
            confirmation_received=True,
            verification_passed=True,
            rollback_available=True,
            idempotency_result=idempotency,
        )
    except BackfillApplyError as exc:
        if exc.result is not None:
            raise
        if replaced and backup_path is not None and backup_hash is not None:
            _raise_post_replace_failure(
                preparation,
                applied,
                backup_path,
                backup_hash,
                cause=exc,
            )
        if backup_path is None or backup_hash is None:
            _raise_backup_failure(preparation, backup_path, backup_hash, cause=exc)
        result = _base_result(
            payload,
            preparation.started_at,
            status=BackfillApplyStatus.ABORTED_UNEXPECTED_CHANGE,
            items=tuple(applied),
            backup_path=backup_path,
            original_hash=preparation.workbook_hash,
            backup_hash=backup_hash,
            unexpected_changes=1,
            confirmation_received=True,
            error_code="TEMPORARY_WORKBOOK_FAILED",
            message="Falha antes da substituição atômica.",
        )
        raise BackfillApplyError(result.message or "Falha no temporário.", result=result) from exc
    except Exception as exc:
        if replaced and backup_path is not None and backup_hash is not None:
            _raise_post_replace_failure(
                preparation,
                applied,
                backup_path,
                backup_hash,
                cause=exc,
            )
        if backup_path is None or backup_hash is None:
            _raise_backup_failure(preparation, backup_path, backup_hash, cause=exc)
        result = _base_result(
            payload,
            preparation.started_at,
            status=BackfillApplyStatus.ABORTED_UNEXPECTED_CHANGE,
            items=tuple(applied),
            backup_path=backup_path,
            original_hash=preparation.workbook_hash,
            backup_hash=backup_hash,
            unexpected_changes=1,
            confirmation_received=True,
            error_code="TEMPORARY_WORKBOOK_FAILED",
            message="Falha antes da substituição atômica.",
        )
        raise BackfillApplyError(result.message or "Falha no temporário.", result=result) from exc
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def apply_backfill_plan(
    workbook_path: Path,
    payload: dict[str, Any],
    *,
    settings: Any,
    confirmation: str,
    preflight_runner=run_preflight,
) -> BackfillApplyResult:
    preparation = prepare_backfill_application(
        workbook_path,
        payload,
        settings=settings,
        preflight_runner=preflight_runner,
    )
    return execute_backfill_application(preparation, confirmation)


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _base_result(
    payload: dict[str, Any],
    started_at: str,
    *,
    status: BackfillApplyStatus,
    **values: Any,
) -> BackfillApplyResult:
    metadata: dict[str, Any] = {
        "status": status,
        "execution_id": payload.get("execution_id"),
        "plan_hash": payload.get("plan_hash"),
        "plan_version": payload.get("plan_version"),
        "technical_processing_format_version": payload.get(
            "technical_processing_format_version"
        ),
        "equipment_rules_version": payload.get("equipment_rules_version"),
        "started_at": started_at,
        "finished_at": _utc_now(),
        "planned_updates": int(payload.get("total_updates") or 0),
    }
    metadata.update(values)
    return BackfillApplyResult(**metadata)


def _raise_abort(
    payload: dict[str, Any],
    started_at: str,
    error_code: str,
    message: str,
    *,
    cause: Exception | None = None,
    **values: Any,
) -> NoReturn:
    result = _base_result(
        payload,
        started_at,
        status=BackfillApplyStatus.ABORTED_PRE_FLIGHT,
        error_code=error_code,
        message=message,
        **values,
    )
    error = BackfillApplyError(message, result=result)
    if cause is None:
        raise error
    raise error from cause


def _validate_production_environment(
    settings: Any, payload: dict[str, Any], started_at: str
) -> None:
    if getattr(settings, "APP_ENV", None) != "production":
        _raise_abort(
            payload,
            started_at,
            "PRODUCTION_ENVIRONMENT_REQUIRED",
            "Backfill real exige APP_ENV=production explicitamente.",
        )
    checks = (
        (not bool(getattr(settings, "DRY_RUN", True)), "DRY_RUN_MUST_BE_FALSE"),
        (bool(getattr(settings, "APPLY_EXCEL", False)), "APPLY_EXCEL_REQUIRED"),
        (bool(getattr(settings, "BACKUP_EXCEL", False)), "BACKUP_EXCEL_REQUIRED"),
    )
    for valid, code in checks:
        if not valid:
            _raise_abort(
                payload,
                started_at,
                code,
                "Configuração operacional incompatível com aplicação real.",
            )


def _load_prepared_updates(
    path: Path, payload: dict[str, Any]
) -> tuple[list[dict[str, Any]], list[BackfillApplyItemResult]]:
    try:
        workbook = load_workbook(path, data_only=False)
    except (OSError, ValueError, KeyError, InvalidFileException) as exc:
        raise BackfillApplyError(
            "Não foi possível abrir a planilha com segurança para o backfill."
        ) from exc
    try:
        prepared, conflicts = _prepare_updates(workbook, payload)
        prepared.sort(
            key=lambda item: (
                str(item["workbook_sheet"]),
                int(item["workbook_row"]),
                str(item.get("protocol") or ""),
            )
        )
        return prepared, conflicts
    finally:
        workbook.close()


def _validate_prepared_updates(
    payload: dict[str, Any],
    prepared: list[dict[str, Any]],
    started_at: str,
    workbook_hash: str,
) -> None:
    expected = int(payload["total_updates"])
    protocols = [str(item.get("protocol") or "") for item in prepared]
    locations = [
        (str(item["workbook_sheet"]), int(item["workbook_row"])) for item in prepared
    ]
    if (
        len(prepared) != expected
        or len(set(protocols)) != len(protocols)
        or len(set(locations)) != len(locations)
        or any(not protocol for protocol in protocols)
    ):
        _raise_abort(
            payload,
            started_at,
            "UPDATE_SET_INVALID",
            "O conjunto de atualizações não é único e completo.",
            original_hash=workbook_hash,
        )


def _raise_conflicts(
    payload: dict[str, Any],
    started_at: str,
    workbook_hash: str,
    conflicts: list[BackfillApplyItemResult],
    *,
    confirmation_received: bool = False,
) -> NoReturn:
    result = _base_result(
        payload,
        started_at,
        status=BackfillApplyStatus.ABORTED_FINGERPRINT_CONFLICT,
        items=tuple(conflicts),
        original_hash=workbook_hash,
        conflicts=len(conflicts),
        fingerprint_conflicts=len(conflicts),
        confirmation_received=confirmation_received,
        error_code="ROW_FINGERPRINT_MISMATCH",
        message="Fingerprint ou quality gate de linha divergente.",
    )
    raise BackfillApplyError(result.message or "Conflito de linha.", result=result)


def _create_validated_backup(
    path: Path,
    original_hash: str,
    equipment_rules_version: str | None,
) -> Path:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    rules_suffix = rules_artifact_suffix(equipment_rules_version)
    backup = path.with_name(f"Planilha_pre_backfill_{rules_suffix}_{timestamp}.xlsx")
    sequence = 1
    while backup.exists():
        backup = path.with_name(
            f"Planilha_pre_backfill_{rules_suffix}_{timestamp}_{sequence}.xlsx"
        )
        sequence += 1
    atomic_copy_file(path, backup, private=True)
    if not _validate_backup(backup, original_hash):
        backup.unlink(missing_ok=True)
        raise OSError("Backup criado, mas sua integridade não foi confirmada.")
    return backup


def _validate_backup(path: Path, original_hash: str) -> bool:
    try:
        if file_sha256(path) != original_hash:
            return False
        workbook = load_workbook(path, read_only=True, data_only=False)
        workbook.close()
        return True
    except (OSError, ValueError, KeyError, InvalidFileException):
        return False


def _raise_backup_failure(
    preparation: BackfillApplicationPreparation,
    backup_path: Path | None,
    backup_hash: str | None,
    *,
    cause: Exception | None = None,
) -> NoReturn:
    result = _base_result(
        preparation.payload,
        preparation.started_at,
        status=BackfillApplyStatus.ABORTED_PRE_FLIGHT,
        backup_path=backup_path,
        original_hash=preparation.workbook_hash,
        backup_hash=backup_hash,
        confirmation_received=True,
        error_code="BACKUP_VALIDATION_FAILED",
        message="Backup integral não pôde ser validado.",
    )
    error = BackfillApplyError(result.message or "Backup inválido.", result=result)
    if cause is None:
        raise error
    raise error from cause


def _save_temporary_workbook(workbook: Any, original_path: Path) -> Path:
    descriptor, name = tempfile.mkstemp(
        prefix=".backfill-apply-", suffix=".xlsx", dir=original_path.parent
    )
    os.close(descriptor)
    temporary = Path(name)
    try:
        workbook.save(temporary)
        _preserve_core_properties(original_path, temporary)
        _sync_file(temporary)
        return temporary
    except Exception:
        temporary.unlink(missing_ok=True)
        raise


def _validate_temporary_workbook(
    path: Path,
    prepared: list[dict[str, Any]],
    preservation_before: str,
) -> bool:
    try:
        return _verify_saved_workbook(path, prepared, preservation_before)
    except (OSError, ValueError, KeyError, InvalidFileException):
        return False


def _sync_file(path: Path) -> None:
    with path.open("r+b") as handle:
        handle.flush()
        os.fsync(handle.fileno())


def _validate_idempotency(path: Path, prepared: list[dict[str, Any]]) -> str:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for item in prepared:
                worksheet = workbook[item["workbook_sheet"]]
                if worksheet.cell(
                    item["workbook_row"], item["module_column"]
                ).value != item["proposed_module_text"]:
                    return "ADDITIONAL_CHANGES_REQUIRED"
                if worksheet.cell(
                    item["workbook_row"], item["inverter_column"]
                ).value != item["proposed_inverter_text"]:
                    return "ADDITIONAL_CHANGES_REQUIRED"
        finally:
            workbook.close()
    except (OSError, ValueError, KeyError, InvalidFileException):
        return "VALIDATION_FAILED"
    return "NO_ADDITIONAL_CHANGES"


def _restore_validated_backup(
    backup_path: Path, workbook_path: Path, original_hash: str
) -> bool:
    try:
        atomic_copy_file(backup_path, workbook_path, private=True)
        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        workbook.close()
        return file_sha256(workbook_path) == original_hash
    except (OSError, ValueError, KeyError, InvalidFileException):
        return False


def _raise_post_replace_failure(
    preparation: BackfillApplicationPreparation,
    applied: list[BackfillApplyItemResult],
    backup_path: Path,
    backup_hash: str,
    *,
    cause: Exception | None = None,
) -> NoReturn:
    restored = _restore_validated_backup(
        backup_path, preparation.workbook_path, preparation.workbook_hash
    )
    status = (
        BackfillApplyStatus.FAILED_ROLLED_BACK
        if restored
        else BackfillApplyStatus.FAILED_ROLLBACK_UNCONFIRMED
    )
    result = _base_result(
        preparation.payload,
        preparation.started_at,
        status=status,
        items=tuple(applied),
        backup_path=backup_path,
        original_hash=preparation.workbook_hash,
        backup_hash=backup_hash,
        final_hash=_safe_file_hash(preparation.workbook_path),
        applied_updates=0,
        updates_by_sheet=preparation.updates_by_sheet,
        confirmation_received=True,
        verification_passed=False,
        rollback_available=True,
        rollback_executed=True,
        rollback_status="CONFIRMED" if restored else "UNCONFIRMED",
        idempotency_result="NOT_CHECKED",
        error_code="POST_WRITE_VALIDATION_FAILED",
        message="Validação pós-gravação falhou; rollback executado.",
    )
    error = BackfillApplyError(result.message or "Falha pós-gravação.", result=result)
    if cause is None:
        raise error
    raise error from cause


def _prepare_updates(workbook, payload: dict[str, Any]):
    prepared: list[dict[str, Any]] = []
    conflicts: list[BackfillApplyItemResult] = []
    for item in payload["items"]:
        if item["action"] != BackfillAction.UPDATE_EQUIPMENT.value:
            continue
        sheet = str(item["workbook_sheet"])
        row = int(item["workbook_row"])
        if sheet not in workbook.sheetnames:
            conflicts.append(_conflict(item, "WORKSHEET_NOT_FOUND"))
            continue
        ws = workbook[sheet]
        header = _find_header(ws)
        if header is None:
            conflicts.append(_conflict(item, "WORKBOOK_HEADER_NOT_FOUND"))
            continue
        _, columns = header
        protocol = str(ws.cell(row, columns["Protocolo"]).value or "").strip() or None
        current_module = str(ws.cell(row, columns["Placa"]).value or "").strip()
        current_inverter = str(ws.cell(row, columns["Inversor"]).value or "").strip()
        current_fingerprint = row_fingerprint(
            sheet, row, protocol, current_module, current_inverter
        )
        if (
            protocol != item.get("protocol")
            or current_fingerprint != item["expected_row_fingerprint"]
        ):
            conflicts.append(
                _conflict(
                    item,
                    "ROW_FINGERPRINT_MISMATCH",
                    current_module=current_module,
                    current_inverter=current_inverter,
                )
            )
            continue
        if not _row_quality_gate(item, current_module, current_inverter):
            conflicts.append(
                _conflict(
                    item,
                    "SEMANTIC_QUALITY_GATE_FAILED",
                    current_module=current_module,
                    current_inverter=current_inverter,
                )
            )
            continue
        prepared.append(
            {
                **item,
                "module_column": columns["Placa"],
                "inverter_column": columns["Inversor"],
                "current_module_text": current_module,
                "current_inverter_text": current_inverter,
            }
        )
    return prepared, conflicts


def _conflict(
    item: dict[str, Any],
    reason: str,
    *,
    current_module: str | None = None,
    current_inverter: str | None = None,
) -> BackfillApplyItemResult:
    module = current_module if current_module is not None else ""
    inverter = current_inverter if current_inverter is not None else ""
    return BackfillApplyItemResult(
        workbook_sheet=str(item["workbook_sheet"]),
        workbook_row=int(item["workbook_row"]),
        protocol=item.get("protocol"),
        action=BackfillAction.ROW_CONFLICT,
        previous_module_text=module,
        final_module_text=module,
        previous_inverter_text=inverter,
        final_inverter_text=inverter,
        reasons=(reason,),
    )


def _write_prepared_updates(workbook, prepared: list[dict[str, Any]]):
    results: list[BackfillApplyItemResult] = []
    for item in prepared:
        ws = workbook[item["workbook_sheet"]]
        current_module = str(
            ws.cell(item["workbook_row"], item["module_column"]).value or ""
        ).strip()
        current_inverter = str(
            ws.cell(item["workbook_row"], item["inverter_column"]).value or ""
        ).strip()
        if not _row_quality_gate(item, current_module, current_inverter):
            raise BackfillApplyError(
                "Quality gate semântico bloqueou a linha antes da gravação."
            )
        ws.cell(item["workbook_row"], item["module_column"]).value = item[
            "proposed_module_text"
        ]
        ws.cell(item["workbook_row"], item["inverter_column"]).value = item[
            "proposed_inverter_text"
        ]
        results.append(
            BackfillApplyItemResult(
                workbook_sheet=item["workbook_sheet"],
                workbook_row=item["workbook_row"],
                protocol=item.get("protocol"),
                action=BackfillAction.UPDATE_EQUIPMENT,
                previous_module_text=item["current_module_text"],
                final_module_text=item["proposed_module_text"],
                previous_inverter_text=item["current_inverter_text"],
                final_inverter_text=item["proposed_inverter_text"],
                reasons=tuple(item.get("reasons") or ()),
            )
        )
    return results


def _row_quality_gate(
    item: dict[str, Any], current_module: str, current_inverter: str
) -> bool:
    if (
        item.get("technical_validation_status") != "approved"
        or item.get("semantic_validation_status") != "approved"
        or item.get("blocking_violations")
    ):
        return False
    try:
        proposed = canonical_collection_from_dict(
            item.get("proposed_canonical_collection")
        )
        comparison = compare_equipment_row_transition(
            current_module,
            current_inverter,
            item.get("proposed_module_text") or "",
            item.get("proposed_inverter_text") or "",
            proposed,
        )
    except (TypeError, ValueError):
        return False
    proposed_module, proposed_inverter = format_canonical_collection(proposed)
    return bool(
        comparison.status == "approved"
        and not comparison.errors
        and proposed_module == item.get("proposed_module_text")
        and proposed_inverter == item.get("proposed_inverter_text")
    )


def _preserve_core_properties(original: Path, saved: Path) -> None:
    """Preserve the workbook core-properties member byte for byte."""
    member = "docProps/core.xml"
    with zipfile.ZipFile(original) as source:
        original_core = source.read(member)
    fd, archive_name = tempfile.mkstemp(
        prefix=f".{saved.stem}.metadata.", suffix=saved.suffix, dir=saved.parent
    )
    os.close(fd)
    archive_path = Path(archive_name)
    try:
        with zipfile.ZipFile(saved) as source, zipfile.ZipFile(archive_path, "w") as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == member:
                    content = original_core
                target.writestr(info, content)
        os.replace(archive_path, saved)
    finally:
        archive_path.unlink(missing_ok=True)


def _verify_saved_workbook(
    path: Path,
    prepared: list[dict[str, Any]],
    preservation_before: str,
) -> bool:
    workbook = load_workbook(path, data_only=False)
    try:
        target_values = {
            (item["workbook_sheet"], item["workbook_row"], column)
            for item in prepared
            for column in (item["module_column"], item["inverter_column"])
        }
        if _preservation_fingerprint(workbook, target_values) != preservation_before:
            return False
        for item in prepared:
            ws = workbook[item["workbook_sheet"]]
            if ws.cell(item["workbook_row"], item["module_column"]).value != item[
                "proposed_module_text"
            ]:
                return False
            if ws.cell(item["workbook_row"], item["inverter_column"]).value != item[
                "proposed_inverter_text"
            ]:
                return False
        return True
    finally:
        workbook.close()


def _preservation_fingerprint(workbook, excluded_values: set[tuple[str, int, int]]) -> str:
    payload: dict[str, Any] = {
        "sheetnames": list(workbook.sheetnames),
        "properties": _safe_properties(workbook.properties),
        "active_sheet": workbook.index(workbook.active),
        "defined_names": sorted(
            (name, str(definition))
            for name, definition in workbook.defined_names.items()
        ),
        "calculation": str(workbook.calculation),
        "protection": str(workbook.security),
        "external_links": len(workbook._external_links),
        "sheets": [],
    }
    for ws in workbook.worksheets:
        cells = []
        for row in ws.iter_rows():
            for cell in row:
                is_backfill_target = (ws.title, cell.row, cell.column) in excluded_values
                cells.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": (
                            "<BACKFILL_TARGET>"
                            if is_backfill_target
                            else _json_value(cell.value)
                        ),
                        "data_type": (
                            "<BACKFILL_TARGET>" if is_backfill_target else cell.data_type
                        ),
                        "style": str(cell._style),
                        "number_format": cell.number_format,
                        "hyperlink": str(cell.hyperlink.target) if cell.hyperlink else None,
                        "comment": cell.comment.text if cell.comment else None,
                    }
                )
        payload["sheets"].append(
            {
                "title": ws.title,
                "state": ws.sheet_state,
                "freeze_panes": str(ws.freeze_panes or ""),
                "auto_filter": ws.auto_filter.ref,
                "merged_cells": sorted(str(item) for item in ws.merged_cells.ranges),
                "tables": sorted(
                    (table.name, table.ref, str(table.tableStyleInfo))
                    for table in ws.tables.values()
                ),
                "data_validations": sorted(
                    (str(item.sqref), str(item))
                    for item in ws.data_validations.dataValidation
                ),
                "conditional_formatting": sorted(
                    (str(key), tuple(str(rule) for rule in rules))
                    for key, rules in ws.conditional_formatting._cf_rules.items()
                ),
                "column_dimensions": sorted(
                    (key, dimension.width, dimension.hidden, dimension.outlineLevel)
                    for key, dimension in ws.column_dimensions.items()
                ),
                "row_dimensions": sorted(
                    (key, dimension.height, dimension.hidden, dimension.outlineLevel)
                    for key, dimension in ws.row_dimensions.items()
                ),
                "page_margins": str(ws.page_margins),
                "page_setup": str(ws.page_setup),
                "print_options": str(ws.print_options),
                "print_area": str(ws.print_area),
                "print_title_rows": str(ws.print_title_rows),
                "print_title_cols": str(ws.print_title_cols),
                "sheet_properties": str(ws.sheet_properties),
                "sheet_protection": str(ws.protection),
                "views": tuple(str(view) for view in ws.views.sheetView),
                "charts": len(ws._charts),
                "images": len(ws._images),
                "cells": cells,
            }
        )
    canonical = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _safe_properties(properties) -> dict[str, Any]:
    return {
        name: _json_value(getattr(properties, name, None))
        for name in (
            "title",
            "subject",
            "creator",
            "keywords",
            "description",
            "lastModifiedBy",
            "category",
            "contentStatus",
            "identifier",
            "language",
            "version",
            "revision",
            "created",
            "modified",
            "lastPrinted",
        )
    }


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _safe_file_hash(path: Path) -> str | None:
    try:
        return file_sha256(path)
    except OSError:
        return None

