from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from enum import StrEnum
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from automacao_gd.domain.models import PortalSolicitation
from automacao_gd.infrastructure.persistence.atomic import atomic_write_json, atomic_write_text


class ReconciliationIssueType(StrEnum):
    PORTAL_COMPLETED_MISSING_IN_WORKBOOK = "PORTAL_COMPLETED_MISSING_IN_WORKBOOK"
    DUPLICATE_PROTOCOL_IN_WORKBOOK = "DUPLICATE_PROTOCOL_IN_WORKBOOK"
    PROTOCOL_IN_WRONG_YEAR_SHEET = "PROTOCOL_IN_WRONG_YEAR_SHEET"
    YEAR_SHEET_VALIDATION_UNAVAILABLE = "YEAR_SHEET_VALIDATION_UNAVAILABLE"
    COMPLETION_EMPTY = "COMPLETION_EMPTY"
    COMPLETION_OPEN = "COMPLETION_OPEN"
    COMPLETION_VALID_DATE = "COMPLETION_VALID_DATE"
    COMPLETION_INVALID_VALUE = "COMPLETION_INVALID_VALUE"
    MODULE_FIELD_EMPTY = "MODULE_FIELD_EMPTY"
    INVERTER_FIELD_EMPTY = "INVERTER_FIELD_EMPTY"
    BOTH_EQUIPMENT_FIELDS_EMPTY = "BOTH_EQUIPMENT_FIELDS_EMPTY"
    EQUIPMENT_EMPTY_EXPECTED_BY_STATUS = "EQUIPMENT_EMPTY_EXPECTED_BY_STATUS"
    EQUIPMENT_EMPTY_REQUIRES_REVIEW = "EQUIPMENT_EMPTY_REQUIRES_REVIEW"
    INCOMPLETE_WORKBOOK_RECORD = "INCOMPLETE_WORKBOOK_RECORD"
    TRAILING_MATERIALIZED_EMPTY_ROW = "TRAILING_MATERIALIZED_EMPTY_ROW"
    WORKBOOK_PROTOCOL_NOT_IN_CURRENT_PORTAL_COMPLETED_SET = (
        "WORKBOOK_PROTOCOL_NOT_IN_CURRENT_PORTAL_COMPLETED_SET"
    )
    INVALID_WORKBOOK_PROTOCOL = "INVALID_WORKBOOK_PROTOCOL"
    INVALID_PORTAL_PROTOCOL = "INVALID_PORTAL_PROTOCOL"


@dataclass(frozen=True)
class ProtocolNormalization:
    raw: object
    normalized: str | None
    cell_type: str
    valid: bool
    reason: str | None = None


@dataclass(frozen=True)
class PortalProtocolIndex:
    records: list[dict[str, Any]]
    unique_protocols: set[str]
    invalid_records: list[dict[str, Any]]
    duplicate_records: list[dict[str, Any]]
    raw_count: int
    unique_count: int


@dataclass(frozen=True)
class WorkbookProtocolRecord:
    sheet_name: str
    row_number: int
    protocol_raw: object
    protocol_normalized: str | None
    protocol_cell_type: str
    client_present: bool
    ingress_date_raw: object
    ingress_date_normalized: str | None
    completion_raw: object
    completion_type: str
    parecer_raw: object
    placa_raw: object
    inversor_raw: object
    row_fingerprint: str

    def public_dict(self) -> dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "row_number": self.row_number,
            "protocol_raw": str(self.protocol_raw or ""),
            "protocol_normalized": self.protocol_normalized,
            "protocol_cell_type": self.protocol_cell_type,
            "client_present": self.client_present,
            "ingress_date_raw": _public_cell(self.ingress_date_raw),
            "ingress_date_normalized": self.ingress_date_normalized,
            "completion_raw": _public_cell(self.completion_raw),
            "completion_type": self.completion_type,
            "parecer_raw": _public_cell(self.parecer_raw),
            "placa_present": bool(str(self.placa_raw or "").strip()),
            "inversor_present": bool(str(self.inversor_raw or "").strip()),
            "row_fingerprint": self.row_fingerprint,
        }


@dataclass(frozen=True)
class WorkbookProtocolIndex:
    records: list[WorkbookProtocolRecord]
    records_by_protocol: dict[str, list[WorkbookProtocolRecord]]
    invalid_protocol_rows: list[dict[str, Any]]
    structural_findings: list[dict[str, Any]]
    sheet_classification: dict[str, str]
    project_rows: int
    rows_without_valid_protocol: int
    duplicate_extra_rows: int


@dataclass
class ReconciliationResult:
    metadata: dict[str, Any]
    portal_summary: dict[str, Any]
    workbook_summary: dict[str, Any]
    set_reconciliation: dict[str, Any]
    missing_in_workbook: list[dict[str, Any]]
    duplicate_workbook_protocols: list[dict[str, Any]]
    wrong_year_sheet: list[dict[str, Any]]
    completion_audit: list[dict[str, Any]]
    equipment_field_audit: list[dict[str, Any]]
    incomplete_records: list[dict[str, Any]]
    workbook_only_protocols: list[dict[str, Any]]
    structural_findings: list[dict[str, Any]]
    pagination: dict[str, Any]
    metrics_scope: str
    authoritative_status: dict[str, Any]
    safety: dict[str, Any]
    operational_batch: dict[str, Any]
    review: dict[str, Any]
    decision: str
    workbook_index: WorkbookProtocolIndex = field(repr=False)

    def to_public_dict(self) -> dict[str, Any]:
        return {
            "metadata": self.metadata,
            "portal_summary": self.portal_summary,
            "workbook_summary": self.workbook_summary,
            "set_reconciliation": self.set_reconciliation,
            "missing_in_workbook": self.missing_in_workbook,
            "duplicate_workbook_protocols": self.duplicate_workbook_protocols,
            "wrong_year_sheet": self.wrong_year_sheet,
            "completion_audit": self.completion_audit,
            "equipment_field_audit": self.equipment_field_audit,
            "incomplete_records": self.incomplete_records,
            "workbook_only_protocols": self.workbook_only_protocols,
            "structural_findings": self.structural_findings,
            "pagination": self.pagination,
            "metrics_scope": self.metrics_scope,
            "authoritative_status": self.authoritative_status,
            "safety": self.safety,
            "operational_batch": self.operational_batch,
            "review": self.review,
            "decision": self.decision,
        }


def normalize_protocol(value: object) -> ProtocolNormalization:
    if value is None or isinstance(value, bool):
        return ProtocolNormalization(value, None, "empty", False, "empty")
    if isinstance(value, int):
        return ProtocolNormalization(value, str(value), "numeric", True)
    if isinstance(value, float):
        if value.is_integer():
            return ProtocolNormalization(value, str(int(value)), "numeric", True)
        return ProtocolNormalization(value, None, "numeric", False, "fractional_numeric")
    text = str(value).strip()
    if not text:
        return ProtocolNormalization(value, None, "empty", False, "empty")
    compact = re.sub(r"\s+", "", text)
    if re.search(r"[eE][+-]?\d+", compact) or not compact.isdigit():
        return ProtocolNormalization(value, None, "text", False, "invalid_characters")
    return ProtocolNormalization(value, compact, "text", True)


def reconcile_portal_workbook(
    portal_records: list[PortalSolicitation],
    workbook_path: Path | str,
    *,
    operational_selected_protocols: list[str] | None = None,
    pagination: dict[str, Any] | None = None,
    started_at: datetime | None = None,
) -> ReconciliationResult:
    started_at = started_at or datetime.now()
    workbook_path = Path(workbook_path)
    sha_before = _sha256_file(workbook_path)
    portal_index = build_portal_protocol_index(portal_records)
    workbook_index = build_workbook_protocol_index(workbook_path)
    sha_after = _sha256_file(workbook_path)

    portal_set = portal_index.unique_protocols
    workbook_set = set(workbook_index.records_by_protocol)
    matched = sorted(portal_set & workbook_set)
    missing = sorted(portal_set - workbook_set)
    workbook_only = sorted(workbook_set - portal_set)

    duplicates = _duplicate_issues(workbook_index)
    wrong_year, year_findings = _year_sheet_issues(workbook_index.records)
    completion = _completion_issues(matched, workbook_index)
    equipment = _equipment_issues(workbook_index.records)
    incomplete = _incomplete_record_issues(workbook_index.records, equipment, completion)
    missing_issues = [_missing_issue(protocol, portal_index) for protocol in missing]
    workbook_only_issues = [
        {
            "issue_type": ReconciliationIssueType.WORKBOOK_PROTOCOL_NOT_IN_CURRENT_PORTAL_COMPLETED_SET,
            "protocol": protocol,
            "occurrences": [
                _record_location(record)
                for record in workbook_index.records_by_protocol.get(protocol, [])
            ],
        }
        for protocol in workbook_only
    ]
    structural = [
        *workbook_index.invalid_protocol_rows,
        *workbook_index.structural_findings,
        *year_findings,
    ]
    duplicate_extra_rows = workbook_index.duplicate_extra_rows
    rows_without_valid_protocol = workbook_index.rows_without_valid_protocol
    workbook_project_rows = workbook_index.project_rows
    metrics_ok = workbook_project_rows == (
        len(workbook_set) + duplicate_extra_rows + rows_without_valid_protocol
    ) and len(portal_set) == (len(matched) + len(missing))
    readonly_ok = sha_before == sha_after
    non_blocking_findings = _has_non_blocking_findings(
        missing_issues=missing_issues,
        duplicates=duplicates,
        wrong_year=wrong_year,
        completion=completion,
        equipment=equipment,
        incomplete=incomplete,
        workbook_only=workbook_only_issues,
        structural=structural,
    )
    pagination_status = _pagination_status(pagination)
    if not pagination_status["set_reconciliation_authoritative"]:
        decision = "STAGE2_BLOCKED - PORTAL_PAGINATION_INCOMPLETE"
    elif not metrics_ok or not readonly_ok:
        decision = "STAGE2_REJECTED - RECONCILIATION_METRICS_INCONSISTENT"
    elif non_blocking_findings:
        decision = "STAGE2_PARTIAL - PAGINATION_COMPLETE_WITH_NON_BLOCKING_DATA_FINDINGS"
    else:
        decision = "STAGE2_COMPLETE - FULL_PORTAL_RECONCILIATION_APPROVED"
    result = ReconciliationResult(
        metadata={
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "source": "option_5_completed_filter",
            "workbook_file": workbook_path.name,
            "workbook_reference": "configured_official_workbook",
            "duration_seconds": round((datetime.now() - started_at).total_seconds(), 3),
        },
        portal_summary={
            "portal_concluded_raw": portal_index.raw_count,
            "portal_concluded_unique": portal_index.unique_count,
            "portal_invalid_protocols": len(portal_index.invalid_records),
            "portal_duplicate_listing_rows": len(portal_index.duplicate_records),
        },
        workbook_summary={
            "workbook_project_rows": workbook_project_rows,
            "workbook_unique_protocols": len(workbook_set),
            "workbook_duplicate_extra_rows": duplicate_extra_rows,
            "workbook_rows_without_valid_protocol": rows_without_valid_protocol,
            "workbook_invalid_protocol_rows": len(workbook_index.invalid_protocol_rows),
            "sheet_classification": workbook_index.sheet_classification,
        },
        set_reconciliation={
            "matched_unique": len(matched),
            "missing_in_workbook_unique": len(missing),
            "workbook_only_unique": len(workbook_only),
            "portal_equation_valid": len(portal_set) == (len(matched) + len(missing)),
            "workbook_equation_valid": workbook_project_rows
            == (len(workbook_set) + duplicate_extra_rows + rows_without_valid_protocol),
        },
        missing_in_workbook=missing_issues,
        duplicate_workbook_protocols=duplicates,
        wrong_year_sheet=wrong_year,
        completion_audit=completion,
        equipment_field_audit=equipment,
        incomplete_records=incomplete,
        workbook_only_protocols=workbook_only_issues,
        structural_findings=structural,
        pagination=pagination_status["pagination"],
        metrics_scope=pagination_status["metrics_scope"],
        authoritative_status={
            "set_reconciliation_authoritative": pagination_status[
                "set_reconciliation_authoritative"
            ]
        },
        safety={
            "read_only": True,
            "workbook_sha_before_reconciliation": sha_before,
            "workbook_sha_after_reconciliation": sha_after,
            "workbook_sha_preserved": readonly_ok,
            "workbook_save_called": False,
            "backup_created": False,
            "os_replace_called": False,
            "pdf_downloads": 0,
            "detail_pages_opened": 0,
            "archives": 0,
        },
        operational_batch={
            "reconciliation_protocols_audited": len(portal_set),
            "operational_protocols_selected": len(operational_selected_protocols or []),
            "operational_protocols_processed": len(operational_selected_protocols or []),
        },
        review={"P0": 0, "P1": 0, "P2": 0, "P3": 0},
        decision=decision,
        workbook_index=workbook_index,
    )
    return result


def _has_non_blocking_findings(
    *,
    missing_issues: list[dict[str, Any]],
    duplicates: list[dict[str, Any]],
    wrong_year: list[dict[str, Any]],
    completion: list[dict[str, Any]],
    equipment: list[dict[str, Any]],
    incomplete: list[dict[str, Any]],
    workbook_only: list[dict[str, Any]],
    structural: list[dict[str, Any]],
) -> bool:
    completion_findings = [
        item
        for item in completion
        if item["issue_type"] != ReconciliationIssueType.COMPLETION_VALID_DATE
    ]
    return any(
        (
            missing_issues,
            duplicates,
            wrong_year,
            completion_findings,
            equipment,
            incomplete,
            workbook_only,
            structural,
        )
    )


def _pagination_status(pagination: dict[str, Any] | None) -> dict[str, Any]:
    if not pagination:
        normalized: dict[str, Any] = {
            "pagination_complete": True,
            "last_page_confirmed": True,
            "last_page_number": None,
            "next_page_available_after_stop": False,
            "pagination_stop_reason": "not_provided_unit_context",
            "pagination_safety_cap": None,
            "pages_visited": [],
        }
        return {
            "pagination": normalized,
            "metrics_scope": "global",
            "set_reconciliation_authoritative": True,
        }
    normalized = {
        "pagination_complete": bool(pagination.get("pagination_complete")),
        "last_page_confirmed": bool(pagination.get("last_page_confirmed")),
        "last_page_number": pagination.get("last_page_number"),
        "next_page_available_after_stop": bool(
            pagination.get("next_page_available_after_stop")
        ),
        "pagination_stop_reason": pagination.get("pagination_stop_reason"),
        "pagination_safety_cap": pagination.get("pagination_safety_cap"),
        "pages_visited": list(pagination.get("pages_visited") or []),
    }
    authoritative = (
        normalized["pagination_complete"]
        and normalized["last_page_confirmed"]
        and not normalized["next_page_available_after_stop"]
    )
    return {
        "pagination": normalized,
        "metrics_scope": "global" if authoritative else "partial",
        "set_reconciliation_authoritative": authoritative,
    }


def build_portal_protocol_index(records: list[PortalSolicitation]) -> PortalProtocolIndex:
    public_records: list[dict[str, Any]] = []
    invalid: list[dict[str, Any]] = []
    duplicates: list[dict[str, Any]] = []
    unique: set[str] = set()
    seen: set[str] = set()
    for record in records:
        normalized = normalize_protocol(record.protocol)
        public = {
            "protocol": normalized.normalized,
            "protocol_raw": str(record.protocol or ""),
            "status_raw": record.status,
            "entry_date_raw": record.entry_date,
            "page_number": record.page_number,
            "row_index": record.row_index,
        }
        if not normalized.valid or not normalized.normalized:
            public["issue_type"] = ReconciliationIssueType.INVALID_PORTAL_PROTOCOL
            public["reason"] = normalized.reason
            invalid.append(public)
            continue
        public["protocol"] = normalized.normalized
        if normalized.normalized in seen:
            duplicate = dict(public)
            duplicate["issue_type"] = "DUPLICATE_PROTOCOL_IN_PORTAL_LISTING"
            duplicates.append(duplicate)
            continue
        seen.add(normalized.normalized)
        unique.add(normalized.normalized)
        public_records.append(public)
    return PortalProtocolIndex(
        records=public_records,
        unique_protocols=unique,
        invalid_records=invalid,
        duplicate_records=duplicates,
        raw_count=len(records),
        unique_count=len(unique),
    )


def build_workbook_protocol_index(workbook_path: Path | str) -> WorkbookProtocolIndex:
    workbook_path = Path(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    records: list[WorkbookProtocolRecord] = []
    invalid_rows: list[dict[str, Any]] = []
    structural: list[dict[str, Any]] = []
    sheet_classification: dict[str, str] = {}
    rows_without_valid_protocol = 0
    project_rows = 0
    try:
        for sheet in workbook.worksheets:
            classification = _sheet_classification(sheet.title)
            sheet_classification[sheet.title] = classification
            header = _find_header(sheet)
            if header is None:
                if classification == "unknown":
                    structural.append(
                        {
                            "issue_type": "UNKNOWN_SHEET_WITHOUT_PROJECT_HEADER",
                            "sheet_name": sheet.title,
                        }
                    )
                continue
            row_items: list[tuple[int, dict[str, Any], bool]] = []
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header["row"] + 1, values_only=True),
                start=header["row"] + 1,
            ):
                values = _row_values_from_tuple(row, header)
                row_items.append((row_number, values, _row_has_project_material(values)))
            material_rows = [
                row_number for row_number, _, has_material in row_items if has_material
            ]
            last_material_row = max(
                max(material_rows, default=header["row"]),
                sheet.max_row,
            )
            for row_number, values, has_material in row_items:
                if not has_material:
                    if row_number <= last_material_row:
                        structural.append(
                            {
                                "issue_type": ReconciliationIssueType.TRAILING_MATERIALIZED_EMPTY_ROW,
                                "sheet_name": sheet.title,
                                "row_number": row_number,
                            }
                        )
                    continue
                project_rows += 1
                normalized = normalize_protocol(values["protocol"])
                if not normalized.valid or not normalized.normalized:
                    rows_without_valid_protocol += 1
                    invalid_rows.append(
                        {
                            "issue_type": ReconciliationIssueType.INVALID_WORKBOOK_PROTOCOL,
                            "sheet_name": sheet.title,
                            "row_number": row_number,
                            "protocol_raw": str(values["protocol"] or ""),
                            "reason": normalized.reason,
                        }
                    )
                    continue
                record = _workbook_record(sheet.title, row_number, values, normalized)
                records.append(record)
    finally:
        workbook.close()
    records_by_protocol: dict[str, list[WorkbookProtocolRecord]] = {}
    for record in records:
        if record.protocol_normalized:
            records_by_protocol.setdefault(record.protocol_normalized, []).append(record)
    duplicate_extra_rows = sum(max(len(items) - 1, 0) for items in records_by_protocol.values())
    return WorkbookProtocolIndex(
        records=records,
        records_by_protocol=records_by_protocol,
        invalid_protocol_rows=invalid_rows,
        structural_findings=structural,
        sheet_classification=sheet_classification,
        project_rows=project_rows,
        rows_without_valid_protocol=rows_without_valid_protocol,
        duplicate_extra_rows=duplicate_extra_rows,
    )


def save_reconciliation_reports(
    result: ReconciliationResult,
    logs_dir: Path | str,
    *,
    timestamp: str | None = None,
    filename_prefix: str = "portal_workbook_reconciliation",
) -> tuple[Path, Path]:
    logs_dir = Path(logs_dir)
    timestamp = timestamp or datetime.now().strftime("%Y%m%dT%H%M%SZ")
    json_path = logs_dir / f"{filename_prefix}_{timestamp}.json"
    markdown_path = logs_dir / f"{filename_prefix}_{timestamp}.md"
    public = result.to_public_dict()
    atomic_write_json(json_path, public, private=True)
    atomic_write_text(markdown_path, _markdown_report(public), private=True)
    return json_path, markdown_path


def reconciliation_summary(result: ReconciliationResult, report_path: str | None = None) -> dict[str, Any]:
    return {
        "metrics_scope": result.metrics_scope,
        "set_reconciliation_authoritative": result.authoritative_status[
            "set_reconciliation_authoritative"
        ],
        "pagination_complete": result.pagination["pagination_complete"],
        "last_page_confirmed": result.pagination["last_page_confirmed"],
        "last_page_number": result.pagination["last_page_number"],
        "next_page_available_after_stop": result.pagination[
            "next_page_available_after_stop"
        ],
        "pagination_stop_reason": result.pagination["pagination_stop_reason"],
        "pages_read": result.portal_summary.get("pages_read", 0),
        "portal_concluded_unique": result.portal_summary["portal_concluded_unique"],
        "workbook_unique_protocols": result.workbook_summary["workbook_unique_protocols"],
        "matched_unique": result.set_reconciliation["matched_unique"],
        "missing_in_workbook_unique": result.set_reconciliation["missing_in_workbook_unique"],
        "workbook_only_unique": result.set_reconciliation["workbook_only_unique"],
        "duplicate_workbook_protocols": len(result.duplicate_workbook_protocols),
        "wrong_year_sheet": len(result.wrong_year_sheet),
        "completion_empty": sum(
            1
            for item in result.completion_audit
            if item["issue_type"] == ReconciliationIssueType.COMPLETION_EMPTY
        ),
        "equipment_empty_requires_review": sum(
            1
            for item in result.equipment_field_audit
            if item["issue_type"] == ReconciliationIssueType.EQUIPMENT_EMPTY_REQUIRES_REVIEW
        ),
        "incomplete_records": len(result.incomplete_records),
        "report_path": report_path,
    }


def _workbook_record(
    sheet_name: str,
    row_number: int,
    values: dict[str, Any],
    normalized: ProtocolNormalization,
) -> WorkbookProtocolRecord:
    ingress = _normalize_date(values["ingress_date"])
    payload = [
        sheet_name,
        row_number,
        normalized.normalized,
        _public_cell(values["ingress_date"]),
        _public_cell(values["completion"]),
        str(values["placa"] or ""),
        str(values["inversor"] or ""),
    ]
    fingerprint = hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    ).hexdigest()
    return WorkbookProtocolRecord(
        sheet_name=sheet_name,
        row_number=row_number,
        protocol_raw=values["protocol"],
        protocol_normalized=normalized.normalized,
        protocol_cell_type=normalized.cell_type,
        client_present=bool(str(values["client"] or "").strip()),
        ingress_date_raw=values["ingress_date"],
        ingress_date_normalized=ingress.isoformat() if ingress else None,
        completion_raw=values["completion"],
        completion_type=_completion_type(values["completion"]),
        parecer_raw=values["parecer"],
        placa_raw=values["placa"],
        inversor_raw=values["inversor"],
        row_fingerprint=fingerprint,
    )


def _find_header(sheet) -> dict[str, int] | None:
    for row_number, row in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=min(sheet.max_row, 20),
            max_col=min(sheet.max_column, 40),
            values_only=True,
        ),
        start=1,
    ):
        headers: dict[str, int] = {}
        for column, value in enumerate(row, start=1):
            if value is None:
                continue
            headers[_normalize_header(str(value))] = column
        if "protocolo" in headers and "conclusao" in headers:
            return {
                "row": row_number,
                "client": headers.get("cliente", 1),
                "protocol": headers["protocolo"],
                "ingress_date": headers.get("data_de_ingresso", headers.get("data_ingresso", 3)),
                "completion": headers["conclusao"],
                "parecer": headers.get("parecer", 5),
                "placa": headers.get("placa", 6),
                "inversor": headers.get("inversor", 7),
            }
    return None


def _row_values_from_tuple(row: tuple[Any, ...], header: dict[str, int]) -> dict[str, Any]:
    return {
        "client": _tuple_cell(row, header["client"]),
        "protocol": _tuple_cell(row, header["protocol"]),
        "ingress_date": _tuple_cell(row, header["ingress_date"]),
        "completion": _tuple_cell(row, header["completion"]),
        "parecer": _tuple_cell(row, header["parecer"]),
        "placa": _tuple_cell(row, header["placa"]),
        "inversor": _tuple_cell(row, header["inversor"]),
    }


def _tuple_cell(row: tuple[Any, ...], one_based_column: int) -> Any:
    index = one_based_column - 1
    if index < 0 or index >= len(row):
        return None
    return row[index]


def _normalize_header(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    normalized = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    normalized = re.sub(r"\s+", "_", normalized.strip().casefold())
    return normalized


def _row_values(sheet, header: dict[str, int], row_number: int) -> dict[str, Any]:
    return {
        "client": sheet.cell(row_number, header["client"]).value,
        "protocol": sheet.cell(row_number, header["protocol"]).value,
        "ingress_date": sheet.cell(row_number, header["ingress_date"]).value,
        "completion": sheet.cell(row_number, header["completion"]).value,
        "parecer": sheet.cell(row_number, header["parecer"]).value,
        "placa": sheet.cell(row_number, header["placa"]).value,
        "inversor": sheet.cell(row_number, header["inversor"]).value,
    }


def _row_has_project_material(values: dict[str, Any]) -> bool:
    return any(str(value or "").strip() for value in values.values())


def _last_material_row(sheet, header: dict[str, int]) -> int:
    last = header["row"]
    for row_number in range(header["row"] + 1, sheet.max_row + 1):
        if _row_has_project_material(_row_values(sheet, header, row_number)):
            last = row_number
    return max(last, sheet.max_row)


def _sheet_classification(sheet_name: str) -> str:
    name = sheet_name.strip()
    if re.fullmatch(r"\d{4}", name) or re.fullmatch(r"\d{4}\s*-\s*\d{4}", name):
        return "valid_year_sheet"
    if name.casefold() in {"config", "auxiliar", "modelo"}:
        return "auxiliary_sheet"
    return "unknown"


def _allowed_years_for_sheet(sheet_name: str) -> set[int] | None:
    name = sheet_name.strip()
    if re.fullmatch(r"\d{4}", name):
        return {int(name)}
    match = re.fullmatch(r"(\d{4})\s*-\s*(\d{4})", name)
    if match:
        start = int(match.group(1))
        end = int(match.group(2))
        return set(range(start, end + 1))
    return None


def _normalize_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def _completion_type(value: object) -> str:
    if value is None or not str(value).strip():
        return "empty"
    if isinstance(value, (datetime, date)):
        return "valid_date"
    text = str(value).strip().casefold()
    if text == "em aberto":
        return "open"
    if _normalize_date(value):
        return "valid_date"
    return "invalid"


def _duplicate_issues(index: WorkbookProtocolIndex) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for protocol, records in sorted(index.records_by_protocol.items()):
        if len(records) <= 1:
            continue
        issues.append(
            {
                "issue_type": ReconciliationIssueType.DUPLICATE_PROTOCOL_IN_WORKBOOK,
                "protocol": protocol,
                "occurrences": [_record_location(record) for record in records],
            }
        )
    return issues


def _year_sheet_issues(
    records: list[WorkbookProtocolRecord],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    wrong: list[dict[str, Any]] = []
    unavailable: list[dict[str, Any]] = []
    for record in records:
        allowed = _allowed_years_for_sheet(record.sheet_name)
        if allowed is None:
            continue
        ingress = _normalize_date(record.ingress_date_raw)
        if ingress is None:
            unavailable.append(
                {
                    "issue_type": ReconciliationIssueType.YEAR_SHEET_VALIDATION_UNAVAILABLE,
                    "protocol": record.protocol_normalized,
                    "sheet_name": record.sheet_name,
                    "row_number": record.row_number,
                    "row_fingerprint": record.row_fingerprint,
                }
            )
            continue
        if ingress.year not in allowed:
            wrong.append(
                {
                    "issue_type": ReconciliationIssueType.PROTOCOL_IN_WRONG_YEAR_SHEET,
                    "protocol": record.protocol_normalized,
                    "sheet_name": record.sheet_name,
                    "row_number": record.row_number,
                    "ingress_year": ingress.year,
                    "expected_sheet": str(ingress.year),
                    "row_fingerprint": record.row_fingerprint,
                }
            )
    return wrong, unavailable


def _completion_issues(
    matched_protocols: list[str],
    index: WorkbookProtocolIndex,
) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for protocol in matched_protocols:
        for record in index.records_by_protocol.get(protocol, []):
            completion_type = _completion_type(record.completion_raw)
            issue_type = {
                "empty": ReconciliationIssueType.COMPLETION_EMPTY,
                "open": ReconciliationIssueType.COMPLETION_OPEN,
                "valid_date": ReconciliationIssueType.COMPLETION_VALID_DATE,
                "invalid": ReconciliationIssueType.COMPLETION_INVALID_VALUE,
            }[completion_type]
            issues.append(
                {
                    "issue_type": issue_type,
                    "protocol": protocol,
                    "sheet_name": record.sheet_name,
                    "row_number": record.row_number,
                    "completion_type": completion_type,
                    "row_fingerprint": record.row_fingerprint,
                }
            )
    return issues


def _equipment_issues(records: list[WorkbookProtocolRecord]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    for record in records:
        placa_empty = not str(record.placa_raw or "").strip()
        inversor_empty = not str(record.inversor_raw or "").strip()
        if not placa_empty and not inversor_empty:
            continue
        if placa_empty and inversor_empty:
            base_type = ReconciliationIssueType.BOTH_EQUIPMENT_FIELDS_EMPTY
        elif placa_empty:
            base_type = ReconciliationIssueType.MODULE_FIELD_EMPTY
        else:
            base_type = ReconciliationIssueType.INVERTER_FIELD_EMPTY
        status_issue = (
            ReconciliationIssueType.EQUIPMENT_EMPTY_EXPECTED_BY_STATUS
            if _parecer_allows_empty_equipment(record.parecer_raw)
            else ReconciliationIssueType.EQUIPMENT_EMPTY_REQUIRES_REVIEW
        )
        issues.append(
            {
                "issue_type": status_issue,
                "field_issue_type": base_type,
                "protocol": record.protocol_normalized,
                "sheet_name": record.sheet_name,
                "row_number": record.row_number,
                "row_fingerprint": record.row_fingerprint,
            }
        )
    return issues


def _incomplete_record_issues(
    records: list[WorkbookProtocolRecord],
    equipment_issues: list[dict[str, Any]],
    completion_issues: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    equipment_requires_review = {
        item["protocol"]
        for item in equipment_issues
        if item["issue_type"] == ReconciliationIssueType.EQUIPMENT_EMPTY_REQUIRES_REVIEW
    }
    invalid_completion = {
        item["protocol"]
        for item in completion_issues
        if item["issue_type"] == ReconciliationIssueType.COMPLETION_INVALID_VALUE
    }
    issues: list[dict[str, Any]] = []
    for record in records:
        missing: list[str] = []
        if _normalize_date(record.ingress_date_raw) is None:
            missing.append("ingress_date")
        if record.parecer_raw is None or not str(record.parecer_raw).strip():
            missing.append("parecer")
        if record.protocol_normalized in invalid_completion:
            missing.append("completion")
        if record.protocol_normalized in equipment_requires_review:
            missing.append("equipment")
        if missing:
            issues.append(
                {
                    "issue_type": ReconciliationIssueType.INCOMPLETE_WORKBOOK_RECORD,
                    "protocol": record.protocol_normalized,
                    "sheet_name": record.sheet_name,
                    "row_number": record.row_number,
                    "missing_fields": missing,
                    "row_fingerprint": record.row_fingerprint,
                }
            )
    return issues


def _missing_issue(protocol: str, portal_index: PortalProtocolIndex) -> dict[str, Any]:
    record = next((item for item in portal_index.records if item["protocol"] == protocol), {})
    return {
        "issue_type": ReconciliationIssueType.PORTAL_COMPLETED_MISSING_IN_WORKBOOK,
        "protocol": protocol,
        "page_number": record.get("page_number"),
        "row_index": record.get("row_index"),
        "status_raw": record.get("status_raw"),
        "entry_date_raw": record.get("entry_date_raw"),
    }


def _record_location(record: WorkbookProtocolRecord) -> dict[str, Any]:
    return {
        "sheet_name": record.sheet_name,
        "row_number": record.row_number,
        "row_fingerprint": record.row_fingerprint,
    }


def _parecer_allows_empty_equipment(value: object) -> bool:
    text = _normalize_text(value)
    return any(token in text for token in ("cancelado", "negativo", "indeferido"))


def _normalize_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.casefold()).strip()


def _public_cell(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _markdown_report(payload: dict[str, Any]) -> str:
    def count(section: str) -> int:
        value = payload.get(section)
        return len(value) if isinstance(value, list) else 0

    lines = [
        "# Reconciliação Portal GD × planilha",
        "",
        "## 1. Resumo executivo",
        "",
        f"- Decisão: {payload['decision']}",
        f"- Escopo das métricas: {payload['metrics_scope']}",
        "- Reconciliação autoritativa: "
        f"{payload['authoritative_status']['set_reconciliation_authoritative']}",
        f"- Concluídos únicos no Portal: {payload['portal_summary']['portal_concluded_unique']}",
        f"- Protocolos únicos na planilha: {payload['workbook_summary']['workbook_unique_protocols']}",
        f"- Encontrados nas duas fontes: {payload['set_reconciliation']['matched_unique']}",
        f"- Ausentes na planilha: {payload['set_reconciliation']['missing_in_workbook_unique']}",
        "",
        "## 2. Fontes analisadas",
        "",
        f"- Fonte Portal: {payload['metadata']['source']}",
        f"- Planilha: {payload['metadata']['workbook_file']}",
        f"- Referência da planilha: {payload['metadata']['workbook_reference']}",
        "",
        "## 2.1. Paginação",
        "",
        f"- Completa: {payload['pagination']['pagination_complete']}",
        f"- Última página confirmada: {payload['pagination']['last_page_confirmed']}",
        f"- Última página: {payload['pagination']['last_page_number']}",
        f"- Próxima página após parada: {payload['pagination']['next_page_available_after_stop']}",
        f"- Motivo da parada: {payload['pagination']['pagination_stop_reason']}",
        f"- Teto defensivo: {payload['pagination']['pagination_safety_cap']}",
        f"- Páginas visitadas: {payload['pagination']['pages_visited']}",
        "",
        "## 3. Contagens do Portal",
        "",
        f"- Protocolos brutos: {payload['portal_summary']['portal_concluded_raw']}",
        f"- Protocolos únicos: {payload['portal_summary']['portal_concluded_unique']}",
        f"- Duplicações na listagem: {payload['portal_summary']['portal_duplicate_listing_rows']}",
        "",
        "## 4. Contagens da planilha",
        "",
        f"- Linhas de projeto: {payload['workbook_summary']['workbook_project_rows']}",
        f"- Protocolos únicos válidos: {payload['workbook_summary']['workbook_unique_protocols']}",
        f"- Linhas sem protocolo válido: {payload['workbook_summary']['workbook_rows_without_valid_protocol']}",
        "",
        "## 5. Reconciliação dos conjuntos",
        "",
        f"- P ∩ W: {payload['set_reconciliation']['matched_unique']}",
        f"- P - W: {payload['set_reconciliation']['missing_in_workbook_unique']}",
        f"- W - P: {payload['set_reconciliation']['workbook_only_unique']}",
        f"- Equação Portal válida: {payload['set_reconciliation']['portal_equation_valid']}",
        f"- Equação planilha válida: {payload['set_reconciliation']['workbook_equation_valid']}",
    ]
    sections = [
        ("6. Protocolos ausentes", "missing_in_workbook"),
        ("7. Duplicações", "duplicate_workbook_protocols"),
        ("8. Abas anuais incorretas", "wrong_year_sheet"),
        ("9. Auditoria da coluna Conclusão", "completion_audit"),
        ("10. Auditoria de Placa e Inversor", "equipment_field_audit"),
        ("11. Registros incompletos", "incomplete_records"),
        ("12. Achados estruturais", "structural_findings"),
    ]
    for title, section in sections:
        lines.extend(["", f"## {title}", ""])
        if count(section) == 0:
            lines.append("Nenhum achado.")
        else:
            lines.append(f"Total: {count(section)}")
    lines.extend(
        [
            "",
            "## 13. Segurança e SHA",
            "",
            f"- SHA antes: {payload['safety']['workbook_sha_before_reconciliation']}",
            f"- SHA depois: {payload['safety']['workbook_sha_after_reconciliation']}",
            f"- Planilha preservada: {payload['safety']['workbook_sha_preserved']}",
            "",
            "## 14. Lote operacional",
            "",
            f"- Protocolos auditados: {payload['operational_batch']['reconciliation_protocols_audited']}",
            f"- Protocolos selecionados: {payload['operational_batch']['operational_protocols_selected']}",
            "",
            "## 15. Decisão",
            "",
            payload["decision"],
        ]
    )
    return "\n".join(lines) + "\n"
