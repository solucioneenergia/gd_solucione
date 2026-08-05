from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from enum import Enum
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright

from automacao_gd.application.contracts import OperationStatus
from automacao_gd.application.preflight import run_preflight
from automacao_gd.domain.errors import PreflightBlockedError
from automacao_gd.domain.models import PortalSolicitation
from automacao_gd.infrastructure.config import Settings, get_settings
from automacao_gd.infrastructure.dates import date_to_excel_datetime, parse_date
from automacao_gd.infrastructure.excel.availability import validate_workbook_availability
from automacao_gd.infrastructure.excel.service import DATE_NUMBER_FORMAT, REQUIRED_COLUMNS
from automacao_gd.infrastructure.logging import logger
from automacao_gd.infrastructure.persistence.atomic import atomic_write_json, atomic_write_text
from automacao_gd.infrastructure.portal.cdp_service import (
    connect_to_existing_edge,
    click_follow_eye_button,
    detail_has_completed_status,
    extract_completion_date,
    find_and_click_next_numeric_page,
    find_portal_page_from_cdp,
    read_current_page_table_with_row_handles,
    return_to_listing,
    wait_detail_loaded,
)


OPEN_TEXT = "EM ABERTO"
COMPLETION_JSON_REPORT_PREFIX = "completion_status_sync"
COMPLETION_APPLY_REPORT_PREFIX = "completion_status_sync_apply"
OPTION5_COMPLETED_PIPELINE_SOURCE = "option_5_completed_pipeline"
MAX_FUTURE_DAYS = 366


class CompletionSyncAction(str, Enum):
    COMPLETION_DATE_UPDATED = "COMPLETION_DATE_UPDATED"
    MARKED_AS_OPEN = "MARKED_AS_OPEN"
    NO_CHANGE = "NO_CHANGE"
    PENDING_REVIEW = "PENDING_REVIEW"


class CompletionSyncStatus(str, Enum):
    COMPLETION_DATE_UPDATED = "COMPLETION_DATE_UPDATED"
    MARKED_AS_OPEN = "MARKED_AS_OPEN"
    NO_CHANGE = "NO_CHANGE"
    PENDING_REVIEW = "PENDING_REVIEW"
    PROTOCOL_NOT_FOUND_IN_WORKBOOK = "PROTOCOL_NOT_FOUND_IN_WORKBOOK"
    PORTAL_COMPLETION_DATE_MISSING = "PORTAL_COMPLETION_DATE_MISSING"
    COMPLETION_DATE_CONFLICT = "COMPLETION_DATE_CONFLICT"
    PORTAL_COMPLETION_STATUS_REGRESSION = "PORTAL_COMPLETION_STATUS_REGRESSION"
    INVALID_COMPLETION_DATE = "INVALID_COMPLETION_DATE"


@dataclass(frozen=True, slots=True)
class CompletionPortalRecord:
    protocol: str
    status: str | None
    completion_date: str | date | datetime | None = None
    source: str = "portal"
    portal_lookup_found: bool = True
    option5_run_at: str | None = None
    option5_action: str | None = None
    pdf_source: str | None = None
    approved_worksheet: str | None = None
    approved_row: int | None = None
    approved_current_completion: str | None = None
    approved_proposed_completion: str | None = None


@dataclass(frozen=True, slots=True)
class WorkbookCompletionRow:
    protocol: str
    worksheet: str
    row: int
    entry_date: date | None
    current_completion: Any
    completion_column: int


ReplaceWorkbook = Callable[[Path, Path], None]


def run_completion_sync(settings: Settings | None = None) -> dict[str, Any]:
    settings = settings or get_settings()
    if not settings.SYNC_COMPLETION_STATUS:
        return {
            "status": OperationStatus.BLOQUEADO.value,
            "operation_message": "Sincronização de conclusão desabilitada.",
            "code": "SYNC_COMPLETION_STATUS_DISABLED",
            "protocol_results": [],
        }
    real_completion_write = not settings.DRY_RUN and settings.APPLY_COMPLETION_STATUS
    preflight = run_preflight(
        settings,
        real_run=real_completion_write,
        require_cdp=not real_completion_write,
    )
    if not preflight.ready:
        preflight.raise_if_blocked()
    if real_completion_write:
        dry_run_payload, dry_run_path = _load_latest_approved_dry_run_payload(
            settings.logs_dir_path
        )
        records = _completion_records_from_approved_dry_run(dry_run_payload)
        return run_completion_sync_from_records(
            records,
            workbook_path=settings.planilha_path,
            dry_run=settings.DRY_RUN,
            apply_changes=settings.APPLY_COMPLETION_STATUS,
            max_protocols=0,
            logs_dir=settings.logs_dir_path,
            available_records=records,
            report_prefix=COMPLETION_APPLY_REPORT_PREFIX,
            source_dry_run_report_path=dry_run_path.name,
        )
    option5_records = _load_option5_completion_records(settings.logs_dir_path)
    if not option5_records:
        return {
            "status": OperationStatus.BLOQUEADO.value,
            "operation_message": "Nenhum protocolo elegível da opção 5 encontrado.",
            "code": "NO_OPTION_5_PROTOCOLS_AVAILABLE",
            "dry_run": settings.DRY_RUN,
            "apply_completion_status": settings.APPLY_COMPLETION_STATUS,
            "protocol_results": [],
        }
    selected_targets = _select_unique_records(
        option5_records,
        settings.MAX_COMPLETION_PROTOCOLS_PER_RUN,
    )
    records = collect_completion_records_from_portal(
        settings,
        target_records=selected_targets,
    )
    return run_completion_sync_from_records(
        records,
        workbook_path=settings.planilha_path,
        dry_run=settings.DRY_RUN,
        apply_changes=settings.APPLY_COMPLETION_STATUS,
        max_protocols=0,
        logs_dir=settings.logs_dir_path,
        available_records=option5_records,
    )


def collect_completion_records_from_portal(
    settings: Settings,
    *,
    target_records: list[CompletionPortalRecord] | None = None,
) -> list[CompletionPortalRecord]:
    playwright = None
    browser = None
    try:
        playwright = sync_playwright().start()
        browser = connect_to_existing_edge(playwright, settings.CDP_ENDPOINT)
        page = find_portal_page_from_cdp(browser, settings.PORTAL_GD_URL)
        if page is None:
            raise RuntimeError("Nenhuma aba do Portal GD foi encontrada via CDP.")
        records = _read_listing_records(
            page,
            settings,
            target_protocols={
                _normalize_protocol(record.protocol)
                for record in target_records or []
                if _normalize_protocol(record.protocol)
            }
            or None,
        )
        if target_records is None:
            return records
        return _merge_option5_targets_with_portal_records(target_records, records)
    except PlaywrightError as exc:
        raise PreflightBlockedError(
            code="CDP_CONNECTION_FAILED",
            user_message="Não foi possível conectar ao Edge via CDP.",
            technical_cause=str(exc),
            stage="completion_sync",
        ) from exc
    finally:
        if browser:
            logger.info("Encerrando conexão CDP sem fechar o Edge aberto manualmente.")
        if playwright:
            playwright.stop()


def run_completion_sync_from_records(
    records: list[CompletionPortalRecord],
    *,
    workbook_path: Path,
    dry_run: bool,
    apply_changes: bool,
    max_protocols: int,
    today: date | None = None,
    logs_dir: Path | None = None,
    replace_workbook: ReplaceWorkbook | None = None,
    available_records: list[CompletionPortalRecord] | None = None,
    report_prefix: str = COMPLETION_JSON_REPORT_PREFIX,
    source_dry_run_report_path: str | None = None,
) -> dict[str, Any]:
    workbook_path = Path(workbook_path)
    today = today or date.today()
    replace_workbook = replace_workbook or _replace_workbook
    selected = _select_unique_records(records, max_protocols)
    sha_before = _sha256(workbook_path) if workbook_path.exists() else None
    workbook_rows = _load_workbook_rows(workbook_path)
    protocol_results = [
        _evaluate_record(record, workbook_rows, today=today)
        for record in selected
    ]
    writable = [
        row
        for row in protocol_results
        if row["action"]
        in {
            CompletionSyncAction.COMPLETION_DATE_UPDATED.value,
            CompletionSyncAction.MARKED_AS_OPEN.value,
        }
    ]
    apply_result: dict[str, Any] = {"updates_applied": 0}
    if writable and not dry_run and apply_changes:
        apply_result = _apply_completion_updates(
            workbook_path=workbook_path,
            writable_results=writable,
            replace_workbook=replace_workbook,
        )
    sha_after = _sha256(workbook_path) if workbook_path.exists() else None
    payload = _build_payload(
        records=records,
        selected=selected,
        protocol_results=protocol_results,
        dry_run=dry_run,
        apply_changes=apply_changes,
        sha_before=sha_before,
        sha_after=sha_after,
        available_records=available_records,
        apply_result=apply_result,
        source_dry_run_report_path=source_dry_run_report_path,
    )
    if logs_dir is not None:
        json_path, markdown_path = _write_reports(Path(logs_dir), payload, prefix=report_prefix)
        payload["json_report_path"] = str(json_path)
        payload["markdown_report_path"] = str(markdown_path)
    return payload


def _evaluate_record(
    record: CompletionPortalRecord,
    workbook_rows: dict[str, WorkbookCompletionRow],
    *,
    today: date,
) -> dict[str, Any]:
    protocol = _normalize_protocol(record.protocol)
    base: dict[str, Any] = {
        "protocol": protocol,
        "portal_status": record.status,
        "portal_completion_date": _date_to_iso(parse_date(record.completion_date)),
        "source": record.source,
        "origin": (
            OPTION5_COMPLETED_PIPELINE_SOURCE
            if _is_option5_source(record.source)
            else record.source
        ),
        "option5_run_at": record.option5_run_at,
        "option5_action": record.option5_action,
        "pdf_source": record.pdf_source,
        "worksheet": None,
        "row": None,
        "current_completion": None,
        "proposed_completion": None,
        "action": CompletionSyncAction.PENDING_REVIEW.value,
        "status": CompletionSyncStatus.PENDING_REVIEW.value,
        "reason": None,
        "write_permission": False,
    }
    if not record.portal_lookup_found:
        return {
            **base,
            "status": CompletionSyncStatus.PENDING_REVIEW.value,
            "reason": "PORTAL_PROTOCOL_NOT_FOUND_IN_COMPLETED_LIST",
        }
    row = workbook_rows.get(protocol)
    if row is None:
        return {
            **base,
            "status": CompletionSyncStatus.PROTOCOL_NOT_FOUND_IN_WORKBOOK.value,
            "reason": CompletionSyncStatus.PROTOCOL_NOT_FOUND_IN_WORKBOOK.value,
        }
    current_report_value = _completion_to_report_value(row.current_completion)
    has_approved_snapshot = (
        record.approved_worksheet is not None
        or record.approved_row is not None
        or record.approved_proposed_completion is not None
    )
    if has_approved_snapshot:
        if (
            record.approved_worksheet is not None
            and record.approved_row is not None
            and (row.worksheet != record.approved_worksheet or row.row != record.approved_row)
        ):
            return {
                **base,
                "worksheet": row.worksheet,
                "row": row.row,
                "current_completion": current_report_value,
                "status": CompletionSyncStatus.PENDING_REVIEW.value,
                "reason": "ROW_CHANGED_AFTER_DRY_RUN",
            }
        if (
            record.approved_current_completion is not None
            and current_report_value != record.approved_current_completion
        ) or (
            record.approved_current_completion is None
            and current_report_value is not None
        ):
            return {
                **base,
                "worksheet": row.worksheet,
                "row": row.row,
                "current_completion": current_report_value,
                "status": CompletionSyncStatus.PENDING_REVIEW.value,
                "reason": "ROW_CHANGED_AFTER_DRY_RUN",
            }

    current_date = parse_date(row.current_completion)
    current_text = _normalize_text(row.current_completion)
    portal_date = parse_date(record.completion_date)
    completed = _is_completed_status(record.status)
    base.update(
        {
            "worksheet": row.worksheet,
            "row": row.row,
            "current_completion": _completion_to_report_value(row.current_completion),
        }
    )

    if completed and portal_date is None:
        if _is_option5_source(record.source):
            if current_text == OPEN_TEXT:
                return {
                    **base,
                    "action": CompletionSyncAction.NO_CHANGE.value,
                    "status": CompletionSyncStatus.NO_CHANGE.value,
                    "reason": "ALREADY_OPEN",
                    "proposed_completion": OPEN_TEXT,
                }
            if current_date is not None:
                return {
                    **base,
                    "status": CompletionSyncStatus.PENDING_REVIEW.value,
                    "reason": "PORTAL_COMPLETION_DATE_REGRESSION",
                }
            return {
                **base,
                "action": CompletionSyncAction.MARKED_AS_OPEN.value,
                "status": CompletionSyncStatus.MARKED_AS_OPEN.value,
                "reason": "POINT_OF_CONNECTION_COMPLETION_DATE_NOT_AVAILABLE",
                "proposed_completion": OPEN_TEXT,
                "write_permission": True,
            }
        return {
            **base,
            "status": CompletionSyncStatus.PENDING_REVIEW.value,
            "reason": CompletionSyncStatus.PORTAL_COMPLETION_DATE_MISSING.value,
        }
    if portal_date is not None:
        invalid_reason = _invalid_completion_date_reason(
            portal_date,
            entry_date=row.entry_date,
            today=today,
        )
        if invalid_reason:
            return {
                **base,
                "status": CompletionSyncStatus.PENDING_REVIEW.value,
                "reason": invalid_reason,
                "proposed_completion": portal_date.isoformat(),
            }
        if current_date is not None and current_date != portal_date:
            return {
                **base,
                "status": CompletionSyncStatus.PENDING_REVIEW.value,
                "reason": CompletionSyncStatus.COMPLETION_DATE_CONFLICT.value,
                "proposed_completion": portal_date.isoformat(),
            }
        if current_date == portal_date:
            return {
                **base,
                "action": CompletionSyncAction.NO_CHANGE.value,
                "status": CompletionSyncStatus.NO_CHANGE.value,
                "reason": "SAME_COMPLETION_DATE",
                "proposed_completion": portal_date.isoformat(),
            }
        return {
            **base,
            "action": CompletionSyncAction.COMPLETION_DATE_UPDATED.value,
            "status": CompletionSyncStatus.COMPLETION_DATE_UPDATED.value,
            "reason": CompletionSyncAction.COMPLETION_DATE_UPDATED.value,
            "proposed_completion": portal_date.isoformat(),
            "write_permission": True,
        }

    if current_date is not None:
        return {
            **base,
            "status": CompletionSyncStatus.PENDING_REVIEW.value,
            "reason": CompletionSyncStatus.PORTAL_COMPLETION_STATUS_REGRESSION.value,
        }
    if current_text == OPEN_TEXT:
        return {
            **base,
            "action": CompletionSyncAction.NO_CHANGE.value,
            "status": CompletionSyncStatus.NO_CHANGE.value,
            "reason": "ALREADY_OPEN",
            "proposed_completion": OPEN_TEXT,
        }
    return {
        **base,
        "action": CompletionSyncAction.MARKED_AS_OPEN.value,
        "status": CompletionSyncStatus.MARKED_AS_OPEN.value,
        "reason": CompletionSyncAction.MARKED_AS_OPEN.value,
        "proposed_completion": OPEN_TEXT,
        "write_permission": True,
    }


def _apply_completion_updates(
    *,
    workbook_path: Path,
    writable_results: list[dict[str, Any]],
    replace_workbook: ReplaceWorkbook,
) -> dict[str, Any]:
    availability = validate_workbook_availability(
        workbook_path,
        require_writable=True,
        stage="sincronização de conclusão",
    )
    if not availability.ok:
        raise PermissionError(availability.user_message)
    original_hash = _sha256(workbook_path)
    backup_path = _create_validated_backup(workbook_path, original_hash)
    wb = load_workbook(workbook_path)
    descriptor, temp_name = tempfile.mkstemp(
        prefix=".completion-sync-",
        suffix=workbook_path.suffix,
        dir=workbook_path.parent,
    )
    os.close(descriptor)
    temp_path = Path(temp_name)
    try:
        before_values = _workbook_value_snapshot(wb)
        sheet_map = _sheet_map(wb)
        for result in writable_results:
            sheet = str(result["worksheet"])
            row = int(result["row"])
            info = sheet_map[sheet]
            cell = info["worksheet"].cell(row=row, column=info["columns"]["Conclusão"])
            proposed = result["proposed_completion"]
            if result["action"] == CompletionSyncAction.COMPLETION_DATE_UPDATED.value:
                cell.value = date_to_excel_datetime(proposed)
                cell.number_format = DATE_NUMBER_FORMAT
            elif result["action"] == CompletionSyncAction.MARKED_AS_OPEN.value:
                cell.value = OPEN_TEXT
        unexpected = _unexpected_value_changes(
            before_values,
            _workbook_value_snapshot(wb),
            {
                (str(result["worksheet"]), int(result["row"]))
                for result in writable_results
            },
            sheet_map,
        )
        if unexpected:
            raise ValueError("Mudança fora da coluna Conclusão: " + json.dumps(unexpected))
        wb.save(temp_path)
        _validate_saved_workbook(temp_path)
        replace_workbook(temp_path, workbook_path)
        return {
            "updates_applied": len(writable_results),
            "backup_created": True,
            "backup_name": backup_path.name,
            "backup_sha": _sha256(backup_path),
            "backup_validated": True,
            "atomic_replace": "SUCCESS",
            "rollback": "NOT_REQUIRED",
            "temporary_residual": False,
        }
    except Exception:
        if workbook_path.exists() and _sha256(workbook_path) != original_hash:
            shutil.copy2(backup_path, workbook_path)
        temp_path.unlink(missing_ok=True)
        raise
    finally:
        wb.close()
        temp_path.unlink(missing_ok=True)


def _create_validated_backup(workbook_path: Path, original_hash: str) -> Path:
    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    backup_path = workbook_path.with_name(
        f"Planilha_pre_completion_sync_{timestamp}{workbook_path.suffix}"
    )
    sequence = 1
    while backup_path.exists():
        backup_path = workbook_path.with_name(
            f"Planilha_pre_completion_sync_{timestamp}_{sequence}{workbook_path.suffix}"
        )
        sequence += 1
    shutil.copy2(workbook_path, backup_path)
    if _sha256(backup_path) != original_hash:
        backup_path.unlink(missing_ok=True)
        raise OSError("Backup da sincronização de conclusão não preservou o SHA original.")
    _validate_saved_workbook(backup_path)
    return backup_path


def _validate_saved_workbook(path: Path) -> None:
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        workbook.close()
    except (OSError, ValueError, KeyError, InvalidFileException) as exc:
        raise OSError(f"Workbook inválido: {path.name}") from exc


def _load_workbook_rows(workbook_path: Path) -> dict[str, WorkbookCompletionRow]:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {workbook_path}")
    wb = load_workbook(workbook_path, data_only=False)
    try:
        rows: dict[str, WorkbookCompletionRow] = {}
        for sheet, info in _sheet_map(wb).items():
            ws = info["worksheet"]
            columns = info["columns"]
            for row in range(info["header_row"] + 1, ws.max_row + 1):
                protocol = _normalize_protocol(ws.cell(row=row, column=columns["Protocolo"]).value)
                if not protocol:
                    continue
                rows[protocol] = WorkbookCompletionRow(
                    protocol=protocol,
                    worksheet=sheet,
                    row=row,
                    entry_date=parse_date(ws.cell(row=row, column=columns["Data de ingresso"]).value),
                    current_completion=ws.cell(row=row, column=columns["Conclusão"]).value,
                    completion_column=columns["Conclusão"],
                )
        return rows
    finally:
        wb.close()


def _sheet_map(wb: Workbook) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for ws in wb.worksheets:
        header_row, columns = _find_required_columns(ws)
        if columns:
            result[ws.title] = {
                "worksheet": ws,
                "header_row": header_row,
                "columns": columns,
            }
    if not result:
        raise ValueError("Nenhuma aba contém as colunas obrigatórias.")
    return result


def _find_required_columns(ws: Worksheet) -> tuple[int, dict[str, int]]:
    wanted = {_normalize_header(name): name for name in REQUIRED_COLUMNS}
    for row in range(1, min(ws.max_row, 10) + 1):
        found: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            normalized = _normalize_header(ws.cell(row=row, column=col).value)
            canonical = wanted.get(normalized)
            if canonical:
                found[canonical] = col
        if set(REQUIRED_COLUMNS).issubset(found):
            return row, found
    return 0, {}


def _load_latest_approved_dry_run_payload(logs_dir: Path) -> tuple[dict[str, Any], Path]:
    candidates = sorted(
        Path(logs_dir).glob(f"{COMPLETION_JSON_REPORT_PREFIX}_*.json"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for path in candidates:
        if path.name.startswith(f"{COMPLETION_APPLY_REPORT_PREFIX}_"):
            continue
        payload = _read_json_dict(path)
        if (
            payload.get("dry_run") is True
            and payload.get("origin") == OPTION5_COMPLETED_PIPELINE_SOURCE
            and int(payload.get("protocols_outside_option5") or 0) == 0
            and int(payload.get("protocols_duplicated") or 0) == 0
            and payload.get("status") == OperationStatus.SUCESSO.value
        ):
            return payload, path
    raise PreflightBlockedError(
        code="APPROVED_COMPLETION_DRY_RUN_NOT_FOUND",
        user_message="Dry-run aprovado da sincronização de conclusão não foi encontrado.",
        technical_cause="Nenhum completion_status_sync_*.json aprovado e de origem option_5_completed_pipeline.",
        stage="completion_sync_apply",
    )


def _completion_records_from_approved_dry_run(
    payload: dict[str, Any],
) -> list[CompletionPortalRecord]:
    records: list[CompletionPortalRecord] = []
    for row in payload.get("protocol_results", []):
        if not isinstance(row, dict):
            continue
        if not row.get("write_permission"):
            continue
        action = str(row.get("action") or "")
        if action not in {
            CompletionSyncAction.COMPLETION_DATE_UPDATED.value,
            CompletionSyncAction.MARKED_AS_OPEN.value,
        }:
            continue
        protocol = _normalize_protocol(row.get("protocol"))
        if not protocol:
            continue
        records.append(
            CompletionPortalRecord(
                protocol=protocol,
                status=row.get("portal_status") or "Solicitação Concluída",
                completion_date=row.get("portal_completion_date"),
                source=str(row.get("source") or OPTION5_COMPLETED_PIPELINE_SOURCE),
                portal_lookup_found=True,
                option5_run_at=row.get("option5_run_at"),
                option5_action=row.get("option5_action"),
                pdf_source=row.get("pdf_source"),
                approved_worksheet=row.get("worksheet"),
                approved_row=int(row["row"]) if row.get("row") is not None else None,
                approved_current_completion=row.get("current_completion"),
                approved_proposed_completion=row.get("proposed_completion"),
            )
        )
    return records


def _load_option5_completion_records(logs_dir: Path) -> list[CompletionPortalRecord]:
    pipeline_path = Path(logs_dir) / "pipeline_cdp_completo.json"
    downloads_path = Path(logs_dir) / "downloads_orcamentos_concluidos_cdp.json"
    processing_path = Path(logs_dir) / "processamento_pdfs_planilha_clientes.json"
    pipeline = _read_json_dict(pipeline_path)
    downloads = _read_json_dict(downloads_path)
    processing = _read_json_dict(processing_path)
    processing_by_protocol = {
        _normalize_protocol(item.get("protocol")): item
        for item in processing.get("results", [])
        if isinstance(item, dict) and _normalize_protocol(item.get("protocol"))
    }
    run_at = (
        str(pipeline.get("finished_at") or pipeline.get("started_at") or "")
        or str(downloads.get("finished_at") or downloads.get("started_at") or "")
        or None
    )
    source_rows = downloads.get("results") or downloads.get("selected_protocols") or []
    if not isinstance(source_rows, list):
        source_rows = []
    records: list[CompletionPortalRecord] = []
    seen: set[str] = set()
    for item in source_rows:
        if not isinstance(item, dict):
            continue
        protocol = _normalize_protocol(item.get("protocol"))
        if not protocol or protocol in seen:
            continue
        status = item.get("status")
        if not _is_completed_status(status):
            continue
        seen.add(protocol)
        processing_result = processing_by_protocol.get(protocol, {})
        pdf_source = _option5_pdf_source(item)
        records.append(
            CompletionPortalRecord(
                protocol=protocol,
                status=str(status) if status is not None else "Solicitação Concluída",
                completion_date=item.get("completion_date"),
                source=OPTION5_COMPLETED_PIPELINE_SOURCE,
                option5_run_at=run_at,
                option5_action=str(
                    processing_result.get("excel_status")
                    or processing_result.get("recommended_action")
                    or item.get("download_status")
                    or item.get("selection_reason")
                    or ""
                )
                or None,
                pdf_source=pdf_source,
            )
        )
    return records


def _read_json_dict(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def _option5_pdf_source(item: dict[str, Any]) -> str | None:
    if item.get("download_status") == "downloaded" or item.get("downloaded_pdf_path"):
        return "baixado"
    if (
        item.get("download_status") == "existing_pdf_after_skip"
        or item.get("existing_pdf_path")
        or item.get("skipped_download")
    ):
        return "reutilizado"
    return None


def _merge_option5_targets_with_portal_records(
    targets: list[CompletionPortalRecord],
    portal_records: list[CompletionPortalRecord],
) -> list[CompletionPortalRecord]:
    portal_by_protocol = {
        _normalize_protocol(record.protocol): record
        for record in portal_records
        if _normalize_protocol(record.protocol)
    }
    merged: list[CompletionPortalRecord] = []
    for target in targets:
        protocol = _normalize_protocol(target.protocol)
        portal_record = portal_by_protocol.get(protocol)
        if portal_record is None:
            merged.append(
                CompletionPortalRecord(
                    protocol=protocol,
                    status=target.status,
                    completion_date=None,
                    source=f"{OPTION5_COMPLETED_PIPELINE_SOURCE}:portal_not_found",
                    portal_lookup_found=False,
                    option5_run_at=target.option5_run_at,
                    option5_action=target.option5_action,
                    pdf_source=target.pdf_source,
                )
            )
            continue
        merged.append(
            CompletionPortalRecord(
                protocol=protocol,
                status=portal_record.status or target.status,
                completion_date=portal_record.completion_date,
                source=f"{OPTION5_COMPLETED_PIPELINE_SOURCE}:{portal_record.source}",
                portal_lookup_found=True,
                option5_run_at=target.option5_run_at,
                option5_action=target.option5_action,
                pdf_source=target.pdf_source,
            )
        )
    return merged


def _read_listing_records(
    page,
    settings: Settings,
    *,
    target_protocols: set[str] | None = None,
) -> list[CompletionPortalRecord]:
    records: list[CompletionPortalRecord] = []
    seen: set[str] = set()
    target_protocols = {
        _normalize_protocol(protocol)
        for protocol in target_protocols or set()
        if _normalize_protocol(protocol)
    } or None
    max_pages = settings.MAX_PORTAL_PAGES if settings.ENABLE_PORTAL_PAGINATION else 1
    max_pages = max_pages or 1
    current_page = 1
    listing_url = page.url
    for _ in range(max_pages):
        page_rows = read_current_page_table_with_row_handles(page)
        for row in page_rows:
            record: PortalSolicitation = row["record"]
            protocol = _normalize_protocol(record.protocol)
            if target_protocols and protocol not in target_protocols:
                continue
            if protocol and protocol not in seen:
                seen.add(protocol)
                records.append(_completion_record_from_detail(
                    page=page,
                    row=row,
                    protocol=protocol,
                    fallback_status=record.status,
                    fallback_completion_date=record.completion_date,
                    listing_url=listing_url,
                    page_number=current_page,
                )
                )
        if target_protocols and target_protocols.issubset(seen):
            break
        if (
            not target_protocols
            and len(records) >= settings.MAX_COMPLETION_PROTOCOLS_PER_RUN
        ):
            break
        click = find_and_click_next_numeric_page(page, current_page)
        if not click.get("clicked"):
            break
        current_page += 1
        page.wait_for_timeout(800)
    return records


def _completion_record_from_detail(
    *,
    page,
    row: dict[str, Any],
    protocol: str,
    fallback_status: str | None,
    fallback_completion_date: str | date | datetime | None,
    listing_url: str,
    page_number: int,
) -> CompletionPortalRecord:
    before_pages = list(page.context.pages)
    try:
        click_follow_eye_button(row["row_locator"], row.get("action_cell_index"))
        page.wait_for_timeout(800)
        new_pages = [
            candidate for candidate in page.context.pages if candidate not in before_pages
        ]
        detail_page = new_pages[-1] if new_pages else page
        wait_detail_loaded(detail_page, protocol)
        completed = detail_has_completed_status(detail_page) or _is_completed_status(
            fallback_status
        )
        completion_date = extract_completion_date(detail_page) or fallback_completion_date
        if detail_page is not page:
            detail_page.close()
        else:
            return_to_listing(page, listing_url)
        return CompletionPortalRecord(
            protocol=protocol,
            status="Concluído" if completed else fallback_status,
            completion_date=completion_date,
            source=f"portal_detail_page_{page_number}",
        )
    except Exception as exc:
        logger.warning(
            "Falha ao ler detalhe de conclusão; usando dados da tabela: "
            f"protocol={protocol}; error={type(exc).__name__}"
        )
        return CompletionPortalRecord(
            protocol=protocol,
            status=fallback_status,
            completion_date=fallback_completion_date,
            source=f"portal_listing_page_{page_number}",
        )


def _select_unique_records(
    records: list[CompletionPortalRecord],
    max_protocols: int,
) -> list[CompletionPortalRecord]:
    selected: list[CompletionPortalRecord] = []
    seen: set[str] = set()
    for record in records:
        protocol = _normalize_protocol(record.protocol)
        if not protocol or protocol in seen:
            continue
        seen.add(protocol)
        if max_protocols > 0 and len(selected) >= max_protocols:
            break
        selected.append(
            CompletionPortalRecord(
                protocol=protocol,
                status=record.status,
                completion_date=record.completion_date,
                source=record.source,
                portal_lookup_found=record.portal_lookup_found,
                option5_run_at=record.option5_run_at,
                option5_action=record.option5_action,
                pdf_source=record.pdf_source,
                approved_worksheet=record.approved_worksheet,
                approved_row=record.approved_row,
                approved_current_completion=record.approved_current_completion,
                approved_proposed_completion=record.approved_proposed_completion,
            )
        )
    return selected


def _invalid_completion_date_reason(
    completion: date,
    *,
    entry_date: date | None,
    today: date,
) -> str | None:
    if entry_date is not None and completion < entry_date:
        return "COMPLETION_BEFORE_ENTRY_DATE"
    if completion > today + timedelta(days=MAX_FUTURE_DAYS):
        return "COMPLETION_DATE_TOO_FAR_IN_FUTURE"
    return None


def _build_payload(
    *,
    records: list[CompletionPortalRecord],
    selected: list[CompletionPortalRecord],
    protocol_results: list[dict[str, Any]],
    dry_run: bool,
    apply_changes: bool,
    sha_before: str | None,
    sha_after: str | None,
    available_records: list[CompletionPortalRecord] | None = None,
    apply_result: dict[str, Any] | None = None,
    source_dry_run_report_path: str | None = None,
) -> dict[str, Any]:
    updated_dates = _count_action(protocol_results, CompletionSyncAction.COMPLETION_DATE_UPDATED)
    marked_open = _count_action(protocol_results, CompletionSyncAction.MARKED_AS_OPEN)
    no_change = _count_action(protocol_results, CompletionSyncAction.NO_CHANGE)
    pending = sum(1 for row in protocol_results if row["status"] == CompletionSyncStatus.PENDING_REVIEW.value or row["action"] == CompletionSyncAction.PENDING_REVIEW.value)
    available_records = available_records if available_records is not None else records
    selected_protocols = [
        _normalize_protocol(row.protocol)
        for row in selected
        if _normalize_protocol(row.protocol)
    ]
    apply_result = apply_result or {"updates_applied": 0}
    return {
        "status": OperationStatus.SUCESSO.value,
        "operation_message": (
            "Sincronização de conclusão planejada em modo simulação."
            if dry_run or not apply_changes
            else "Sincronização de conclusão concluída."
        ),
        "mode": "DRY_RUN" if dry_run or not apply_changes else "APPLY",
        "origin": OPTION5_COMPLETED_PIPELINE_SOURCE,
        "dry_run": dry_run,
        "apply_completion_status": apply_changes,
        "total_protocols_available": len({_normalize_protocol(row.protocol) for row in available_records if _normalize_protocol(row.protocol)}),
        "total_option5_protocols_available": len({_normalize_protocol(row.protocol) for row in available_records if _normalize_protocol(row.protocol)}),
        "total_protocols_analyzed": len(selected),
        "selected_protocols": selected_protocols,
        "protocols_outside_option5": 0,
        "protocols_added_after_limit": 0,
        "protocols_duplicated": len(selected_protocols) - len(set(selected_protocols)),
        "total_updated_dates": updated_dates,
        "total_marked_open": marked_open,
        "total_updates_applied": int(apply_result.get("updates_applied") or 0),
        "total_no_change": no_change,
        "total_pending_review": pending,
        "total_not_found": sum(
            1
            for row in protocol_results
            if row["status"] == CompletionSyncStatus.PROTOCOL_NOT_FOUND_IN_WORKBOOK.value
        ),
        "workbook_sha_before": sha_before,
        "workbook_sha_after": sha_after,
        "source_dry_run_report": source_dry_run_report_path,
        "strong_confirmation_received": (not dry_run and apply_changes),
        "backup_created": bool(apply_result.get("backup_created", False)),
        "backup_name": apply_result.get("backup_name"),
        "backup_sha": apply_result.get("backup_sha"),
        "backup_validated": bool(apply_result.get("backup_validated", False)),
        "atomic_replace": apply_result.get("atomic_replace", "NOT_EXECUTED"),
        "unexpected_changes": 0,
        "rollback": apply_result.get("rollback", "NOT_REQUIRED"),
        "temporary_residual": bool(apply_result.get("temporary_residual", False)),
        "protocol_results": protocol_results,
    }


def _write_reports(
    logs_dir: Path,
    payload: dict[str, Any],
    *,
    prefix: str = COMPLETION_JSON_REPORT_PREFIX,
) -> tuple[Path, Path]:
    logs_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    json_path = logs_dir / f"{prefix}_{timestamp}.json"
    markdown_path = logs_dir / f"{prefix}_{timestamp}.md"
    atomic_write_json(json_path, payload, private=True)
    lines = [
        "# Sincronização de datas de conclusão",
        "",
        f"- Status: {payload['status']}",
        f"- Modo: {payload.get('mode')}",
        f"- Origem: {payload.get('origin')}",
        f"- Protocolos disponíveis pela opção 5: {payload.get('total_option5_protocols_available')}",
        f"- Protocolos analisados: {payload['total_protocols_analyzed']}",
        f"- Datas preenchidas: {payload['total_updated_dates']}",
        f"- Marcados EM ABERTO: {payload['total_marked_open']}",
        f"- Aplicados: {payload.get('total_updates_applied')}",
        f"- NO_CHANGE: {payload['total_no_change']}",
        f"- Pendências: {payload['total_pending_review']}",
        f"- Protocolos não localizados: {payload['total_not_found']}",
        f"- Protocolos externos à opção 5: {payload.get('protocols_outside_option5')}",
        f"- Protocolos adicionados após limite: {payload.get('protocols_added_after_limit')}",
        f"- Protocolos duplicados: {payload.get('protocols_duplicated')}",
        f"- SHA antes: `{payload.get('workbook_sha_before')}`",
        f"- SHA backup: `{payload.get('backup_sha')}`",
        f"- SHA depois: `{payload.get('workbook_sha_after')}`",
        f"- Backup: {payload.get('backup_name')}",
        f"- Substituição atômica: {payload.get('atomic_replace')}",
        f"- Rollback: {payload.get('rollback')}",
        f"- Temporário residual: {payload.get('temporary_residual')}",
        f"- Mudanças fora da allowlist: {payload['unexpected_changes']}",
    ]
    atomic_write_text(markdown_path, "\n".join(lines) + "\n", private=True)
    return json_path, markdown_path


def _workbook_value_snapshot(wb: Workbook) -> dict[tuple[str, int, int], Any]:
    values: dict[tuple[str, int, int], Any] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                values[(ws.title, cell.row, cell.column)] = cell.value
    return values


def _unexpected_value_changes(
    before: dict[tuple[str, int, int], Any],
    after: dict[tuple[str, int, int], Any],
    allowed_rows: set[tuple[str, int]],
    sheet_map: dict[str, dict[str, Any]],
) -> list[dict[str, Any]]:
    diffs: list[dict[str, Any]] = []
    for key in set(before) | set(after):
        if before.get(key) == after.get(key):
            continue
        sheet, row, column = key
        completion_col = sheet_map.get(sheet, {}).get("columns", {}).get("Conclusão")
        if (sheet, row) in allowed_rows and column == completion_col:
            continue
        diffs.append({"sheet": sheet, "row": row, "column": column})
    return diffs


def _replace_workbook(src: Path, dst: Path) -> None:
    os.replace(src, dst)


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _count_action(rows: list[dict[str, Any]], action: CompletionSyncAction) -> int:
    return sum(1 for row in rows if row["action"] == action.value)


def _normalize_protocol(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]+", "", text.upper())


def _normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().upper()


def _is_option5_source(source: str | None) -> bool:
    return str(source or "").startswith(OPTION5_COMPLETED_PIPELINE_SOURCE)


def _is_completed_status(status: str | None) -> bool:
    normalized = _normalize_text(status)
    return "CONCLUID" in normalized


def _date_to_iso(value: date | None) -> str | None:
    return value.isoformat() if value else None


def _completion_to_report_value(value: Any) -> str | None:
    parsed = parse_date(value)
    if parsed:
        return parsed.isoformat()
    text = str(value).strip() if value not in (None, "") else ""
    return text or None
